"""Mise à jour des logigrammes du FPPN (« FLOW_CHART_COTISATION_ET_PRODUITS_DEPARGNE_1.pptx ») par Adenka.

Sur chaque logigramme : une pastille rouge numérotée sur chaque étape où un contrôle manque ou un écart est constaté.
Après chaque logigramme : une diapositive « Mise à jour Adenka » (n°, étape, constat, source, action proposée).
Le fichier d'origine n'est pas modifié.

Usage : python3 annoter_logigrammes_fppn.py <logigrammes_fppn.pptx> <sortie.pptx>
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

SRC, SORTIE = sys.argv[1:3]
ROUGE, NOIR, GRIS = "C00000", "060606", "F2F2F2"
ENT = "Entretien du 19/10/2026 avec le responsable des cotisations"
# diapositive : [(texte de la forme à marquer, étape, constat, source, action proposée)]
NOTES = {
    0: [("Création du mécano", "Création du mécano et ajout à la fiche",
         "Nouvelle recrue sans mécano : le premier précompte peut être décalé d'un mois (déclaration non encore documentée).", ENT,
         "Dater chaque souscription et vérifier le mois du premier précompte."),
        ("Création du fichier de précompte", "Création du fichier de précompte (service du personnel)",
         "Écart : l'entretien indique que le responsable des cotisations du FPPN produit le fichier aller de l'EFC et du PEP et le dépose sur la plateforme de partage.", ENT,
         "Préciser qui produit le fichier aller et qui le transmet à la Solde."),
        ("Réception du fichier", "Réception du fichier retour numérique",
         "Contrôle absent : aucun rapprochement du fichier aller et du fichier retour (rejets, écarts de montant).", ENT + " ; rapport, partie 3",
         "Rapprochement mensuel signé, avec l'état des rejets."),
        ("Comptabilisation des opérations", "Comptabilisation des opérations de prélèvement",
         "Contrôle absent : aucun intérêt calculé ni comptabilisé (3 278 913 238 FCFA dus au 31/08/2026), aucun relevé par épargnant, aucun rapprochement avec les comptes 41195100 et 41195200.",
         "Balances 2020 au 31/08/2026 ; rapport, partie 6", "Calcul annuel des intérêts au taux arrêté et relevé annuel par épargnant.")],
    1: [("Calcul des cotations", "Calcul des cotations",
         "La convention d'apporteur d'affaires avec le GNA n'a pas été transmise.", "PBC du 23/10/2026", "Transmettre la convention et les décomptes annuels du GNA."),
        ("Fichier retour", "Fichier retour (Trésor)",
         "Fichier retour de janvier 2026 manquant dans les fichiers reçus.", "Fichiers de précomptes reçus le 15/10/2026 ; rapport, partie 3", "Transmettre le fichier retour de janvier 2026."),
        ("Comptabilisation du pourcentage encaissé", "Comptabilisation du pourcentage encaissé (20 %)",
         "Contrôle absent : l'assiette de la commission de 20 % (primes précomptées seules ou toutes les primes) n'est pas documentée.", "Rapport, question au FPPN n° 19",
         "Rapprocher chaque année la commission des primes collectées.")],
    2: [("Adhésion sur", "Adhésion sur fiche d'adhésion",
         "La durée de la garantie santé (à vie ou limitée) n'est pas documentée. La date de naissance ne figure que dans les fichiers de février à avril 2024 (105 retraités).",
         "CONVENTION DE PARTENARIAT ENTRE IPS-CGRAE ET LE FPPN.pdf ; fichiers retour de la CGRAE", "Fixer la durée de la garantie et transmettre la date de naissance de chaque retraité."),
        ("Création du fichier de précompte", "Création du fichier de précompte et calcul des cotisations",
         "Barème reçu le 23/10/2026 : 2,5 % de (pension de base - MUGEFCI) + 5 000 à 15 000 FCFA selon le grade. Le prélèvement est limité à la quotité cessible : les retraités sans quotité ne sont pas précomptés, et leur liste n'a pas été transmise.",
         "« BAREME_CGRAE.xlsx » ; convention IPS-CGRAE", "Ajouter au logigramme l'étape « quotité cessible suffisante ? » et suivre les souscripteurs non précomptés."),
        ("Fichier retour", "Fichier retour numérique",
         "Fichiers retour de janvier et de juin 2024 manquants.", "Fichiers de précomptes reçus ; PBC du 23/10/2026", "Transmettre les deux fichiers."),
        ("Comptabilisation du", "Comptabilisation du virement",
         "Contrôle absent : le virement de la CGRAE n'est pas rapproché du fichier retour.", ENT, "Rapprocher chaque virement du fichier retour du mois.")],
    3: [("Carte de santé", "Carte de santé FPPN ?",
         "Le fichier des cartes de santé actives au 31/08/2026 n'a pas été transmis.", "PBC du 23/10/2026, n° 25", "Transmettre le fichier des cartes actives, ayants droit compris."),
        ("Création du fichier de précompte", "Création du fichier de précompte (service du personnel)",
         "Taux retenus : 3,77 % (maladie) et 1,23 % (décès), au lieu de 3,5 % et 1,5 % fixés par les statuts (art. 17). Le fichier aller n'est pas transmis au FPPN : le taux de recouvrement (au moins 90 %) ne peut pas être mesuré.",
         "Fichiers de précomptes 936 et 8Z6 ; statuts du FPPN édités le 15/06/2023, art. 17 ; Décision n° 05/2022", "Obtenir chaque mois le fichier aller et la décision qui fixe la clé de 5 %."),
        ("Fichier retour papier", "Fichier retour papier et numérique",
         "Décembre 2024 : fichiers « 122024_effectif_police_936_DECEMBRE 2024.xlsx » et « FPPN_ALLOC_DECES_8Z6_DECEMBRE_2024.xlsx » vides (0 octet).",
         "Fichiers reçus le 15/10/2026", "Transmettre les fichiers de décembre 2024.")],
}


def formes(shapes, ox=0, oy=0, sx=1.0, sy=1.0):
    """Toutes les formes, avec leur position absolue (les groupes sont dépliés)."""
    for sh in shapes:
        if sh.shape_type == 6:
            x_ = sh._element.grpSpPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
            off, ext = x_.find("{*}off"), x_.find("{*}ext")
            choff, chext = x_.find("{*}chOff"), x_.find("{*}chExt")
            fx = int(ext.get("cx")) / max(1, int(chext.get("cx")))
            fy = int(ext.get("cy")) / max(1, int(chext.get("cy")))
            nx, ny = ox + (int(off.get("x")) - 0) * sx, oy + (int(off.get("y")) - 0) * sy
            yield from formes(sh.shapes, nx - int(choff.get("x")) * fx * sx, ny - int(choff.get("y")) * fy * sy, sx * fx, sy * fy)
        else:
            if sh.left is None:
                continue
            yield sh, ox + sh.left * sx, oy + sh.top * sy, sh.width * sx


def texte(cell, t, taille, gras=False, coul=NOIR, align=PP_ALIGN.LEFT):
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = ""
    p.alignment = align
    r = p.add_run()
    r.text = t
    r.font.name = "Helvetica"
    r.font.size = Pt(taille)
    r.font.bold = gras
    r.font.color.rgb = RGBColor.from_string(coul)


prs = Presentation(SRC)
blank = min(prs.slide_layouts, key=lambda l: len(l.placeholders))
ids = list(prs.slides._sldIdLst)
nouvelles = []
for k, sl in enumerate(list(prs.slides)):
    titre = next((sh.text_frame.text.strip() for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip().startswith("Flow chart")), f"Logigramme {k + 1}")
    lignes = []
    for n, (cherche, etape, constat, source, action) in enumerate(NOTES.get(k, []), 1):
        cands = [(sh, x, y, w) for sh, x, y, w in formes(sl.shapes) if sh.has_text_frame and sh.text_frame.text.strip().replace("\n", " ").startswith(cherche)]
        if cherche == "Fichier retour":
            cands = [c_ for c_ in cands if c_[0].text_frame.text.strip() in ("Fichier retour", "Fichier retour numérique")] or cands
        assert cands, (k, cherche)
        sh, x, y, w = cands[0]
        d = Cm(0.55)
        o = sl.shapes.add_shape(MSO_SHAPE.OVAL, int(x + w - d / 2), int(y - d / 2), d, d)
        o.fill.solid()
        o.fill.fore_color.rgb = RGBColor.from_string(ROUGE)
        o.line.color.rgb = RGBColor.from_string("FFFFFF")
        o.shadow.inherit = False
        o.name = f"Adenka point {n}"
        tf = o.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        texte(o, str(n), 9, True, "FFFFFF", PP_ALIGN.CENTER)
        lignes.append((str(n), etape, constat, source, action))
    if not lignes:
        continue
    ns = prs.slides.add_slide(blank)
    for ph in list(ns.placeholders):
        ph._element.getparent().remove(ph._element)
    tb = ns.shapes.add_textbox(Cm(1.2), Cm(0.6), Cm(31.4), Cm(1.2))
    texte(tb, "Mise à jour Adenka du 23/10/2026 : " + titre.replace("Flow chart du processus de comptabilisation ", "comptabilisation "), 16, True)
    st = ns.shapes.add_textbox(Cm(1.2), Cm(1.8), Cm(31.4), Cm(0.8))
    texte(st, "Pastilles rouges du logigramme précédent : contrôle absent ou écart constaté sur les documents reçus. Constats détaillés dans le rapport actuariel et statistique du 23/10/2026.", 10, False, "53565A")
    larg = [1.2, 5.4, 11.4, 6.6, 6.8]
    t = ns.shapes.add_table(len(lignes) + 1, 5, Cm(1.2), Cm(3.0), Cm(sum(larg)), Cm(1.0 * (len(lignes) + 1))).table
    for j, w in enumerate(larg):
        t.columns[j].width = Cm(w)
    for j, e in enumerate(["N°", "Étape du logigramme", "Constat ou contrôle absent", "Source", "Action proposée"]):
        c = t.cell(0, j)
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor.from_string(NOIR)
        texte(c, e, 10, True, "FFFFFF")
    for i, lg in enumerate(lignes, 1):
        for j, v in enumerate(lg):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string(GRIS if i % 2 == 0 else "FFFFFF")
            texte(c, v, 9, j == 0, ROUGE if j == 0 else NOIR, PP_ALIGN.CENTER if j == 0 else PP_ALIGN.LEFT)
    nouvelles.append((ids[k], prs.slides._sldIdLst[-1]))
# chaque diapositive de mise à jour suit son logigramme
lst = prs.slides._sldIdLst
for apres, nouv in nouvelles:
    lst.remove(nouv)
    apres.addnext(nouv)
prs.save(SORTIE)
print(SORTIE, "diapositives :", len(prs.slides), "| points :", sum(len(v) for v in NOTES.values()))
