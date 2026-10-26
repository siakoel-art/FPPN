"""Graphiques Word natifs (modifiables : clic droit > Modifier les données) pour python-docx.

Le graphique est construit avec python-pptx (XML DrawingML du graphique et classeur Excel des données), puis rangé
dans le paquet Word : pièce /word/charts/chartN.xml, classeur incorporé, relation depuis le document et image en ligne.
"""
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml as docx_parse
from docx.shared import Cm
from lxml import etree
from pptx.chart.chart import Chart
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.oxml import parse_xml
from pptx.util import Pt

CT_CHART = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
CT_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TYPES = {"barres": XL_CHART_TYPE.COLUMN_CLUSTERED, "lignes": XL_CHART_TYPE.LINE_MARKERS, "barres_h": XL_CHART_TYPE.BAR_CLUSTERED}
_compteur = {"n": 0}


def _chart_xml(spec):
    cd = CategoryChartData(number_format=spec.get("format", "0.00"))
    cd.categories = spec["categories"]
    for s in spec["series"]:
        cd.add_series(s["nom"], s["valeurs"])
    cs = parse_xml(cd.xml_bytes(TYPES[spec["type"]]))
    ch = Chart(cs, None)
    ch.font.name = "Helvetica"
    ch.font.size = Pt(spec.get("taille_police", 9))
    ch.has_title = False
    ch._chartSpace.chart.get_or_add_autoTitleDeleted().set("val", "1")
    ch.has_legend = spec.get("legende", True)
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    plot = ch.plots[0]
    if spec["type"] != "lignes":
        plot.gap_width = spec.get("ecart", 60)
        plot.overlap = spec.get("chevauchement", 0)
    for se, s in zip(plot.series, spec["series"]):
        coul = RGBColor.from_string(s["couleur"])
        if spec["type"] == "lignes":
            se.format.line.color.rgb = coul
            se.format.line.width = Pt(s.get("epaisseur", 2.25))
            se.smooth = False
            if s.get("tirets"):
                from pptx.enum.dml import MSO_LINE_DASH_STYLE
                se.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            if s.get("sans_marqueur"):
                from pptx.enum.chart import XL_MARKER_STYLE
                se.marker.style = XL_MARKER_STYLE.NONE
            else:
                se.marker.format.fill.solid()
                se.marker.format.fill.fore_color.rgb = coul
                se.marker.format.line.color.rgb = coul
        else:
            se.format.fill.solid()
            se.format.fill.fore_color.rgb = coul
        if s.get("etiquettes", spec.get("etiquettes", True)):
            dl = se.data_labels
            dl.number_format = spec.get("format_etiquettes", spec.get("format", "0.00"))
            dl.number_format_is_linked = False
            dl.font.size = Pt(spec.get("taille_etiquettes", 7.5))
            dl.show_value = True
            if spec["type"] == "lignes":
                dl.position = XL_LABEL_POSITION.ABOVE
            else:
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
    va.tick_labels.font.size = Pt(8)
    va.tick_labels.number_format = spec.get("format_axe", "0")
    va.tick_labels.number_format_is_linked = False
    va.format.line.fill.background()
    if "axe_min" in spec:
        va.minimum_scale = spec["axe_min"]
    if "axe_max" in spec:
        va.maximum_scale = spec["axe_max"]
    if spec.get("titre_axe"):
        va.has_title = True
        va.axis_title.text_frame.text = spec["titre_axe"]
        va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        va.axis_title.text_frame.paragraphs[0].runs[0].font.bold = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(8)
    ca.format.line.color.rgb = RGBColor(0x06, 0x06, 0x06)
    if spec.get("etiquettes_axe_bas"):
        from pptx.enum.chart import XL_TICK_LABEL_POSITION
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    return cs, cd.xlsx_blob


def graphique(r, spec, titre, source, largeur_cm=16.0, hauteur_cm=8.0):
    """Ajoute au rapport (RapportAdenka) une figure numérotée avec un graphique Word natif."""
    _compteur["n"] += 1
    n = _compteur["n"]
    r.n_fig += 1
    r._legende("Figure", r.n_fig, titre)
    pkg = r.doc.part.package
    cs, xlsx = _chart_xml(spec)
    xlsx_part = Part(PackURI(f"/word/embeddings/Microsoft_Excel_Worksheet{n}.xlsx"), CT_XLSX, xlsx, pkg)
    chart_part = Part(PackURI(f"/word/charts/chart{n}.xml"), CT_CHART, b"", pkg)
    rid_x = chart_part.relate_to(xlsx_part, RT.PACKAGE)
    ext = cs.get_or_add_externalData()
    ext.rId = rid_x
    chart_part._blob = etree.tostring(cs, xml_declaration=True, encoding="UTF-8", standalone=True)
    rid_c = r.doc.part.relate_to(chart_part, RT.CHART)
    cx, cy = int(Cm(largeur_cm)), int(Cm(hauteur_cm))
    p = r.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    inline = docx_parse(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:docPr id="{9000 + n}" name="Graphique {n}"/><wp:cNvGraphicFramePr/>'
        '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{rid_c}"/>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')
    p._p.append(inline)
    r._source(source)
    return p
