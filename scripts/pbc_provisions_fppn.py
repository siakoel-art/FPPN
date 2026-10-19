# -*- coding: utf-8 -*-
"""
FPPN - Revue des provisions (rapport v5) : items 71 à 79 de la PBC générale et section I du guide d'entretien
(onglet « 2. Provisionnements », chef comptable).
Utilisation : python pbc_provisions_fppn.py <pbc_generale.xlsx> <guide.xlsx>   (fichiers modifiés en place)
"""
import datetime as dt
import sys
from copy import copy

import openpyxl
from openpyxl.worksheet.cell_range import MultiCellRange

PBC, GUIDE = sys.argv[1:3]
D_DEM, D_LIM = dt.datetime(2026, 10, 19), dt.datetime(2026, 10, 26)
COMMENT = "Revue des provisions (rapport v5) : analyse en attente de ce document."
ITEMS = [
    ("Bases des prestations maladie et décès sur le plus grand nombre d'années disponibles (au moins depuis 2019) jusqu'au 31/08/2026, une ligne par dossier : date de survenance (soins ou décès), date de déclaration, date de paiement, montant, statut (payé, en attente, rejeté). Nécessaires à la PSAP dossier par dossier et à la méthode CIMA des sinistres tardifs.", "Urgent", "Electronique"),
    ("Note de calcul des provisions pour prestations à payer maladie et obsèques (comptes 15801000 et 15802000), liste des dossiers qui les composent, et raison de l'absence de mise à jour depuis 2020.", "Urgent", None),
    ("État des impayés du Grenier par tranche de retard au 31/12/2025 et au 31/08/2026, et politique de dépréciation des prêts en retard (décision, taux appliqués).", "Moyen", "Electronique"),
    ("Politique de crédit et de recouvrement du Grenier : règles d'octroi, recouvrement des agents sortis ou non précomptés, traitement des prêts à mécano anormal.", "Moyen", None),
    ("Détail des mouvements du compte 47111106 « GNA ASSURANCE » de 2023 au 31/08/2026, avec les pièces justificatives.", "Moyen", "Electronique"),
    ("Fichier du personnel du FPPN au 31/12/2025 et au 31/08/2026 : date de naissance, date d'embauche, salaire mensuel, catégorie, âge de départ prévu.", "Moyen", "Electronique"),
    ("Convention collective et accord d'entreprise applicables au personnel du FPPN, articles fixant l'indemnité de départ à la retraite.", "Moyen", None),
    ("Note de calcul de la provision pour indemnités de départ à la retraite (compte 196) : méthode, hypothèses, liste des salariés.", "Moyen", None),
    ("Règles de la garantie santé des retraités : durée (à vie ou limitée), conditions, cotisation, décision qui la fonde.", "Urgent", None),
]
POINTS = [
    ("Provisions figées", "Les provisions maladie et obsèques ont le même montant de 2020 au 31/08/2026, sans dotation ni reprise (comptes 6901, 7901 absents).", "Transmettre la note de calcul et la liste des dossiers (item 72)."),
    ("Compte GNA 47111106", "Compte mouvementé au débit depuis 2023, alors que les primes vont directement à GNA selon l'entretien du 19/10/2026.", "Expliquer la nature des mouvements (item 75)."),
    ("Provision pour départ à la retraite", "Provision reprise et dotée en totalité chaque année, méthode non communiquée.", "Transmettre la note de calcul et le fichier du personnel (items 76 à 78)."),
]
QUESTIONS = [
    ("33", "Provisions", "Comment les provisions maladie et obsèques ont-elles été calculées, et pourquoi n'ont-elles pas changé depuis 2020 ?", "Les montants sont identiques au franc près de 2020 au 31/08/2026.",
     "Attendu : note de calcul, liste des dossiers. Alerte : aucun calcul, montant hérité.", "Dossier par dossier, puis sinistres tardifs (méthode CIMA)", "PSAP = dossiers connus + tardifs", "Règlement d'exécution n° 004/2022/COM/UEMOA, art. 10", "Note de calcul (PBC, item 72)", "Incontournable"),
    ("34", "Provisions", "Les factures des cliniques enregistrées en dettes (compte 40110000) sont-elles aussi comprises dans la provision maladie ?", "Éviter un double compte entre dettes et provision.",
     "Attendu : non, la provision couvre les soins non encore facturés.", "", "", "Règlement d'exécution n° 004/2022, art. 10 ; PCMS", "", "Incontournable"),
    ("35", "Grenier", "Le FPPN déprécie-t-il les prêts du Grenier en retard ? Selon quels taux ?", "Aucune dépréciation passée depuis 2022 alors que des prêts ont plus d'un an de retard.",
     "Attendu : politique écrite. Sinon, barème BCEAO des SFD en référence.", "Encours par tranche × taux", "Dépréciation = Σ encours(k) × taux(k)", "PCMS, principe de prudence ; BCEAO (SFD)", "Politique de dépréciation (PBC, item 73)", "Incontournable"),
    ("36", "Flotte automobile", "Quelle est la nature des mouvements du compte 47111106 « GNA ASSURANCE » ?", "Les primes iraient directement à GNA, mais le compte est mouvementé depuis 2023.",
     "Attendu : avances, remboursements de trop-perçus ou primes.", "", "", "Convention avec GNA", "Détail du compte (PBC, item 75)", "Si le temps le permet"),
    ("37", "Personnel", "Comment la provision pour départ à la retraite (compte 196) est-elle calculée ?", "Provision reprise et dotée en totalité chaque année ; méthode inconnue.",
     "Attendu : barème de la convention collective, fichier du personnel, hypothèses.", "Unités de crédit projetées", "PIDR = Σ IDR × probabilité de présence × actualisation × prorata d'ancienneté", "SYSCOHADA révisé ; convention collective interprofessionnelle", "Note de calcul (PBC, items 76 à 78)", "Si le temps le permet"),
    ("38", "Retraités", "La garantie santé des retraités est-elle accordée à vie ?", "Si oui, une provision mathématique est due.",
     "Attendu : règlement de la garantie.", "Valeur actuelle des soins futurs", "PM = Σ coût × probabilité de survie × actualisation", "Règlement d'exécution n° 004/2022, art. 11", "Règles de la garantie (PBC, item 79)", "Incontournable"),
]

# PBC générale ---------------------------------------------------------------------------------------
wb = openpyxl.load_workbook(PBC)
ws = wb["PBC Actuariat FPPN"]
num = [r for r in range(11, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith("=")]
dern, COLS = num[-1], range(2, 14)
assert len(num) == 70, len(num)
sect = next(r for r in range(dern, 9, -1) if ws.cell(r, 3).value and not ws.cell(r, 2).value)
st_s = {c: copy(ws.cell(sect, c)._style) for c in COLS}
st_i = {c: copy(ws.cell(dern, c)._style) for c in COLS}
r = dern + 2
for c in COLS:
    ws.cell(r, c)._style = copy(st_s[c])
ws.cell(r, 3, "Provisions et engagements (revue du rapport v5, 19/10/2026)")
ws.row_dimensions[r].height = ws.row_dimensions[sect].height
prec = dern
for texte, prio, support in ITEMS:
    r += 1
    for c in COLS:
        ws.cell(r, c)._style = copy(st_i[c])
    ws.cell(r, 2, f"=+B{prec}+1")
    ws.cell(r, 3, texte)
    ws.cell(r, 5, prio)
    ws.cell(r, 7, D_DEM)
    ws.cell(r, 8, D_LIM)
    ws.cell(r, 10, support)
    ws.cell(r, 11, COMMENT)
    ws.cell(r, 13, "Non Reçu")
    ws.row_dimensions[r].height = 36 if len(texte) < 230 else 51
    prec = r
plage = f"$M$11:$M${prec}"
ws["M7"] = f'=IFERROR(COUNTIF({plage},"Reçu")/(COUNTA({plage})-COUNTIF({plage},"Non Applicable")),0)'
for dv in ws.data_validations.dataValidation:
    col = str(dv.sqref).split(":")[0][0]
    dv.sqref = MultiCellRange(f"{col}11:{col}{prec}")
sy = wb["Synthèse du pointage"]
for k, st in ((6, "Reçu"), (7, "Partiellement Reçu"), (8, "Non Reçu")):
    sy.cell(k, 3, f"=COUNTIF('PBC Actuariat FPPN'!{plage},\"{st}\")")
sy["D8"] = sy["D8"].value + ", " + ", ".join(str(n) for n in range(71, 80))
sy["D9"] = "1 à 79"
pc = wb["Points à clarifier"]
d = pc.max_row
for k, (sujet, constat, demande) in enumerate(POINTS, 1):
    rr = d + k
    for c in range(2, 6):
        pc.cell(rr, c)._style = copy(pc.cell(d, c)._style)
    pc.row_dimensions[rr].height = pc.row_dimensions[d].height
    pc.cell(rr, 2, int(pc.cell(d, 2).value) + k)
    pc.cell(rr, 3, sujet)
    pc.cell(rr, 4, constat)
    pc.cell(rr, 5, demande)
wb.save(PBC)

# Guide : section I de l'onglet « 2. Provisionnements » ------------------------------------------------
wb = openpyxl.load_workbook(GUIDE)
ws = wb["2. Provisionnements"]
clot = next(r for r in range(9, ws.max_row + 1) if ws.cell(r, 2).value == "Clôture de l'entretien")
titre_h = next(r for r in range(9, ws.max_row + 1) if str(ws.cell(r, 2).value).startswith("H. "))
q_mod = titre_h + 1
fus = [m for m in ws.merged_cells.ranges if m.min_row >= clot - 1]
for m in fus:
    ws.unmerge_cells(str(m))
haut = {x: ws.row_dimensions[x].height for x in range(clot - 1, ws.max_row + 1)}
n = len(QUESTIONS) + 1
ws.insert_rows(clot - 1, n)
for x, h in haut.items():
    ws.row_dimensions[x + n].height = h
for m in fus:
    ws.merge_cells(start_row=m.min_row + n, end_row=m.max_row + n, start_column=m.min_col, end_column=m.max_col)
ligne = clot - 1
for c in range(1, 13):
    ws.cell(ligne, c)._style = copy(ws.cell(titre_h, c)._style)
ws.merge_cells(start_row=ligne, end_row=ligne, start_column=2, end_column=12)
ws.cell(ligne, 2, "I. Provisions et engagements (revue du rapport v5, 19/10/2026)")
ws.row_dimensions[ligne].height = ws.row_dimensions[titre_h].height
for k, q in enumerate(QUESTIONS, 1):
    for c in range(1, 13):
        ws.cell(ligne + k, c)._style = copy(ws.cell(q_mod, c)._style)
    ws.row_dimensions[ligne + k].height = 46
    for j, v in enumerate(q):
        ws.cell(ligne + k, 2 + j, v or None)
ws.data_validations.dataValidation[0].sqref = MultiCellRange(f"K8:K{ligne + len(QUESTIONS)}")
wb.save(GUIDE)
print(f"PBC : items 71 à 79 jusqu'à la ligne {prec} ; guide : section I lignes {ligne}-{ligne + len(QUESTIONS)}")
