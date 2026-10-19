# -*- coding: utf-8 -*-
"""
FPPN - Ajout des items 56 à 70 (entretien cotisations du 19/10/2026) à la PBC générale pointée,
et des points à clarifier 26 à 29.
Utilisation : python pbc_generale_cotisations_fppn.py <pbc_generale.xlsx> <pbc_complementaire_cotisations.xlsx> <sortie.xlsx>
"""
import sys
from copy import copy

import openpyxl
from openpyxl.worksheet.cell_range import MultiCellRange

GEN, COMP, SORTIE = sys.argv[1:4]
COMMENT = "Demande issue de l'entretien du 19/10/2026 avec le responsable des cotisations (PBC complémentaire cotisations)."
POINTS = [
    ("Délai de trois mois avant les prestations", "Le responsable des cotisations cite un délai d'au moins trois mois de cotisation avant de bénéficier des prestations (19/10/2026). L'entretien du 16/10/2026 indiquait l'absence de délai de carence pour le régime obligatoire.",
     "Transmettre le texte qui fixe ce délai et préciser les produits concernés (item 59)."),
    ("Circuit maladie 936 et décès 8Z6", "Les fichiers aller sont transmis directement par la direction du personnel de la Police à la Solde, sans plateforme de partage. Le FPPN ne reçoit que le retour, par un canal non précisé.",
     "Transmettre les fichiers aller (item 61) et préciser le canal de réception des retours (item 62)."),
    ("MEDOC 8Z7", "Produit en pause depuis environ deux mois selon le responsable des cotisations. Les fichiers retour vont de décembre 2025 à juillet 2026, sauf février et avril 2026 (retours vides).",
     "Préciser la période de fonctionnement et la date de suspension (item 69)."),
    ("Types de prêts du Grenier", "Trois types de prêts cités pendant l'entretien du 19/10/2026, huit types de crédit dans les fichiers de production du Grenier.",
     "Transmettre la liste des types de prêts et leurs conditions (item 70)."),
]

wb = openpyxl.load_workbook(GEN)
ws = wb["PBC Actuariat FPPN"]
dern_item = max(r for r in range(11, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith("="))
sect_modele = next(r for r in range(dern_item, 9, -1) if ws.cell(r, 3).value and not ws.cell(r, 2).value)
COLS = range(2, 14)
st_sect = {c: copy(ws.cell(sect_modele, c)._style) for c in COLS}
st_item = {c: copy(ws.cell(dern_item, c)._style) for c in COLS}

src = openpyxl.load_workbook(COMP)["PBC Cotisations"]
r = dern_item + 2
for c in COLS:
    ws.cell(r, c)._style = copy(st_sect[c])
ws.cell(r, 3, "Cotisations et précomptes (demande complémentaire du 19/10/2026)")
ws.row_dimensions[r].height = ws.row_dimensions[sect_modele].height
prec, premier = dern_item, None
for i in range(11, src.max_row + 1):
    if not str(src.cell(i, 2).value or "").startswith("="):
        continue
    r += 1
    for c in COLS:
        ws.cell(r, c)._style = copy(st_item[c])
    ws.cell(r, 2, f"=+B{prec}+1")
    for c in (3, 5, 7, 8, 10, 13):
        ws.cell(r, c, src.cell(i, c).value)
    # Les onglets cités sont ceux de la PBC complémentaire cotisations
    for o in ("Fichiers manquants", "Modèle ayants droit", "Modèle matricule mécano"):
        ws.cell(r, 3).value = ws.cell(r, 3).value.replace(f"l'onglet « {o} »", f"l'onglet « {o} » de la PBC complémentaire cotisations")
    n = 56 + (r - dern_item - 3)
    ws.cell(r, 11, COMMENT + (" Complète l'item 51 : fichiers détenus par la direction du personnel de la Police." if n == 61 else ""))
    ws.row_dimensions[r].height = src.row_dimensions[i].height
    premier = premier or r
    prec = r
fin = prec

plage = f"$M$11:$M${fin}"
ws["M7"] = f'=IFERROR(COUNTIF({plage},"Reçu")/(COUNTA({plage})-COUNTIF({plage},"Non Applicable")),0)'
for dv in ws.data_validations.dataValidation:
    col = str(dv.sqref).split(":")[0][0]
    dv.sqref = MultiCellRange(f"{col}11:{col}{fin}")

sy = wb["Synthèse du pointage"]
for k, st in ((6, "Reçu"), (7, "Partiellement Reçu"), (8, "Non Reçu")):
    sy.cell(k, 3, f"=COUNTIF('PBC Actuariat FPPN'!{plage},\"{st}\")")
sy["D8"] = sy["D8"].value + ", " + ", ".join(str(n) for n in range(56, 71))
sy["D9"] = "1 à 70"

pc = wb["Points à clarifier"]
d = pc.max_row
for k, (sujet, constat, demande) in enumerate(POINTS, 1):
    rr = d + k
    for c in range(2, 6):
        pc.cell(rr, c)._style = copy(pc.cell(d, c)._style)
    pc.row_dimensions[rr].height = pc.row_dimensions[d].height
    pc.cell(rr, 2, pc.cell(d, 2).value + k if isinstance(pc.cell(d, 2).value, int) else str(int(pc.cell(d, 2).value) + k))
    pc.cell(rr, 3, sujet)
    pc.cell(rr, 4, constat)
    pc.cell(rr, 5, demande)
wb.save(SORTIE)
print(f"items 56 à 70 en lignes {premier}-{fin} ; points à clarifier jusqu'à la ligne {pc.max_row}")
