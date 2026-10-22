# -*- coding: utf-8 -*-
"""
FPPN - Mise en forme pédagogique des onglets de calcul (lecteur non technique)
==============================================================================
Chaque onglet commence par un bloc d'explication en mots simples :
  « À quoi sert cet onglet », « D'où viennent les chiffres » (fichiers exacts), « Comment le calcul est fait »,
  « Comment lire le résultat », « Valable pour », « Limites ».
Les tableaux portent des formules Excel vivantes ; une colonne « Explication » commente chaque ligne quand c'est utile ;
les paramètres modifiables sont en orange.
"""
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

POLICE = "Helvetica"
NOIR, GRIS1, GRIS2, ORANGE, BLEU, GRIS_TXT = "060606", "CCCCCC", "E8E8E8", "EFAC62", "112A40", "53565A"
FIN = Side(style="thin", color="BBBCBC")
RUBRIQUES = ("À quoi sert cet onglet", "D'où viennent les chiffres", "Comment le calcul est fait", "Comment lire le résultat", "Valable pour", "Limites")


def onglet(wb, nom, titre, textes, largeurs):
    """Crée l'onglet, écrit le titre et le bloc d'explication. textes : dict rubrique -> texte. Renvoie (ws, ligne suivante)."""
    ws = wb.create_sheet(nom)
    ws.sheet_view.showGridLines = False
    for j, w in enumerate(largeurs, 1):
        ws.column_dimensions[L(j)].width = w
    ws["A1"] = titre
    ws["A1"].font = Font(name=POLICE, size=14, bold=True, color=NOIR)
    derniere = L(max(len(largeurs), 4))
    r = 3
    for rub in RUBRIQUES:
        if not textes.get(rub):
            continue
        a = ws.cell(r, 1, rub)
        a.font = Font(name=POLICE, size=9, bold=True, color="FFFFFF")
        a.fill = PatternFill("solid", fgColor=BLEU)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(f"B{r}:{derniere}{r}")
        b = ws.cell(r, 2, textes[rub])
        b.font = Font(name=POLICE, size=9, color=NOIR)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        largeur_txt = sum(largeurs[1:]) or 80
        ws.row_dimensions[r].height = max(15, 12.5 * (len(textes[rub]) // max(int(largeur_txt * 1.1), 40) + 1 + textes[rub].count("\n")))
        r += 1
    return ws, r + 1


def sous_titre(ws, r, texte, ncol):
    c = ws.cell(r, 1, texte)
    c.font = Font(name=POLICE, size=11, bold=True, color=BLEU)
    return r + 1


def tableau(ws, r0, entetes, lignes, formats=None, orange=(), commentaires=None, total=None):
    """lignes : listes de valeurs ; une chaîne commençant par « = » est une formule, « {r} » y est remplacé par le numéro de ligne.
    formats : liste de formats par colonne ; orange : colonnes (indices 0..) modifiables ; commentaires : {(i, j): texte}.
    total : liste de valeurs ou formules pour une ligne de total ({d} et {f} : première et dernière ligne de données).
    Renvoie (première ligne de données, dernière ligne de données)."""
    for j, t in enumerate(entetes, 1):
        c = ws.cell(r0, j, t)
        c.font = Font(name=POLICE, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NOIR)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 30
    d = r0 + 1
    for i, lig in enumerate(lignes):
        r = d + i
        for j, v in enumerate(lig, 1):
            if isinstance(v, str) and v.startswith("="):
                v = v.replace("{r}", str(r))
            c = ws.cell(r, j, v)
            c.font = Font(name=POLICE, size=9, bold=(j == 1), color=NOIR)
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and not str(v).startswith("="))
            c.border = Border(bottom=FIN)
            fond = GRIS1 if j == 1 else (GRIS2 if i % 2 else "FFFFFF")
            if (j - 1) in orange:
                fond = ORANGE
            c.fill = PatternFill("solid", fgColor=fond)
            if formats and j - 1 < len(formats) and formats[j - 1]:
                c.number_format = formats[j - 1]
            if commentaires and (i, j - 1) in commentaires:
                c.comment = Comment(commentaires[(i, j - 1)], "Adenka")
    f = d + len(lignes) - 1
    if total:
        r = f + 1
        for j, v in enumerate(total, 1):
            if isinstance(v, str) and v.startswith("="):
                v = v.replace("{d}", str(d)).replace("{f}", str(f))
            c = ws.cell(r, j, v)
            c.font = Font(name=POLICE, size=9, bold=True, color=NOIR)
            c.fill = PatternFill("solid", fgColor=GRIS1)
            if formats and j - 1 < len(formats) and formats[j - 1]:
                c.number_format = formats[j - 1]
    return d, f


def note(ws, r, texte, ncol=6):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(ncol, 2))
    c = ws.cell(r, 1, texte)
    c.font = Font(name=POLICE, size=9, italic=True, color=GRIS_TXT)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = max(15, 12.5 * (len(texte) // 150 + 1))
    return r + 1


FCFA, PCT, NB1, NB2, ENT = "#,##0", "0.00%", "0.0", "0.00", "0"
