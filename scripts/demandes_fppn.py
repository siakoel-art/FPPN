# -*- coding: utf-8 -*-
"""
FPPN - Liste des questions et des documents à transmettre par le FPPN (document destiné au client)
=================================================================================================
Onglet 1 : questions simples et précises, avec la pièce qui permet d'y répondre.
Onglet 2 : documents à transmettre, tirés des items ouverts de la PBC (non reçus ou partiellement reçus),
           avec, pour un item partiellement reçu, la seule partie qui reste à transmettre.
Numéros : PBC actuarielle du 23/10/2026.
Utilisation : python demandes_fppn.py <pbc_23.xlsx> <correspondance.json> <sortie.xlsx>
"""
import json
import re
import sys

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

PBC, MAPJ, SORTIE = sys.argv[1:4]
M = {int(k): v for k, v in json.load(open(MAPJ))["mapping"].items()}
ws = openpyxl.load_workbook(PBC).worksheets[0]
ITEMS, sec, n = [], None, 0
for r in range(10, ws.max_row + 1):
    if ws.cell(r, 2).value is not None:
        n += 1
        ITEMS.append({"n": n, "texte": ws.cell(r, 3).value, "prio": ws.cell(r, 5).value, "statut": ws.cell(r, 13).value, "section": sec})
    elif ws.cell(r, 3).value:
        sec = re.sub(r"\s*\(.*\)", "", ws.cell(r, 3).value).replace("Dossier « 08-Nouveaux Dossiers Compta » reçu le 22/10/2026", "Allocations décès, clinique et projet de Cocody").replace("Documents reçus le 22/10/2026", "Comptes, personnel et retraités").replace("Demandes issues de la vérification de cohérence du 23/10/2026", "Compléments du 23/10/2026")


def nouveau(texte):
    return next(i["n"] for i in ITEMS if i["texte"].startswith(texte))


def ref(*anciens):
    return ", ".join(str(x) for x in sorted({M[a] for a in anciens}))


RESTE = {  # partie restant à transmettre pour un item partiellement reçu (ancienne numérotation)
    1: "Statuts et règlement intérieur applicables du 01/01/2020 au 30/06/2021 (la version adoptée le 30/06/2021 est reçue).",
    4: "Avenants à la convention IPS-CGRAE, son annexe 2 (bulletin de souscription), barème de la cotisation des retraités et liste des retraités souscripteurs non précomptés.",
    7: "Toute étude actuarielle ou note de calcul des engagements, dont l'étude des risques longs déposée à l'AIRMS (les rapports de gestion reçus ne contiennent pas de calcul actuariel).",
    9: "Fichier des mutualistes actifs avec date de naissance, sexe, date d'adhésion, statut et date de sortie, au 31/12 de 2019 à 2025 et au 31/08/2026 (extraction du logiciel Easy Money).",
    12: "Cotisations appelées et encaissées par mois et par catégorie de membres, de 2020 à 2023.",
    14: "Liste des dossiers non payés au 31/12 de 2019 à 2025 et au 31/08/2026, toutes garanties. Si cet état n'est pas tenu, l'extraire de la base des prestations.",
    17: "Soldes d'ouverture et de clôture, intérêts crédités et retraits de chaque épargnant PEP et EFC, de 2020 au 31/08/2026, et rapprochement avec les comptes 41195100 et 41195200.",
    21: "Échéancier de remboursement par emprunteur au 31/08/2026 (les balances âgées sont reçues).",
    23: "Valeur au bilan, dernière évaluation et revenus annuels de chaque bien immobilier et de la clinique.",
    26: "Nom, fonction et contact des responsables des prestations, de l'épargne, des placements et des données des mutualistes.",
    27: "Détail mensuel des prestations maladie suivies par le service, de 2019 au 31/08/2026.",
    28: "États des dossiers décès de juillet et août 2026, et date d'enregistrement de chaque dossier dans les états 2022, 2024 et 2025.",
    30: "Barèmes décès, frais funéraires et invalidité appliqués de 2001 à 2019, avec les décisions qui les ont adoptés.",
    34: "Versions signées des manuels de procédures (adhésions, cotisations, prestations, crédits).",
    35: "Rapport de contrôle de l'AIRMS, plan de redressement en cours et correspondances avec l'AIRMS.",
    42: "Fichier des cartes de santé actives au 31/08/2026, ayants droit compris.",
    44: "Décision d'agrément et liste des risques autorisés (l'attestation d'immatriculation n° 1D/0052016/CI du 05/04/2017 est reçue).",
    47: "Rapports du commissaire aux comptes de 2020 à 2025 et états financiers par risque long (les états financiers de 2020 au 31/08/2026 sont reçus).",
    53: "Procès-verbaux du conseil d'administration, de l'organe de contrôle et des Assemblées générales de 2020 à 2025, dont l'approbation du reclassement des provisions en 2024.",
    56: "Document qui notifie à la Solde la répartition du taux de 5 %, et clé appliquée en janvier et février 2024.",
    57: "Règlement de la flotte automobile et date de démarrage de chaque produit.",
    71: "Base des prestations maladie datée : date des soins, date de dépôt de la facture, date de paiement, pour chaque facture.",
    72: "Pièce d'origine des provisions maladie et obsèques, ou décision qui permettra de les reprendre ou de les justifier.",
    73: "État des impayés du Grenier par tranche de retard et décision fixant les taux de dépréciation.",
    79: "Durée de la garantie santé des retraités (à vie ou limitée), cotisation et ayants droit couverts, et décision qui la fonde.",
    80: "Conventions, taux, échéances et relevés des DAT BOA de 2020 au 31/08/2026.",
    82: "Inventaire et tableau de variation de tous les placements de 2016 au 31/08/2026.",
    83: "Correspondance avec Orabank sur l'arrêt des prélèvements et le remboursement du DAT.",
    84: "Situation et valeur de chaque terrain (projet, attribution, litige).",
    85: "Encaissements des loyers des immeubles d'Aboisso.",
    86: "Rapport de réévaluation du siège et décision de comptabilisation au compte 10620000.",
    89: "Convention d'apporteur d'affaires avec GNA, décomptes annuels des commissions et états des primes par adhérent.",
    91: "État des chèques en suspens, dont ceux du compte BOA EFC émis depuis décembre 2021.",
    92: "Fichiers des décès transmis par la direction du personnel, de 2020 à 2025.",
    97: "Déclarations mensuelles CNPS et déclarations de 2026.",
    99: "Décision fixant le taux servi sur le PEP et sur l'EFC.",
    103: "Statut civil ou policier de chaque agent de la Direction Générale.",
    109: "Délibération approuvant le cadre organique des emplois et la grille des salaires.",
    119: "Comptabilité séparée de la clinique et résultats annuels de 2020 à 2025.",
}
RESTE_N = {}
for old, txt in RESTE.items():
    RESTE_N.setdefault(M[old], txt)

QUESTIONS = [
    ("Cotisations", "Pourquoi la Solde précompte-t-elle 3,77 % du salaire de base pour la maladie et 1,23 % pour le décès, alors que les statuts (art. 17) fixent 3,5 % et 1,5 % ?", ref(56), "Urgent"),
    ("Cotisations", "Le délai d'au moins trois mois de cotisation avant les prestations s'applique-t-il à la maladie et au décès, ou seulement à l'EFC et au PEP ?", ref(59), "Urgent"),
    ("Cotisations", "Quelle décision a porté l'EFC de 6 000 FCFA à 10 000 FCFA en 2021, et quelle règle s'applique aux anciens souscripteurs ?", ref(58), "Moyen"),
    ("Cotisations", "Que représentent les références PTI, RPEP, EP, TI, R-PEP et PTQ précomptées au poste 998 (174 461 860 FCFA de 2020 à août 2026) ?", str(nouveau("Nature des références PTI")), "Moyen"),
    ("Cotisations", "Le FPPN rapproche-t-il les fichiers aller et retour de chaque produit ? Si oui, qui le fait et selon quelle procédure écrite ?", str(nouveau("Procédure écrite du circuit")), "Faible"),
    ("Épargne", "Quel taux est servi sur l'EFC : 3,5 %, 5,5 % ou 14 % ? Quelle décision signée le fixe ?", ref(99), "Urgent"),
    ("Épargne", "Le dépôt à terme du compte 52210200 est-il un placement du FPPN ou un placement affecté à l'EFC, comme l'indique l'intitulé du compte ?", ref(80), "Urgent"),
    ("Provisions", "Quelle décision permettra de reprendre ou de justifier les provisions maladie (2 867 602 878 FCFA) et obsèques (1 961 695 412 FCFA), passées vers 2004 et inchangées depuis ?", ref(72), "Urgent"),
    ("Provisions", "Pourquoi les sinistres décès 2021 valent-ils 156 020 000 FCFA dans l'étude des ratios et 401 856 666 FCFA dans la balance 2021 (compte 60210000) ?", str(nouveau("Explication de l'écart sur les sinistres décès 2021")), "Moyen"),
    ("Provisions", "À quel grade correspond l'allocation décès de 2 200 000 FCFA des états de fin d'année, et quel texte fonde l'exclusion du suicide ?", ref(116), "Moyen"),
    ("Retraités", "La garantie santé des policiers retraités est-elle accordée à vie ou pour une durée limitée ? Quelle est leur cotisation ?", ref(79), "Urgent"),
    ("Personnel", "Quel âge de départ à la retraite le FPPN applique-t-il à son personnel civil ? Cinq agents de 55 ans ou plus sont toujours en poste en 2025, dont un de 60 ans.", ref(113), "Urgent"),
    ("Personnel", "Pourquoi douze dates d'embauche diffèrent-elles entre les DISA 2021 à 2023 et la DISA 2025 ?", ref(113), "Urgent"),
    ("Personnel", "À quelle entité (Direction Générale, clinique, Grenier) sont rattachés les dix-neuf agents recrutés en 2025, et combien d'enfants à charge a chaque agent ?", ref(113), "Urgent"),
    ("Placements", "Quelle est la nature de la créance de 3 407 519 850 FCFA sur l'hôpital de la Police (compte 47111100), inchangée depuis 2020, et un remboursement est-il prévu ?", ref(87), "Moyen"),
    ("Placements", "Le compte 10620000 enregistre-t-il la réévaluation du terrain du siège ou de matériels, comme l'indique son intitulé ?", ref(86), "Urgent"),
    ("Placements", "Quel est le budget du projet de Cocody, son plan de financement et la part financée par l'épargne des mutualistes ?", ref(118), "Urgent"),
    ("Grenier", "Pourquoi la plus grande part de l'encours du Grenier était-elle en retard fin janvier 2026, et comment sont réglées les échéances non précomptées par la Solde ?", str(nouveau("Grenier : explication de la hausse")), "Moyen"),
    ("Flotte automobile", "Sur quelle assiette la commission de 20 % versée par GNA est-elle calculée (primes précomptées seules, ou toutes les primes) ?", ref(89), "Moyen"),
    ("MEDOC", "Sur quelle période le produit MEDOC a-t-il fonctionné, et à quelle date a-t-il été suspendu ?", ref(69), "Faible"),
    ("Gouvernance", "Quelle version des comptes 2023 l'Assemblée générale a-t-elle approuvée, et pourquoi le résultat au 31/08/2026 diffère-t-il entre les états financiers (-437 148 460 FCFA) et la balance (-441 256 500 FCFA) ?", ref(110), "Urgent"),
    ("Gouvernance", "Le directeur général a-t-il été désigné à nouveau ou prolongé après la mandature 2020-2023 ? Qui autorise les placements ?", ref(117), "Urgent"),
    ("Conformité", "Quels sont le numéro, la date et les risques autorisés de l'agrément du FPPN par l'AIRMS ?", ref(44), "Urgent"),
    ("Conformité", "Le FPPN adhère-t-il au fonds de garantie FOMUS-CI ? Quel montant de cotisation a-t-il payé depuis 2022 ?", ref(49), "Moyen"),
    ("Conformité", "Une étude actuarielle des risques longs ou une certification des comptes a-t-elle été réalisée et déposée à l'AIRMS ?", ref(48, 47), "Moyen"),
    ("Conformité", "Quel est l'effectif total de la Police nationale au 31/12 de 2020 à 2025 et au 31/08/2026 ?", str(nouveau("Effectif total de la Police nationale")), "Moyen"),
]

NOIR = "000000"
CP = {"Urgent": "FF0000", "Moyen": "FFC000", "Faible": "00B050", "Non reçu": "FF0000", "Partiellement reçu": "FFC000"}
fin = Side(style="thin", color="808080")
BORD = Border(left=fin, right=fin, top=fin, bottom=fin)


def feuille(wb, nom, titre, sous, ent, larg, lignes, colorees, legendes):
    w = wb.create_sheet(nom)
    w.sheet_view.showGridLines = False
    w["A1"], w["A2"] = titre, sous
    w["A1"].font = Font(name="Helvetica", size=14, bold=True)
    w["A2"].font = Font(name="Helvetica", size=9, italic=True)
    for j, (e, l) in enumerate(zip(ent, larg), 1):
        c = w.cell(4, j, e)
        c.font = Font(name="Helvetica", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NOIR)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORD
        w.column_dimensions[c.column_letter].width = l
    for i, row in enumerate(lignes, 5):
        for j, v in enumerate(row, 1):
            c = w.cell(i, j, v)
            c.font = Font(name="Helvetica", size=9)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BORD
            if j in colorees and v in CP:
                c.fill = PatternFill("solid", fgColor=CP[v])
                c.font = Font(name="Helvetica", size=9, bold=True, color="000000" if CP[v] == "FFC000" else "FFFFFF")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    w.freeze_panes = "A5"
    w.auto_filter.ref = f"A4:{w.cell(4, len(ent)).column_letter}{4 + len(lignes)}"
    w.page_setup.orientation = "landscape"
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0
    w.print_title_rows = "4:4"
    r = 6 + len(lignes)
    for t, items in legendes:
        w.cell(r, 2, t).font = Font(name="Helvetica", size=10, bold=True)
        r += 1
        for code, sens in items:
            c = w.cell(r, 2, code)
            c.fill = PatternFill("solid", fgColor=CP[code])
            c.font = Font(name="Helvetica", size=9, bold=True, color="000000" if CP[code] == "FFC000" else "FFFFFF")
            c.border = BORD
            d = w.cell(r, 3, sens)
            d.font = Font(name="Helvetica", size=9)
            d.border = BORD
            r += 1
        r += 1


LP = ("Code couleur : priorité", [("Urgent", "Nécessaire pour conclure sur les provisions, les engagements ou la conformité."),
                                  ("Moyen", "Précise un résultat ou un indicateur."), ("Faible", "Complète la documentation.")])
LS = ("Code couleur : statut", [("Non reçu", "Aucun document reçu pour cette demande."), ("Partiellement reçu", "Seule la partie indiquée reste à transmettre.")])
wb = Workbook()
wb.remove(wb.active)
feuille(wb, "1_Questions", "Questions au FPPN", "Réponse écrite attendue, avec la pièce qui l'établit. N° PBC : liste des documents demandés du 23/10/2026.",
        ["N°", "Thème", "Question", "N° PBC", "Priorité", "Réponse du FPPN", "Pièce jointe", "Date de réponse"], [5, 16, 70, 9, 10, 50, 30, 14],
        [(i, t, q, p, pr, "", "", "") for i, (t, q, p, pr) in enumerate(QUESTIONS, 1)], {5}, [LP])
ouverts = [x for x in ITEMS if x["statut"] in ("Non Reçu", "Partiellement Reçu")]
lignes = []
for k, x in enumerate(ouverts, 1):
    st = "Non reçu" if x["statut"] == "Non Reçu" else "Partiellement reçu"
    demande = RESTE_N.get(x["n"], x["texte"]) if st == "Partiellement reçu" else x["texte"]
    lignes.append((k, x["n"], x["section"], demande, st, x["prio"], "", ""))
feuille(wb, "2_Documents", "Documents à transmettre par le FPPN", "Demandes ouvertes de la liste des documents demandés (PBC) du 23/10/2026. Pour une demande partiellement reçue, seule la partie qui manque est indiquée.",
        ["N°", "N° PBC", "Thème", "Document à transmettre", "Statut", "Priorité", "Date de transmission", "Commentaire du FPPN"], [5, 8, 28, 80, 15, 10, 14, 40],
        lignes, {5, 6}, [LP, LS])
wb.save(SORTIE)
print(SORTIE, "questions :", len(QUESTIONS), "documents :", len(lignes), "partiels sans précision :", [x["n"] for x in ouverts if x["statut"] == "Partiellement Reçu" and x["n"] not in RESTE_N])
