# -*- coding: utf-8 -*-
"""
FPPN - PBC actuarielle du 23/10/2026 : doublons fusionnés, numérotation continue, pointage à jour
================================================================================================
1. Fusion des doublons : l'item conservé reprend la demande et le pointage des items fusionnés
   (statut combiné : Reçu si tous reçus, Non Reçu si aucun, sinon Partiellement Reçu).
2. Mise à jour du pointage avec la vérification de cohérence du 23/10/2026 (grands livres renvoyés lisibles,
   fichiers vides de décembre 2024, réponses déjà présentes dans les documents).
3. Ajout des demandes nouvelles issues de la vérification.
4. Renumérotation continue, onglet « Correspondance des numéros » (ancien → nouveau), mise à jour des numéros
   d'items dans les onglets annexes et de la synthèse du pointage.
Utilisation : python maj_pbc_23_fppn.py <pbc_22.xlsx> <pbc_23.xlsx> <correspondance.json>
"""
import datetime as dt
import json
import re
import sys
from copy import copy

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SRC, DST, MAP_JSON = sys.argv[1:4]
J = dt.datetime(2026, 10, 23)
D = dt.datetime(2026, 10, 30)
ST = {"Reçu": 2, "Partiellement Reçu": 1, "Non Reçu": 0}

FUSION = {  # item conservé : items fusionnés (ancienne numérotation)
    61: [51], 66: [10], 71: [13, 40], 14: [29], 99: [18], 102: [16, 46], 56: [3, 45], 100: [2, 33],
    30: [93, 101], 86: [50], 103: [76], 104: [77], 78: [98], 79: [108], 89: [37, 68], 94: [32],
    57: [60], 119: [88], 72: [6], 7: [48], 17: [19], 82: [20], 54: [24], 35: [8], 53: [5, 106, 107],
}
TEXTE = {  # demandes reformulées après fusion ou vérification
    62: "Fichiers retour maladie (936) et décès (8Z6) manquants ou vides : « 122024_effectif_police_936_DECEMBRE 2024.xlsx » et « FPPN_ALLOC_DECES_8Z6_DECEMBRE_2024.xlsx » "
        "(reçus vides, 0 octet) ; retours 8Z6 de mai à octobre 2024 et d'avril 2025 ; retours 936 et 8Z6 de janvier 2020 à décembre 2023.",
    63: "Fichiers manquants des autres produits : retour Grenier 9X4 de décembre 2024, retour EFC 870 de mars 2025, retour CGRAE de janvier et de juin 2024 (le fichier reçu pour juin 2024 "
        "est un fichier aller), retours MEDOC 8Z7 de 2024 à 2026, fichiers aller et retour de la flotte 9W3 manquants (aller : 03, 05, 06 et 08/2025, 03 et 08/2026 ; retour : 04, 06 et 08/2025, 01/2026).",
    71: "Base des prestations, une ligne par facture ou par dossier, du 01/01/2020 au 31/08/2026 : matricule du mutualiste, bénéficiaire, prestataire, numéro de facture, date des soins, "
        "date de dépôt de la facture au FPPN, date de paiement, montant facturé, montant payé, motif d'un rejet (maladie). Les dossiers décès sont demandés à l'item des sinistres décès.",
    56: "Document qui notifie à la Solde la répartition du taux de 5 % entre maladie et décès (la Solde précompte 3,77 % et 1,23 % du salaire de base, les statuts, art. 17, fixent 3,5 % et 1,5 %), "
        "avec la clé appliquée de 2020 à 2026.",
    30: "Barèmes des capitaux décès, frais funéraires, allocations aux ayants droit et allocations d'invalidité appliqués de 2001 à 2019, avec les décisions qui les ont adoptés (le barème en vigueur depuis 2020 est reçu).",
    80: "DAT BOA (comptes 52210200 et 52210300) : conventions, taux, échéances et relevés de 2020 au 31/08/2026, et réponse écrite à la question : le DAT du compte 52210200 est-il un placement "
        "du FPPN ou un placement affecté à l'EFC, comme l'indique l'intitulé du compte ?",
    72: "Toute pièce sur l'origine des provisions pour prestations à payer maladie (2 867 602 878 FCFA) et obsèques (1 961 695 412 FCFA), inchangées depuis au moins le 01/09/2020 "
        "et passées vers 2004 selon le chef comptable, et décision qui permettra de les reprendre ou de les justifier.",
    53: "Procès-verbaux des réunions du conseil d'administration, de l'organe de contrôle et de l'Assemblée générale de 2020 à 2025, dont : décisions modifiant les cotisations, les prestations ou les taux "
        "servis sur l'épargne ; affectation des résultats (statuts, art. 90) ; approbation du reclassement des provisions en comptes 15801000 et 15802000 en 2024 (statuts, art. 75).",
}
POINTAGE = {  # nouveau commentaire de pointage (ajouté) et statut éventuel
    62: (None, "Vérification du 23/10/2026 : les deux fichiers de décembre 2024 sont vides (0 octet, même empreinte) et ne figurent dans aucun envoi suivant (18/10 et 22/10/2026). "
               "Le canal de réception n'est plus demandé : les fichiers retour 936 et 8Z6 ont été transmis le 15/10/2026."),
    92: (None, "Vérification du 23/10/2026 : le calcul de fin d'année est comptabilisé au compte 41591000 (2024 : 140 800 000 FCFA ; 2025 : 130 600 000 FCFA, égal au crédit de l'année). Point résolu sur le compte utilisé."),
    72: (None, "Vérification du 23/10/2026 : montants identiques du report à nouveau du 01/09/2020 au 31/08/2026 ; seul mouvement, le reclassement de 2024 au même montant. Ce sont les provisions héritées décrites par le chef comptable."),
    70: (None, "Les huit types de crédit des fichiers de production sont : ACTIVITES GENERATRICE DE REVENUS, CREDIT DE SANTE, CREDIT GRENIER SHOP, CREDIT HABITAT, CREDIT MARIAGE, CRÉDIT FLASH, CRÉDIT SCOLAIRE, EQUIPEMENT DOMESTIQUE."),
    55: (None, "Sigle CMP : Clinique médicale du policier."),
    30: (None, "Barème en vigueur lu dans « CARTOGRAPHIE DES RISQUES COUVERTS ET DES PRESTATIONS OFFERTES.docx » : officier 1 600 000 FCFA, sous-officier 1 400 000 FCFA."),
}
NOUVEAUX = [
    ("Nature des références PTI, RPEP, EP, TI, R-PEP et PTQ précomptées au poste 998 de 2020 à août 2026 (174 461 860 FCFA à côté des références PEP) : remboursements de prêts PEP (comptes 27110000 et 27120000) ou autre produit.", "Moyen"),
    ("Explication de l'écart sur les sinistres décès 2021 : 156 020 000 FCFA dans « ETUDE DES RATIO MALADIE DECES 2.xlsx », contre 401 856 666 FCFA au compte 60210000 de la balance 2021.", "Moyen"),
    ("Validation, ligne par ligne, des montants décès supérieurs à 1 000 000 FCFA présents dans les fichiers retour 8Z6 (liste jointe par Adenka), et montant correct de chaque ligne.", "Faible"),
    ("Procédure écrite du circuit des fichiers aller et retour de chaque produit (maladie, décès, EFC, PEP, Grenier, MEDOC, flotte, retraités) : qui produit, qui transmet, qui reçoit, qui rapproche.", "Faible"),
    ("Effectif total de la Police nationale au 31/12 de 2020 à 2025 et au 31/08/2026 (taux de pénétration, Décision n° 05/2022).", "Moyen"),
    ("Grenier : explication de la hausse des retards de fin janvier 2026, et mode de règlement des échéances non précomptées par la Solde (paiement direct, compensation, relance).", "Moyen"),
]


def comb(a, b):
    if a == b:
        return a
    if "Non Applicable" in (a, b):
        return b if a == "Non Applicable" else a
    return "Partiellement Reçu"


wb = openpyxl.load_workbook(SRC)
ws = wb["PBC Actuariat FPPN"]
last = ws.max_row
rows, n = [], 0
for r in range(10, last + 1):
    b, c = ws.cell(r, 2).value, ws.cell(r, 3).value
    if b is not None:
        n += 1
        rows.append({"type": "item", "old": n, "v": {k: ws.cell(r, k).value for k in range(2, 15)}, "h": ws.row_dimensions[r].height})
    elif c:
        rows.append({"type": "section", "v": {3: c}, "h": ws.row_dimensions[r].height})
STY_ITEM = {k: copy(ws.cell(11, k)._style) for k in range(1, 15)}
STY_SEC = {k: copy(ws.cell(10, k)._style) for k in range(1, 15)}
items = {x["old"]: x for x in rows if x["type"] == "item"}

# 1. fusion
fusionnes = {}
for keep, others in FUSION.items():
    t = items[keep]["v"]
    for o in others:
        ov = items[o]["v"]
        fusionnes[o] = keep
        t[13] = comb(t[13] or "Non Reçu", ov[13] or "Non Reçu")
        if ov[9] and (not t[9] or (isinstance(t[9], dt.datetime) and isinstance(ov[9], dt.datetime) and ov[9] > t[9])):
            t[9] = ov[9]
        t[11] = (t[11] or "") + f"\nFusion du 23/10/2026 : reprend l'ancien item {o} (« {str(ov[3])[:160]}{'…' if len(str(ov[3])) > 160 else ''} »), statut « {ov[13]} »."
        if ST.get(ov[5]) is None and ov[5] == "Urgent":
            t[5] = "Urgent"
        if ov[5] == "Urgent":
            t[5] = "Urgent"
# 2. pointage et textes
for k, txt in TEXTE.items():
    items[k]["v"][3] = txt
for k, (st, com) in POINTAGE.items():
    if st:
        items[k]["v"][13] = st
    items[k]["v"][11] = (items[k]["v"][11] or "") + "\n" + com

# 3. nouvelle liste
new_rows = []
for x in rows:
    if x["type"] == "item" and x["old"] in fusionnes:
        continue
    new_rows.append(x)
new_rows.append({"type": "section", "v": {3: "Demandes issues de la vérification de cohérence du 23/10/2026"}, "h": None})
for txt, prio in NOUVEAUX:
    new_rows.append({"type": "item", "old": None, "v": {3: txt, 5: prio, 7: J, 8: D, 11: "Demande issue de la vérification de cohérence des chiffres du 23/10/2026.", 13: "Non Reçu"}, "h": None})
# supprime les sections vides
propre = []
for i, x in enumerate(new_rows):
    if x["type"] == "section":
        suiv = next((y for y in new_rows[i + 1:] if True), None)
        if suiv is None or suiv["type"] == "section":
            continue
    propre.append(x)

# 4. réécriture
for r in range(10, last + 30):
    ws.row_dimensions[r].height = None
    for k in range(1, 15):
        c = ws.cell(r, k)
        c.value = None
        c._style = copy(ws.cell(200, 20)._style)
mapping, r, n, prev = {}, 10, 0, None
for x in propre:
    sty = STY_SEC if x["type"] == "section" else STY_ITEM
    for k in range(1, 15):
        ws.cell(r, k)._style = copy(sty[k])
    if x["type"] == "section":
        ws.cell(r, 3).value = x["v"][3]
    else:
        n += 1
        if x["old"]:
            mapping[x["old"]] = n
        for k, v in x["v"].items():
            if k != 2:
                ws.cell(r, k).value = v
        ws.cell(r, 2).value = "=1" if prev is None else f"=+B{prev}+1"
        prev = r
    if x["h"]:
        ws.row_dimensions[r].height = x["h"]
    r += 1
fin = r - 1
for o, k in fusionnes.items():
    mapping[o] = mapping[k]
ws["C5"] = "DEMANDE DES DOCUMENTS (LISTE COMPLÈTE) ET POINTAGE AU 23/10/2026"
ws["M7"] = f'=IFERROR(COUNTIF($M$11:$M${fin},"Reçu")/(COUNTA($M$11:$M${fin})-COUNTIF($M$11:$M${fin},"Non Applicable")),0)'
ws.data_validations.dataValidation = []
for col, f in (("E", '"Urgent,Moyen,Faible"'), ("J", '"Physique,Electronique,Physique et Electronique"'), ("M", '"Reçu,Non Reçu,Non Applicable,Partiellement Reçu"')):
    dv = DataValidation(type="list", formula1=f, allow_blank=True)
    dv.add(f"{col}11:{col}{fin}")
    ws.add_data_validation(dv)
ws.print_area = f"A1:N{fin}"


# 5. numéros dans les onglets annexes
def remap_txt(s):
    def f(m):
        return str(mapping.get(int(m.group(0)), m.group(0)))
    return re.sub(r"\b\d{1,3}\b", f, s)


for name in ("Pointage autres équipes", "Pointage nouveaux fichiers", "Pointage lot 1 (contenu)", "Pointage lots 2 et 3 (contenu)", "Pointage dossier compta 22-10"):
    w = wb[name]
    for rr in range(6, w.max_row + 1):
        c = w.cell(rr, 1)
        if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.strip().isdigit()):
            o = int(c.value)
            c.value = f"{mapping.get(o, o)} (ancien {o})" if mapping.get(o, o) != o else mapping.get(o, o)
w = wb["Couverture mensuelle"]
for rr in range(7, w.max_row + 1):
    c = w.cell(rr, 3)
    if isinstance(c.value, str) and re.search(r"\d", c.value):
        c.value = remap_txt(c.value)
    elif isinstance(c.value, (int, float)):
        c.value = mapping.get(int(c.value), c.value)
w = wb["Synthèse du pointage"]
for row in w.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("=") and "$M$11:$M$" in c.value:
            c.value = re.sub(r"\$M\$11:\$M\$\d+", f"$M$11:$M${fin}", c.value)
statuts = {}
for rr in range(11, fin + 1):
    if ws.cell(rr, 2).value is not None:
        num = sum(1 for q in range(11, rr + 1) if ws.cell(q, 2).value is not None)
        statuts.setdefault(ws.cell(rr, 13).value, []).append(num)
for rr in range(1, w.max_row + 1):
    lab = w.cell(rr, 2).value
    if lab in statuts:
        w.cell(rr, 4).value = ", ".join(str(x) for x in statuts[lab])
    elif lab == "Total":
        w.cell(rr, 4).value = f"1 à {n}"
w["A3"] = None
w["B3"] = "Pointage au 23/10/2026 : envois des 15/10, 18/10, 19/10, 20/10, 21/10 et 22/10/2026, vérification de cohérence du 23/10/2026. Doublons fusionnés (onglet « Correspondance des numéros »)."

# 5 bis. points à clarifier : statut au 23/10/2026
REPONSES = {
    7: "Fichiers texte relus (montant et date de chaque ligne) : totaux identiques à la base, mois par mois.",
    12: "Exports relus : nombre de prêts, débloqué et encours identiques aux totaux imprimés par le logiciel, pour les 37 balances âgées.",
    15: "Tous les fichiers endommagés ont été renvoyés lisibles le 22/10/2026 (dossier « 02-Nouveaux_GL et Balances »), dont « GRAND LIVRE DES COMPTES 2022.xlsx » (année complète, 192 comptes).",
    29: "Les huit types sont dans les fichiers de production : ACTIVITES GENERATRICE DE REVENUS, CREDIT DE SANTE, CREDIT GRENIER SHOP, CREDIT HABITAT, CREDIT MARIAGE, CRÉDIT FLASH, CRÉDIT SCOLAIRE, EQUIPEMENT DOMESTIQUE. La politique de crédit est reçue.",
    30: "Ce sont les provisions héritées décrites par le chef comptable : montants identiques depuis le report à nouveau du 01/09/2020, seul mouvement le reclassement de 2024. Reste demandée toute pièce sur leur origine.",
    33: "Le grand livre 2024 montre un dépôt d'un mois renouvelé chaque mois et augmenté d'environ 40 000 000 FCFA, puis viré au FPPN en décembre pour payer l'EFC (« INTERET/ DAT EFC »). Reste demandée la convention.",
    36: "Calcul de fin d'année comptabilisé au compte 41591000 : 140 800 000 FCFA en 2024 et 130 600 000 FCFA en 2025, égaux aux états de fin d'année.",
    37: "Barème lu dans « CARTOGRAPHIE DES RISQUES COUVERTS ET DES PRESTATIONS OFFERTES.docx » : officier de police 1 600 000 FCFA, sous-officier 1 400 000 FCFA.",
    56: "Le barème des états (1 400 000 et 1 600 000 FCFA) est celui de la cartographie des risques reçue le 22/10/2026. Seul le montant de 2 200 000 FCFA reste à rattacher à un grade.",
}
wp = wb["Points à clarifier"]
c = wp.cell(5, 5, "Statut au 23/10/2026")
c._style = copy(wp.cell(5, 4)._style)
c2 = wp.cell(5, 6, "Réponse établie par les documents")
c2._style = copy(wp.cell(5, 4)._style)
wp.column_dimensions["E"].width, wp.column_dimensions["F"].width = 16, 60
for rr in range(6, wp.max_row + 1):
    num = wp.cell(rr, 2).value
    if num is None:
        continue
    num = int(num)
    st = "Résolu" if num in REPONSES and num not in (30, 33, 56) else ("Expliqué en partie" if num in REPONSES else "Ouvert")
    cs = wp.cell(rr, 5, st)
    cs.fill = PatternFill("solid", fgColor={"Résolu": "00B050", "Expliqué en partie": "FFC000", "Ouvert": "FF0000"}[st])
    cs.font = Font(name="Helvetica", size=9, bold=True, color="FFFFFF" if st != "Expliqué en partie" else "000000")
    cs.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cr = wp.cell(rr, 6, REPONSES.get(num, ""))
    cr.font = Font(name="Helvetica", size=9)
    cr.alignment = Alignment(wrap_text=True, vertical="top")
rl = wp.max_row + 2
wp.cell(rl, 4, "Code couleur : statut").font = Font(name="Helvetica", size=10, bold=True)
for k, (lab, col, txt) in enumerate((("Résolu", "00B050", "La réponse figure dans les documents reçus, point clos."),
                                     ("Expliqué en partie", "FFC000", "Le constat est expliqué, une pièce reste demandée."),
                                     ("Ouvert", "FF0000", "Réponse ou document du FPPN attendu.")), rl + 1):
    a = wp.cell(k, 4, lab)
    a.fill = PatternFill("solid", fgColor=col)
    a.font = Font(name="Helvetica", size=9, bold=True, color="000000" if col == "FFC000" else "FFFFFF")
    wp.cell(k, 5, txt).font = Font(name="Helvetica", size=9)

# 6. onglet de correspondance
wc = wb.create_sheet("Correspondance des numéros", 1)
wc["A2"] = "Correspondance des numéros d'items"
wc["A2"].font = Font(name="Helvetica", size=14, bold=True)
wc["A3"] = "Numérotation de la PBC du 22/10/2026 et numérotation du 23/10/2026. Un item fusionné renvoie à l'item qui reprend sa demande."
wc["A3"].font = Font(name="Helvetica", size=9, italic=True)
for j, (h, wdt) in enumerate((("Ancien n°", 12), ("Nouveau n°", 12), ("Traitement", 60)), 1):
    c = wc.cell(5, j, h)
    c.font = Font(name="Helvetica", size=10, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="000000")
    wc.column_dimensions[c.column_letter].width = wdt
for i, o in enumerate(sorted(mapping), 6):
    wc.cell(i, 1, o)
    wc.cell(i, 2, mapping[o])
    t = wc.cell(i, 3, f"Fusionné avec l'ancien item {fusionnes[o]} (doublon)" if o in fusionnes else "Conservé")
    if o in fusionnes:
        t.fill = PatternFill("solid", fgColor="FFC000")
    for k in (1, 2, 3):
        wc.cell(i, k).font = Font(name="Helvetica", size=9)
        wc.cell(i, k).alignment = Alignment(horizontal="center" if k < 3 else "left")
r0 = len(mapping) + 7
wc.cell(r0, 2, "Code couleur").font = Font(name="Helvetica", size=10, bold=True)
wc.cell(r0 + 1, 2).fill = PatternFill("solid", fgColor="FFC000")
wc.cell(r0 + 1, 3, "Item fusionné avec un autre item qui demandait la même pièce")
wb.save(DST)
json.dump({"mapping": mapping, "fusionnes": fusionnes, "derniere_ligne": fin, "items": n}, open(MAP_JSON, "w"), indent=1)
print(DST, "items :", n, "fusionnés :", len(fusionnes), "dernière ligne :", fin)
