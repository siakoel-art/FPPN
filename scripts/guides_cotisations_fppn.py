# -*- coding: utf-8 -*-
"""
FPPN - Mise à jour des guides d'entretien après l'entretien cotisations du 19/10/2026
- onglet « 4. Cotisations » : réponses relevées (colonne Notes) et section G de relances ;
- onglet « 2. Provisionnements » : questions 24 à 27 rétablies (vides dans la version du 19/10/2026)
  et section H pour le chef comptable (paiements manuels, clé maladie et décès).
Utilisation : python guides_cotisations_fppn.py <guide.xlsx> <questions_24_27.json> <sortie.xlsx>
"""
import json
import sys
from copy import copy

import openpyxl

GUIDE, Q2427, SORTIE = sys.argv[1:4]
E19 = "Entretien du 19/10/2026 (déclaration à confirmer par document) : "
SANS = "Entretien du 19/10/2026 : pas de réponse relevée sur ce point."

NOTES = {
    1: E19 + "maladie et décès en pourcentage du salaire de base ; autres produits en montants fixes (EFC, PEP, MEDOC, prêts, flotte). Barème à déposer sur la plateforme de partage, non reçu (PBC, item 56).",
    2: E19 + "la répartition du taux entre maladie et décès relève de la comptabilité. Question reportée au chef comptable (onglet 2, question 32).",
    3: E19 + "fichier aller des retraités produit par le responsable des cotisations et déposé sur la plateforme de partage ; fichier retour déposé par la CGRAE.",
    4: E19 + "l'EFC et le PEP sont présentés aux élèves policiers comme obligatoires dès l'école, sans texte. Délai d'au moins trois mois de cotisation avant les prestations (à recouper avec l'entretien du 16/10/2026).",
    5: E19 + "précompte sur salaire. Maladie et décès : fichier aller transmis directement par la direction du personnel de la Police à la Solde, sans plateforme de partage ; retour reçu en direct par le FPPN, canal à préciser. Produits complémentaires, CGRAE et GNA : fichiers aller produits et déposés sur la plateforme de partage par le responsable des cotisations, retours déposés par l'entité concernée.",
    6: SANS,
    7: SANS,
    8: SANS,
    9: E19 + "le fichier aller contient le mécano, le nom, le code du poste, une référence et le montant.",
    10: E19 + "les élèves policiers souscrivent avec leur seul matricule ; le mécano est reporté en masse sur les fiches après son attribution, et le premier précompte peut être décalé d'un mois.",
    11: E19 + "aucun rapprochement entre fichiers aller et retour, aucun état des cotisations, aucun suivi mutualiste par mutualiste.",
    12: SANS,
    13: E19 + "fichier des ayants droit attendu, sans les âges (PBC, item 66).",
    14: SANS,
    15: E19 + "EFC : 6 000 FCFA avant 2021, 10 000 FCFA pour les nouvelles souscriptions depuis 2021 ; les anciens restent à 6 000 FCFA mais peuvent passer à 10 000 FCFA ; aucun document formel. PEP : à partir de 10 000 FCFA, par paliers de 5 000 FCFA ; durée à vérifier.",
    16: SANS,
    17: E19 + "aucun contrôle des retours. Un état séparé des motifs de rejet n'est reçu que depuis environ deux mois. Paiements manuels possibles, enregistrés avec date et heure.",
    18: E19 + "aucun suivi des écarts entre cotisations demandées et précomptées.",
    19: E19 + "l'interlocuteur est le responsable informatique ; il gère les cotisations parce qu'il maîtrise le logiciel qui produit les fichiers. Nom du logiciel non relevé.",
    20: SANS,
}

SRC_CR = "Compte rendu d'entretien du 19/10/2026"
RELANCES = [
    ("21", "Circuit", "Par quel canal le FPPN reçoit-il les fichiers retour maladie (936) et décès (8Z6), et qui les reçoit ?", "Compléter le circuit des précomptes obligatoires.", "Attendu : courriel, support physique ou dépôt par la Solde, avec le nom du destinataire au FPPN.", "", "", SRC_CR, "Fichiers retour manquants (PBC, item 62)", "Incontournable"),
    ("22", "Rejets", "Depuis quelle date exacte la Solde transmet-elle l'état des motifs de rejet ?", "Savoir à partir de quel mois les écarts peuvent être expliqués.", "Attendu : une date et les états reçus. L'interlocuteur cite environ deux mois.", "", "", SRC_CR, "États des motifs de rejet (PBC, item 64)", "Incontournable"),
    ("23", "PEP", "Quelle est la durée d'un plan PEP et quelles sont les conditions de retrait ?", "Établir l'échéancier de l'épargne.", "Attendu : règlement du PEP. La durée n'a pas été précisée le 19/10/2026.", "", "", SRC_CR, "Documentation des produits (PBC, item 57)", "Incontournable"),
    ("24", "Prestations", "Le délai d'au moins trois mois de cotisation s'applique-t-il à tous les produits, y compris la maladie et le décès ? Quel texte le fixe ?", "Lever la contradiction avec l'entretien du 16/10/2026, qui indiquait l'absence de délai de carence pour le régime obligatoire.", "Attendu : un texte et la liste des produits concernés.", "", "", SRC_CR + " ; compte rendu du 16/10/2026", "Règle des trois mois (PBC, item 59)", "Incontournable"),
    ("25", "Grenier", "Quels sont les types de prêts du Grenier et leurs conditions ?", "Rapprocher les trois types cités des huit types de crédit présents dans les fichiers de production.", "Attendu : liste des types de prêts, montants, durées, taux.", "", "", SRC_CR + " ; fichiers de production du Grenier reçus le 15/10/2026", "Types de prêts (PBC, item 70)", "Si le temps le permet"),
    ("26", "MEDOC", "Sur quelle période MEDOC a-t-il fonctionné, et quand a-t-il été suspendu ?", "Expliquer les mois sans fichier du poste 8Z7.", "Attendu : date de démarrage, date de suspension, conventions avec les pharmacies.", "", "", SRC_CR, "Documentation MEDOC (PBC, item 69)", "Si le temps le permet"),
    ("27", "Barème", "Quel document notifie à la Solde les taux maladie et décès appliqués ?", "Valider la répartition 3,77 % et 1,23 % mesurée dans les fichiers retour.", "Attendu : note ou courrier de notification des taux à la Solde.", "", "", SRC_CR, "Barème et notification (PBC, item 56)", "Incontournable"),
    ("28", "Produits", "À quelle date chaque produit a-t-il démarré : EFC, PEP, MEDOC, prêts du Grenier, flotte automobile, retraités ?", "Distinguer les mois sans fichier des mois où le produit n'existait pas.", "Attendu : une date par produit.", "", "", SRC_CR, "Dates de démarrage (PBC, item 60)", "Incontournable"),
]
PROV = [
    ("31", "Cotisations", "Comment les paiements manuels de cotisations sont-ils encaissés, comptabilisés et suivis ?", "Mesurer les cotisations qui échappent aux fichiers de précomptes.", "Le responsable des cotisations indique que des paiements manuels sont possibles et enregistrés avec date et heure (entretien du 19/10/2026). Attendu : compte d'encaissement, état des paiements.", "", "", SRC_CR, "Liste des paiements manuels (PBC, item 65)", "Incontournable"),
    ("32", "Cotisations", "Quelle est la clé de répartition du taux de cotisation entre maladie (poste 936) et décès (poste 8Z6) ?", "Le responsable des cotisations renvoie cette répartition à la comptabilité.", "Les fichiers retour montrent 3,77 % et 1,23 % du salaire de base. Attendu : la clé et le texte qui la fixe.", "", "Taux maladie + taux décès = 5,00 % du salaire de base (fichiers retour)", SRC_CR, "Barème (PBC, item 56)", "Incontournable"),
]


def style(ws, src, dst, ncol=12):
    for c in range(1, ncol + 1):
        ws.cell(dst, c)._style = copy(ws.cell(src, c)._style)
    ws.row_dimensions[dst].height = ws.row_dimensions[src].height


def inserer(ws, ligne, n):
    """Insère n lignes avant `ligne` en décalant les fusions et les validations."""
    fus = [r for r in ws.merged_cells.ranges if r.min_row >= ligne]
    for r in fus:
        ws.unmerge_cells(str(r))
    haut = {r: ws.row_dimensions[r].height for r in range(ligne, ws.max_row + 1)}
    ws.insert_rows(ligne, n)
    for r, h in haut.items():
        ws.row_dimensions[r + n].height = h
    for r in fus:
        ws.merge_cells(start_row=r.min_row + n, end_row=r.max_row + n, start_column=r.min_col, end_column=r.max_col)
    for dv in ws.data_validations.dataValidation:
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(
            " ".join(f"{c.coord.split(':')[0][0]}{c.min_row}:{c.coord.split(':')[0][0]}{c.max_row + n}" for c in dv.sqref.ranges))


def section(ws, ligne, titre, questions, src_titre, src_q):
    inserer(ws, ligne, len(questions) + 1)
    style(ws, src_titre, ligne)
    ws.merge_cells(start_row=ligne, end_row=ligne, start_column=2, end_column=12)
    ws.cell(ligne, 2, titre)
    for k, q in enumerate(questions, 1):
        style(ws, src_q, ligne + k)
        ws.row_dimensions[ligne + k].height = 46
        for j, v in enumerate(q):
            ws.cell(ligne + k, 2 + j, v or None)


wb = openpyxl.load_workbook(GUIDE)

# 4. Cotisations
ws = wb["4. Cotisations"]
for r in range(9, ws.max_row + 1):
    n = ws.cell(r, 2).value
    if str(n).isdigit() and int(n) in NOTES:
        ws.cell(r, 12, NOTES[int(n)])
        ws.cell(r, 12).alignment = copy(ws.cell(r, 6).alignment)
        ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 60)
ws["B3"] = ws["B3"].value + " ; entretien tenu avec le responsable des cotisations, également responsable informatique (nom à compléter)"
clot = next(r for r in range(9, ws.max_row + 1) if ws.cell(r, 2).value == "Clôture de l'entretien")
section(ws, clot - 1, "G. Relances après l'entretien du 19/10/2026", RELANCES, 31, 32)

# 2. Provisionnements : questions 24 à 27 rétablies, puis section H
ws = wb["2. Provisionnements"]
q2427 = json.load(open(Q2427, encoding="utf-8"))
for q in q2427:
    r = next(r for r in range(9, ws.max_row + 1) if str(ws.cell(r, 2).value) == q[0])
    for m in [m for m in ws.merged_cells.ranges if m.min_row == r]:
        ws.unmerge_cells(str(m))
    style(ws, 37, r)
    for j, v in enumerate(q):
        ws.cell(r, 2 + j, v or None)
clot = next(r for r in range(9, ws.max_row + 1) if ws.cell(r, 2).value == "Clôture de l'entretien")
titre_g = next(r for r in range(9, ws.max_row + 1) if str(ws.cell(r, 2).value).startswith("G. "))
section(ws, clot - 1, "H. Cotisations : questions issues de l'entretien du 19/10/2026 (chef comptable)", PROV, titre_g, 37)

wb.save(SORTIE)
print("guides mis à jour")
