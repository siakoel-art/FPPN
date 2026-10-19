# -*- coding: utf-8 -*-
"""
FPPN - Outils de mise en page partagés (charte Adenka) : texte, analyse justifiée, tableaux natifs, sigles
=========================================================================================================

Fonctions reprises du rapport v3 (rapport_v3_fppn.py) pour les comptes rendus d'entretien et le rapport v4.
"""
import math
import re

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

NOIR, GRIS, BLEU, ORANGE, ROUGE, GRIS_F, GRIS_C = "060606", "F2F2F2", "112A40", "EFAC62", "C00000", "53565A", "BBBCBC"
BAS_CONTENU = 16.3          # aucun contenu sous cette ordonnée (cm) : place des sigles et de la source

SIGLES = [
    ("FPPN", r"\bFPPN\b", "Fonds de prévoyance de la Police nationale"),
    ("AIRMS", r"\bAIRMS\b", "Agence ivoirienne de régulation de la mutualité sociale"),
    ("UEMOA", r"\bUEMOA\b", "Union économique et monétaire ouest-africaine"),
    ("PCMS", r"\bPCMS\b", "plan comptable des mutuelles sociales"),
    ("CIMA", r"\bCIMA\b", "Conférence interafricaine des marchés d'assurances"),
    ("CGRAE", r"\bCGRAE\b", "Caisse générale de retraite des agents de l'État"),
    ("BCEAO", r"\bBCEAO\b", "Banque centrale des États de l'Afrique de l'Ouest"),
    ("PEP", r"\bPEP\b", "Plan Épargne Plus"),
    ("EFC", r"\bEFC\b", "entraide fin de carrière"),
    ("CMP", r"\bCMP\b", "sigle des comptes du FPPN, signification à confirmer"),
    ("FCP", r"\bFCP\b", "fonds commun de placement"),
    ("DAT", r"\bDAT\b", "dépôt à terme"),
    ("DAF", r"\bDAF\b", "directeur administratif et financier"),
    ("CNPS", r"\bCNPS\b", "Caisse nationale de prévoyance sociale"),
    ("IS", r"\bIS\b", "impôt sur les salaires (part patronale)"),
    ("FDFP", r"\bFDFP\b", "Fonds de développement de la formation professionnelle"),
    ("SGI", r"\bSGI\b", "société de gestion et d'intermédiation"),
    ("S/C", r"S/C", "prestations / cotisations"),
    ("TCAM", r"\bTCAM\b", "taux de croissance annuel moyen"),
    ("IC", r"\bIC\b", "intervalle de confiance"),
    ("PAR30", r"\bPAR ?30\b", "part de l'encours en retard de plus de 30 jours"),
    ("PSAP", r"\bPSAP\b", "provision pour sinistres à payer"),
    ("IBNR", r"\bIBNR\b", "sinistres survenus non encore déclarés"),
    ("A/E", r"\bA/E\b", "décès observés / décès attendus"),
    ("PBC", r"\bPBC\b", "liste des documents demandés au client"),
    ("FNG", r"\bFNG\b", "fonds national de garantie"),
    ("FOMUS-CI", r"FOMUS-CI", "fonds de garantie des mutuelles sociales de Côte d'Ivoire"),
    ("MEPS", r"\bMEPS\b", "ministère de l'Emploi et de la Protection sociale"),
    ("GNA", r"\bGNA\b", "assureur de la flotte automobile (compte 47111106)"),
    ("MEDOC", r"\bMEDOC\b", "produit santé complémentaire du poste 8Z7"),
    ("TAFIRE", r"\bTAFIRE\b", "tableau financier des ressources et des emplois"),
    ("CAC", r"\bCAC\b", "commissaire aux comptes"),
    ("HAO", r"\bHAO\b", "hors activités ordinaires"),
    ("CMU", r"\bCMU\b", "couverture maladie universelle"),
    ("SHA-256", r"SHA-256", "algorithme d'empreinte numérique"),
    ("CSV", r"\bCSV\b", "fichier texte de tableau"),
    ("OCR", r"\bOCR\b", "reconnaissance optique de caractères"),
    ("FANAF", r"\bFANAF\b", "Fédération des sociétés d'assurances de droit national africaines"),
    ("CPM", r"\bCPM\b", "Comité de politique monétaire"),
    ("SCI", r"\bSCI\b", "société civile immobilière"),
    ("AG", r"\bAG\b", "assemblée générale"),
    ("R²", r"R²", "coefficient de détermination"),
    ("qx", r"\bqx\b", "probabilité de décès entre x et x + 1 ans"),
    ("Lx", r"\bLx\b", "survivants à l'âge x"),
    ("e_x", r"\be_?x\b|\be5[05]\b|\be60\b|\be0\b", "espérance de vie à l'âge x"),
    ("ä_x", r"ä_?x|ä5[05]|ä6[05]", "valeur d'une rente viagère de 1 par an à l'âge x"),
    ("TD, TV", r"\bT[DV] CIMA\b|\(T[DV]\)", "anciennes tables CIMA de décès et de survie"),
    ("936", r"(?:poste |postes |maladie |Maladie |\(|^)936\b", "poste maladie"),
    ("8Z6", r"\b8Z6\b", "poste décès"),
    ("870", r"(?:poste |EFC |\(|^)870\b", "poste EFC"),
    ("998", r"(?:poste |PEP |\(|^)998\b", "poste PEP"),
    ("9X4", r"\b9X4\b", "poste Grenier"),
    ("9W3", r"\b9W3\b", "poste flotte automobile"),
    ("8Z7", r"\b8Z7\b", "poste MEDOC"),
]


def sigles_du_texte(texte):
    trouves = []
    for nom, motif, defi in SIGLES:
        if re.search(motif, texte, flags=re.M):
            trouves.append(f"{nom} : {defi}")
    return trouves



def txt(slide, x, y, l, h, texte, taille=14, gras=False, couleur=NOIR, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(l), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.15)
    for i, ligne in enumerate(str(texte).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ligne
        r.font.name = "Helvetica"
        r.font.size = Pt(taille)
        r.font.bold = gras
        r.font.color.rgb = RGBColor.from_string(couleur)
    return tb


def hauteur_texte(paras, largeur, taille, espace=6):
    cpl = max(1.0, (largeur - 0.3) / (0.0181 * taille))
    lignes = sum(math.ceil(len(p.replace("**", "")) / cpl) for p in paras)
    return lignes * taille * 1.22 * 0.03528 + len(paras) * espace * 0.03528


def corps(slide, x, y, l, h, paras, tmax=14, tmin=9):
    """Texte d'analyse justifié, **gras** accepté ; la taille est la plus grande qui tient dans la zone."""
    if not paras or h < 0.6:
        return None
    # Charte retenue le 19/10/2026 : le paragraphe d'implication reste séparé, sans libellé
    paras = [re.sub(r"^\*\*Implication pour (le FPPN|la mission)\.\*\* ", "", p_) for p_ in paras]
    taille = tmin
    for t in (14, 13, 12, 11.5, 11, 10.5, 10, 9.5, 9):
        if t <= tmax and hauteur_texte(paras, l, t) <= h:
            taille = t
            break
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(l), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Cm(0.15)
    for i, bloc in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.JUSTIFY
        p.space_after = Pt(6)
        for morceau in re.split(r"(\*\*.+?\*\*)", bloc):
            if not morceau:
                continue
            r = p.add_run()
            gras = morceau.startswith("**") and morceau.endswith("**")
            r.text = morceau[2:-2] if gras else morceau
            r.font.name = "Helvetica"
            r.font.size = Pt(taille)
            r.font.bold = gras
            r.font.color.rgb = RGBColor.from_string(NOIR)
    return tb


def boite(slide, x, y, l, h, texte, fond=BLEU, couleur="FFFFFF", taille=12, gras=True, forme=MSO_SHAPE.ROUNDED_RECTANGLE, contour=None):
    sh = slide.shapes.add_shape(forme, Cm(x), Cm(y), Cm(l), Cm(h))
    if fond:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fond)
    else:
        sh.fill.background()
    if contour:
        sh.line.color.rgb = RGBColor.from_string(contour)
        sh.line.width = Pt(1.5)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Cm(0.15)
    for i, ligne in enumerate(str(texte).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = ligne
        r.font.name = "Helvetica"
        r.font.size = Pt(taille if i == 0 else max(taille - 2, 9))
        r.font.bold = gras and i == 0
        r.font.color.rgb = RGBColor.from_string(couleur)
    return sh


def fleche(slide, x1, y1, x2, y2, couleur=GRIS_F, epaisseur=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = RGBColor.from_string(couleur)
    c.line.width = Pt(epaisseur)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


COUL_STATUT = [("Conforme", BLEU, "FFFFFF"), ("Non conforme", ROUGE, "FFFFFF"), ("Au-dessus", ORANGE, NOIR),
               ("En attente", GRIS_C, NOIR), ("À apprécier", GRIS_C, NOIR)]


def hauteurs_lignes(spec, taille):
    hs = []
    for i, vals in enumerate([spec["entetes"]] + spec["lignes"]):
        h = 0
        for v, w in zip(vals, spec["largeurs"]):
            cpl = max(1.0, (w - 0.3) / (0.0186 * taille))
            n = max(1, math.ceil(len(str(v)) / cpl))
            h = max(h, n * taille * 1.2 * 0.03528 + 0.16)
        hs.append(max(h, spec.get("hauteur_ligne") or 0.5))
    return hs


def tableau(slide, spec, y0=3.85, x0=1.4, taille=None, hauteur_max=None):
    """Tableau natif ; renvoie le bas du tableau (cm). La police diminue si le tableau dépasse la hauteur disponible."""
    y = y0
    if spec.get("titre"):
        txt(slide, x0, y, sum(spec["largeurs"]), 0.6, spec["titre"], 12, False, ORANGE, PP_ALIGN.CENTER)
        y += 0.75
    lignes, larg = spec["lignes"], spec["largeurs"]
    taille = taille or spec.get("taille") or (11 if len(lignes) <= 8 else 10)
    hmax = (hauteur_max or (BAS_CONTENU - y))
    hs = hauteurs_lignes(spec, taille)
    while sum(hs) > hmax and taille > 7.5:
        taille -= 0.5
        hs = hauteurs_lignes(spec, taille)
    n = len(lignes) + 1
    shape = slide.shapes.add_table(n, len(larg), Cm(x0), Cm(y), Cm(sum(larg)), Cm(sum(hs)))
    t = shape.table
    tblPr = t._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblPr.set(attr, "0")
    st = tblPr.find(qn("a:tableStyleId"))
    if st is not None:
        st.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"
    for j, w in enumerate(larg):
        t.columns[j].width = Cm(w)
    for i in range(n):
        t.rows[i].height = Cm(hs[i])
        vals = spec["entetes"] if i == 0 else lignes[i - 1]
        for j, v in enumerate(vals):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Cm(0.12)
            c.margin_top = c.margin_bottom = Cm(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            p.text = ""
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", "0")
            pPr.set("indent", "0")
            pPr.set("lvl", "0")
            for bn in pPr.findall(qn("a:buNone")):
                pPr.remove(bn)
            pPr.insert(0, pPr.makeelement(qn("a:buNone"), {}))
            r = p.add_run()
            r.text = str(v)
            r.font.name = "Helvetica"
            r.font.size = Pt(taille)
            r.font.bold = (i == 0) or (j == 0)
            fond, coul = (NOIR, "FFFFFF") if i == 0 else ((GRIS if i % 2 == 0 else "FFFFFF"), NOIR)
            if i > 0 and spec.get("statut_col") == j:
                for debut, f_, c_ in COUL_STATUT:
                    if str(v).startswith(debut):
                        fond, coul = f_, c_
                        r.font.bold = True
                        break
            r.font.color.rgb = RGBColor.from_string(coul)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string(fond)
    return y + sum(hs)




def ligne_sigles(slide, textes):
    """Ajoute en bas de page la définition des sigles présents dans les textes de la diapositive."""
    sig_ = sigles_du_texte("\n".join(textes))
    if sig_:
        ligne_s = "Sigles : " + " | ".join(sig_)
        txt(slide, 1.02, 16.48, 31.75, 0.66, ligne_s, 7 if len(ligne_s) < 330 else 6.5, False, "595959", PP_ALIGN.LEFT, MSO_ANCHOR.BOTTOM)
