# -*- coding: utf-8 -*-
"""
FPPN - Mise à jour des guides d'entretien après les entretiens du 19/10/2026 avec M. Koffi, chef comptable
- onglet « 3. Placements » : réponses relevées (colonne Notes) et section G de relances ;
- onglet « 2. Provisionnements » : réponses aux questions 31 à 38 et section J de relances ;
- onglet « 4. Cotisations » : nom de l'interlocuteur (M. Kouamé Attuoman).
Utilisation : python guides_placements_provisions_fppn.py <guide.xlsx>   (fichier modifié en place)
"""
import sys
from copy import copy

import openpyxl
from openpyxl.worksheet.cell_range import MultiCellRange

GUIDE = sys.argv[1]
EP = "Entretien du 19/10/2026 à 14h avec M. Koffi (déclaration à confirmer par document) : "
EV = "Entretien du 19/10/2026 à 16h avec M. Koffi (déclaration à confirmer par document) : "
SANS = "Entretiens du 19/10/2026 : pas de réponse relevée sur ce point."
SRC = "Compte rendu d'entretien du 19/10/2026"
NOTES_P = {
    1: EP + "DAT BOA alimenté chaque mois (60 millions par mois en 2026) et débouclé en fin d'année ; FCP de 500 millions géré par United Capital (SGI) ; terrains. Aucune obligation ni action citée.",
    2: EP + "BOA pour le DAT, United Capital (SGI) pour le FCP ; Orabank par le passé.",
    3: EP + "le DAT court jusqu'au 30 ou 31 décembre ; pas d'échéance citée pour le FCP.",
    4: SANS,
    5: EP + "FCP : rendement annoncé de 8 à 10 %, non garanti. Taux du DAT non communiqué.",
    6: EP + "aucune politique écrite ; toutes les décisions viennent de la direction générale ; ni comité ni avis du conseil d'administration.",
    7: EP + "le montant du DAT est fixé selon les départs à la retraite prévus (liste de la direction du personnel) et une simulation de la comptabilité.",
    8: SANS,
    9: EP + "liste des départs à la retraite fournie par la direction du personnel, rapprochée des souscripteurs, puis simulation par la comptabilité.",
    10: SANS,
    11: EP + "le Trésor vire parfois partiellement ; le reste est versé plus tard ; le DAF relance.",
    12: EP + "le Grenier tient sa propre comptabilité de prêts dans un logiciel de microfinance ; impayés non abordés.",
    13: EP + "le terrain du siège a été réévalué par un expert immobilier (compte 10620000) ; date non précisée.",
    14: EP + "deux villas et un autre bâtiment à Aboisso sont loués à usage d'habitation ; pas de programme de location-vente cité.",
    15: SANS,
    16: EP + "les placements sont suivis sur les relevés bancaires ; les rapprochements bancaires ne sont plus supervisés.",
    17: SANS,
    18: EP + "aucun placement bloqué cité ; incident passé avec Orabank (versements arrêtés, DAT remboursé au terme).",
}
NOTES_V = {
    31: SANS,
    32: EV + "3,5 % pour la maladie et 1,5 % pour le décès, sans document qui fixe la répartition.",
    33: EV + "montants hérités d'environ 2004, avant le PCMS ; jamais recalculés ; archives perdues lors d'un sinistre en 2010. Aucune provision maladie ; calcul de fin d'année des allocations décès restant à payer.",
    34: EV + "les factures reçues jusqu'en mars sont rattachées à l'exercice (dettes) ; aucune provision maladie n'est calculée.",
    35: EV + "le Grenier tient sa propre comptabilité de prêts ; la dépréciation n'a pas été abordée.",
    36: EP + "le FPPN est apporteur d'affaires de GNA et perçoit 20 % des primes collectées, en fin d'année.",
    37: EV + "convention collective appliquée à l'ancienneté ; risques professionnels à 2 % ; fichiers de calcul et déclarations CNPS demandés.",
    38: SANS,
}
REL_P = [
    ("19", "DAT", "Quel est le taux du DAT BOA, et pourquoi le compte 52210200 s'intitule-t-il « Entraide fin de carrière » ?", "Confirmer la nature et le rendement du principal placement.", "Attendu : convention et relevés.", "", "", SRC, "PBC, item 80", "Incontournable"),
    ("20", "FCP", "Quelle est la valeur liquidative du FCP au 31/08/2026 et son rendement réel depuis la souscription ?", "Le FCP n'est pas revalorisé dans les comptes.", "Attendu : relevés du gestionnaire.", "", "Rendement = (valeur finale - valeur initiale) / valeur initiale", SRC, "PBC, item 81", "Incontournable"),
    ("21", "Orabank", "En quelle année l'incident Orabank a-t-il eu lieu ? Les balances montrent un DAT ouvert en 2022 et soldé en 2023.", "Dater l'incident cité.", "Attendu : correspondance avec la banque.", "", "", SRC, "PBC, item 83", "Si le temps le permet"),
    ("22", "GNA", "Sur quelle assiette la commission de 20 % est-elle calculée ? En 2024, la commission comptabilisée ne vaut pas 20 % des précomptes 9W3.", "Rapprocher les commissions des primes.", "Attendu : convention et décomptes.", "", "Commission = 20 % × primes collectées", SRC, "PBC, item 89", "Si le temps le permet"),
    ("23", "Terrains", "Quels sont les terrains du FPPN, leur situation et les projets prévus ?", "Évaluer les actifs immobiliers.", "Attendu : inventaire et titres.", "", "", SRC, "PBC, item 84", "Si le temps le permet"),
    ("24", "Hôpital", "Où en sont les démarches auprès de l'État sur la créance de l'hôpital de la Police ?", "Apprécier le recouvrement de 3 407 519 850 FCFA.", "Attendu : correspondances et décision.", "", "", SRC, "PBC, item 87", "Incontournable"),
]
REL_V = [
    ("39", "Décès", "Dans quel compte le calcul de fin d'année des allocations décès restant à payer est-il comptabilisé ?", "Le compte 15802000 ne bouge pas depuis 2020.", "Attendu : compte et calculs 2020 à 2025.", "", "", SRC, "PBC, item 92", "Incontournable"),
    ("40", "Décès", "Quel est le montant de l'allocation décès d'un officier ?", "Montant non lisible dans la transcription.", "Attendu : barème officiel.", "", "", SRC, "PBC, item 93", "Incontournable"),
    ("41", "Provisions", "Quand et par quel courrier l'AIRMS a-t-elle demandé le reclassement des provisions en compte 158 ?", "Le reclassement est visible dans la balance 2024.", "Attendu : courrier de l'AIRMS.", "", "", SRC, "PBC, item 95", "Si le temps le permet"),
]


def style(ws, src, dst, ncol=12):
    for c in range(1, ncol + 1):
        ws.cell(dst, c)._style = copy(ws.cell(src, c)._style)
    ws.row_dimensions[dst].height = ws.row_dimensions[src].height


def section(ws, titre, questions):
    clot = next(r for r in range(9, ws.max_row + 1) if ws.cell(r, 2).value == "Clôture de l'entretien")
    titre_mod = max(r for r in range(9, clot) if ws.cell(r, 2).value and not str(ws.cell(r, 2).value).isdigit())
    q_mod = titre_mod + 1
    ligne, n = clot - 1, len(questions) + 1
    fus = [m for m in ws.merged_cells.ranges if m.min_row >= ligne]
    for m in fus:
        ws.unmerge_cells(str(m))
    haut = {x: ws.row_dimensions[x].height for x in range(ligne, ws.max_row + 1)}
    ws.insert_rows(ligne, n)
    for x, h in haut.items():
        ws.row_dimensions[x + n].height = h
    for m in fus:
        ws.merge_cells(start_row=m.min_row + n, end_row=m.max_row + n, start_column=m.min_col, end_column=m.max_col)
    style(ws, titre_mod, ligne)
    ws.merge_cells(start_row=ligne, end_row=ligne, start_column=2, end_column=12)
    ws.cell(ligne, 2, titre)
    for k, q in enumerate(questions, 1):
        style(ws, q_mod, ligne + k)
        ws.row_dimensions[ligne + k].height = 46
        for j, v in enumerate(q):
            ws.cell(ligne + k, 2 + j, v or None)
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(f"K8:K{ligne + len(questions)}")


def notes(ws, d):
    for r in range(9, ws.max_row + 1):
        n = ws.cell(r, 2).value
        if str(n).isdigit() and int(n) in d:
            ws.cell(r, 12, d[int(n)])
            ws.cell(r, 12).alignment = copy(ws.cell(r, 6).alignment)
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 60)


wb = openpyxl.load_workbook(GUIDE)
ws = wb["3. Placements"]
notes(ws, NOTES_P)
ws["B3"] = ws["B3"].value + " ; entretien tenu le 19/10/2026 à 14h avec M. Koffi, chef comptable"
section(ws, "G. Relances après l'entretien du 19/10/2026", REL_P)
ws = wb["2. Provisionnements"]
notes(ws, NOTES_V)
section(ws, "J. Relances après l'entretien du 19/10/2026 (16h)", REL_V)
ws = wb["4. Cotisations"]
ws["B3"] = ws["B3"].value.replace("(nom à compléter)", "(M. Kouamé Attuoman)")
wb.save(GUIDE)
print("guides mis à jour")
