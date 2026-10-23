# -*- coding: utf-8 -*-
"""
FPPN - Classeur des calculs : barre de navigation, volets figés, tableaux Excel, module VBA
==========================================================================================
1. Décale chaque onglet d'une ligne vers le bas (formules, mises en forme conditionnelles, fusions, hauteurs ajustées) et
   écrit en ligne 1 une barre de boutons-liens : Guide, onglet précédent, onglet suivant, contrôle, onglets sources.
   Les liens marchent sans macro, sur Windows, Mac et Excel en ligne.
2. Fige seulement la barre de navigation et le titre (les blocs d'explication défilent).
3. Transforme chaque tableau (en-tête noir) en tableau Excel : en défilant, Excel remplace les lettres des colonnes par
   les titres du tableau, qui restent donc toujours visibles.
4. Écrit le module VBA « FPPN_navigation.bas » (boutons macros facultatifs), à importer en deux minutes.
"""
import re
from copy import copy

import openpyxl
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

POLICE = "Helvetica"
BLEU, ORANGE, GRIS = "112A40", "EFAC62", "53565A"
REF = re.compile(r"(?<![A-Za-z0-9_.!])((?:'[^']+'|[A-Za-z0-9_]+)!)?(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(A-Za-z_!])")
REF_FEUILLE = re.compile(r"(?<![A-Za-z0-9_.])('[^']+'|[A-Za-z][A-Za-z0-9_]*|[0-9][A-Za-z0-9_]*)!")


def decaler_formule(f, n=1):
    """Ajoute n à toutes les lignes des références de cellule d'une formule (hors chaînes entre guillemets)."""
    morceaux = f.split('"')
    for i in range(0, len(morceaux), 2):
        morceaux[i] = REF.sub(lambda m: f"{m.group(1) or ''}{m.group(2)}{m.group(3)}{m.group(4)}{int(m.group(5)) + n}", morceaux[i])
    return '"'.join(morceaux)


def decaler_plage(sqref, n=1):
    return " ".join(re.sub(r"([A-Z]{1,3})(\d+)", lambda m: f"{m.group(1)}{int(m.group(2)) + n}", p) for p in str(sqref).split())


def feuilles_citees(f):
    noms = set()
    for m in REF_FEUILLE.finditer(f):
        noms.add(m.group(1).strip("'"))
    return noms


def copier_decale(src, dst, n=1):
    for row in src.iter_rows():
        for c in row:
            if c.value is None and not c.has_style:
                continue
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                v = decaler_formule(v, n)
            d = dst.cell(row=c.row + n, column=c.column, value=v)
            if c.has_style:
                d._style = copy(c._style)
            if c.comment:
                d.comment = copy(c.comment)
            if c.hyperlink:
                d.hyperlink = copy(c.hyperlink)
                d.hyperlink.ref = d.coordinate
    for k, dim in src.column_dimensions.items():
        dst.column_dimensions[k].width = dim.width
        dst.column_dimensions[k].hidden = dim.hidden
    for k, dim in src.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[k + n].height = dim.height
    for m in src.merged_cells.ranges:
        dst.merge_cells(start_row=m.min_row + n, start_column=m.min_col, end_row=m.max_row + n, end_column=m.max_col)
    for cf in src.conditional_formatting:
        for rule in cf.rules:
            r2 = copy(rule)
            r2.formula = [decaler_formule(x, n) for x in (rule.formula or [])]
            dst.conditional_formatting.add(decaler_plage(cf.sqref, n), r2)
    for dv in src.data_validations.dataValidation:
        d2 = copy(dv)
        d2.sqref = openpyxl.worksheet.cell_range.MultiCellRange(decaler_plage(dv.sqref, n))
        dst.add_data_validation(d2)
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines
    dst.sheet_properties.tabColor = src.sheet_properties.tabColor
    dst.page_setup.orientation = src.page_setup.orientation
    dst.page_setup.fitToWidth, dst.page_setup.fitToHeight = src.page_setup.fitToWidth, src.page_setup.fitToHeight
    dst.sheet_properties.pageSetUpPr = copy(src.sheet_properties.pageSetUpPr)


def ligne_entete_noire(ws, r):
    """Vrai si la ligne r commence (colonne A) une ligne d'en-tête de tableau (fond noir ou bleu nuit)."""
    c = ws.cell(r, 1)
    rgb = str(c.fill.fgColor.rgb) if c.fill and c.fill.fill_type == "solid" and c.fill.fgColor is not None else ""
    return rgb[-6:] in ("000000", "060606", "112A40") and isinstance(c.value, str) and c.value.strip() != ""


def tableaux_excel(ws, prefixe, deja):
    """Ajoute un tableau Excel sur chaque tableau détecté (en-tête noir en colonne A) si ses titres sont des textes uniques."""
    fusion = set()
    for m in ws.merged_cells.ranges:
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                fusion.add((r, c))
    n = 0
    r = 3          # ligne 1 : barre de navigation ; ligne 2 : titre
    while r <= ws.max_row:
        if not ligne_entete_noire(ws, r) or (r, 1) in fusion:
            r += 1
            continue
        c2 = 1
        while c2 + 1 <= ws.max_column and isinstance(ws.cell(r, c2 + 1).value, str) and ws.cell(r, c2 + 1).value.strip() and (r, c2 + 1) not in fusion:
            c2 += 1
        titres = [str(ws.cell(r, c).value).strip() for c in range(1, c2 + 1)]
        f = r
        while f + 1 <= ws.max_row:
            v = ws.cell(f + 1, 1).value
            ligne = [ws.cell(f + 1, c).value for c in range(1, c2 + 1)]
            if all(x is None or x == "" for x in ligne):
                break
            if isinstance(v, str) and v.startswith(("Total", "Moyenne", "Croissance annuelle moyenne", "Source", "Nombre et âge")):
                break
            if any((f + 1, c) in fusion for c in range(1, c2 + 1)):
                break
            f += 1
        ok = f > r and len(set(t.lower() for t in titres)) == len(titres) and all(titres) and c2 >= 2
        if ok:
            for c in range(1, c2 + 1):
                ws.cell(r, c).value = titres[c - 1]
            n += 1
            nom = f"{prefixe}_{n}"
            while nom in deja:
                nom += "b"
            deja.add(nom)
            t = Table(displayName=nom, ref=f"A{r}:{L(c2)}{f}")
            t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False, showLastColumn=False, showRowStripes=False, showColumnStripes=False)
            t.tableColumns = [TableColumn(id=i + 1, name=titres[i]) for i in range(c2)]
            t.autoFilter = None
            ws.add_table(t)
        r = f + 1
    return n


def finaliser(wb, guide="0_guide", controle="0_controle"):
    noms = list(wb.sheetnames)
    sources = {}
    for nom in noms:
        cit = set()
        for row in wb[nom].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    cit |= feuilles_citees(c.value)
        sources[nom] = [x for x in sorted(cit) if x in noms and x != nom]
    deja = set()
    for i, nom in enumerate(noms):
        src = wb[nom]
        gel = src.freeze_panes
        filtre = src.auto_filter.ref
        dst = wb.create_sheet(f"tmp{i}", i)
        copier_decale(src, dst, 1)
        del wb[nom]
        dst.title = nom
        # barre de navigation
        boutons = [("⌂ Guide", guide), ("◀ Précédent", noms[i - 1] if i > 0 else None), ("Suivant ▶", noms[i + 1] if i + 1 < len(noms) else None), ("✓ Contrôle", controle)]
        j = 1
        for lib, cible in boutons:
            if not cible or cible == nom:
                continue
            c = dst.cell(1, j, lib)
            c.hyperlink = f"#'{cible}'!A1"
            c.font = Font(name=POLICE, size=9, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=BLEU)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = Border(*(Side(style="thin", color="FFFFFF"),) * 4)
            j += 1
        if sources[nom]:
            c = dst.cell(1, j, "Chiffres tirés de :")
            c.font = Font(name=POLICE, size=9, italic=True, color=GRIS)
            c.alignment = Alignment(horizontal="right", vertical="center")
            j += 1
            for s in sources[nom][:6]:
                c = dst.cell(1, j, s)
                c.hyperlink = f"#'{s}'!A1"
                c.font = Font(name=POLICE, size=9, bold=True, color=BLEU, underline="single")
                c.fill = PatternFill("solid", fgColor=ORANGE)
                c.alignment = Alignment(horizontal="center", vertical="center")
                j += 1
        dst.row_dimensions[1].height = 20
        # volets : barre de navigation et titre seulement
        dst.freeze_panes = "A3"
        if filtre:
            dst.auto_filter.ref = decaler_plage(filtre, 1)
        n = tableaux_excel(dst, "T" + re.sub(r"[^A-Za-z0-9]", "", nom)[:20], deja)
        if n and dst.auto_filter.ref:
            dst.auto_filter.ref = None
    return sources


VBA = r'''Attribute VB_Name = "FPPN_navigation"
' FPPN - Classeur des calculs : boutons de navigation (facultatifs)
' Mode d'emploi (2 minutes) :
'   1. Alt + F11 (éditeur VBA) ; menu Fichier > Importer un fichier... ; choisir FPPN_navigation.bas ; fermer l'éditeur.
'   2. Alt + F8 ; choisir InstallerBoutons ; Exécuter.
'   3. Fichier > Enregistrer sous > type « Classeur Excel (prenant en charge les macros) (*.xlsm) ».
' Les liens bleus de la ligne 1 de chaque onglet marchent sans ces macros.
Option Explicit

Public Sub AllerAuGuide()
    On Error Resume Next
    ThisWorkbook.Worksheets("0_guide").Activate
    ThisWorkbook.Worksheets("0_guide").Range("A1").Select
End Sub

Public Sub AllerAuControle()
    On Error Resume Next
    ThisWorkbook.Worksheets("0_controle").Activate
End Sub

Public Sub OngletSuivant()
    Dim i As Long
    i = ActiveSheet.Index
    Do While i < ThisWorkbook.Sheets.Count
        i = i + 1
        If ThisWorkbook.Sheets(i).Visible = xlSheetVisible Then ThisWorkbook.Sheets(i).Activate: Exit Sub
    Loop
End Sub

Public Sub OngletPrecedent()
    Dim i As Long
    i = ActiveSheet.Index
    Do While i > 1
        i = i - 1
        If ThisWorkbook.Sheets(i).Visible = xlSheetVisible Then ThisWorkbook.Sheets(i).Activate: Exit Sub
    Loop
End Sub

Public Sub ListeDesOnglets()
    ' Affiche la liste de tous les onglets (clic droit sur les flèches en bas à gauche, version Windows)
    On Error Resume Next
    Application.CommandBars("Workbook tabs").ShowPopup
    If Err.Number <> 0 Then AllerAuGuide
End Sub

Public Sub InstallerBoutons()
    Dim ws As Worksheet, b As Object, x As Double, k As Long
    Dim libelles As Variant, macros As Variant
    libelles = Array("Guide", "< Précédent", "Suivant >", "Tous les onglets", "Contrôle")
    macros = Array("AllerAuGuide", "OngletPrecedent", "OngletSuivant", "ListeDesOnglets", "AllerAuControle")
    For Each ws In ThisWorkbook.Worksheets
        For Each b In ws.Buttons
            If Left(b.Name, 4) = "nav_" Then b.Delete
        Next b
        x = ws.Range("L1").Left
        For k = 0 To UBound(libelles)
            Set b = ws.Buttons.Add(x + k * 92, ws.Range("A1").Top + 1, 88, 17)
            b.Name = "nav_" & k
            b.Caption = libelles(k)
            b.OnAction = macros(k)
            b.Font.Size = 8
        Next k
    Next ws
    MsgBox "Boutons installés sur " & ThisWorkbook.Worksheets.Count & " onglets. Enregistrez le classeur au format .xlsm.", vbInformation
End Sub

Public Sub SupprimerBoutons()
    Dim ws As Worksheet, b As Object
    For Each ws In ThisWorkbook.Worksheets
        For Each b In ws.Buttons
            If Left(b.Name, 4) = "nav_" Then b.Delete
        Next b
    Next ws
End Sub
'''


def ecrire_vba(chemin):
    with open(chemin, "w", encoding="cp1252", newline="\r\n") as f:
        f.write(VBA)
