# -*- coding: utf-8 -*-
"""
FPPN - Demande complémentaire après les entretiens placements et provisions du 19/10/2026 (items 80 à 99),
reportée dans la PBC générale (sections et points à clarifier).
Utilisation : python pbc_placements_provisions_fppn.py <pbc_modele.xlsx> <pbc_generale.xlsx> <sortie_complementaire.xlsx>
(la PBC générale est modifiée en place)
"""
import datetime as dt
import sys
from copy import copy

import openpyxl
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation

MODELE, GENERALE, SORTIE = sys.argv[1:4]
D_DEM, D_LIM = dt.datetime(2026, 10, 19), dt.datetime(2026, 10, 26)
SECTIONS = [
    ("Placements et autres actifs (entretien de 14h)", [
        ("Conventions, taux, échéances et relevés du DAT BOA (compte 52210200) de 2020 au 31/08/2026.", "Urgent", "Electronique"),
        ("Bulletin de souscription au FCP géré par United Capital (SGI), valeurs liquidatives et relevés de 2025 au 31/08/2026.", "Urgent", "Electronique"),
        ("Inventaire et tableau de variation de tous les placements de 2016 au 31/08/2026 : souscriptions, intérêts, débouclages, par banque ou gestionnaire.", "Urgent", "Electronique"),
        ("Correspondance avec Orabank sur l'arrêt des prélèvements mensuels et le remboursement du DAT.", "Moyen", None),
        ("Titres de propriété ou actes d'acquisition des terrains, situation de chaque terrain (projet, attribution, litige).", "Moyen", None),
        ("Baux et encaissements des immeubles d'Aboisso.", "Faible", None),
        ("Rapport de réévaluation du siège et décision de comptabilisation au compte 10620000.", "Urgent", None),
        ("Dossier de la créance sur l'hôpital de la Police (compte 47111100) : pièces, correspondances avec l'État, décision d'autonomie de l'hôpital.", "Moyen", None),
        ("Bilan d'ouverture de la clinique du FPPN (CMP) et détail des comptes 47111104 et 47111105.", "Moyen", "Electronique"),
        ("Convention d'apporteur d'affaires avec GNA et décomptes annuels des commissions (20 % des primes collectées).", "Moyen", None),
        ("Liste des comptes bancaires, actifs et inactifs, avec leurs signataires.", "Moyen", "Electronique"),
        ("État des chèques en suspens et des chèques annulés lors des changements de signataires.", "Faible", "Electronique"),
    ]),
    ("Provisions (entretien de 16h)", [
        ("Fichiers des décès transmis par la direction du personnel et calcul de fin d'année des allocations restant à payer, de 2020 à 2025, avec le compte utilisé.", "Urgent", "Electronique"),
        ("Barème officiel des allocations décès par grade (sous-officier, officier, commissaire).", "Urgent", None),
        ("Contrats et décomptes de gestion de la garantie décès avec LMAI et Sanlam.", "Moyen", None),
        ("Courrier de l'AIRMS demandant le reclassement des provisions en compte 158.", "Moyen", None),
        ("Notification du taux de risques professionnels (2 %) par la CNPS.", "Moyen", None),
        ("Déclarations CNPS de 2020 au 31/08/2026.", "Moyen", "Electronique"),
        ("Fichiers de calcul de la provision pour départ à la retraite de 2020 à 2026 (complète l'item 78).", "Moyen", "Electronique"),
        ("Décision fixant le taux servi de 3,5 % sur le PEP et l'EFC.", "Urgent", None),
    ]),
]
POINTS = [
    ("DAT BOA", "Le chef comptable le décrit comme un placement du FPPN qui sert aussi à payer l'EFC ; le compte s'intitule « DAT BOA Entraide fin de carrière ».", "Confirmer la nature et l'affectation du DAT (item 80)."),
    ("Compte 10620000", "Réévaluation du terrain du siège selon le chef comptable ; l'intitulé du compte parle de matériels.", "Transmettre le rapport de réévaluation (item 86)."),
    ("Commissions GNA", "Commission déclarée de 20 % des primes collectées ; en 2024, la commission comptabilisée ne vaut pas 20 % des précomptes 9W3.", "Transmettre la convention et les décomptes (item 89)."),
    ("Allocations décès", "Calcul de fin d'année des allocations restant à payer ; compte de comptabilisation non précisé.", "Indiquer le compte et transmettre les calculs (item 92)."),
    ("Barème décès des officiers", "Montant de l'allocation des officiers non confirmé.", "Transmettre le barème (item 93)."),
]
COLS = range(2, 14)


# ---------------------------------------------------------------- PBC complémentaire
wb = openpyxl.load_workbook(MODELE)
ws = wb.worksheets[0]
ws.title = "PBC Placements et provisions"
ws["C4"] = "AUDIT DU FPPN 2020 – 31 AOÛT 2026 : DEMANDE COMPLÉMENTAIRE PLACEMENTS ET PROVISIONS"
ws["C5"] = "SUITE AUX ENTRETIENS DU 19/10/2026 AVEC LE CHEF COMPTABLE (PLACEMENTS À 14H, PROVISIONS À 16H)"
st_vide = {c: copy(ws.cell(9, c)._style) for c in COLS}
st_sect = {c: copy(ws.cell(10, c)._style) for c in COLS}
st_item = {c: copy(ws.cell(11, c)._style) for c in COLS}
for r in range(10, ws.max_row + 1):
    for c in COLS:
        ws.cell(r, c).value = None
        ws.cell(r, c)._style = copy(st_vide[c])
    ws.row_dimensions[r].height = 13.5
ws.data_validations.dataValidation = []
r, prec, premier = 10, None, None
for titre, items in SECTIONS:
    for c in COLS:
        ws.cell(r, c)._style = copy(st_sect[c])
    ws.cell(r, 3, titre)
    ws.row_dimensions[r].height = 18
    r += 1
    for texte, prio, support in items:
        for c in COLS:
            ws.cell(r, c)._style = copy(st_item[c])
        ws.cell(r, 2, "=80" if prec is None else f"=+B{prec}+1")
        ws.cell(r, 3, texte)
        ws.cell(r, 5, prio)
        ws.cell(r, 7, D_DEM)
        ws.cell(r, 8, D_LIM)
        ws.cell(r, 10, support)
        ws.cell(r, 13, "Non Reçu")
        ws.row_dimensions[r].height = 21 if len(texte) < 110 else 36
        premier = premier or r
        prec = r
        r += 1
    r += 1
ws["M7"] = f'=IFERROR(COUNTIF($M${premier}:$M${prec},"Reçu")/(COUNTA($M${premier}:$M${prec})-COUNTIF($M${premier}:$M${prec},"Non Applicable")),0)'
for col, liste in (("E", '"Urgent,Moyen,Faible"'), ("J", '"Physique,Electronique,Physique et Electronique"'),
                   ("M", '"Reçu,Non Reçu,Non Applicable,Partiellement Reçu"')):
    dv = DataValidation(type="list", formula1=liste, allow_blank=True)
    dv.add(f"{col}{premier}:{col}{prec}")
    ws.add_data_validation(dv)
for nom in [n for n in wb.sheetnames if n != ws.title]:
    del wb[nom]
wb.save(SORTIE)

# ---------------------------------------------------------------- PBC générale
wb = openpyxl.load_workbook(GENERALE)
ws = wb["PBC Actuariat FPPN"]
num = [x for x in range(11, ws.max_row + 1) if str(ws.cell(x, 2).value or "").startswith("=")]
assert len(num) == 79, len(num)
dern = num[-1]
sect = next(x for x in range(dern, 9, -1) if ws.cell(x, 3).value and not ws.cell(x, 2).value)
st_s = {c: copy(ws.cell(sect, c)._style) for c in COLS}
st_i = {c: copy(ws.cell(dern, c)._style) for c in COLS}
r, prec = dern + 1, dern
for titre, items in SECTIONS:
    r += 1
    for c in COLS:
        ws.cell(r, c)._style = copy(st_s[c])
    ws.cell(r, 3, titre.replace("(entretien", "(entretiens du 19/10/2026,").replace("de 14h)", "14h)").replace("de 16h)", "16h)"))
    ws.row_dimensions[r].height = ws.row_dimensions[sect].height
    for texte, prio, support in items:
        r += 1
        for c in COLS:
            ws.cell(r, c)._style = copy(st_i[c])
        ws.cell(r, 2, f"=+B{prec}+1")
        ws.cell(r, 3, texte)
        ws.cell(r, 5, prio)
        ws.cell(r, 7, D_DEM)
        ws.cell(r, 8, D_LIM)
        ws.cell(r, 10, support)
        ws.cell(r, 11, "Demande issue des entretiens du 19/10/2026 avec le chef comptable (PBC complémentaire placements et provisions).")
        ws.cell(r, 13, "Non Reçu")
        ws.row_dimensions[r].height = 21 if len(texte) < 110 else 36
        prec = r
    r += 1
plage = f"$M$11:$M${prec}"
ws["M7"] = f'=IFERROR(COUNTIF({plage},"Reçu")/(COUNTA({plage})-COUNTIF({plage},"Non Applicable")),0)'
for dv in ws.data_validations.dataValidation:
    col = str(dv.sqref).split(":")[0][0]
    dv.sqref = MultiCellRange(f"{col}11:{col}{prec}")
sy = wb["Synthèse du pointage"]
for k, st in ((6, "Reçu"), (7, "Partiellement Reçu"), (8, "Non Reçu")):
    sy.cell(k, 3, f"=COUNTIF('PBC Actuariat FPPN'!{plage},\"{st}\")")
sy["D8"] = sy["D8"].value + ", " + ", ".join(str(n) for n in range(80, 100))
sy["D9"] = "1 à 99"
pc = wb["Points à clarifier"]
d = pc.max_row
for k, (sujet, constat, demande) in enumerate(POINTS, 1):
    rr = d + k
    for c in range(2, 6):
        pc.cell(rr, c)._style = copy(pc.cell(d, c)._style)
    pc.row_dimensions[rr].height = pc.row_dimensions[d].height
    pc.cell(rr, 2, int(pc.cell(d, 2).value) + k)
    pc.cell(rr, 3, sujet)
    pc.cell(rr, 4, constat)
    pc.cell(rr, 5, demande)
wb.save(GENERALE)
print(f"complémentaire : items 80 à {80 + sum(len(i) for _, i in SECTIONS) - 1} ; générale : jusqu'à la ligne {prec}, points jusqu'à {pc.max_row}")
