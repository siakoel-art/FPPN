# -*- coding: utf-8 -*-
"""
FPPN - Rapport actuariel et statistique, version 3 (PowerPoint, charte Adenka, Timesaver)
=========================================================================================

Entrées : resultats.json, complements.json, ratios.json (analyses), dossier des images de code, skill adenka-powerpoint.
Sorties : deck .pptx (graphiques et tableaux natifs, schémas modifiables, notes du présentateur) et
          fichier de traçabilité (chaque chiffre écrit : diapositive, valeur exacte, texte affiché).

Nouveautés de la version 3 :
- partie « Cadre réglementaire et conformité prudentielle » (Décision n° 05/2022/COM/UEMOA, règlement
  d'exécution n° 004/2022/COM/UEMOA, règlement n° 07/2009/CM/UEMOA, PCMS, arrêté n° 00115/MEPS/CAB) ;
- sigles de chaque diapositive définis en bas de page ;
- analyse en trois temps sous chaque tableau, graphique ou illustration : constat en gras, preuves et
  explication, puis un paragraphe d'implication pour le FPPN et la mission ;
- gabarits Timesaver pour les méthodes, la chaîne de traitement et les diapositives « À retenir » ;
- extraits de code mis en forme comme dans VS Code sous Windows.

Règles de présentation : pourcentages à deux décimales ; montants exacts en FCFA, sans décimale, dans
les textes et les tableaux ; graphiques : données exactes en FCFA, affichées en millions.
Utilisation : python rapport_v3_fppn.py <resultats.json> <complements.json> <ratios.json> <dossier_images> <skill> <sortie.pptx>
"""
import io
import json
import math
import os
import re
import sys

RES, COMP, RAT, IMG, SKILL, SORTIE = sys.argv[1:7]
sys.path.insert(0, os.path.join(SKILL, "scripts"))
from pptx_tools import selection, remplir, finaliser  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_TICK_LABEL_POSITION  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Cm, Pt  # noqa: E402

R = json.load(open(RES, encoding="utf-8"))
C = json.load(open(COMP, encoding="utf-8"))
Q = json.load(open(RAT, encoding="utf-8"))
TIMESAVER = os.path.join(SKILL, "assets", "Timesaver_Adenka.pptx")
NOIR, GRIS, BLEU, ORANGE, ROUGE, GRIS_F, GRIS_C = "060606", "F2F2F2", "112A40", "EFAC62", "C00000", "53565A", "BBBCBC"
BAS_CONTENU = 16.3          # aucun contenu sous cette ordonnée (cm) : place des sigles et de la source

# ---------------------------------------------------------------------------
# Formats et traçabilité
# ---------------------------------------------------------------------------
TRACE = []
SL = []
COURANTE = None  # diapositive en cours de post-traitement (sinon : la prochaine à créer)
FIG, TAB = [], []


def nb(x, dec=0):
    return f"{x:,.{dec}f}".replace(",", " ").replace(".", ",")


def V(label, valeur, texte):
    """Enregistre le chiffre écrit (diapositive courante) pour le contrôle automatique."""
    TRACE.append({"diapositive": COURANTE or len(SL) + 1, "indicateur": label, "valeur_exacte": valeur, "texte": texte})
    return texte


def pc(label, x):
    return V(label, x, nb(100 * x, 2) + " %")


def pcv(label, x):
    return V(label, x, nb(x, 2) + " %")


def F(label, x):
    return V(label, x, nb(round(x)) + " FCFA")


def N(label, x, dec=0):
    return V(label, x, nb(x, dec))


def fig(titre):
    FIG.append([len(FIG) + 1, titre, len(SL) + 1])
    return f"Figure {len(FIG)} : {titre}"


def tab(titre):
    TAB.append([len(TAB) + 1, titre, len(SL) + 1])
    return f"Tableau {len(TAB)} : {titre}"


def an(liste, a):
    return next(r for r in liste if r.get("exercice", r.get("annee")) == a)


SIN, BIL = R["sinistralite"], R["bilan"]
s = {a: an(SIN, a) for a in range(2020, 2027)}
b = {a: an(BIL, a) for a in range(2020, 2027)}
GP, GE, EFF, CG, T936 = R["grenier_par"], R["grenier_encours"], R["effectifs_mensuels"], R["cgrae"], R["taux_936"]
EX = C["extraits"]
TC = {t["indicateur"]: t for t in C["tcam"]}
MORT = {m["table"]: m for m in C["mortalite"]}
IC = C["ic_grenier"]
SC_ = C["sensibilite_cle"]
INF = {int(k): v for k, v in C["inflation"]["valeurs"].items()}
RT = {x["code"]: x for x in Q["ratios"]}
for x in Q["ratios"]:
    for cle in ("valeurs", "valeurs_retraitees", "statuts", "statuts_retraites", "num", "den"):
        if cle in x:
            x[cle] = {(int(k) if str(k).isdigit() else k): v for k, v in x[cle].items()}
AG = {int(k): v for k, v in Q["agregats"].items()}
MS = {m["exercice"]: m for m in Q["marge_solvabilite"]}
AFF = {a["exercice_affectation"]: a for a in Q["affectation"]}
EXR = {e["exercice"]: e for e in Q["excedent_par_risque"]}
MOB = Q["mobilisation"]
cats_ex = [str(a) if a < 2026 else "2026 (8 mois)" for a in range(2020, 2027)]
yrs_bil = [str(a) if a < 2026 else "31/08/2026" for a in range(2020, 2027)]
FMT_M = "# ##0,,"           # données exactes en FCFA, affichage en millions

SRC_PREC = "Source : fichiers de précomptes de la Solde, de la CGRAE et du FPPN reçus le 15/10/2026 ; traitements Adenka"
SRC_COMPTA = "Source : balances des comptes du FPPN (Sage 100), exercices 2020 à 2025 et du 01/01 au 31/08/2026 (provisoire) ; traitements Adenka"
SRC_GRENIER = "Source : exports du logiciel de prêts du Grenier du Policier (balances âgées, production) reçus le 15/10/2026 ; traitements Adenka"
SRC_CIMA = "Source : tables de mortalité CIMA H, CIMA F, TD CIMA et TV CIMA (fichiers fournis le 18/10/2026) ; table ronde CIMA FANAF 2011 ; calculs Adenka"
SRC_BM = f"Source : Banque mondiale, FP.CPI.TOTL.ZG, mise à jour du {C['inflation']['mise_a_jour'][8:10]}/{C['inflation']['mise_a_jour'][5:7]}/{C['inflation']['mise_a_jour'][:4]} ; " + SRC_COMPTA[9:]
DEC = "Décision n° 05/2022/COM/UEMOA du 1er juin 2022, annexe"
REX = "règlement d'exécution n° 004/2022/COM/UEMOA du 1er juin 2022"
R07 = "règlement n° 07/2009/CM/UEMOA du 26 juin 2009"
SRC_REG = "Source : " + DEC + " ; " + REX + " ; " + R07 + " ; balances du FPPN ; calculs Adenka"
BAREME = "barème indiqué par le responsable des sinistres et prestations lors de l'entretien du 16/10/2026"

# ---------------------------------------------------------------------------
# Sigles : définis en bas de chaque diapositive où ils apparaissent
# ---------------------------------------------------------------------------
SIGLES = [
    ("FPPN", r"\bFPPN\b", "Fonds de prévoyance de la Police nationale"),
    ("AIRMS", r"\bAIRMS\b", "Agence ivoirienne de régulation de la mutualité sociale"),
    ("UEMOA", r"\bUEMOA\b", "Union économique et monétaire ouest-africaine"),
    ("PCMS", r"\bPCMS\b", "plan comptable des mutuelles sociales"),
    ("CIMA", r"\bCIMA\b", "Conférence interafricaine des marchés d'assurances"),
    ("CGRAE", r"\bCGRAE\b", "Caisse générale de retraite des agents de l'État"),
    ("BCEAO", r"\bBCEAO\b", "Banque centrale des États de l'Afrique de l'Ouest"),
    ("PEP", r"\bPEP\b", "Plan Épargne Plus"),
    ("EFC", r"\bEFC\b", "entraide fin de carrière"),
    ("CMP", r"\bCMP\b", "sigle des comptes du FPPN, signification à confirmer"),
    ("FCP", r"\bFCP\b", "fonds commun de placement"),
    ("DAT", r"\bDAT\b", "dépôt à terme"),
    ("S/C", r"S/C", "prestations / cotisations"),
    ("TCAM", r"\bTCAM\b", "taux de croissance annuel moyen"),
    ("IC", r"\bIC\b", "intervalle de confiance"),
    ("PAR30", r"\bPAR ?30\b", "part de l'encours en retard de plus de 30 jours"),
    ("PSAP", r"\bPSAP\b", "provision pour sinistres à payer"),
    ("IBNR", r"\bIBNR\b", "sinistres survenus non encore déclarés"),
    ("A/E", r"\bA/E\b", "décès observés / décès attendus"),
    ("PBC", r"\bPBC\b", "liste des documents demandés au client"),
    ("FNG", r"\bFNG\b", "fonds national de garantie"),
    ("FOMUS-CI", r"FOMUS-CI", "fonds de garantie des mutuelles sociales de Côte d'Ivoire"),
    ("MEPS", r"\bMEPS\b", "ministère de l'Emploi et de la Protection sociale"),
    ("GNA", r"\bGNA\b", "assureur de la flotte automobile (compte 47111106)"),
    ("MEDOC", r"\bMEDOC\b", "produit santé complémentaire du poste 8Z7"),
    ("TAFIRE", r"\bTAFIRE\b", "tableau financier des ressources et des emplois"),
    ("CAC", r"\bCAC\b", "commissaire aux comptes"),
    ("HAO", r"\bHAO\b", "hors activités ordinaires"),
    ("CMU", r"\bCMU\b", "couverture maladie universelle"),
    ("SHA-256", r"SHA-256", "algorithme d'empreinte numérique"),
    ("CSV", r"\bCSV\b", "fichier texte de tableau"),
    ("OCR", r"\bOCR\b", "reconnaissance optique de caractères"),
    ("FANAF", r"\bFANAF\b", "Fédération des sociétés d'assurances de droit national africaines"),
    ("CPM", r"\bCPM\b", "Comité de politique monétaire"),
    ("SCI", r"\bSCI\b", "société civile immobilière"),
    ("AG", r"\bAG\b", "assemblée générale"),
    ("R²", r"R²", "coefficient de détermination"),
    ("qx", r"\bqx\b", "probabilité de décès entre x et x + 1 ans"),
    ("Lx", r"\bLx\b", "survivants à l'âge x"),
    ("e_x", r"\be_?x\b|\be5[05]\b|\be60\b|\be0\b", "espérance de vie à l'âge x"),
    ("ä_x", r"ä_?x|ä5[05]|ä6[05]", "valeur d'une rente viagère de 1 par an à l'âge x"),
    ("TD, TV", r"\bT[DV] CIMA\b|\(T[DV]\)", "anciennes tables CIMA de décès et de survie"),
    ("936", r"(?:poste |postes |maladie |Maladie |\(|^)936\b", "poste maladie"),
    ("8Z6", r"\b8Z6\b", "poste décès"),
    ("870", r"(?:poste |EFC |\(|^)870\b", "poste EFC"),
    ("998", r"(?:poste |PEP |\(|^)998\b", "poste PEP"),
    ("9X4", r"\b9X4\b", "poste Grenier"),
    ("9W3", r"\b9W3\b", "poste flotte automobile"),
    ("8Z7", r"\b8Z7\b", "poste MEDOC"),
]


def sigles_du_texte(texte):
    trouves = []
    for nom, motif, defi in SIGLES:
        if re.search(motif, texte, flags=re.M):
            trouves.append(f"{nom} : {defi}")
    return trouves


def contenu(entete, notes, **k):
    SL.append({"type": "contenu", "entete": entete, "notes": notes, **k})


def separateur(num, titre, sous_titre, notes=""):
    SL.append({"type": "separateur", "num": num, "titre": titre, "sous_titre": sous_titre, "notes": notes})


def gab(num, entete, remplissage, notes, source=None, analyse=None):
    SL.append({"type": "gab", "gab": num, "entete": entete, "remplir": remplissage, "notes": notes, "source": source,
               "analyse": analyse})


def retenir(theme, constat, preuves, implication, action, notes):
    """Diapositive « À retenir » sur le gabarit Timesaver n° 13 (4 cartes) : constat, preuves, implication, action."""
    gab(13, {"titre": f"À retenir : {theme}", "sous_titre": constat[0], "synthese": constat[1]},
        {"cartes": [("Constat", constat[2]), ("Preuves", preuves), ("Implication", implication), ("Action", action)]},
        notes, source="Source : analyses Adenka")


def methode_gab(entete, items, notes, source="Source : Adenka", attente=False):
    """Méthode sur le gabarit Timesaver n° 9 (arborescence à 4 blocs) : (titre, texte) x 4."""
    gab(9, entete, {"blocs": items, "attente": attente}, notes, source=source)

# ===========================================================================
# CONTENU
# ===========================================================================
NB_NC = sum(1 for x in Q["ratios"] if x["statut_global"] == "Non conforme")
NB_C = sum(1 for x in Q["ratios"] if x["statut_global"] == "Conforme")
NB_ATT = sum(1 for x in Q["ratios"] if x["statut_global"] in ("En attente", "À apprécier"))
sig = lambda c, a=2025: RT[c]["valeurs"].get(a)  # noqa: E731

SL.append({"type": "couverture", "notes": "Rapport d'étape au 15/10/2026, version 3. Données arrêtées à l'envoi du FPPN reçu le 15/10/2026, complétées le 18/10/2026 (grands livres PDF, inflation Banque mondiale, tables CIMA, documentation publique de l'AIRMS)."})
SL.append({"type": "sommaire", "notes": "Le rapport suit l'ordre : résumé exécutif, compréhension du FPPN, démarche, données, analyses par thème, conformité prudentielle, tables de mortalité, méthodes à venir, outils, analyses en attente, annexes."})

# ---------------------------------------------------------------- RÉSUMÉ EXÉCUTIF
separateur("R", "Résumé exécutif", "Constats, chiffres clés, risques et actions",
           "Cette partie peut être présentée seule : elle reprend les messages essentiels du rapport, avec les chiffres et les renvois aux parties détaillées.")
contenu({"titre": "Résumé exécutif", "sous_titre": "Le FPPN dégage des excédents mais reste fragile, peu documenté et hors normes prudentielles",
         "synthese": f"Fonds propres négatifs, épargne couverte à {pc('Couverture épargne 31/08/2026', b[2026]['couverture_epargne'])} par des actifs dédiés, {NB_NC} ratios prudentiels non conformes sur 36"},
        "Questions probables. Pourquoi « fragile » malgré des excédents ? Parce que les excédents annuels ne suffisent pas encore à effacer les pertes passées (fonds propres négatifs) et que l'épargne due aux mutualistes n'est pas couverte par des actifs disponibles. Pourquoi « peu documenté » ? Aucune base de prestations, pas de date de naissance des mutualistes, pas de barème officiel transmis. Pourquoi « hors normes » ? Les ratios de la Décision n° 05/2022/COM/UEMOA, appliqués aux balances, ne sont pas respectés pour la réserve obligatoire, les fonds propres, la liquidité et les frais. Chaque ligne renvoie à une partie détaillée du rapport.",
        tableau={"titre": tab("Principaux constats au 15/10/2026"), "entetes": ["Domaine", "Constat", "Chiffre clé"],
                 "largeurs": [4.2, 17.2, 9.8], "lignes": [
                     ["Données", "Les précomptes mensuels ne commencent qu'en 2023 ou 2024 selon le poste, et aucune base de prestations n'a été reçue.", N("Lignes de précomptes contrôlées", 4454416) + " lignes de précomptes contrôlées"],
                     ["Cotisations", "Le précompte total vaut 5 % du salaire de base, mais la Solde le répartit en 3,77 % maladie et 1,23 % décès, au lieu du " + BAREME + " (3,5 % et 1,5 %).", pcv("Taux total observé", 5.0) + " = " + pcv("Taux maladie observé", 100 * T936["mediane"]) + " + 1,23 %"],
                     ["Maladie", "Les prestations maladie comptabilisées augmentent bien plus vite que les cotisations maladie.", "S/C " + pc("S/C maladie 2021", s[2021]["sc_maladie"]) + " (2021), " + pc("S/C maladie 2025", s[2025]["sc_maladie"]) + " (2025)"],
                     ["Provisions", "Les provisions pour prestations à payer gardent le même montant de 2020 à 2026, alors que les prestations ont fortement augmenté.", F("Provisions techniques", b[2026]["provisions_techniques"])],
                     ["Épargne", "L'épargne due aux mutualistes (PEP, EFC) dépasse largement les actifs qui lui sont réservés.", F("Épargne totale 31/08/2026", b[2026]["epargne_totale"])],
                     ["Fonds propres", "Les fonds propres restent négatifs à chaque clôture, même si les excédents réduisent l'écart.", F("Fonds propres 2025", b[2025]["fonds_propres"]) + " (2025)"],
                     ["Conformité", "Les ratios de la Décision n° 05/2022 ne sont pas respectés pour la réserve obligatoire, les fonds propres, la liquidité et les frais de fonctionnement.", f"{NB_NC} non conformes, {NB_C} conformes, {NB_ATT} en attente"],
                     ["Grenier", "Le portefeuille de prêts croît très vite, et le FPPN le finance en grande partie.", F("Encours Grenier 31/08/2026", GP["encours"]) + " d'encours"]]},
        disposition="dessous",
        analyse=["**Les fragilités sont surtout financières et documentaires, pas d'exploitation.** Le FPPN encaisse plus qu'il ne dépense chaque année depuis 2020, sauf sur les huit premiers mois de 2026."],
        implication="Les trois chantiers prioritaires sont la transmission des données manquantes, le recalcul des provisions et un plan de mise en conformité prudentielle à présenter à l'AIRMS.",
        source="Source : analyses Adenka sur les documents reçus le 15/10/2026 et la documentation publique de l'AIRMS")

gab(49, {"titre": "Quatre indicateurs clés", "sous_titre": "Quatre ratios résument la situation technique et financière du FPPN",
         "synthese": "Sinistralité maladie, couverture de l'épargne, précompte des prêts et taux de cotisation, mesurés sur les données reçues"},
    {"titres": ["Sinistralité maladie 2025", "Couverture de l'épargne", "Précompte des prêts", "Taux de cotisation"],
     "valeurs": [pc("S/C maladie 2025 (gab)", s[2025]["sc_maladie"]), pc("Couverture épargne (gab)", b[2026]["couverture_epargne"]),
                 pc("Précomptes / échéances août 2026", [e for e in R["grenier_echeances"] if e["mois"] == "août 26"][0]["precompte_9x4"] / [e for e in R["grenier_echeances"] if e["mois"] == "août 26"][0]["echeances_balance"]),
                 pcv("Taux total (gab)", 5.0)],
     "commentaires": ["Les prestations maladie comptabilisées représentent cette part des cotisations maladie de 2025. Ce ratio était de " + pc("S/C 2021 (gab)", s[2021]["sc_maladie"]) + " en 2021.",
                      "Les banques, dépôts à terme et fonds réservés à l'épargne couvrent cette part de l'épargne due au 31/08/2026.",
                      "Les précomptes du Grenier (poste 9X4) couvrent cette part des échéances des prêts en cours en août 2026.",
                      "Le précompte total sur le salaire de base se répartit en 3,77 % maladie et 1,23 % décès, sur " + N("Lignes taux", T936["lignes"]) + " lignes."]},
    "Comment lire : chaque ligne donne un indicateur, sa valeur et sa définition. Les quatre valeurs sont calculées sur les fichiers du FPPN et détaillées dans les parties 5, 6, 7 et 4. Le précompte des prêts mesure la part des échéances effectivement retenue sur salaire ; l'écart peut être payé autrement (caisse, banque), ce que le Grenier doit confirmer.",
    source=SRC_PREC)

contenu({"titre": "Fonctionnement du FPPN", "sous_titre": "Les cotisations et l'épargne financent prestations, prêts et investissements",
         "synthese": f"En 2025 : {F('Cotisations 2025', s[2025]['cot_maladie'] + s[2025]['cot_deces'])} de cotisations, {F('Prestations 2025', s[2025]['prest_maladie'] + s[2025]['prest_deces'])} de prestations"},
        "Ce schéma résume les flux d'argent. Les mutualistes cotisent par précompte sur salaire (Solde) ou sur pension (CGRAE). Le FPPN paie les prestations maladie (cliniques, pharmacies) et décès (familles), gère l'épargne (PEP, EFC), finance les prêts du Grenier et des investissements (CMP, terrains). Les montants sont ceux de la comptabilité 2025 et du bilan au 31/08/2026.",
        schema="fonctionnement",
        analyse=["**L'épargne des mutualistes est la ressource la plus importante du FPPN, devant les cotisations.** Elle atteint " + F("Épargne (schéma texte)", b[2026]["epargne_totale"]) + " au 31/08/2026, soit plus de quatre fois les cotisations de 2025."],
        implication="Le FPPN emploie cette épargne en partie dans des actifs peu liquides (CMP, terrains, prêts) : sa capacité à rembourser les épargnants est le premier enjeu de la mission.",
        source=SRC_COMPTA)

gab(23, {"titre": "Six constats majeurs", "sous_titre": "Les fragilités portent sur les données, les provisions, l'épargne et les normes",
         "synthese": "Chaque constat est démontré dans la partie indiquée, à partir des fichiers du FPPN et des textes de l'UEMOA"},
    {"blocs": [("Données incomplètes", "Le FPPN n'a transmis ni base de prestations, ni dates de naissance. Les analyses actuarielles centrales restent en attente (parties 3 et 14)."),
               ("Sinistralité en hausse", "Le S/C maladie passe de " + pc("S/C 2021 (6c)", s[2021]["sc_maladie"]) + " en 2021 à " + pc("S/C 2025 (6c)", s[2025]["sc_maladie"]) + " en 2025 (partie 5)."),
               ("Provisions figées", "Les provisions gardent le même montant depuis 2020, alors que les prestations ont été multipliées par " + N("Multiplicateur prestations", (s[2025]["prest_maladie"] + s[2025]["prest_deces"]) / (s[2021]["prest_maladie"] + s[2021]["prest_deces"]), 2) + " (partie 9)."),
               ("Épargne non adossée", "Des actifs réservés couvrent seulement " + pc("Couverture (6c)", b[2026]["couverture_epargne"]) + " de l'épargne due (partie 6)."),
               ("Fonds propres négatifs", "Les fonds propres valent " + F("FP 2025 (6c)", b[2025]["fonds_propres"]) + " fin 2025, et l'actif repose sur des emplois peu liquides (partie 9)."),
               ("Normes non respectées", f"{NB_NC} ratios de la Décision n° 05/2022 sur 36 ne sont pas respectés, dont la réserve obligatoire et la liquidité (partie 10).")]},
    "Les six constats sont classés du plus général (données) au plus réglementaire (normes). Les chiffres sont exacts et proviennent des balances comptables et des précomptes. Le constat « actifs peu liquides » repose sur la nature des comptes (investissements, terrains, créances) ; leur valeur réelle n'a pas été expertisée.",
    source="Source : analyses Adenka")

contenu({"titre": "Graphiques clés (1/2)", "sous_titre": "La sinistralité maladie monte pendant que l'épargne reste peu couverte",
         "synthese": f"S/C maladie 2025 : {pc('S/C 2025 (gc)', s[2025]['sc_maladie'])} ; couverture de l'épargne au 31/08/2026 : {pc('Couverture (gc)', b[2026]['couverture_epargne'])}"},
        "Graphique de gauche : rapport entre prestations et cotisations maladie comptabilisées, par exercice. Graphique de droite : épargne due aux mutualistes (passif) et actifs qui lui sont réservés. Les deux graphiques sont détaillés dans les parties 5 et 6.",
        graphiques=[{"titre": fig("Ratio S/C maladie (prestations / cotisations)"), "type": "lignes", "categories": cats_ex,
                     "series": [{"nom": "S/C maladie", "valeurs": [round(100 * s[a]["sc_maladie"], 2) for a in range(2020, 2027)], "couleur": BLEU}],
                     "format": "0.00", "legende": None, "unite": "En %"},
                    {"titre": fig("Épargne due et actifs réservés"), "type": "barres", "categories": yrs_bil,
                     "series": [{"nom": "Épargne PEP et EFC", "valeurs": [round(b[a]["epargne_totale"]) for a in range(2020, 2027)], "couleur": BLEU},
                                {"nom": "Actifs réservés", "valeurs": [round(b[a]["actifs_dedies_epargne"]) for a in range(2020, 2027)], "couleur": ORANGE}],
                     "format": FMT_M, "unite": "En millions de FCFA"}],
        bandeau="**Les deux courbes vont dans le mauvais sens pour la solvabilité.** Les prestations consomment une part croissante des cotisations, et l'épargne grossit plus vite que les actifs qui la garantissent. Implication : le FPPN dispose de moins en moins de marge pour faire face à un choc de prestations ou de retraits.",
        source=SRC_COMPTA)

contenu({"titre": "Graphiques clés (2/2)", "sous_titre": "Les fonds propres restent négatifs et l'encours de prêts explose",
         "synthese": f"Fonds propres 2025 : {F('FP 2025 (gc)', b[2025]['fonds_propres'])} ; encours du Grenier au 31/08/2026 : {F('Encours (gc)', GP['encours'])}"},
        "Gauche : fonds propres (résultat compris) à chaque clôture. Droite : encours des prêts du Grenier en fin d'année et au 31/08/2026. Détails dans les parties 9 et 7.",
        graphiques=[{"titre": fig("Fonds propres"), "type": "barres", "categories": yrs_bil,
                     "series": [{"nom": "Fonds propres", "valeurs": [round(b[a]["fonds_propres"]) for a in range(2020, 2027)], "couleur": BLEU}],
                     "format": FMT_M, "legende": None, "unite": "En millions de FCFA"},
                    {"titre": fig("Encours des prêts du Grenier"), "type": "barres",
                     "categories": [g["arrete"] for g in GE if g["arrete"].startswith("31/12") or g["arrete"] == "31/08/2026"],
                     "series": [{"nom": "Encours", "valeurs": [round(g["encours"]) for g in GE if g["arrete"].startswith("31/12") or g["arrete"] == "31/08/2026"], "couleur": ORANGE}],
                     "format": FMT_M, "legende": None, "unite": "En millions de FCFA"}],
        bandeau="**Le redressement des fonds propres s'interrompt en 2026, au moment où l'encours des prêts atteint son plus haut niveau.** Implication : le FPPN finance une activité de crédit en croissance rapide sans coussin de fonds propres pour absorber d'éventuelles pertes.",
        source=SRC_COMPTA + " ; " + SRC_GRENIER[9:])

cles = ["SIN", "MSE", "ROB", "FEP", "LIQ", "EQS", "CHF", "RL1"]
STATUT_TXT = {"Conforme": "Conforme", "Non conforme": "Non conforme", "En attente": "En attente", "À apprécier": "À apprécier"}


def val_ratio(x, a, retraite=False):
    v = (x.get("valeurs_retraitees") if retraite else x["valeurs"]).get(a)
    if v is None:
        return "n.d."
    lab = f"{x['code']} {a}{' retraité' if retraite else ''}"
    if x["unite"] == "%":
        return pc(lab, v)
    if x["unite"] == "mois":
        return N(lab, v, 2) + " mois"
    if x["unite"] == "jours":
        return N(lab, v, 0) + " jours"
    if x["unite"] == "FCFA":
        return F(lab, v)
    return N(lab, v, 2)


contenu({"titre": "Conformité prudentielle en bref", "sous_titre": "Les ratios de solvabilité et de liquidité de l'UEMOA ne sont pas respectés",
         "synthese": f"Sur 36 indicateurs de la Décision n° 05/2022 : {NB_NC} non conformes, {NB_C} conformes, {NB_ATT} en attente de données ou à apprécier"},
        "Le statut est jugé sur les exercices 2022 à 2025, la Décision étant entrée en vigueur le 1er juin 2022. « Comme le FPPN » : comptes pris tels qu'ils sont tenus. « Retraité » : épargne des adhérents classée en ressource longue et avances au CMP, au Grenier et à l'Hôpital de la Police classées en emplois longs. Détail dans la partie 10.",
        tableau={"titre": tab("Huit ratios prudentiels clés, exercice 2025"), "entetes": ["Ratio", "Norme", "Comme le FPPN", "Retraité", "Statut 2022 à 2025"],
                 "largeurs": [7.4, 3.6, 3.0, 2.8, 3.2], "statut_col": 4, "lignes": [
                     [RT[c]["nom"].replace(" (risques longs)", "").replace("Ratio de solvabilité de la branche épargne", "Solvabilité de l'épargne"),
                      RT[c]["norme"].replace("Au moins 1 (marge disponible au moins égale à la marge exigée)", "Au moins 1").replace(" (tableau 1) ; 9 mois au § 1.2.2.5.4", ""),
                      val_ratio(RT[c], 2025), val_ratio(RT[c], 2025, True) if "valeurs_retraitees" in RT[c] else "",
                      RT[c]["statut_global"]] for c in cles]},
        disposition="droite",
        analyse=["**Les ratios qui mesurent la solidité financière sont non conformes, ceux qui mesurent l'activité courante le sont souvent.** L'autonomie financière et le financement propre sont respectés, car les cotisations couvrent les charges de l'année.",
                 "**Le classement de l'épargne change le verdict sur la liquidité.** En dette à court terme, comme dans le plan comptable, la liquidité 2025 vaut " + val_ratio(RT["LIQ"], 2025) + ". En ressource longue, elle vaut " + val_ratio(RT["LIQ"], 2025, True) + "."],
        implication="Le règlement n° 07/2009 prévoit le retrait d'agrément après deux exercices de non-respect des règles prudentielles (art. 84). Plusieurs ratios sont non conformes plusieurs années de suite : un plan de mise en conformité doit être préparé.",
        source=SRC_REG)

gab(107, {"titre": "Quatre taux à surveiller", "sous_titre": "Une faible part des engagements est couverte par des ressources disponibles",
          "synthese": "Couverture de l'épargne, sinistralité, trésorerie rapportée à l'épargne et prêts en retard de plus de 30 jours"},
    {"anneaux": [(b[2026]["couverture_epargne"], "Épargne couverte par des actifs réservés (31/08/2026)", "Couverture anneau"),
                 (s[2025]["sc_maladie"], "Sinistralité maladie (S/C 2025)", "S/C anneau"),
                 (C["retrait_max_couvert"]["par_tresorerie"], "Trésorerie totale / épargne due (31/08/2026)", "Trésorerie/épargne anneau"),
                 (IC["p30"], "Prêts en retard de plus de 30 jours (en nombre, 31/08/2026)", "Prêts retard anneau")]},
    "Chaque anneau donne un pourcentage. La trésorerie totale (toutes banques et caisses) ne représente qu'une petite part de l'épargne due. Les prêts en retard sont mesurés ici en nombre de prêts ; en montant d'encours, la part est plus faible (voir partie 7).",
    source=SRC_COMPTA + " ; " + SRC_GRENIER[9:],
    analyse="**Trois anneaux sur quatre montrent une couverture faible des engagements.** La trésorerie de toutes les banques ne rembourserait qu'environ un cinquième de l'épargne due. Implication : en cas de retraits importants, le FPPN doit céder ou mobiliser des actifs peu liquides.")

gab(83, {"titre": "Carte des risques", "sous_titre": "Quatre causes identifiées produisent quatre risques pour le FPPN",
         "synthese": "Les causes relèvent de la gestion des données et des choix d'emploi des ressources ; les effets restent à chiffrer"},
    {"causes": ["Pas de base de prestations ni de dates de naissance", "Provisions jamais recalculées", "Épargne utilisée pour d'autres emplois", "Écart entre échéances et précomptes"],
     "effets": ["Coût réel des sinistres inconnu", "Engagements mal mesurés au bilan", "Manque de liquidités en cas de retraits", "Retards et pertes sur les prêts"]},
    "Lecture : chaque colonne relie une cause (en haut) à son effet (en bas). Les causes sont démontrées par les données ; les effets sont des risques, c'est-à-dire des conséquences possibles, à confirmer par les analyses en attente.",
    source="Source : analyses Adenka")

contenu({"titre": "Recommandations prioritaires", "sous_titre": "Sept actions conditionnent la fiabilité des engagements et la conformité",
         "synthese": "Les trois premières actions relèvent de la transmission de données, les quatre suivantes des décisions de gestion"},
        "Ordre de priorité : sans les données (actions 1 à 3), les provisions et les engagements ne peuvent pas être calculés. Les actions 4 à 7 relèvent de la direction du FPPN et du conseil d'administration. Les numéros entre parenthèses renvoient à la liste des documents demandés (PBC).",
        tableau={"titre": tab("Actions recommandées"), "entetes": ["N°", "Action", "Objet", "Priorité"], "largeurs": [1.0, 12.4, 14.4, 3.4],
                 "lignes": [["1", "Transmettre les bases de prestations décès et santé 2020 à 2026 (13, 28, 40).", "Calculer les provisions, la sinistralité par année de survenance et les délais.", "Urgente"],
                            ["2", "Transmettre le fichier des mutualistes avec date de naissance et sexe (9).", "Mesurer la mortalité, projeter l'équilibre technique et calculer les engagements.", "Urgente"],
                            ["3", "Transmettre les comptes individuels PEP et EFC et leurs règlements (16, 17).", "Évaluer la dette envers chaque épargnant et ses échéances.", "Urgente"],
                            ["4", "Recalculer les provisions techniques à chaque clôture et les reclasser en compte 190.", "Respecter l'article 7 du règlement d'exécution et le plan comptable.", "Haute"],
                            ["5", "Présenter à l'AIRMS un plan de mise en conformité des ratios prudentiels.", "Réserve obligatoire, fonds propres, liquidité, frais de fonctionnement.", "Haute"],
                            ["6", "Adosser l'épargne à des placements liquides et séparer ses comptes (cantonnement).", "Couvrir l'épargne exigible, encadrer le financement du Grenier et du CMP.", "Haute"],
                            ["7", "Faire réaliser l'étude actuarielle quinquennale des risques longs.", "Article 14 du règlement d'exécution n° 004/2022.", "Moyenne"]]},
        disposition="dessous",
        analyse=["**Les actions 1 à 3 lèvent les limites de ce rapport ; les actions 4 à 7 corrigent les écarts déjà démontrés.**"],
        implication="Le plan de mise en conformité (action 5) dépend des données des actions 1 à 3 : les deux chantiers doivent être menés ensemble.",
        source="Source : analyses Adenka ; numéros de la liste des documents demandés (PBC)")

gab(39, {"titre": "Feuille de route proposée", "sous_titre": "Quatre étapes pour passer du constat à l'évaluation des engagements",
         "synthese": "Chaque étape dépend de la précédente ; la première est la transmission des données manquantes"},
    {"etapes": [("Collecter", "Le FPPN transmet les bases de prestations, les mutualistes avec date de naissance, les comptes d'épargne et les barèmes."),
                ("Calculer", "Adenka calcule les provisions (PSAP, Chain Ladder), la mortalité observée et les engagements d'épargne."),
                ("Vérifier", "Adenka rapproche les résultats de la comptabilité et des ratios de l'UEMOA, puis les discute avec le FPPN."),
                ("Projeter", "Adenka projette l'équilibre technique sur 5 à 10 ans et chiffre les recommandations.")]},
    "La feuille de route ne fixe pas de dates : elles dépendent du délai de transmission des données par le FPPN. Chaque étape produit une mise à jour de ce rapport.",
    source="Source : Adenka")

# ---------------------------------------------------------------- PARTIE 1 : COMPRENDRE LE FPPN
separateur(1, "Comprendre le FPPN", "Cadre, garanties, circuits et notions de base",
           "Cette partie donne au lecteur les clés pour comprendre la suite, sans connaissance préalable en actuariat.")
contenu({"titre": "Cadre juridique", "sous_titre": "Le FPPN est une mutuelle sociale agréée, régie par le droit de l'UEMOA",
         "synthese": "Le règlement n° 07/2009/CM/UEMOA fixe les règles, l'AIRMS contrôle, le Code CIMA sert de référence technique"},
        "Une mutuelle sociale n'est pas une compagnie d'assurance : elle est régie par le règlement UEMOA de 2009, pas par le Code CIMA. Le Code CIMA sert ici seulement de référence technique (tables de mortalité, taux d'actualisation), faute de règle équivalente dans les textes de la mutualité. Le répertoire de l'AIRMS mis à jour le 15/10/2026 cite le FPPN au n° 05 des mutuelles agréées et immatriculées ; le numéro et la date de son agrément restent à obtenir.",
        tableau={"titre": tab("Textes et organismes de référence"), "entetes": ["Texte ou organisme", "Rôle pour le FPPN", "Usage dans ce rapport"], "largeurs": [8.0, 11.4, 11.8],
                 "lignes": [["Règlement n° 07/2009/CM/UEMOA du 26/06/2009", "Il fixe l'agrément, les obligations, les sanctions et le contrôle des mutuelles sociales.", "Cadre général (partie 10)"],
                            ["Règlement d'exécution n° 004/2022/COM/UEMOA du 01/06/2022", "Il fixe les règles prudentielles : provisions, réserves, risques longs, placements.", "Exigences prudentielles (partie 10)"],
                            ["Décision n° 05/2022/COM/UEMOA du 01/06/2022", "Elle fixe les indicateurs et ratios prudentiels et leurs normes.", "Tableau de bord des ratios (partie 10)"],
                            ["Règlement n° 03/2012/CM/UEMOA et PCMS", "Ils imposent le plan comptable des mutuelles sociales depuis l'exercice 2014.", "Lecture et contrôle des comptes"],
                            ["AIRMS", "L'Agence ivoirienne de régulation de la mutualité sociale agrée et contrôle les mutuelles.", "Destinataire des états et des ratios"],
                            ["Code CIMA, article 338", "Il fixe les tables de mortalité et le taux d'intérêt maximal (3,5 %) des assureurs vie.", "Référence technique (partie 11)"]]},
        disposition="dessous",
        analyse=["**Le FPPN doit respecter des normes chiffrées depuis le 1er juin 2022.** Avant cette date, les textes ne fixaient pas de ratios prudentiels précis pour les mutuelles."],
        implication="Les ratios de la Décision n° 05/2022 s'appliquent pleinement aux exercices 2022 à 2025 de la période d'analyse : la partie 10 les calcule un par un.",
        source="Source : textes cités (documentation publique de l'AIRMS fournie le 18/10/2026) ; répertoire AIRMS des mutuelles agréées, mise à jour du 15/10/2026")

contenu({"titre": "Garanties et produits", "sous_titre": "Huit postes de précompte portent l'ensemble des produits du FPPN",
         "synthese": "Maladie et décès sont les garanties de base ; épargne, prêts, flotte et MEDOC sont des produits complémentaires"},
        "Chaque produit a un code de poste sur la fiche de paie (Solde) : 936 maladie, 8Z6 décès, 870 EFC, 998 PEP, 9X4 Grenier, 9W3 flotte, 8Z7 MEDOC. Les retraités sont précomptés par la CGRAE sur leur pension. Les montants moyens sont observés dans les fichiers de précomptes, pas tirés d'un barème (non reçu).",
        tableau={"titre": tab("Produits du FPPN et montants observés"), "entetes": ["Poste", "Produit", "Nature", "Montant moyen observé"], "largeurs": [2.2, 5.4, 7.8, 5.2],
                 "lignes": [["936", "Maladie", "Garantie santé des actifs, 3,77 % du salaire de base", F("Moyen 936 août 2026", [e for e in EFF if e["annee"] == 2026 and e["mois"] == 8][0]["montant"] / [e for e in EFF if e["annee"] == 2026 and e["mois"] == 8][0]["cotisants"]) + " par mois (août 2026)"],
                            ["8Z6", "Décès", "Capital décès, frais funéraires, allocations, 1,23 % du salaire de base", "voir partie 4"],
                            ["870", "EFC", "Épargne versée en fin de carrière", F("Moyen EFC", an(R["efc_versements"], 2025)["versements"] / an(R["efc_versements"], 2025)["epargnants"] / an(R["efc_versements"], 2025)["mois"]) + " par mois (2025)"],
                            ["998", "PEP", "Épargne volontaire", "voir partie 6"],
                            ["9X4", "Grenier du Policier", "Remboursement de prêts", "voir partie 7"],
                            ["9W3", "Flotte automobile", "Assurance des véhicules, assureur GNA", "environ 10 000 FCFA"],
                            ["8Z7", "MEDOC", "Produit santé complémentaire, nature à préciser", "5 000 FCFA"],
                            ["CGRAE", "Retraités", "Garantie santé, précompte sur pension", F("Moyen CGRAE", sum(c["moyen"] for c in CG) / len(CG))]]},
        disposition="droite",
        analyse=["**Le FPPN couvre à la fois des risques courts et des risques longs.** La maladie se règle dans l'année ; l'EFC et le PEP sont versés des années après les cotisations.",
                 "**Plusieurs produits n'ont pas de document qui fixe leurs règles.** Le barème, le règlement de l'EFC et du PEP, et la nature du produit MEDOC n'ont pas été transmis."],
        implication="Les risques longs (EFC, PEP) relèvent de règles prudentielles plus strictes : étude actuarielle tous les cinq ans, certification des comptes, réserve de trésorerie (partie 10).",
        source=SRC_PREC)

contenu({"titre": "Circuit des précomptes", "sous_titre": "Les cotisations sont retenues à la source, puis rapprochées par le FPPN",
         "synthese": "Chaque mois, le FPPN envoie un fichier « aller » ; la Solde ou la CGRAE renvoie un fichier « retour » des sommes retenues"},
        "Le fichier aller dit ce que le FPPN demande de retenir ; le fichier retour dit ce qui a été réellement retenu. L'écart entre les deux révèle des rejets (agent sorti, montant erroné, plafond atteint). Ce rapprochement est la base du contrôle des cotisations (partie 3).",
        schema="precomptes",
        analyse=["**Le FPPN ne reçoit ses cotisations qu'à travers les fichiers retour de la Solde et de la CGRAE.** Il ne peut donc contrôler ses encaissements qu'en comparant les deux fichiers."],
        implication="Les fichiers aller des postes maladie et décès n'ont pas été transmis : le taux de recouvrement des cotisations exigé par la Décision n° 05/2022 (au moins 90 %) ne peut pas être calculé pour ces postes.",
        source="Source : fichiers aller et retour reçus le 15/10/2026 ; Adenka")

contenu({"titre": "Flux financiers 2025", "sous_titre": "Les ressources de l'année couvrent les prestations, les frais et dégagent un excédent",
         "synthese": f"Produits {F('Produits 2025', s[2025]['produits'])}, charges {F('Charges 2025', s[2025]['charges'])}, résultat {F('Résultat 2025', s[2025]['resultat'])}"},
        "Le résultat de l'exercice est la différence entre produits et charges (y compris les éléments exceptionnels, dits hors activités ordinaires). En 2025, il reste positif, mais il diminue depuis 2021 (partie 5).",
        graphique={"titre": fig("Produits, charges et résultat 2025"), "type": "barres",
                   "categories": ["Cotisations maladie", "Cotisations décès", "Prestations maladie", "Prestations décès", "Frais de gestion", "Résultat"],
                   "series": [{"nom": "2025", "valeurs": [round(s[2025]["cot_maladie"]), round(s[2025]["cot_deces"]), round(s[2025]["prest_maladie"]), round(s[2025]["prest_deces"]), round(s[2025]["frais_gestion"]), round(s[2025]["resultat"])], "couleur": BLEU}],
                   "format": FMT_M, "legende": None, "unite": "En millions de FCFA"},
        analyse=[f"**Les cotisations sont la ressource principale du FPPN.** Elles atteignent {F('Cot mal 2025', s[2025]['cot_maladie'])} en maladie et {F('Cot dec 2025', s[2025]['cot_deces'])} en décès.",
                 f"**Les prestations maladie absorbent la plus grande part des ressources.** Elles valent {F('Prest mal 2025', s[2025]['prest_maladie'])}, contre {F('Prest dec 2025', s[2025]['prest_deces'])} en décès.",
                 f"**Les frais de gestion pèsent {pc('Taux frais 2025', s[2025]['taux_frais'])} des cotisations.** Ils couvrent le personnel, les indemnités des organes, le fonctionnement et les amortissements."],
        implication="La garantie décès dégage un excédent qui compense en partie la garantie maladie : le résultat global masque un déséquilibre de la maladie, analysé par risque dans la partie 10.",
        source=SRC_COMPTA)

contenu({"titre": "Notions de base", "sous_titre": "Huit notions suffisent pour lire l'ensemble du rapport",
         "synthese": "Chaque notion est définie ici une fois, puis utilisée avec le même sens dans tout le rapport"},
        "Ces définitions sont volontairement simples. Les formules exactes figurent dans les diapositives « méthode » de chaque partie et dans l'annexe des méthodes.",
        tableau={"titre": tab("Notions utilisées dans le rapport"), "entetes": ["Notion", "Définition simple", "Exemple dans le rapport"], "largeurs": [4.8, 16.2, 10.2],
                 "lignes": [["Cotisation", "Somme que le mutualiste verse pour être couvert.", "5 % du salaire de base"],
                            ["Prestation", "Somme que le FPPN paie quand le risque survient (soins, décès).", "Factures des cliniques, capital décès"],
                            ["Sinistralité (S/C)", "Part des cotisations consommée par les prestations.", "S/C maladie 2025"],
                            ["Provision technique", "Somme mise de côté au bilan pour payer des prestations déjà nées mais pas encore payées.", "Provision pour prestations à payer"],
                            ["Engagement", "Ce que le FPPN devra payer plus tard à ses mutualistes.", "Épargne PEP et EFC à restituer"],
                            ["Adossement", "Correspondance entre les dettes et les actifs disponibles pour les payer.", "Épargne et placements réservés"],
                            ["Fonds propres", "Ce qui reste quand on retire toutes les dettes de tous les actifs.", "Négatifs de 2020 à 2026"],
                            ["Ratio prudentiel", "Rapport chiffré qu'une norme impose de respecter pour protéger les adhérents.", "Réserve obligatoire de 6 à 9 mois"]]},
        disposition="dessous",
        analyse=["**Ces huit notions portent toute l'analyse.** Les quatre premières décrivent l'activité de l'année ; les quatre dernières décrivent la capacité du FPPN à tenir ses promesses dans la durée."],
        source="Source : Adenka ; " + DEC)

# ---------------------------------------------------------------- PARTIE 2 : DÉMARCHE
separateur(2, "Contexte et démarche", "Mission, périmètre, chaîne de traitement")
contenu({"titre": "Mission et périmètre", "sous_titre": "Le volet actuariel répond à quatre questions sur les engagements et les placements",
         "synthese": "Période d'analyse du 1er janvier 2020 au 31 août 2026 ; données arrêtées à l'envoi du FPPN reçu le 15/10/2026"},
        "Les quatre questions viennent du dossier de préparation de la mission ; les trois axes complémentaires (données, conformité, équilibre) en découlent. La colonne « état » dit honnêtement ce qui est fait et ce qui attend des données.",
        tableau={"titre": tab("Périmètre de la mission actuarielle et statistique"), "entetes": ["Axe", "Question traitée", "État au 15/10/2026"], "largeurs": [6.2, 15.0, 10.0],
                 "lignes": [["1. Engagements", "Combien le FPPN doit-il à ses adhérents et ayants droit ?", "Partiel : épargne seulement"],
                            ["2. Suffisance des provisions", "Les provisions au bilan suffisent-elles à payer les prestations dues ?", "Constat fait, recalcul en attente"],
                            ["3. Adossement actif-passif", "Les actifs arrivent-ils en caisse au bon moment pour payer ?", "Réalisé sur les balances"],
                            ["4. Adéquation des placements", "Les placements sont-ils sûrs, liquides et rentables ?", "Réalisé sur les balances"],
                            ["Contrôle des données", "Les données sont-elles complètes et cohérentes avec la comptabilité ?", "Réalisé sur les données reçues"],
                            ["Conformité prudentielle", "Les ratios de la Décision n° 05/2022 sont-ils respectés ?", "Réalisé : 36 indicateurs examinés"],
                            ["Équilibre technique", "Les cotisations futures couvrent-elles les prestations futures ?", "En attente des âges et des dossiers"]]},
        disposition="dessous",
        analyse=["**Trois des quatre questions de la mission peuvent recevoir une première réponse ; la première attend les données individuelles.**"],
        implication="Le calcul des engagements (question 1) conditionne la réponse définitive aux questions 2 et 3 : il reste la priorité dès réception des fichiers.",
        source="Source : dossier de préparation de la mission (13/10/2026) et liste des documents demandés (PBC)")
gab(41, {"titre": "Chaîne de traitement", "sous_titre": "Chaque chiffre est tracé du fichier du FPPN jusqu'au rapport",
         "synthese": "Six étapes, chacune avec son contrôle automatique ; le détail de chaque étape figure dans la partie 13"},
    {"etapes": [("Réception", "Le FPPN envoie 4 287 fichiers ; Adenka en dresse l'inventaire et repère doublons et fichiers illisibles."),
                ("Extraction", "Les tableurs sont convertis sans perte en CSV, avec une empreinte SHA-256 pour chaque fichier."),
                ("Base", "Les 4,45 millions de lignes sont réunies ; les lignes de total et les montants doublés sont écartés."),
                ("Contrôles", "Les totaux des logiciels sont retrouvés, les fichiers aller et retour sont rapprochés."),
                ("Analyses", "Les indicateurs, tests statistiques et ratios sont calculés et écrits dans le classeur des calculs."),
                ("Rapport", "Les graphiques sont natifs et chaque chiffre affiché est relu et comparé à sa valeur exacte.")]},
    "Ce schéma est modifiable dans PowerPoint et sera complété au fil des entretiens. Chaque étape correspond à un script Python conservé dans le dossier de la mission ; le contrôle indiqué est automatique.",
    source="Source : scripts Python Adenka (dossier scripts de la mission)")

# ---------------------------------------------------------------- PARTIE 3 : DONNÉES
separateur(3, "Données et qualité", "Couverture, contrôles, anomalies détectées")
methode_gab({"titre": "Méthode : contrôle des données", "sous_titre": "Trois contrôles garantissent que la base reproduit fidèlement les fichiers",
             "synthese": "Exhaustivité, cohérence interne et rapprochement entre sources indépendantes, puis une limite assumée"},
            [("Exhaustivité", "Adenka calcule l'empreinte SHA-256 de chaque fichier et compte lignes et cellules avant et après conversion : 218 fichiers copiés et 491 feuilles converties sont identiques."),
             ("Cohérence interne", "Les totaux imprimés par les logiciels sont retrouvés dans les 37 balances âgées du Grenier, et les débits égalent les crédits dans les 7 exercices comptables."),
             ("Rapprochement", "Deux sources qui décrivent le même flux doivent concorder : le taux de précompte rapporte le montant retour au montant aller, poste par poste et mois par mois."),
             ("Limite", "Ces contrôles prouvent que la base est fidèle aux fichiers reçus. Ils ne prouvent pas que les fichiers du FPPN sont eux-mêmes exacts.")],
            "Exhaustivité : rien n'est perdu entre le fichier reçu et la base. Cohérence : les totaux imprimés par les logiciels sont retrouvés. Rapprochement : deux sources qui décrivent le même flux doivent concorder. Limite : la fidélité aux fichiers n'est pas l'exactitude des fichiers.")
cov = {(c["poste"], c["sens"]): c["mois_recus"] for c in R["couverture"]}
lib_flux = [("Comptabilité (balances)", 80), ("PEP 998 retour", cov[("998", "retour")]), ("EFC 870 retour", cov[("870", "retour")]),
            ("Grenier 9X4 retour", cov[("9X4", "retour")]), ("Flotte 9W3 retour", cov[("9W3", "retour")]), ("Maladie 936", cov[("936", "retour")]),
            ("Retraités CGRAE retour", cov[("CGRAE", "retour")]), ("Décès 8Z6", cov[("8Z6", "retour")]), ("MEDOC 8Z7 retour", cov[("8Z7", "retour")]),
            ("Prestations décès et santé", 0)]
contenu({"titre": "Couverture de la période", "sous_titre": "Les données mensuelles ne couvrent pas les années 2020 à 2022",
         "synthese": "Sur 80 mois, seuls la comptabilité et le PEP sont complets ; aucune donnée de prestations n'a été reçue"},
        "Un mois est compté comme reçu quand au moins un fichier lisible le couvre. Les balances comptables couvrent l'année entière en totaux annuels ; le détail des écritures de janvier à août 2020 manque.",
        graphique={"titre": fig("Mois reçus par flux de données, sur 80 mois"), "type": "barres_h", "categories": [x[0] for x in lib_flux][::-1],
                   "series": [{"nom": "Mois reçus", "valeurs": [x[1] for x in lib_flux][::-1], "couleur": BLEU}], "format": "0",
                   "unite": "En nombre de mois (janvier 2020 à août 2026)"},
        analyse=["**Seules la comptabilité et l'épargne PEP couvrent toute la période d'analyse.** La comptabilité donne des totaux annuels ; les fichiers texte de la Solde donnent le PEP par mutualiste et par mois de 2020 à 2022.",
                 "**Les précomptes mensuels commencent en 2023 ou en 2024.** Les fichiers maladie et décès débutent en janvier 2024, ceux des retraités en février 2024 et ceux de l'EFC en janvier 2023.",
                 "**Aucune base de prestations n'a été transmise.** Le coût des soins et des décès n'est connu que par les comptes, sans date de survenance ni bénéficiaire."],
        implication="Les analyses mensuelles portent sur 2024 à 2026, et les années 2020 à 2023 ne peuvent être analysées qu'en totaux comptables. Les conclusions sur l'évolution de la sinistralité restent donc fragiles tant que les bases de prestations ne sont pas reçues.",
        source=SRC_PREC)
rt = R["lignes_total_ecartees"]
lt = EX["ligne_total"]
contenu({"titre": "Calcul pas à pas : lignes de total", "sous_titre": "Une ligne de total cachée aurait doublé les cotisations maladie du mois",
         "synthese": f"Fichier de février 2025 : la dernière ligne porte {F('Total février 2025', float(lt['lignes'][-1][5]))}, soit la somme des {N('Lignes fichier février 2025', lt['nb_lignes_fichier'] - 1)} lignes"},
        "Lecture : (1) la colonne matricule de la dernière ligne contient un nombre de lignes, pas un matricule ; (2) le montant est la somme de toutes les lignes au-dessus ; (3) compter cette ligne aurait doublé le total du mois. Règle appliquée : une ligne dont le montant égale la somme des autres (à 1 % près) est écartée. Au total, " + N("Lignes total écartées (notes)", rt["n"]) + " lignes de ce type ont été écartées. Noms et matricules masqués.",
        pas_a_pas={"entetes": lt["entete"], "lignes": lt["lignes"], "largeurs": [2.0, 2.6, 1.6, 3.4, 3.6, 3.6, 3.4, 2.4],
                   "rouge": [(len(lt["lignes"]) - 1, 3), (len(lt["lignes"]) - 1, 5)],
                   "explications": ["La colonne « matricule » de la dernière ligne contient le nombre de lignes du fichier.",
                                    f"Le montant de {F('Total février (pas)', float(lt['lignes'][-1][5]))} est la somme de toutes les lignes au-dessus.",
                                    f"Adenka écarte cette ligne et les {N('Lignes total écartées', rt['n'])} lignes du même type, ce qui évite {F('Montant total écarté', rt['montant'])} de double compte."],
                   "fichier": lt["fichier"]},
        analyse=["**Une simple somme des fichiers aurait gonflé les cotisations.** Les lignes de total ont la même forme que les lignes des mutualistes et ne se repèrent que par le calcul."],
        implication="Le FPPN doit retirer ces lignes de total de ses fichiers, ou les isoler, avant tout traitement automatique.",
        source="Source : " + lt["fichier"] + " (identifiant d'extraction " + lt["id"] + ") ; noms et matricules masqués")
md = EX["montants_doubles"]
contenu({"titre": "Calcul pas à pas : montants doublés", "sous_titre": "Des montants saisis deux fois à la suite gonflaient les précomptes décès",
         "synthese": f"Avril 2024 : {N('Montants corrigés', md['total_corriges'])} montants de plus d'un million de FCFA sont des montants écrits deux fois"},
        "Lecture : (1) le montant source répète deux fois la même suite de chiffres ; (2) un précompte décès de plusieurs milliards de FCFA par agent est impossible (médiane du poste : environ 2 300 FCFA) ; (3) le montant retenu est la première moitié. Correction appliquée seulement au poste décès, au-delà d'un million de FCFA, et signalée au FPPN.",
        pas_a_pas={"entetes": ["Matricule", "Nom", "Montant dans le fichier", "Montant corrigé"], "lignes": md["lignes"], "largeurs": [4.0, 4.0, 8.0, 6.0],
                   "rouge": [(0, 2), (0, 3)],
                   "explications": ["Le montant du fichier répète deux fois la suite « 47924 ».", "Un précompte décès de plus d'un million de FCFA par agent est impossible.",
                                    "Adenka retient la première moitié du montant et signale la ligne au FPPN."], "fichier": md["fichier"]},
        analyse=["**L'erreur vient de la saisie du fichier source, pas d'un calcul.** Les deux moitiés du montant sont identiques, ce qui signe une double frappe."],
        implication="La correction est limitée au poste décès et à un seuil d'un million de FCFA pour ne pas modifier de montants réels. Le FPPN doit confirmer chaque ligne signalée.",
        source="Source : " + md["fichier"] + " ; noms et matricules masqués")
ar = R["aller_retour"]
nomp = {"870": "EFC 870", "998": "PEP 998", "9X4": "Grenier 9X4", "9W3": "Flotte 9W3", "CGRAE": "Retraités CGRAE", "8Z7": "MEDOC 8Z7"}
contenu({"titre": "Rapprochement aller et retour", "sous_titre": "La Solde précompte en moyenne 97 à 100 % des montants demandés",
         "synthese": "Sur les mois communs, l'écart moyen est faible, mais certains mois s'écartent fortement"},
        "Taux de précompte = montant retour / montant aller sur les mois où les deux fichiers existent. Un taux supérieur à 100 % signifie que le fichier aller était incomplet (cas du PEP en juillet 2024).",
        graphique={"titre": fig("Montant précompté rapporté au montant demandé, mois communs"), "type": "barres", "categories": [nomp[a["poste"]] for a in ar],
                   "series": [{"nom": "Taux moyen", "valeurs": [round(100 * a["taux_recouvrement"], 2) for a in ar], "couleur": BLEU},
                              {"nom": "Mois le plus bas", "valeurs": [round(100 * a["taux_min_mensuel"], 2) for a in ar], "couleur": ORANGE}],
                   "format": "0.00", "unite": "En % du montant demandé"},
        analyse=[f"**La Solde et la CGRAE retiennent presque tout ce que le FPPN demande.** Le taux moyen va de {pc('Taux min moyen', min(a['taux_recouvrement'] for a in ar))} à {pc('Taux max moyen', max(a['taux_recouvrement'] for a in ar))} selon le poste.",
                 f"**Certains mois s'écartent nettement de la moyenne.** Le mois le plus bas atteint {pc('Min Grenier', [a for a in ar if a['poste'] == '9X4'][0]['taux_min_mensuel'])} pour le Grenier et {pc('Min PEP', [a for a in ar if a['poste'] == '998'][0]['taux_min_mensuel'])} pour le PEP.",
                 "**Ces écarts traduisent des rejets ou des fichiers incomplets.** Le FPPN doit pouvoir les expliquer mois par mois."],
        implication="Pour les retraités, le taux mesuré sert d'indicateur de recouvrement des cotisations : il respecte la norme d'au moins 90 % de la Décision n° 05/2022. Pour les cotisations maladie et décès des actifs, le calcul attend les fichiers aller.",
        source=SRC_PREC)
apm, apx = EX["aller_retour_pep"]["min"], EX["aller_retour_pep"]["max"]
contenu({"titre": "Calcul pas à pas : écarts du PEP", "sous_titre": "Les deux mois extrêmes du PEP s'expliquent par le contenu des fichiers",
         "synthese": f"Janvier 2025 : {pc('Taux PEP janv 2025', apm['taux'])} ; juillet 2024 : {pc('Taux PEP juil 2024', apx['taux'])}"},
        "Janvier 2025 : le fichier retour compte moins de lignes que le fichier aller, des précomptes n'ont pas été faits. Juillet 2024 : le fichier retour compte plus de lignes que le fichier aller, le fichier aller n'incluait pas les nouveaux épargnants (le fichier aller d'août 2024 en compte nettement plus).",
        pas_a_pas={"entetes": ["Mois", "Fichier aller", "Lignes aller", "Montant aller", "Lignes retour", "Montant retour", "Taux"],
                   "lignes": [["janv. 2025", apm["fichier_a"], N("Lignes aller janv25", apm["lignes_a"]), F("Montant aller janv25", apm["montant_a"]), N("Lignes retour janv25", apm["lignes_r"]), F("Montant retour janv25", apm["montant_r"]), pc("Taux janv25 (t)", apm["taux"])],
                              ["juil. 2024", apx["fichier_a"], N("Lignes aller juil24", apx["lignes_a"]), F("Montant aller juil24", apx["montant_a"]), N("Lignes retour juil24", apx["lignes_r"]), F("Montant retour juil24", apx["montant_r"]), pc("Taux juil24 (t)", apx["taux"])]],
                   "largeurs": [2.4, 6.6, 2.6, 5.4, 2.6, 5.4, 3.0], "rouge": [(0, 4), (1, 4)],
                   "explications": ["En janvier 2025, " + N("Écart lignes janv25", apm["lignes_a"] - apm["lignes_r"]) + " épargnants demandés n'ont pas été précomptés.",
                                    "En juillet 2024, le fichier retour compte " + N("Écart lignes juil24", apx["lignes_r"] - apx["lignes_a"]) + " lignes de plus que le fichier aller.",
                                    "Le taux mensuel dépend donc de la qualité des deux fichiers, et le FPPN doit tenir un état des rejets."],
                   "fichier": "FPPN-SOLDE\\PEP\\FICHIER ALLER et FICHIER RETOUR"},
        analyse=["**Un taux supérieur à 100 % ne signifie pas un trop-perçu.** Il révèle un fichier aller incomplet, qui n'inclut pas les nouveaux épargnants."],
        implication="Sans état des rejets, le FPPN ne peut pas prouver que chaque épargnant a été précompté du bon montant : c'est une condition pour reconstituer les comptes individuels.",
        source=SRC_PREC)
contenu({"titre": "Fichiers vides ou endommagés", "sous_titre": "Onze fichiers sont vides, illisibles ou mal classés et doivent être renvoyés",
         "synthese": "Les fichiers comptables endommagés ont tous un double lisible, sauf le grand livre des comptes 2022"},
        "Un fichier vide (0 octet) ne contient aucune donnée ; un fichier endommagé a une partie de ses données détruite (souvent lors d'une copie ou d'un transfert). Aucun de ces défauts ne vient du traitement d'Adenka : les empreintes prouvent que les fichiers sont identiques à ceux reçus.",
        tableau={"titre": tab("Fichiers à renvoyer ou à corriger"), "entetes": ["Fichier", "Défaut", "Conséquence"], "largeurs": [11.4, 6.2, 13.6],
                 "lignes": [["Décès 8Z6 et maladie 936, décembre 2024", "Vides (0 octet)", "Le mois de décembre 2024 n'a aucune donnée."],
                            ["MEDOC 8Z7, retours de février et avril 2026", "Aucune ligne", "Les précomptes MEDOC de ces deux mois sont inconnus."],
                            ["Balances des tiers 2021 et 2025 (Excel)", "Endommagées", "Adenka a utilisé les PDF lisibles."],
                            ["Grands livres des comptes 2021 et 2024 (Excel)", "Endommagés", "Adenka a utilisé les PDF lisibles."],
                            ["Grand livre des comptes 2022 (Excel et PDF)", "Endommagés", "Adenka a récupéré 417 pages sur 657 ; la balance 2022 est intacte."],
                            ["Grands livres des tiers 2023 et 2025 (Excel)", "Endommagés", "Adenka a utilisé les PDF lisibles."],
                            ["Fichiers aller rangés dans les dossiers retour", "Mal classés", "Les retours de CGRAE juin 2024 et GNA janvier 2026 manquent."]]},
        disposition="dessous",
        analyse=["**Les défauts portent sur les fichiers reçus et non sur leur traitement.** Les empreintes SHA-256 prouvent que les fichiers analysés sont identiques à ceux transmis."],
        implication="Le renvoi de ces fichiers complètera décembre 2024 et le détail comptable 2022, sans changer les totaux annuels déjà utilisés.",
        source="Source : journal de contrôle de l'extraction ; note interne des anomalies")
retenir("données", ("Les données reçues sont fidèlement reprises, mais incomplètes", "Base contrôlée à 100 %, sans prestations, sans âges et sans précomptes mensuels de 2020 à 2022",
                    "La base reproduit les fichiers du FPPN sans perte, mais il manque les données nécessaires aux calculs actuariels."),
        f"Les empreintes, les lignes et les totaux des logiciels sont identiques. Adenka a écarté {N('Total écartés (ret)', rt['n'])} lignes de total et corrigé {N('Corrigés (ret)', md['total_corriges'])} montants doublés.",
        "La qualité des analyses dépend désormais de la transmission des données manquantes, et non plus du traitement des fichiers.",
        "Le FPPN doit transmettre les bases de prestations, les dates de naissance et les fichiers vides de décembre 2024.",
        "Message pour le client : la qualité de l'analyse dépend désormais de la transmission des données manquantes, pas du traitement.")

# ---------------------------------------------------------------- PARTIE 4 : EFFECTIFS ET COTISATIONS
separateur(4, "Effectifs et cotisations", "Cotisants, retraités, taux, cotisations en termes réels")
methode_gab({"titre": "Méthode : effectifs et taux", "sous_titre": "Effectifs, flux et taux sont mesurés agent par agent, chaque mois",
             "synthese": "Un cotisant est un matricule précompté au poste maladie 936 dans le mois"},
            [("Effectif", "L'effectif du mois est le nombre de matricules distincts du fichier maladie 936 de ce mois. Trente et un mois sont disponibles."),
             ("Flux", "Une entrée est un matricule présent un mois et absent le mois précédent ; une sortie est l'inverse. Seuls les mois consécutifs sont comparés."),
             ("Taux de cotisation", f"Le taux rapporte le précompte du mois au salaire de base du mois, pour chaque agent : {N('Lignes taux (méthode)', T936['lignes'])} lignes sont calculées."),
             ("Limite", "Les fichiers ne donnent pas le motif de sortie (retraite, décès, radiation), ni l'âge, ni le sexe des agents.")],
            "Effectif du mois m = nombre de matricules distincts du fichier maladie 936 du mois m. Entrée = matricule présent en m et absent en m-1 ; sortie = l'inverse. Taux de cotisation = précompte du mois / salaire de base du mois.")
e0 = [e for e in EFF if e["annee"] == 2024 and e["mois"] == 1][0]
e1 = [e for e in EFF if e["annee"] == 2026 and e["mois"] == 8][0]
tdc = C["tendances"][0]
contenu({"titre": "Effectifs cotisants", "sous_titre": "Les cotisants actifs augmentent d'un tiers en trente et un mois",
         "synthese": f"{N('Cotisants janv 2024', e0['cotisants'])} cotisants en janvier 2024, {N('Cotisants août 2026', e1['cotisants'])} en août 2026"},
        f"La droite de tendance estimée par moindres carrés donne une hausse de {N('Pente cotisants', tdc['pente'], 2)} cotisants par mois (intervalle de confiance à 95 % : {N('IC bas pente', tdc['ic95_bas'], 2)} à {N('IC haut pente', tdc['ic95_haut'], 2)}). La probabilité d'observer une telle pente si la vraie pente était nulle est inférieure à 0,01 % (p-valeur {tdc['p_valeur']:.2e}). La hausse se fait par paliers en juillet, ce qui rend la droite seulement approximative (R² = {nb(tdc['r2'], 2)}).",
        graphique={"titre": fig("Cotisants actifs précomptés (poste 936), par mois"), "type": "lignes", "categories": [e["libelle"] for e in EFF],
                   "series": [{"nom": "Cotisants", "valeurs": [e["cotisants"] for e in EFF], "couleur": BLEU}], "format": "# ##0", "etiquettes": False,
                   "axe_min": 20000, "legende": None, "unite": "En nombre de mutualistes"},
        analyse=[f"**Les cotisants actifs augmentent de {pc('Hausse cotisants', e1['cotisants'] / e0['cotisants'] - 1)} en 31 mois.** Le taux de croissance annuel moyen vaut {pc('TCAM cotisants', TC['Cotisants actifs (maladie 936)']['tcam'])}.",
                 f"**La hausse se fait par paliers en juillet, mois d'intégration des promotions de recrues.** Les fichiers comptent {N('Entrées juil 2024', EX['entrees_juillet']['2024']['entrees'])} entrées en juillet 2024 et {N('Entrées juil 2025', EX['entrees_juillet']['2025']['entrees'])} en juillet 2025.",
                 f"**La tendance est statistiquement significative.** Elle correspond à {N('Pente (texte)', tdc['pente'], 2)} cotisants de plus par mois en moyenne."],
        implication="Les recrues cotisent dès leur entrée, alors que leurs besoins de soins et d'épargne viendront plus tard : la croissance actuelle améliore les comptes de l'année, mais elle prépare des engagements plus importants à long terme, que seule une projection par âge permettra de chiffrer.",
        source=SRC_PREC + " (décembre 2024 : fichier vide)")
fl = R["flux_mensuels"]
contenu({"titre": "Entrées et sorties", "sous_titre": "Les entrées massives de juillet s'opposent à des sorties faibles et régulières",
         "synthese": f"Entrées : {N('Entrées 2024', R['flux_annuels'][0]['entrees'])} en 2024, {N('Entrées 2025', R['flux_annuels'][1]['entrees'])} en 2025 ; sorties : {N('Sorties 2024', R['flux_annuels'][0]['sorties'])} et {N('Sorties 2025', R['flux_annuels'][1]['sorties'])}"},
        "Les mois sans mois précédent disponible (janvier 2024, janvier 2025 après le fichier vide de décembre 2024) n'ont pas de flux calculé. Les sorties de novembre 2024 et janvier 2026 sont plus nombreuses : leur motif doit être demandé au FPPN (départs à la retraite de fin d'année probables, non vérifiés).",
        graphique={"titre": fig("Entrées et sorties de cotisants, par mois"), "type": "barres", "categories": [f["mois"] for f in fl],
                   "series": [{"nom": "Entrées", "valeurs": [f["entrees"] for f in fl], "couleur": BLEU}, {"nom": "Sorties", "valeurs": [f["sorties"] for f in fl], "couleur": ORANGE}],
                   "format": "# ##0", "etiquettes": False, "unite": "En nombre de mutualistes"},
        analyse=["**Les entrées sont concentrées en juillet, mois d'intégration des recrues.** Les autres mois, les entrées et les sorties restent faibles.",
                 "**Les sorties ne dépassent la moyenne qu'en novembre 2024 et en janvier 2026.** Leur motif n'est pas indiqué dans les fichiers.",
                 f"**La fidélisation des adhérents est forte.** Parmi les matricules précomptés en 2024, {pc('Fidélisation 2025 (flux)', RT['FID']['valeurs'][2025])} le sont encore en 2025, au-dessus de la norme d'au moins 80 % de la Décision n° 05/2022."],
        implication="Le motif de chaque sortie est indispensable pour séparer retraites, décès et radiations, et donc pour mesurer la mortalité des policiers et projeter les départs.",
        source=SRC_PREC)
anc = R["anciennete"]
contenu({"titre": "Structure des effectifs", "sous_titre": "Une population jeune en service, recrutée en majorité depuis moins de dix ans",
         "synthese": f"En {anc['mois']}, ancienneté médiane de {N('Ancienneté médiane', anc['mediane'], 2)} ans, salaire de base médian de {F('Salaire base médian', anc['salaire_base_median'])}"},
        "L'ancienneté est calculée à partir de la date de recrutement du fichier détaillé de janvier 2025. Elle ne donne pas l'âge : deux agents de même ancienneté peuvent avoir des âges différents. Pour cette raison, aucune hypothèse d'âge n'a été faite.",
        graphique={"titre": fig(f"Agents par ancienneté de service, {anc['mois']}"), "type": "barres", "categories": anc["categories"],
                   "series": [{"nom": "Agents", "valeurs": anc["valeurs"], "couleur": BLEU}], "format": "# ##0", "unite": "En nombre d'agents"},
        analyse=[f"**{pc('Part moins de 10 ans', sum(anc['valeurs'][:2]) / sum(anc['valeurs']))} des agents ont moins de dix ans de service.** Les recrutements récents dominent la population cotisante.",
                 f"**Une seconde vague a 15 à 25 ans de service.** Elle regroupe {pc('Part 15-25 ans', (anc['valeurs'][3] + anc['valeurs'][4]) / sum(anc['valeurs']))} des agents.",
                 "**Cette seconde vague partira à la retraite dans les années à venir.** Elle deviendra une population de retraités précomptés par la CGRAE, qui consomme davantage de soins."],
        implication="Les départs de cette vague déclencheront le versement de l'EFC et une hausse de la sinistralité des retraités : c'est le premier effet que la projection de l'équilibre devra chiffrer dès réception des âges.",
        source=SRC_PREC + " (fichiers maladie 936 détaillés)")
tx = EX["taux_cotisation"]
contenu({"titre": "Calcul pas à pas : taux de cotisation", "sous_titre": "Pour chaque agent, maladie et décès font 5,00 % du salaire de base",
         "synthese": f"La même répartition, 3,77 % et 1,23 %, se retrouve sur {N('Lignes taux (pas)', T936['lignes'])} lignes des mois détaillés"},
        "Calcul pour le premier agent : 11 354 / 301 159 = 3,77 % (maladie) ; 3 704 / 301 159 = 1,23 % (décès) ; total 5,00 %. Le " + BAREME + " est de 3,5 % et 1,5 %. Le total est donc bon mais la répartition diffère. Matricules masqués.",
        pas_a_pas={"entetes": ["Matricule", "Grade", "Salaire de base", "Maladie 936", "Décès 8Z6", "Maladie %", "Décès %", "Total %"],
                   "lignes": [r[:5] + [r[5].replace(".", ",") + " %", r[6].replace(".", ",") + " %", r[7].replace(".", ",") + " %"] for r in tx["lignes"]],
                   "largeurs": [3.2, 2.2, 4.2, 3.8, 3.6, 3.4, 3.2, 3.2], "rouge": [(0, 5), (0, 6), (0, 7)],
                   "explications": ["Le précompte maladie vaut 11 354 / 301 159 = 3,77 % du salaire de base, au lieu de 3,5 %.", "Le précompte décès vaut 3 704 / 301 159 = 1,23 % du salaire de base, au lieu de 1,5 %.", "Le total de 5,00 % est respecté, mais la répartition entre maladie et décès diffère du barème indiqué."],
                   "fichier": "FPPN-SOLDE\\EFFECTIF POLICE (MALADIE)\\2024\\042024_effectif_police_936_AVRIL 2024.xlsx et FPPN_ALLOC_DECES_8Z6_AVRIL_2024.xlsx"},
        analyse=["**La Solde applique la même clé à tous les agents.** L'écart avec le barème n'est donc pas une erreur ponctuelle, mais une règle de calcul."],
        implication="Le barème officiel (décision, statuts ou règlement intérieur) doit être transmis : sans lui, le résultat de chaque garantie dépend de la source retenue.",
        source=SRC_PREC + " ; matricules masqués")
tot25 = s[2025]["cot_maladie"] + s[2025]["cot_deces"]
contenu({"titre": "Taux de cotisation", "sous_titre": "Trois clés de répartition coexistent pour le même taux de 5 %",
         "synthese": f"Barème indiqué en entretien 3,5 / 1,5 ; Solde 3,77 / 1,23 ; comptabilité 2025 {nb(5 * s[2025]['cot_maladie'] / tot25, 2)} / {nb(5 * s[2025]['cot_deces'] / tot25, 2)}"},
        "La clé comptable est la part des cotisations maladie (compte 70100000) dans le total des cotisations maladie et décès (comptes 70100000 et 70210000) de 2025, multipliée par 5 %. Elle est proche du barème indiqué en entretien, pas de la clé de la Solde.",
        graphique={"titre": fig("Répartition du taux de 5 % selon la source"), "type": "barres_empilees",
                   "categories": ["Barème indiqué en entretien", "Précompte Solde", "Comptabilité 2025"],
                   "series": [{"nom": "Maladie", "valeurs": [3.5, round(100 * T936["mediane"], 2), round(5 * s[2025]["cot_maladie"] / tot25, 2)], "couleur": BLEU},
                              {"nom": "Décès", "valeurs": [1.5, round(5 - 100 * T936["mediane"], 2), round(5 * s[2025]["cot_deces"] / tot25, 2)], "couleur": ORANGE}],
                   "format": "0.00", "unite": "En % du salaire de base"},
        analyse=["**Le total de 5 % est respecté dans les trois sources.** Seule la répartition entre maladie et décès change.",
                 "**La Solde retient 3,77 % au poste maladie et 1,23 % au poste décès.** Le responsable des sinistres et prestations a indiqué 3,5 % et 1,5 % lors de l'entretien du 16/10/2026.",
                 f"**La comptabilité affecte {pc('Clé comptable maladie', SC_['cle_comptable_maladie'])} des cotisations à la maladie**, proche des 70 % du barème indiqué en entretien. Elle ne suit donc pas la Solde."],
        implication="La Décision n° 05/2022 impose de suivre l'excédent par risque : sans clé unique, le FPPN ne peut pas calculer la réserve obligatoire exigée pour chaque risque (au moins 30 % de l'excédent maladie, 50 % des autres).",
        source=SRC_PREC + " ; " + SRC_COMPTA[9:])
contenu({"titre": "Sensibilité : clé de répartition", "sous_titre": "La clé retenue change la sinistralité de chaque garantie de plusieurs points",
         "synthese": f"S/C maladie 2025 : {pc('S/C clé compta', SC_['sc_maladie_compta'])} avec la clé comptable, {pc('S/C clé Solde', SC_['sc_maladie_cle_solde'])} avec la clé de la Solde"},
        "Calcul : S/C maladie = prestations maladie / (cotisations totales × part maladie). Avec la clé de la Solde (75,40 %), la part des cotisations affectée à la maladie augmente, donc le ratio maladie baisse et le ratio décès monte. Les prestations ne changent pas : seule l'affectation des cotisations change.",
        graphique={"titre": fig("Ratio S/C 2025 selon la clé de répartition"), "type": "barres", "categories": ["Clé comptable", "Clé de la Solde (3,77 / 1,23)", "Barème indiqué (3,5 / 1,5)"],
                   "series": [{"nom": "S/C maladie", "valeurs": [round(100 * SC_[k], 2) for k in ("sc_maladie_compta", "sc_maladie_cle_solde", "sc_maladie_cle_bareme")], "couleur": BLEU},
                              {"nom": "S/C décès", "valeurs": [round(100 * SC_[k], 2) for k in ("sc_deces_compta", "sc_deces_cle_solde", "sc_deces_cle_bareme")], "couleur": ORANGE}],
                   "format": "0.00", "unite": "En %"},
        analyse=[f"**Avec la clé de la Solde, le S/C maladie 2025 passe à {pc('S/C solde (t)', SC_['sc_maladie_cle_solde'])}** et le S/C décès monte à {pc('S/C décès solde', SC_['sc_deces_cle_solde'])}.",
                 f"**Avec le barème indiqué en entretien, le S/C maladie vaut {pc('S/C barème', SC_['sc_maladie_cle_bareme'])}.** Les prestations restent les mêmes : seule l'affectation des cotisations change.",
                 "**Le résultat par garantie dépend donc d'une décision de gestion.** Cette décision doit être écrite dans le barème et appliquée de la même façon par la Solde et par la comptabilité."],
        implication="Selon la clé, la maladie se situe dans la fourchette réglementaire de 70 % à 80 % ou en dessous : la conclusion sur l'équilibre de la garantie maladie ne sera ferme qu'avec le barème officiel.",
        source=SRC_COMPTA)
contenu({"titre": "Cotisations en termes réels", "sous_titre": "Hors inflation, les cotisations progressent de 23 % entre 2020 et 2025",
         "synthese": f"Hausse nominale {pc('Croissance nominale cot', C['cot_croissance_nominale_2020_2025'])}, hausse des prix {pc('Croissance prix', C['prix_croissance_2020_2025'])}, hausse réelle {pc('Croissance réelle cot', C['cot_croissance_reelle_2020_2025'])}"},
        "Méthode : indice des prix base 100 en 2020, construit en cumulant l'inflation annuelle de la Banque mondiale (2021 : 4,16 % ; 2022 : 5,23 % ; 2023 : 4,37 % ; 2024 : 3,45 % ; 2025 : 0,13 %). Cotisations réelles = cotisations nominales × 100 / indice. Elles mesurent le pouvoir d'achat des cotisations en francs de 2020.",
        graphique={"titre": fig("Cotisations nominales et réelles (francs de 2020)"), "type": "lignes", "categories": [str(a) for a in range(2020, 2026)],
                   "series": [{"nom": "Nominales", "valeurs": [round(r["cotisations_nominales"]) for r in C["reel"]], "couleur": BLEU},
                              {"nom": "Réelles (francs de 2020)", "valeurs": [round(r["cotisations_reelles_2020"]) for r in C["reel"]], "couleur": ORANGE}],
                   "format": FMT_M, "unite": "En millions de FCFA"},
        analyse=[f"**L'inflation explique une partie de la hausse des cotisations.** Les prix ont augmenté de {pc('Prix (t)', C['prix_croissance_2020_2025'])} entre 2020 et 2025.",
                 f"**La hausse réelle reste nette : {pc('Réelle (t)', C['cot_croissance_reelle_2020_2025'])}.** Elle vient surtout de la hausse des effectifs, le taux de 5 % étant resté le même.",
                 f"**L'inflation de 2025 est très faible ({pcv('Inflation 2025', INF[2025])})**, après {pcv('Inflation 2022', INF[2022])} en 2022."],
        implication="Le coût des soins suit les prix de la santé : une cotisation fixée en pourcentage du salaire ne protège le FPPN que si les salaires suivent ces prix. La projection devra tester cette hypothèse.",
        source=SRC_BM)
cot25 = an(R["cotisations"], 2025)
contenu({"titre": "Cotisations comptabilisées", "sous_titre": "Les cotisations comptabilisées progressent de 45,76 % entre 2020 et 2025",
         "synthese": f"{F('Cot 2020', s[2020]['cot_maladie'] + s[2020]['cot_deces'])} en 2020, {F('Cot 2025 (cc)', tot25)} en 2025 ; taux de croissance annuel moyen {pc('TCAM cot', TC['Cotisations maladie et décès']['tcam'])}"},
        "Rapprochement 2025 : cotisations maladie comptabilisées comparées aux précomptes maladie (12 mois) et retraités CGRAE (12 mois). Les cotisations à recevoir du Trésor (compte 41111000) mesurent ce que l'État doit encore au FPPN à la clôture.",
        graphique={"titre": fig("Cotisations comptabilisées par garantie"), "type": "barres_empilees", "categories": cats_ex,
                   "series": [{"nom": "Maladie", "valeurs": [round(s[a]["cot_maladie"]) for a in range(2020, 2027)], "couleur": BLEU},
                              {"nom": "Décès", "valeurs": [round(s[a]["cot_deces"]) for a in range(2020, 2027)], "couleur": ORANGE}],
                   "format": FMT_M, "unite": "En millions de FCFA"},
        analyse=[f"**En 2025, comptabilité et précomptes maladie diffèrent de {pc('Écart maladie 2025', 1 - s[2025]['cot_maladie'] / (cot25['prec_936'] + cot25['prec_cgrae']))}.** La comptabilité enregistre {F('Compta mal 2025', s[2025]['cot_maladie'])}, contre {F('Précomptes mal 2025', cot25['prec_936'] + cot25['prec_cgrae'])} de précomptes.",
                 "**L'écart vient de la clé de répartition.** La comptabilité affecte au décès une part des cotisations plus grande que la Solde.",
                 f"**Fin 2025, l'État devait {F('Créance Trésor 2025', b[2025]['cotisations_a_recevoir_etat'])} de cotisations.** Le compte est soldé au 31/08/2026."],
        implication="Les cotisations totales concordent, mais leur partage entre maladie et décès diffère selon la source. Tant que cet écart n'est pas réglé, la sinistralité et l'excédent de chaque risque restent incertains.",
        source=SRC_COMPTA)
retenir("effectifs et cotisations", ("Les cotisations augmentent avec les effectifs, mais leur répartition doit être clarifiée", "Effectifs en hausse d'un tiers, taux total de 5 % respecté, trois clés de répartition différentes",
                                     "Le FPPN encaisse davantage chaque année grâce aux recrues, mais il applique trois clés différentes pour partager les cotisations entre maladie et décès."),
        f"Le FPPN compte {N('Cotisants (ret)', e1['cotisants'])} cotisants en août 2026. La Solde retient 3,77 % et 1,23 %, alors que le barème indiqué en entretien prévoit 3,5 % et 1,5 %.",
        "Le résultat de chaque garantie et la réserve obligatoire par risque dépendent d'une clé qui n'est fixée par aucun document transmis.",
        "Le FPPN doit transmettre le barème officiel, les dates de naissance et les motifs de sortie des agents.",
        "La question de la clé de répartition est la plus importante de cette partie : elle change le résultat de chaque garantie.")

# ---------------------------------------------------------------- PARTIE 5 : SINISTRALITÉ
separateur(5, "Sinistralité et résultat", "Prestations, frais de gestion, résultat")
methode_gab({"titre": "Méthode : sinistralité comptable", "sous_titre": "Le ratio S/C rapporte les prestations de l'année aux cotisations de l'année",
             "synthese": "S/C = prestations comptabilisées / cotisations comptabilisées, par garantie et par exercice"},
            [("Formule", "Le S/C d'une garantie rapporte ses prestations comptabilisées de l'année à ses cotisations comptabilisées de la même année, de 2020 à 2026."),
             ("Prestations maladie", "Adenka retient les comptes 601, 6013, 6018 et 60800007, soit douze comptes de prestations maladie en 2025."),
             ("Prestations décès", "Adenka retient les comptes 6021 et 6024. En 2020, la prime d'assurance décès (65180000) s'y ajoute, car un assureur portait alors la garantie."),
             ("Norme et limite", "La Décision n° 05/2022 fixe une sinistralité entre 70 % et 80 %. Le ratio comptable mélange les soins de l'année et des factures d'années antérieures.")],
            "Le ratio comptable mélange les soins de l'année et les factures d'années antérieures payées ou comptabilisées dans l'année. Le vrai ratio (par année de survenance des soins) demande la base des prestations avec la date des soins.")
sc25 = EX["sc_2025"]
contenu({"titre": "Calcul pas à pas : S/C maladie 2025", "sous_titre": "Douze comptes de prestations divisés par un compte de cotisations",
         "synthese": f"{F('Prest mal 2025 (pas)', s[2025]['prest_maladie'])} / {F('Cot mal 2025 (pas)', s[2025]['cot_maladie'])} = {pc('S/C 2025 (pas)', s[2025]['sc_maladie'])}"},
        "Le compte 60120000 « Charges de prestations maladie » porte à lui seul plus de la moitié du total : il regroupe des prestations sans détail de nature, ce qui empêche l'analyse par type de soins. Montants exacts lus dans la balance 2025.",
        pas_a_pas={"entetes": ["Compte", "Intitulé", "Montant 2025"], "lignes": [[c[0], c[1], nb(float(c[2])) + " FCFA"] for c in sc25["prestations"]] + [["Total", "Prestations maladie", F("Total prest (pas)", s[2025]["prest_maladie"])], [sc25["cotisations"][0][0], sc25["cotisations"][0][1], nb(-float(sc25["cotisations"][0][2])) + " FCFA"]],
                   "largeurs": [3.4, 11.0, 6.0], "rouge": [(7, 2), (12, 2), (13, 2)], "petit": True,
                   "explications": ["Le compte global 60120000 porte plus de la moitié des prestations, sans détail de nature.", "Les douze comptes de prestations maladie totalisent ce montant.", "Les cotisations maladie acquises forment le dénominateur du ratio."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2025\\BALANCE FPPN 2025.xlsx"},
        analyse=["**Appliquée à la seule maladie, la fourchette réglementaire de 70 % à 80 % est respectée en 2025.** Le compte global empêche pourtant de savoir quels soins font monter le ratio."],
        implication="Le FPPN doit ventiler le compte 60120000 par nature de soins et par prestataire, ce qui permettra de cibler les mesures de maîtrise des dépenses.",
        source="Source : balance des comptes 2025 (Sage 100) ; traitements Adenka")
contenu({"titre": "Sinistralité comptable", "sous_titre": "Les prestations maladie absorbent presque toutes les cotisations maladie",
         "synthese": f"S/C maladie : {pc('S/C 2021 (s)', s[2021]['sc_maladie'])} en 2021, {pc('S/C 2025 (s)', s[2025]['sc_maladie'])} en 2025, {pc('S/C 2026 (s)', s[2026]['sc_maladie'])} sur huit mois de 2026"},
        f"Test de tendance (moindres carrés sur 2020 à 2025) : le S/C maladie progresse de {N('Pente S/C', C['tendances'][1]['pente'], 2)} points par an (IC 95 % : {N('IC bas S/C', C['tendances'][1]['ic95_bas'], 2)} à {N('IC haut S/C', C['tendances'][1]['ic95_haut'], 2)} ; p-valeur {nb(C['tendances'][1]['p_valeur'], 4)}). Avec six points seulement, la tendance est significative au seuil de 5 % mais son rythme est imprécis.",
        graphique={"titre": fig("Prestations comptabilisées / cotisations (S/C)"), "type": "lignes", "categories": cats_ex,
                   "series": [{"nom": "Maladie", "valeurs": [round(100 * s[a]["sc_maladie"], 2) for a in range(2020, 2027)], "couleur": BLEU},
                              {"nom": "Décès (prime d'assurance incluse en 2020)", "valeurs": [round(100 * s[a]["sc_deces_avec_prime"], 2) for a in range(2020, 2027)], "couleur": ORANGE}],
                   "format": "0.00", "unite": "En % des cotisations de la garantie"},
        analyse=[f"**Le ratio maladie a été multiplié par {N('Mult S/C', s[2025]['sc_maladie'] / s[2021]['sc_maladie'], 2)} entre 2021 et 2025.** Il dépasse la fourchette réglementaire sur les huit premiers mois de 2026.",
                 "**Les ratios très bas de 2020 à 2022 sont suspects.** Des factures de ces années ont probablement été comptabilisées plus tard, ce que seule la base des prestations permettra de vérifier.",
                 f"**En décès, le ratio oscille entre {pc('Min S/C décès', min(s[a]['sc_deces'] for a in range(2022, 2027)))} et {pc('Max S/C décès', max(s[a]['sc_deces'] for a in range(2022, 2027)))} depuis 2022.** La garantie décès reste largement excédentaire."],
        implication="À ce rythme, les prestations maladie dépasseront les cotisations maladie : sans mesure de maîtrise ou révision du barème, la garantie maladie deviendra déficitaire et pèsera sur la réserve obligatoire.",
        source=SRC_COMPTA)
tlp = C["tendances"][2]
contenu({"titre": "Croissance des prestations", "sous_titre": "Les prestations maladie croissent beaucoup plus vite que les cotisations",
         "synthese": f"Taux de croissance annuel moyen 2021 à 2025 : prestations maladie {pc('TCAM prest mal', TC['Prestations maladie']['tcam'])}, cotisations {pc('TCAM cot (cr)', TC['Cotisations maladie et décès']['tcam'])}"},
        f"Le taux de croissance annuel moyen (TCAM) est le taux constant qui mène de la valeur de départ à la valeur d'arrivée : TCAM = (V_fin / V_début)^(1/n) - 1. Une régression du logarithme des prestations sur le temps donne une croissance annuelle moyenne de {pc('Croissance log', tlp['croissance_annuelle'])} (IC 95 % : {pc('IC bas log', tlp['croissance_ic95_bas'])} à {pc('IC haut log', tlp['croissance_ic95_haut'])}), estimation plus robuste aux valeurs extrêmes que le TCAM.",
        tableau={"titre": tab("Taux de croissance annuels moyens"), "entetes": ["Indicateur", "Période", "Valeur de départ", "Valeur d'arrivée", "Multiplié par", "TCAM"], "largeurs": [7.6, 4.6, 5.6, 5.6, 3.6, 4.2],
                 "lignes": [[t["indicateur"], f"{t['debut']} à {t['fin']}",
                             (F(f"V0 {t['indicateur']}", t["v0"]) if t["v0"] > 100000 else N(f"V0 {t['indicateur']}", t["v0"])),
                             (F(f"V1 {t['indicateur']}", t["v1"]) if t["v1"] > 100000 else N(f"V1 {t['indicateur']}", t["v1"])),
                             N(f"Mult {t['indicateur']}", t["multiplicateur"], 2), pc(f"TCAM {t['indicateur']}", t["tcam"])] for t in C["tcam"]]},
        disposition="dessous",
        analyse=["**Les prestations maladie et l'encours du Grenier croissent bien plus vite que les cotisations et les effectifs.** Les ressources ne suivent pas le rythme des emplois."],
        implication="Un écart durable entre la croissance des prestations et celle des cotisations conduit mécaniquement à un déficit technique : la projection devra en fixer l'horizon.",
        source=SRC_COMPTA + " ; " + SRC_PREC[9:] + " ; " + SRC_GRENIER[9:])
nat = R["prestations_maladie_nature"]
contenu({"titre": "Prestations maladie par nature", "sous_titre": "Plus de la moitié des prestations maladie est enregistrée sans détail de nature",
         "synthese": f"En 2025, le compte global porte {F('Compte global 2025', nat[5]['2025'])} sur {F('Prest mal 2025 (n)', s[2025]['prest_maladie'])}"},
        "Lecture : chaque couleur est une nature de soins. La catégorie « autres prestations maladie » correspond aux comptes 6012x, dont le compte global 60120000. Imagerie, analyses et consultations sont les natures détaillées qui progressent le plus.",
        graphique={"titre": fig("Prestations maladie comptabilisées par nature"), "type": "barres_empilees", "categories": cats_ex,
                   "series": [{"nom": n["nature"], "valeurs": [round(n[str(e)]) for e in range(2020, 2027)]} for n in nat],
                   "format": FMT_M, "etiquettes": False, "legende": "droite", "unite": "En millions de FCFA"},
        analyse=[f"**Les prestations maladie passent de {F('Prest 2021', s[2021]['prest_maladie'])} en 2021 à {F('Prest 2025 (nat)', s[2025]['prest_maladie'])} en 2025.**",
                 f"**Imagerie, analyses et consultations représentent {pc('Part img+ana+cons', (nat[4]['2025'] + nat[3]['2025'] + nat[1]['2025']) / s[2025]['prest_maladie'])} des prestations de 2025.** Ce sont les natures détaillées qui progressent le plus.",
                 "**Le compte global empêche le suivi par type de soins et par prestataire.**"],
        implication="Sans détail par nature et par prestataire, le FPPN ne peut ni négocier ses tarifs, ni repérer les abus, ni calculer le taux de prestations indues exigé par la Décision n° 05/2022.",
        source=SRC_COMPTA)
contenu({"titre": "Frais de gestion et résultat", "sous_titre": "Les excédents se réduisent et deviennent un déficit en 2026",
         "synthese": f"Résultat net : {F('Résultat 2020', s[2020]['resultat'])} en 2020, {F('Résultat 2025 (fr)', s[2025]['resultat'])} en 2025, {F('Résultat 2026', s[2026]['resultat'])} sur huit mois de 2026"},
        "Frais de gestion = comptes 61 à 66, hors prime d'assurance décès. Résultat = produits - charges + résultat hors activités ordinaires (classe 8). Les huit mois de 2026 n'intègrent pas les écritures de fin d'année (amortissements, provisions).",
        graphique={"titre": fig("Frais de gestion et résultat net"), "type": "barres", "categories": cats_ex,
                   "series": [{"nom": "Frais de gestion", "valeurs": [round(s[a]["frais_gestion"]) for a in range(2020, 2027)], "couleur": GRIS_C},
                              {"nom": "Résultat net", "valeurs": [round(s[a]["resultat"]) for a in range(2020, 2027)], "couleur": BLEU}],
                   "format": FMT_M, "unite": "En millions de FCFA"},
        analyse=[f"**Les frais de gestion pèsent entre {pc('Min frais', min(s[a]['taux_frais'] for a in range(2021, 2026)))} et {pc('Max frais', max(s[a]['taux_frais'] for a in range(2021, 2026)))} des cotisations depuis 2021.** La Décision n° 05/2022 limite les charges de fonctionnement à 20 % des cotisations.",
                 f"**Le résultat baisse depuis 2021.** Il passe de {F('Résultat 2021', s[2021]['resultat'])} à {F('Résultat 2025 (t)', s[2025]['resultat'])} en 2025.",
                 "**Les huit premiers mois de 2026 sont déficitaires**, avant même les écritures de fin d'année."],
        implication="Les frais pèsent presque deux fois la norme réglementaire. Leur réduction est le levier le plus direct pour restaurer l'excédent pendant que la sinistralité monte.",
        source=SRC_COMPTA)
retenir("sinistralité", ("La sinistralité maladie approche le niveau des cotisations", "Ratio S/C maladie multiplié par plus de douze depuis 2021, résultat en baisse, déficit sur 2026",
                         "Les prestations maladie absorbent une part croissante des cotisations, et le résultat du FPPN baisse depuis 2021."),
        f"La hausse du S/C maladie est statistiquement significative, d'environ {N('Pente S/C (ret)', C['tendances'][1]['pente'], 2)} points par an de 2020 à 2025. Les frais dépassent la norme de 20 % des cotisations.",
        "Sans maîtrise des prestations et des frais, le FPPN deviendra déficitaire et ne pourra plus doter sa réserve obligatoire.",
        "Le FPPN doit détailler le compte global de prestations par nature et par prestataire, et transmettre la base des prestations.",
        "La hausse rapide du ratio maladie est le principal risque technique : à ce rythme, les cotisations maladie ne couvriront plus les prestations.")

# ---------------------------------------------------------------- PARTIE 6 : ÉPARGNE
separateur(6, "Épargne des mutualistes", "Dette d'épargne, actifs réservés, sensibilité aux retraits")
methode_gab({"titre": "Méthode : adossement de l'épargne", "sous_titre": "La couverture compare l'épargne due aux actifs qui lui sont réservés",
             "synthese": "Couverture = (banques, dépôts à terme et fonds réservés) / (comptes PEP et EFC)"},
            [("Passif d'épargne", "L'épargne due correspond aux soldes créditeurs des comptes 41195100 (PEP) et 41195200 (EFC). C'est une dette du FPPN envers les mutualistes."),
             ("Actifs réservés", "Adenka retient les comptes bancaires et dépôts à terme dont l'intitulé désigne l'EFC ou le PEP (5211020x, 5221020x, 5211030x, 5221030x, 526120) et le fonds commun de placement (5330)."),
             ("Couverture", "La couverture rapporte les actifs réservés à l'épargne due, à chaque clôture de 2020 au 31/08/2026."),
             ("Limite", "Le règlement des produits n'a pas été transmis : la part de l'épargne exigible à court terme et l'affectation du fonds commun de placement restent à confirmer.")],
            "L'épargne PEP et EFC appartient aux mutualistes : c'est une dette du FPPN. Un bon adossement suppose des actifs liquides et réservés d'un montant au moins égal à la part exigible.")
ep = EX["epargne_2026"]
contenu({"titre": "Calcul pas à pas : couverture de l'épargne", "sous_titre": "Deux comptes de dette face à quatre comptes d'actifs réservés",
         "synthese": f"{F('Actifs dédiés (pas)', b[2026]['actifs_dedies_epargne'])} / {F('Épargne (pas)', b[2026]['epargne_totale'])} = {pc('Couverture (pas)', b[2026]['couverture_epargne'])} au 31/08/2026"},
        "Montants lus dans la balance au 31/08/2026 (provisoire). Les soldes créditeurs du passif sont affichés en valeur positive. Le fonds commun de placement est compté comme actif réservé bien que sa destination ne soit pas précisée : la couverture réelle est donc au plus de ce niveau.",
        pas_a_pas={"entetes": ["Compte", "Intitulé", "Solde au 31/08/2026"],
                   "lignes": [[c[0], c[1], nb(-float(c[2])) + " FCFA"] for c in ep["passif"]] + [["", "Total épargne due", F("Épargne totale (pas)", b[2026]["epargne_totale"])]] + [[c[0], c[1], nb(float(c[2])) + " FCFA"] for c in ep["actifs"]] + [["", "Total actifs réservés", F("Total actifs (pas)", b[2026]["actifs_dedies_epargne"])]],
                   "largeurs": [3.4, 11.4, 6.6], "rouge": [(2, 2), (7, 2)],
                   "explications": ["Les deux comptes 41195 enregistrent la dette envers les épargnants PEP et EFC.", "Les banques, le dépôt à terme et le fonds commun de placement forment les actifs réservés.", f"La couverture vaut donc {pc('Couverture (expl)', b[2026]['couverture_epargne'])} au 31/08/2026."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2026\\BALANCE FPPN DU 01-01-2026 AU 31-08-2026.xlsx"},
        analyse=["**Ce ratio est aussi le ratio de solvabilité des risques longs de la Décision n° 05/2022** (actifs admis / engagements réglementés), dont la norme est d'au moins 1."],
        implication="Avec une couverture inférieure à un cinquième, le FPPN ne respecte pas cette norme : la partie 10 la reprend dans le tableau de bord sans refaire le calcul.",
        source="Source : balance au 31/08/2026 (impression provisoire)")
contenu({"titre": "Adossement de l'épargne", "sous_titre": "L'épargne des mutualistes n'est couverte qu'à 16,28 % par des actifs réservés",
         "synthese": f"Couverture : entre {pc('Min couverture', min(b[a]['couverture_epargne'] for a in range(2020, 2027)))} et {pc('Max couverture', max(b[a]['couverture_epargne'] for a in range(2020, 2027)))} selon les années"},
        "L'épargne non couverte finance d'autres emplois : investissements au CMP, terrains, créance sur l'Hôpital de la Police, prêts du Grenier. Ces emplois ne sont pas mobilisables rapidement.",
        graphique={"titre": fig("Épargne au passif et actifs réservés"), "type": "barres", "categories": yrs_bil,
                   "series": [{"nom": "Épargne PEP et EFC", "valeurs": [round(b[a]["epargne_totale"]) for a in range(2020, 2027)], "couleur": BLEU},
                              {"nom": "Actifs réservés", "valeurs": [round(b[a]["actifs_dedies_epargne"]) for a in range(2020, 2027)], "couleur": ORANGE}],
                   "format": FMT_M, "unite": "En millions de FCFA"},
        analyse=[f"**La dette d'épargne augmente de {pc('Hausse épargne', b[2026]['epargne_totale'] / b[2020]['epargne_totale'] - 1)} depuis 2020.** Les actifs réservés ne suivent pas cette hausse.",
                 f"**L'épargne EFC triple.** Elle passe de {F('EFC 2020', b[2020]['epargne_efc'])} en 2020 à {F('EFC 2026', b[2026]['epargne_efc'])} au 31/08/2026.",
                 "**L'épargne non couverte finance des emplois peu liquides** : investissements au CMP, terrains, créance sur l'Hôpital de la Police et prêts du Grenier."],
        implication="Le règlement d'exécution n° 004/2022 impose le cantonnement des patrimoines quand une mutuelle couvre à la fois des risques courts et longs (art. 15). L'épargne n'est séparée qu'en partie des autres ressources du FPPN.",
        source=SRC_COMPTA)
pv = [p for p in R["pep_versements"] if p["annee"] >= 2020]
ev = {e["annee"]: e for e in R["efc_versements"]}
contenu({"titre": "Versements d'épargne", "sous_titre": "Les versements au PEP ont été multipliés par quatre depuis 2022",
         "synthese": f"PEP 2025 : {F('PEP 2025', pv[5]['versements'])} versés par {N('Épargnants PEP 2025', pv[5]['epargnants'])} épargnants"},
        "Les versements sont les précomptes retour de la Solde. La différence entre versements de l'année et variation du compte d'épargne mesure les sorties de l'année (échéances, retraits, transferts), sans pouvoir les détailler faute de comptes individuels.",
        graphique={"titre": fig("Versements précomptés par produit d'épargne"), "type": "barres", "categories": [str(p["annee"]) if p["annee"] < 2026 else "2026 (8 mois)" for p in pv],
                   "series": [{"nom": "PEP 998", "valeurs": [round(p["versements"]) for p in pv], "couleur": BLEU},
                              {"nom": "EFC 870", "valeurs": [round(ev[p["annee"]]["versements"]) if p["annee"] in ev else None for p in pv], "couleur": ORANGE}],
                   "format": FMT_M, "unite": "En millions de FCFA (EFC à partir de 2023)"},
        analyse=[f"**Les épargnants PEP passent de {N('PEP 2020', pv[0]['epargnants'])} en 2020 à {N('PEP 2026', pv[-1]['epargnants'])} en 2026.** Le produit attire un nombre croissant de policiers.",
                 f"**En 2025, le PEP reçoit {F('PEP 2025 (t)', pv[5]['versements'])} mais le compte n'augmente que de {F('Var PEP 2025', b[2025]['epargne_pep'] - b[2024]['epargne_pep'])}.** La différence correspond à des sorties d'épargne.",
                 f"**Les sorties d'épargne atteignent {F('Sorties épargne 2025', AG[2025]['sorties_epargne'])} en 2025** (débits des comptes 41195100 et 41195200)."],
        implication="Ces sorties sont les « engagements » d'un trimestre que la réserve de trésorerie des risques longs doit couvrir (art. 19 du règlement d'exécution) : leur suivi mensuel est nécessaire pour fixer le niveau de cette réserve.",
        source=SRC_PREC + " ; " + SRC_COMPTA[9:])
rr = C["sensibilite_retraits"]
contenu({"titre": "Sensibilité : retraits d'épargne", "sous_titre": "Un retrait de plus d'un cinquième de l'épargne épuiserait toute la trésorerie du FPPN",
         "synthese": f"Les actifs réservés couvrent un retrait de {pc('Retrait couvert actifs', C['retrait_max_couvert']['par_actifs_dedies'])} de l'épargne ; la trésorerie totale, {pc('Retrait couvert tréso', C['retrait_max_couvert']['par_tresorerie'])}"},
        "Scénarios théoriques, pas des prévisions : on suppose qu'une part de l'épargne est demandée en une fois et on la compare aux actifs réservés puis à toute la trésorerie (banques et caisses). Les conditions réelles de retrait (délais, pénalités, échéances) ne sont pas connues.",
        graphique={"titre": fig("Besoin de liquidités selon le taux de retrait"), "type": "barres", "categories": [nb(100 * r["taux_retrait"], 2) + " %" for r in rr],
                   "series": [{"nom": "Montant à rembourser", "valeurs": [round(r["besoin"]) for r in rr], "couleur": BLEU},
                              {"nom": "Actifs réservés", "valeurs": [round(r["actifs_dedies"]) for r in rr], "couleur": ORANGE},
                              {"nom": "Trésorerie totale", "valeurs": [round(r["tresorerie_totale"]) for r in rr], "couleur": GRIS_C}],
                   "format": FMT_M, "unite": "En millions de FCFA, au 31/08/2026"},
        analyse=[f"**Les actifs réservés couvrent un retrait de 10 % de l'épargne ({F('Retrait 10%', rr[1]['besoin'])}).**",
                 f"**Au-delà de {pc('Seuil tréso', C['retrait_max_couvert']['par_tresorerie'])} de retraits, toute la trésorerie du FPPN ne suffit plus.**",
                 "**Le FPPN doit alors céder ou mobiliser des actifs peu liquides**, comme les terrains ou les avances au CMP."],
        implication="Le risque de liquidité dépend des conditions de retrait du PEP et de l'EFC : le règlement de ces produits et les âges des épargnants permettront de passer de scénarios théoriques à un échéancier réel.",
        source=SRC_COMPTA)
contenu({"titre": "Poste 998 : autres références", "sous_titre": "Le poste de l'épargne PEP porte aussi d'autres références à expliquer",
         "synthese": "PTI, RPEP, EP et TI apparaissent à côté des références PEP sur la période 2020 à 2026"},
        "La référence est le code qui accompagne chaque précompte (exemple : PEP0522 = PEP souscrit en mai 2022). Les références PTI ont des montants mensuels importants et peu de mutualistes : elles pourraient correspondre à des remboursements de prêts, ce que le FPPN doit confirmer.",
        tableau={"titre": tab("Références du poste 998, 2020 à août 2026"), "entetes": ["Référence", "Lignes", "Montant total", "Mutualistes"], "largeurs": [4.0, 4.0, 7.4, 4.0],
                 "lignes": [[r["produit"], N(f"Lignes {r['produit']}", r["lignes"]), F(f"Montant {r['produit']}", r["montant"]), N(f"Mut {r['produit']}", r["mutualistes"])] for r in R["poste_998_references"]]},
        disposition="droite",
        analyse=["**Le poste 998 ne porte pas que de l'épargne PEP.** D'autres références, dont PTI, apparaissent avec des montants importants et peu de mutualistes.",
                 "**Leur nature n'est documentée nulle part.** Il peut s'agir de remboursements de prêts PEP (comptes 27110000 et 27120000), ce que le FPPN doit confirmer."],
        implication="Tant que ces références ne sont pas expliquées, la dette d'épargne PEP ne peut pas être reconstituée à partir des précomptes.",
        source=SRC_PREC)
retenir("épargne", ("L'épargne due aux mutualistes est très peu couverte par des actifs disponibles", "Couverture de 16,28 % au 31/08/2026 ; un retrait de 21,34 % épuiserait toute la trésorerie",
                    "La dette d'épargne augmente plus vite que les actifs réservés, et l'écart finance des emplois peu liquides."),
        F("Épargne (ret)", b[2026]["epargne_totale"]) + " d'épargne sont dus au 31/08/2026, pour " + F("Actifs réservés (ret)", b[2026]["actifs_dedies_epargne"]) + " d'actifs réservés.",
        "Le FPPN ne peut pas rembourser rapidement une vague de retraits, et il ne respecte pas le ratio de solvabilité des risques longs.",
        "Le FPPN doit transmettre le règlement des produits, les comptes individuels et les conditions de retrait, puis cantonner l'épargne.",
        "L'adossement de l'épargne est le principal risque financier relevé.")

# ---------------------------------------------------------------- PARTIE 7 : GRENIER
separateur(7, "Prêts du Grenier", "Encours, production, retards, précomptes")
methode_gab({"titre": "Méthode : risque de crédit", "sous_titre": "Le risque des prêts se mesure par la part de l'encours en retard",
             "synthese": "Encours en retard de plus de 30 jours / encours total, à chaque date de balance âgée"},
            [("Part en retard (montant)", "Adenka rapporte l'encours des prêts en retard de plus de 30 jours à l'encours total, sur les 37 balances âgées reçues."),
             ("Part en retard (nombre)", "Adenka rapporte le nombre de prêts en retard au nombre de prêts au 31/08/2026, avec un intervalle de confiance de Wilson à 95 %."),
             ("Précompte des échéances", "Adenka rapporte les précomptes 9X4 du mois aux échéances des prêts en cours, pour les mois de 2026."),
             ("Limite", "Le retard est celui que calcule le logiciel de prêts. Il dépend de la date à laquelle la Solde verse les précomptes.")],
            "L'encours est le capital restant dû. Le retard (en jours) est celui imprimé par le logiciel de prêts. On mesure la part en nombre de prêts et en montant d'encours.")
ge_fa = [g for g in GE if g["arrete"].startswith("31/12") or g["arrete"] == "31/08/2026"]
contenu({"titre": "Encours du Grenier", "sous_titre": "L'encours de prêts a été multiplié par 132 depuis fin 2021",
         "synthese": f"Au 31/08/2026 : {N('Prêts', GP['prets'])} prêts, {F('Encours (gr)', GP['encours'])}, dont {F('Financement FPPN', b[2026]['grenier_financement'])} financés par le FPPN"},
        f"Taux de croissance annuel moyen de l'encours de fin 2022 à fin 2025 : {pc('TCAM encours', TC['Encours des prêts du Grenier']['tcam'])}. La créance du FPPN sur le Grenier (compte 47111102) mesure le financement apporté.",
        graphique={"titre": fig("Encours des prêts et financement du FPPN"), "type": "barres", "categories": [g["arrete"] for g in ge_fa],
                   "series": [{"nom": "Encours", "valeurs": [round(g["encours"]) for g in ge_fa], "couleur": BLEU},
                              {"nom": "Financement FPPN", "valeurs": [round(g["financement_fppn_47111102"]) if g["financement_fppn_47111102"] is not None else None for g in ge_fa], "couleur": ORANGE}],
                   "format": FMT_M, "unite": "En millions de FCFA"},
        analyse=[f"**L'encours passe de {F('Encours 2021', ge_fa[0]['encours'])} fin 2021 à {F('Encours 2026 (t)', GP['encours'])} au 31/08/2026**, soit {N('Mult encours', GP['encours'] / ge_fa[0]['encours'], 2)} fois plus.",
                 f"**Le FPPN finance {pc('Part financement', b[2026]['grenier_financement'] / GP['encours'])} de l'encours.** Le Grenier repose donc presque entièrement sur les ressources du FPPN.",
                 f"**Les prêts sont courts et chers.** La durée médiane est de {N('Durée médiane', GP['duree_mediane'])} mois et le taux médian de {pcv('Taux médian', GP['taux_median'])}."],
        implication="Si le Grenier relève des activités génératrices de revenus, le règlement d'exécution n° 004/2022 les autorise à condition qu'elles ne mettent pas en péril l'activité d'assurance et que leurs pertes ne soient pas couvertes par les cotisations ni par les réserves (art. 26) : le financement du Grenier doit être encadré par une décision formelle.",
        source=SRC_GRENIER + " ; " + SRC_COMPTA[9:])
gpr = R["grenier_production"]
contenu({"titre": "Production de prêts", "sous_titre": "Les crédits scolaires et habitat portent la croissance de la production",
         "synthese": f"Montants débloqués : {F('Prod 2022', gpr['total'][1])} en 2022, {F('Prod 2025', gpr['total'][4])} en 2025 ({N('Nb prêts 2025', gpr['nombre'][4])} prêts)"},
        "Production = montants débloqués sur la période (fichiers annuels de 2021 à 2025, mensuels par sous-portefeuille en 2026). La politique de crédit de 2021 est au format Word ancien et doit être renvoyée lisible pour vérifier taux, durées et plafonds.",
        graphique={"titre": fig("Montants débloqués par type de crédit"), "type": "barres_empilees", "categories": [str(a) if a < 2026 else "2026 (8 mois)" for a in gpr["annees"]],
                   "series": [{"nom": t.capitalize(), "valeurs": [round(v) for v in gpr["valeurs"][t]]} for t in gpr["types"]],
                   "format": FMT_M, "etiquettes": False, "legende": "droite", "unite": "En millions de FCFA"},
        analyse=[f"**La production a été multipliée par {N('Mult prod', gpr['total'][4] / gpr['total'][1], 2)} entre 2022 et 2025.**",
                 f"**Les crédits scolaires et habitat font {pc('Part scol+hab', (gpr['valeurs']['CRÉDIT SCOLAIRE'][4] + gpr['valeurs']['CREDIT HABITAT'][4]) / gpr['total'][4])} des montants de 2025.**",
                 "**Le crédit habitat, aux montants unitaires plus importants, allonge la durée du risque.**"],
        implication="Une production qui croît plus vite que les ressources propres du Grenier augmente le besoin de financement du FPPN : la politique de crédit doit fixer un plafond d'encours rapporté aux ressources disponibles.",
        source=SRC_GRENIER)
ph = [p for p in R["grenier_par_historique"] if p["arrete"] != "30/09/2026"]
contenu({"titre": "Retards de remboursement", "sous_titre": "Les retards ont culminé début 2026 avant de revenir à un niveau faible",
         "synthese": f"Encours en retard de plus de 30 jours : {pc('PAR30 janv 2026', [p for p in ph if p['arrete'] == '31/01/2026'][0]['par30'])} fin janvier 2026, {pc('PAR30 août 2026', GP['par30'])} au 31/08/2026"},
        f"Au 31/08/2026, {N('Prêts retard 30', IC['retard30'])} prêts sur {N('Prêts (ic)', IC['prets'])} ont plus de 30 jours de retard, soit {pc('Part prêts retard', IC['p30'])} des prêts (IC de Wilson à 95 % : {pc('IC bas retard', IC['p30_bas'])} à {pc('IC haut retard', IC['p30_haut'])}). En montant, la part est plus faible ({pc('PAR30 (n)', GP['par30'])}) : les prêts en retard sont de petits encours.",
        graphique={"titre": fig("Part de l'encours en retard de plus de 30 et 90 jours"), "type": "lignes", "categories": [p["arrete"] for p in ph],
                   "series": [{"nom": "Plus de 30 jours", "valeurs": [round(100 * p["par30"], 2) for p in ph], "couleur": BLEU},
                              {"nom": "Plus de 90 jours", "valeurs": [round(100 * p["par90"], 2) for p in ph], "couleur": ORANGE}],
                   "format": "0.00", "etiquettes": False, "unite": "En % de l'encours"},
        analyse=["**Fin janvier 2026, la plus grande part de l'encours avait entre 31 et 90 jours de retard.**",
                 "**Ces retards se sont résorbés en quelques mois.** Ce profil fait penser à des échéances non précomptées en fin d'année 2025.",
                 f"**{N('Prêts > 180 j', IC['retard180'])} prêts ont plus de 180 jours de retard.** Ces prêts doivent être provisionnés."],
        implication="Les retards dépendent surtout du circuit de précompte, pas de la solvabilité des emprunteurs : le FPPN doit obtenir de la Solde un calendrier fiable et suivre chaque mois les échéances non précomptées.",
        source=SRC_GRENIER)
pr_ = EX["prets_retard"]
contenu({"titre": "Calcul pas à pas : prêts en retard", "sous_titre": "Les retards les plus anciens dépassent trois ans et révèlent des mécanos anormaux",
         "synthese": "Les quatre prêts les plus en retard au 31/08/2026 ont entre 1 105 et 1 308 jours de retard"},
        "Lecture : (1) le retard est imprimé en jours par le logiciel ; (2) trois de ces prêts ont un « mécano » de 2 ou 3 chiffres (142, 87, 166), alors qu'un mécano en compte 6 : ce sont des anomalies de saisie ; (3) ces prêts anciens ont encore un encours et doivent être provisionnés ou passés en perte. Noms masqués.",
        pas_a_pas={"entetes": ["Mécano", "Nom", "Montant débloqué", "Encours", "Retard (jours)"], "lignes": [[r[0], r[1], nb(float(r[2])) + " FCFA", nb(float(r[3])) + " FCFA", r[4]] for r in pr_["lignes"]],
                   "largeurs": [3.4, 4.0, 6.0, 5.6, 4.6], "rouge": [(0, 0), (0, 4), (1, 0)],
                   "explications": ["Ces prêts ont un mécano de 2 ou 3 chiffres, ce qui révèle une saisie anormale.", "Le plus ancien de ces prêts a plus de trois ans de retard.", "Ces prêts ont encore un encours inscrit, qui doit être provisionné."],
                   "fichier": "AUDIT ADENKA\\SERVICE GRENIER DU POLICIER\\FICHIER ALLER\\GRENIER AOÛT 2026\\BalanceAgeeDetaille AOUT 2026.xls"},
        analyse=["**Les prêts les plus anciens ne peuvent plus être précomptés, car leur mécano est faux.** Sans mécano valide, la Solde ne peut pas retenir l'échéance sur le salaire."],
        implication="Le Grenier doit corriger ces identifiants et décider, pour chaque prêt, d'une provision ou d'un passage en perte.",
        source=SRC_GRENIER + " ; noms masqués")
ech = [e for e in R["grenier_echeances"] if e["precompte_9x4"]]
contenu({"titre": "Échéances et précomptes", "sous_titre": "Les précomptes ne couvrent que 76 à 85 % des échéances des prêts en cours",
         "synthese": f"De janvier à août 2026, l'écart mensuel va de {F('Écart min', min(e['echeances_balance'] - e['precompte_9x4'] for e in ech))} à {F('Écart max', max(e['echeances_balance'] - e['precompte_9x4'] for e in ech))}"},
        "Échéances = somme des montants d'échéance des prêts ayant un encours dans la balance âgée du mois. Précomptes = précomptes 9X4 retour du même mois. L'écart peut être payé hors Solde (caisse, banque) ou correspondre à des rejets.",
        graphique={"titre": fig("Échéances des prêts en cours et précomptes 9X4, 2026"), "type": "barres", "categories": [e["mois"] for e in ech],
                   "series": [{"nom": "Échéances", "valeurs": [round(e["echeances_balance"]) for e in ech], "couleur": BLEU},
                              {"nom": "Précomptes 9X4", "valeurs": [round(e["precompte_9x4"]) for e in ech], "couleur": ORANGE}],
                   "format": FMT_M, "unite": "En millions de FCFA"},
        analyse=[f"**La part précomptée va de {pc('Min précompte/éch', min(e['precompte_9x4'] / e['echeances_balance'] for e in ech))} à {pc('Max précompte/éch', max(e['precompte_9x4'] / e['echeances_balance'] for e in ech))} selon le mois.**",
                 "**L'export du logiciel perd la lettre-clé du mécano** (321100C devient 321100), ce qui complique le rapprochement prêt par prêt.",
                 "**Un rapprochement avec les encaissements bancaires du Grenier est nécessaire** pour savoir si l'écart est payé autrement."],
        implication="Si l'écart n'est pas encaissé par un autre canal, il alimente chaque mois de nouveaux retards : c'est la question à poser en priorité au service du Grenier.",
        source=SRC_GRENIER + " ; " + SRC_PREC[9:])
sp = C["sensibilite_pertes_grenier"]
contenu({"titre": "Sensibilité : pertes sur prêts", "sous_titre": "Même une perte totale sur les prêts très en retard reste limitée en montant",
         "synthese": f"Encours en retard de plus de 180 jours : {F('Encours 180', sp['encours_180j'])} ; de plus de 90 jours : {F('Encours 90', sp['encours_90j'])}"},
        "Trois taux de perte théoriques (25 %, 50 %, 100 %) appliqués à l'encours en retard. Ce ne sont pas des estimations de perte, mais des ordres de grandeur pour fixer une provision. La politique de provisionnement doit être fixée par le FPPN.",
        tableau={"titre": tab("Provision selon le taux de perte"), "entetes": ["Taux de perte", "Sur l'encours > 180 jours", "Sur l'encours > 90 jours"], "largeurs": [5.0, 7.2, 7.2],
                 "lignes": [[pc(f"Perte {p['perte']}", p["perte"]), F(f"Prov180 {p['perte']}", p["provision_180j"]), F(f"Prov90 {p['perte']}", p["provision_90j"])] for p in sp["scenarios"]]},
        disposition="droite",
        analyse=[f"**Une perte totale sur les prêts en retard de plus de 90 jours coûterait {F('Perte max 90', sp['scenarios'][-1]['provision_90j'])}.** Ce montant reste faible au regard de l'encours.",
                 "**Le risque du Grenier n'est donc pas la perte sur prêts.** Il tient au volume de ressources du FPPN immobilisées dans les prêts."],
        implication="Une politique de provisionnement écrite (taux par tranche de retard) suffirait à couvrir ce risque, mais elle ne règle pas la question de la liquidité de l'épargne.",
        source=SRC_GRENIER)
retenir("Grenier", ("Le Grenier croît très vite grâce au financement du FPPN", "Encours multiplié par 132, financé à 87,67 %, retards maîtrisés mais précomptes incomplets",
                    "Le Grenier a multiplié ses prêts en s'appuyant presque entièrement sur les ressources du FPPN."),
        f"L'encours atteint {F('Encours (ret)', GP['encours'])} au 31/08/2026. Les précomptes couvrent 76 à 85 % des échéances, et 126 prêts ont plus de 180 jours de retard.",
        "Le principal risque n'est pas la perte sur prêts, limitée, mais l'immobilisation de l'épargne des mutualistes dans les prêts.",
        "Le FPPN doit transmettre une politique de crédit lisible, les encaissements hors Solde et la décision qui encadre le financement du Grenier.",
        "Le risque principal n'est pas la perte sur prêts, limitée, mais la mobilisation de l'épargne des mutualistes pour financer cette croissance.")

# ---------------------------------------------------------------- PARTIE 8 : RETRAITÉS ET ANNEXES
separateur(8, "Retraités et annexes", "CGRAE, flotte automobile, MEDOC, projets immobiliers")
contenu({"titre": "Retraités CGRAE", "sous_titre": "Les retraités cotisants sont passés de 16 à 610 en deux ans et demi",
         "synthese": f"Précompte moyen de {F('Moyen CGRAE (r)', sum(c['moyen'] for c in CG) / len(CG))} par mois sur la pension ; données depuis février 2024"},
        "Les retraités cotisent sur leur pension par la CGRAE. Leur coût santé, supérieur à celui des actifs, devra être comparé à leur cotisation dès réception des factures. La convention avec la CGRAE n'a pas été reçue.",
        graphique={"titre": fig("Retraités précomptés par la CGRAE, par mois"), "type": "lignes", "categories": [c["libelle"] for c in CG],
                   "series": [{"nom": "Retraités", "valeurs": [c["retraites"] for c in CG], "couleur": BLEU}], "format": "0", "etiquettes": False, "legende": None,
                   "unite": "En nombre de retraités"},
        analyse=[f"**{N('Retraités août 2026', CG[-1]['retraites'])} retraités sont précomptés en août 2026.** Leur nombre augmente chaque mois depuis février 2024.",
                 "**Plusieurs anomalies de gestion apparaissent dans les fichiers** : précomptes en double ou indus remboursés, numéros de décision erronés, précomptes non effectués.",
                 f"**La CGRAE précompte {pc('Recouvrement CGRAE (r)', Q['effectifs']['recouvrement_cgrae'])} des montants demandés** sur {N('Mois CGRAE (r)', Q['effectifs']['recouvrement_cgrae_mois'])} mois communs."],
        implication="Si la garantie santé est maintenue à vie aux retraités pour une cotisation inférieure à son coût, elle devient un engagement viager : sa valeur se calculera avec la table CIMA F (partie 11) dès que les âges et les coûts de soins seront transmis.",
        source=SRC_PREC)
contenu({"titre": "Produits annexes", "sous_titre": "Deux produits sont précomptés sans documentation transmise",
         "synthese": "Flotte automobile (9W3) et MEDOC (8Z7) : nature, garanties et comptabilisation à préciser ; projets immobiliers suivis en compte d'attente"},
        "Ces produits ne relèvent pas directement des garanties de base, mais ils transitent par le FPPN et engagent sa responsabilité vis-à-vis des adhérents. Aucun compte n'est intitulé « SCI » dans les balances ; les projets immobiliers figurent en comptes 4711120x.",
        tableau={"titre": tab("Produits annexes et points à clarifier"), "entetes": ["Produit", "Ce que montrent les données", "À obtenir du FPPN"], "largeurs": [4.8, 13.2, 13.2],
                 "lignes": [["Flotte automobile (9W3)", "La Solde précompte environ 10 000 FCFA par mois et par adhérent de 2023 à 2026 ; le compte 47111106 « GNA ASSURANCE » suit l'assureur.", "La convention avec l'assureur, les primes reversées, les sinistres et le rôle exact du FPPN."],
                            ["MEDOC (8Z7)", "Le FPPN précompte 5 000 FCFA par mois depuis novembre 2025 ; les retours de février et avril 2026 sont vides.", "La nature du produit, les garanties, les bénéficiaires et la comptabilisation."],
                            ["Projets immobiliers", "Les comptes 47111200 et 47111201 portent des projets immobiliers ; aucun compte n'est intitulé SCI.", "Le statut juridique des projets et le lien éventuel avec la SCI « La Prévoyance »."],
                            ["Identifiants « 183FPPN »", "Les fichiers aller PEP et Grenier de 2025 et 2026 contiennent 34 identifiants sans mécano.", "L'identité et le mécano de ces personnes."]]},
        disposition="dessous",
        analyse=["**Ces activités engagent le FPPN sans que leurs règles soient documentées.** Le FPPN encaisse des précomptes pour le compte de tiers ou pour des produits dont la nature n'est pas écrite."],
        implication="Si le FPPN porte lui-même le risque de la flotte automobile, il exerce une activité d'assurance de dommages qui sort du champ des mutuelles sociales : la convention doit le préciser.",
        source=SRC_PREC + " ; " + SRC_COMPTA[9:])

# ---------------------------------------------------------------- PARTIE 9 : BILAN
separateur(9, "Bilan et placements", "Provisions, adossement par échéance, fonds propres, rendement")
methode_gab({"titre": "Méthode : lecture du bilan", "sous_titre": "Le bilan est reconstitué compte par compte à partir des balances",
             "synthese": "Provisions, épargne, actifs et fonds propres sont calculés à chaque clôture, de 2020 au 31/08/2026"},
            [("Provisions techniques", "Adenka lit les comptes 19011000 et 19012000 jusqu'en 2023, puis 15801000 et 15802000 à partir de 2024 (prestations à payer)."),
             ("Fonds propres", "Les fonds propres additionnent les comptes 10, 11 et 12 et le résultat de l'exercice, à chaque clôture."),
             ("Adossement par échéance", "Adenka classe l'actif selon son délai de mobilisation et le passif selon son échéance, puis compare les deux, tranche par tranche."),
             ("Limite", "Le bilan n'a pas été transmis, et le 31/08/2026 est provisoire : il ne comprend ni amortissements ni écritures de fin d'année.")],
            "Le bilan n'a pas été transmis : il est reconstitué à partir des soldes des balances. Fonds propres = réserves + report à nouveau + pertes reportées + résultat de l'exercice.")
pvx = EX["provisions"]
contenu({"titre": "Calcul pas à pas : provisions", "sous_titre": "Les deux comptes de provisions portent exactement le même montant sept ans de suite",
         "synthese": f"Maladie {F('Prov mal', b[2020]['provision_maladie'])} et obsèques {F('Prov obs', b[2020]['provision_obseques'])} à chaque clôture"},
        "Lecture : (1) même montant au franc près de 2020 à 2026 ; (2) seul le numéro de compte change en 2024 (reclassement du compte 19 vers le compte 158) ; (3) une estimation recalculée chaque année varierait avec les dossiers en attente. La provision n'est donc pas une estimation actuelle.",
        pas_a_pas={"entetes": ["Clôture", "Compte maladie", "Provision maladie", "Compte obsèques", "Provision obsèques"],
                   "lignes": [[a if a != "2026" else "31/08/2026", pvx[a][0][0], nb(-float(pvx[a][0][2])) + " FCFA", pvx[a][1][0], nb(-float(pvx[a][1][2])) + " FCFA"] for a in sorted(pvx)],
                   "largeurs": [3.6, 4.0, 6.8, 4.0, 6.8], "rouge": [(0, 2), (6, 2), (4, 1)],
                   "explications": ["La provision maladie vaut le même montant en 2020…", "… et au 31/08/2026, au franc près.", "Seul le numéro de compte change en 2024, du compte 19 vers le compte 158."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2020 à 2026 (balances des comptes)"},
        analyse=["**La provision n'est pas recalculée.** Un montant identique au franc près pendant sept ans ne peut pas résulter d'une évaluation des dossiers en cours."],
        implication="Le règlement d'exécution n° 004/2022 impose une provision pour prestations à payer calculée exercice par exercice, dossier par dossier, augmentée des sinistres non déclarés (art. 10). Le compte 158 est par ailleurs réservé aux provisions réglementées à caractère fiscal.",
        source=SRC_COMPTA)
pm_ = C["provisions_mois"]
contenu({"titre": "Provisions techniques", "sous_titre": "La provision ne couvre plus que 21 mois de prestations, contre 115 en 2021",
         "synthese": f"Provisions / prestations mensuelles : {N('Mois 2021', pm_[1]['mois_couverts'], 2)} mois en 2021, {N('Mois 2025', pm_[5]['mois_couverts'], 2)} en 2025, {N('Mois 2026', pm_[6]['mois_couverts'], 2)} en 2026"},
        "Le nombre de mois couverts n'est pas une norme : c'est un indicateur de cohérence. Une provision figée perd son sens quand les prestations changent : elle était très supérieure aux besoins en 2021 et diminue en proportion depuis. Seul un calcul fondé sur les dossiers dira si elle est suffisante.",
        graphique={"titre": fig("Mois de prestations couverts par les provisions"), "type": "barres", "categories": cats_ex,
                   "series": [{"nom": "Mois couverts", "valeurs": [round(p["mois_couverts"], 2) for p in pm_], "couleur": BLEU}], "format": "0.00", "legende": None,
                   "unite": "En nombre de mois de prestations"},
        analyse=[f"**La provision vaut {F('Prov tot (p)', b[2026]['provisions_techniques'])} chaque année.**",
                 f"**Les prestations mensuelles passent de {F('Prest mens 2021', pm_[1]['prestations_mensuelles'])} en 2021 à {F('Prest mens 2026', pm_[6]['prestations_mensuelles'])} en 2026.** La provision représente donc de moins en moins de mois de prestations.",
                 "**Le calcul de la provision (PSAP, Chain Ladder) suppose les bases de prestations demandées** (partie 12)."],
        implication="Si les délais de paiement s'allongent pendant que les prestations augmentent, la provision deviendra insuffisante : son recalcul est la première tâche dès réception des bases de prestations.",
        source=SRC_COMPTA)
contenu({"titre": "Structure de l'actif", "sous_titre": "L'actif repose sur des créances et des investissements peu liquides",
         "synthese": f"Au 31/08/2026 : CMP {F('CMP (a)', b[2026]['cmp_investissements'])}, terrains {F('Terrains', b[2026]['terrains'])}, trésorerie {F('Trésorerie', b[2026]['tresorerie_totale'])}"},
        "Les montants « CMP » regroupent les comptes 47111101 (« Polyclinique du FPPN »), 47111104 (« Investissements au CMP ») et 47111105 (« Charges de fonctionnement au CMP ») ; la signification du sigle CMP reste à confirmer par le FPPN. La créance sur l'Hôpital de la Police (47111100) est inchangée depuis 2020.",
        graphique={"titre": fig("Principaux postes de l'actif"), "type": "barres_empilees", "categories": yrs_bil,
                   "series": [{"nom": "Investissements CMP", "valeurs": [round(b[a]["cmp_investissements"]) for a in range(2020, 2027)], "couleur": BLEU},
                              {"nom": "Terrains", "valeurs": [round(b[a]["terrains"]) for a in range(2020, 2027)], "couleur": GRIS_F},
                              {"nom": "Créance Hôpital de la Police", "valeurs": [round(b[a]["creance_hopital_police"]) for a in range(2020, 2027)], "couleur": GRIS_C},
                              {"nom": "Financement du Grenier", "valeurs": [round(b[a]["grenier_financement"]) for a in range(2020, 2027)], "couleur": ORANGE},
                              {"nom": "Trésorerie", "valeurs": [round(b[a]["tresorerie_totale"]) for a in range(2020, 2027)], "couleur": NOIR}],
                   "format": FMT_M, "etiquettes": False, "legende": "droite", "unite": "En millions de FCFA"},
        analyse=[f"**Les montants « CMP » passent de {F('CMP 2020', b[2020]['cmp_investissements'])} en 2020 à {F('CMP 2026', b[2026]['cmp_investissements'])} au 31/08/2026.**",
                 f"**La créance sur l'Hôpital de la Police ({F('Hôpital', b[2026]['creance_hopital_police'])}) n'a pas bougé depuis 2020.** Aucun remboursement n'apparaît.",
                 f"**Les dettes envers les cliniques atteignent {F('Fournisseurs cliniques', b[2026]['fournisseurs_cliniques'])} au 31/08/2026.**"],
        implication="Le FPPN finance aussi des charges de fonctionnement « au CMP ». Le règlement d'exécution n° 004/2022 interdit de couvrir les pertes d'activités sanitaires par les cotisations ou les réserves obligatoires (art. 27) : une comptabilité séparée de ces activités doit le démontrer.",
        source=SRC_COMPTA)
mo = {m["cloture"]: m for m in MOB}
m26 = mo["31/08/2026"]
contenu({"titre": "Adossement par échéance", "sous_titre": "L'actif mobilisable en moins d'un an couvre la moitié du passif à moins d'un an",
         "synthese": f"Au 31/08/2026 : actif mobilisable en moins d'un an {F('Actif < 1 an 2026', m26['immediat'] + m26['moins_1_an'])}, passif à moins d'un an hors épargne {F('Passif < 1 an 2026', m26['passif_moins_1_an'])}"},
        "Étapes 1 à 3 de la méthode d'adossement du dossier de préparation : classer l'actif par délai de mobilisation, classer le passif par échéance, comparer tranche par tranche. Actif immédiat : caisse et comptes à vue. Moins d'un an : dépôts à terme, fonds commun de placement, créances à court terme. 1 à 5 ans : prêts PEP et financement du Grenier (durée médiane des prêts de 10 mois). Plus de 5 ans ou incertain : immobilisations, terrains, avances au CMP, créance sur l'Hôpital de la Police. Passif à moins d'un an : dettes à court terme hors épargne et provisions pour prestations à payer. L'épargne est présentée à part : son échéance dépend des âges et des conditions de retrait.",
        graphique={"titre": fig("Actif par délai de mobilisation et passif par échéance"), "type": "barres_empilees",
                   "categories": ["Actif 31/12/2025", "Passif 31/12/2025", "Actif 31/08/2026", "Passif 31/08/2026"],
                   "series": [{"nom": "Actif immédiat", "valeurs": [round(mo["31/12/2025"]["immediat"]), None, round(m26["immediat"]), None], "couleur": NOIR},
                              {"nom": "Actif à moins d'un an", "valeurs": [round(mo["31/12/2025"]["moins_1_an"]), None, round(m26["moins_1_an"]), None], "couleur": GRIS_F},
                              {"nom": "Actif à 1 à 5 ans", "valeurs": [round(mo["31/12/2025"]["1_a_5_ans"]), None, round(m26["1_a_5_ans"]), None], "couleur": ORANGE},
                              {"nom": "Actif à plus de 5 ans ou incertain", "valeurs": [round(mo["31/12/2025"]["plus_5_ans_ou_incertain"]), None, round(m26["plus_5_ans_ou_incertain"]), None], "couleur": GRIS_C},
                              {"nom": "Passif à moins d'un an", "valeurs": [None, round(mo["31/12/2025"]["passif_moins_1_an"]), None, round(m26["passif_moins_1_an"])], "couleur": BLEU},
                              {"nom": "Épargne, échéance à établir", "valeurs": [None, round(mo["31/12/2025"]["epargne_echeance_a_etablir"]), None, round(m26["epargne_echeance_a_etablir"])], "couleur": ROUGE}],
                   "format": FMT_M, "etiquettes": False, "legende": "droite", "unite": "En millions de FCFA"},
        analyse=[f"**Plus des trois quarts de l'actif sont à plus de cinq ans ou de délai incertain.** Cette tranche vaut {F('Actif > 5 ans 2026', m26['plus_5_ans_ou_incertain'])} sur {F('Actif total mobilisation 2026', m26['total_actif'])} au 31/08/2026.",
                 f"**L'actif à moins d'un an ne couvre pas le passif à moins d'un an.** L'écart vaut {F('Écart < 1 an 2026', m26['passif_moins_1_an'] - m26['immediat'] - m26['moins_1_an'])} au 31/08/2026, avant même toute sortie d'épargne.",
                 f"**L'épargne ({F('Épargne mobilisation 2026', m26['epargne_echeance_a_etablir'])}) n'a pas d'échéancier connu.** Chaque départ à la retraite déclenche pourtant le versement de l'EFC."],
        implication="Le passif à moins d'un an comprend les provisions pour prestations à payer, dont le montant n'est pas recalculé. Le FPPN finance des engagements courts avec des actifs longs : c'est le risque de liquidité central, à confirmer avec l'échéancier de l'épargne.",
        source=SRC_COMPTA + " ; dossier de préparation de la mission, § 3.1")
fp = EX["fonds_propres_2025"]
contenu({"titre": "Calcul pas à pas : fonds propres 2025", "sous_titre": "Les pertes reportées dépassent les réserves, d'où des fonds propres négatifs",
         "synthese": f"Fonds propres 2025 = {F('FP hors résultat 2025', b[2025]['fonds_propres_hors_resultat'])} + résultat {F('Résultat 2025 (fp)', b[2025]['resultat_exercice'])} = {F('FP 2025 (pas)', b[2025]['fonds_propres'])}"},
        "Les comptes de capitaux propres ont un solde créditeur (positif pour le FPPN), sauf le report à nouveau débiteur (pertes passées), affiché ici en négatif. Libellés exacts de la balance 2025.",
        pas_a_pas={"entetes": ["Compte", "Intitulé (balance 2025)", "Montant"], "lignes": [[c[0], c[1], nb(-float(c[2])) + " FCFA"] for c in fp["comptes"]] + [["", "Sous-total capitaux propres", F("Sous-total FP", b[2025]["fonds_propres_hors_resultat"])], ["", "Résultat de l'exercice 2025", F("Résultat (pas FP)", b[2025]["resultat_exercice"])], ["", "Fonds propres", F("FP (pas total)", b[2025]["fonds_propres"])]],
                   "largeurs": [3.4, 11.4, 6.6], "rouge": [(3, 2), (6, 2)],
                   "explications": ["Les pertes reportées (compte 129) dépassent l'ensemble des réserves.", "Les fonds propres restent négatifs malgré le résultat positif de 2025.", "Les excédents futurs devront absorber ce solde négatif."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2025\\BALANCE FPPN 2025.xlsx"},
        analyse=["**Les réserves existent, mais les pertes passées les dépassent.** Le compte 10620000 « Réserve disponible » est le plus important des comptes positifs, et sa nature reste à confirmer (partie 10)."],
        implication="La marge de solvabilité de la Décision n° 05/2022 se calcule sur ces réserves : selon que les pertes reportées sont déduites ou non, le FPPN est conforme ou non (partie 10).",
        source="Source : balance des comptes 2025")
contenu({"titre": "Fonds propres", "sous_titre": "Les fonds propres restent négatifs malgré six années d'excédents",
         "synthese": f"{F('FP 2020 (f)', b[2020]['fonds_propres'])} fin 2020, {F('FP 2025 (f)', b[2025]['fonds_propres'])} fin 2025, {F('FP 2026 (f)', b[2026]['fonds_propres'])} au 31/08/2026"},
        "Des fonds propres négatifs signifient que les dettes dépassent les actifs comptables. Les ratios de fonds propres et de marge de solvabilité de la Décision n° 05/2022 sont calculés dans la partie 10.",
        graphique={"titre": fig("Fonds propres et résultat de l'exercice"), "type": "barres", "categories": yrs_bil,
                   "series": [{"nom": "Fonds propres", "valeurs": [round(b[a]["fonds_propres"]) for a in range(2020, 2027)], "couleur": BLEU},
                              {"nom": "Résultat de l'exercice", "valeurs": [round(b[a]["resultat_exercice"]) for a in range(2020, 2027)], "couleur": ORANGE}],
                   "format": FMT_M, "unite": "En millions de FCFA"},
        analyse=["**Les excédents de 2020 à 2025 réduisent l'écart sans le combler.** Les fonds propres passent de " + F("FP 2020 (t)", b[2020]["fonds_propres"]) + " à " + F("FP 2025 (t)", b[2025]["fonds_propres"]) + ".",
                 "**Le déficit des huit premiers mois de 2026 inverse la tendance.** Les fonds propres reculent à nouveau au 31/08/2026.",
                 f"**Le ratio des fonds propres et assimilés vaut {val_ratio(RT['FEP'], 2025)} en 2025, pour une norme d'au moins 1.**"],
        implication="Au rythme des excédents de 2024 et 2025, il faudrait encore au moins un exercice excédentaire pour revenir à des fonds propres positifs. Le déficit de 2026 éloigne cette échéance : le redressement dépend de la maîtrise de la sinistralité et des frais.",
        source=SRC_COMPTA)
rp = R["rendement_placements"]
contenu({"titre": "Rendement des placements", "sous_titre": "Le problème des placements est leur faible volume, pas leur rendement",
         "synthese": f"Rendement moyen apparent 2021 à 2025 : {pc('Rendement moyen', sum(r['produits_financiers'] for r in rp) / sum(r['placements_moyens'] for r in rp))} par an ; taux directeur BCEAO entre 2,00 % et 3,50 %"},
        "Rendement apparent = produits financiers (classe 77) / encours moyen des placements (52 à 53). Il est irrégulier : l'essentiel des intérêts est comptabilisé en 2025 à la clôture d'un dépôt à terme. Taux directeurs : communiqués du Comité de politique monétaire de la BCEAO (validés).",
        graphique={"titre": fig("Taux directeur de la BCEAO"), "type": "lignes", "categories": [t[0] for t in R["externe"]["taux_directeur"]],
                   "series": [{"nom": "Taux directeur", "valeurs": [t[1] for t in R["externe"]["taux_directeur"]], "couleur": BLEU}], "format": "0.00", "legende": None,
                   "unite": "En %, date d'effet"},
        analyse=[f"**Les produits financiers de 2021 à 2025 atteignent {F('Produits fin', sum(r['produits_financiers'] for r in rp))}**, pour un encours moyen de placements d'environ {F('Placements moyens', sum(r['placements_moyens'] for r in rp) / len(rp))}.",
                 "**Le rendement moyen dépasse le taux directeur, mais il est irrégulier.** L'essentiel des intérêts est comptabilisé en 2025, à la clôture d'un dépôt à terme.",
                 f"**Le vrai sujet est le volume placé**, face à {F('Épargne (rend)', b[2026]['epargne_totale'])} d'épargne à restituer. Le solde des comptes courants non rémunérés atteint {F('Comptes courants 2026', AG[2026]['comptes_courants_521'])} au 31/08/2026."],
        implication="La Décision n° 05/2022 demande aux mutuelles qui gèrent des risques longs un solde des comptes courants non rémunérés proche de zéro : une partie de cette trésorerie peut être placée à terme sans nuire à la liquidité.",
        source="Source : BCEAO, communiqués du Comité de politique monétaire ; " + SRC_COMPTA[9:])
retenir("bilan", ("Le bilan combine provisions figées, fonds propres négatifs et actifs peu liquides", "Trois fragilités cumulées limitent la capacité du FPPN à faire face à un choc",
                  "Le FPPN porte des engagements courts et une épargne croissante avec des actifs longs, des provisions non recalculées et des fonds propres négatifs."),
        "Les provisions gardent le même montant de 2020 à 2026, les fonds propres sont négatifs à chaque clôture, et plus des trois quarts de l'actif sont à plus de cinq ans ou de délai incertain.",
        "Le FPPN ne peut absorber ni un choc de prestations, ni une vague de retraits, sans céder des actifs peu liquides.",
        "Le FPPN doit adopter une politique de provisionnement, une politique de placement et un encadrement écrit des investissements dans les structures liées.",
        "Ces constats appellent une décision de gouvernance : politique de provisionnement, politique de placement, encadrement des investissements dans les structures liées.")

# ---------------------------------------------------------------- PARTIE 10 : CONFORMITÉ PRUDENTIELLE
separateur(10, "Conformité prudentielle", "Textes UEMOA, ratios, réserves, plan comptable",
           "Cette partie applique aux comptes du FPPN les normes de la mutualité sociale UEMOA. Les textes viennent de la documentation publique de l'AIRMS fournie le 18/10/2026 et validée comme source.")
contenu({"titre": "Textes applicables", "sous_titre": "Cinq textes fixent les règles prudentielles que le FPPN doit respecter",
         "synthese": "Les articles cités sont lus dans les textes intégraux ; les pages scannées ont été lues par reconnaissance optique de caractères (OCR) et vérifiées à l'image"},
        "Hiérarchie : le règlement de 2009 fixe les principes et renvoie au règlement d'exécution pour les règles prudentielles (art. 58) ; la Décision de 2022 chiffre les normes ; l'arrêté de 2023 fixe les sanctions en Côte d'Ivoire. Le PCMS donne les numéros de comptes que la Décision utilise dans ses formules.",
        tableau={"titre": tab("Textes de la mutualité sociale appliqués au FPPN"), "entetes": ["Texte", "Ce qu'il impose au FPPN", "Articles utilisés"], "largeurs": [8.6, 15.8, 6.8],
                 "lignes": [["Règlement n° 07/2009/CM/UEMOA du 26/06/2009", "Il impose l'agrément, une comptabilité selon le plan comptable, la répartition des excédents en réserves et le placement des fonds en banque.", "Art. 22, 28, 60 à 63, 83, 84"],
                            ["Règlement d'exécution n° 004/2022/COM/UEMOA du 01/06/2022", "Il impose des provisions suffisantes, une réserve légale, des règles propres aux risques longs et le cantonnement des patrimoines.", "Art. 3 à 27, 35"],
                            ["Décision n° 05/2022/COM/UEMOA du 01/06/2022", "Elle fixe 13 indicateurs de gouvernance, 11 indicateurs de contrôle, 9 ratios pour les risques longs, la marge de solvabilité et l'affectation des excédents.", "Annexe, tableaux 1 et 2, § 1.1 à 2.9"],
                            ["Règlement n° 03/2012/CM/UEMOA et PCMS de mai 2012", "Ils imposent le plan comptable des mutuelles sociales à partir de l'exercice ouvert le 1er janvier 2014.", "Art. 2 ; liste des comptes ; fiches 1, 2, 6, 25"],
                            ["Arrêté n° 00115/MEPS/CAB du 12/12/2023", "Il fixe les sanctions administratives et pécuniaires en cas de manquement.", "Annexe, lignes 6 et 13"]]},
        disposition="dessous",
        analyse=["**Ces textes forment un ensemble complet et chiffré depuis juin 2022.** Ils couvrent la comptabilité, les provisions, les réserves, la liquidité et les sanctions."],
        implication="Le FPPN dépasse les seuils du système normal du plan comptable (plus de 10 millions de FCFA de cotisations ou plus de 1 000 adhérents, art. 4 du règlement d'exécution) : il doit tenir une comptabilité d'engagement et produire un TAFIRE (art. 5).",
        source="Source : textes cités, documentation publique de l'AIRMS fournie le 18/10/2026 (validée)")
contenu({"titre": "Risques courts et risques longs", "sous_titre": "Le FPPN couvre des risques courts et, avec l'épargne, des risques longs",
         "synthese": "Un risque est long quand plus d'un an sépare la cotisation de la prestation (règlement d'exécution n° 004/2022, art. 1)"},
        "Définitions de l'article 1 du règlement d'exécution : risque court quand l'écart entre l'encaissement des cotisations et le versement des prestations est d'un an au plus ; risque long au-delà. L'article 3 classe la vieillesse et le décès lié à un mécanisme de capitalisation parmi les risques longs. La qualification retenue ici doit être confirmée par l'AIRMS ou par l'agrément du FPPN, non transmis.",
        tableau={"titre": tab("Qualification des garanties du FPPN"), "entetes": ["Garantie", "Délai entre cotisation et prestation", "Qualification proposée", "Conséquence"], "largeurs": [4.4, 8.4, 6.6, 11.8],
                 "lignes": [["Maladie (936, CGRAE)", "Les soins sont payés dans l'année.", "Risque court", "Ratios du tableau 1 de la Décision n° 05/2022."],
                            ["Décès (8Z6)", "Le capital est payé au décès, sans capitalisation connue.", "Risque court", "Ratios du tableau 1 ; qualification à revoir si un mécanisme de capitalisation existe."],
                            ["EFC (870)", "L'épargne est versée en fin de carrière, des années après les versements.", "Risque long (vieillesse)", "Ratios de la section II, réserve de trésorerie, étude actuarielle, certification."],
                            ["PEP (998)", "L'épargne est restituée à l'échéance du plan, en général après plus d'un an.", "Risque long, à confirmer", "Mêmes obligations que l'EFC si la durée du plan dépasse un an."]]},
        disposition="dessous",
        analyse=["**Avec l'EFC et le PEP, le FPPN entre dans le champ des risques longs.** Ces deux produits représentent " + F("Épargne risques longs", b[2026]["epargne_totale"]) + " d'engagements au 31/08/2026."],
        implication="Une mutuelle qui couvre à la fois des risques courts et longs doit cantonner ses patrimoines et séparer ses opérations (art. 15), faire certifier ses comptes par un commissaire aux comptes (art. 12) et réaliser une étude actuarielle tous les cinq ans (art. 14).",
        source="Source : " + REX + ", art. 1, 3, 12 à 15 ; fichiers de précomptes ; balances du FPPN")
gab(31, {"titre": "Lire un ratio : les tiroirs du bilan", "sous_titre": "Chaque compte est rangé dans un tiroir selon le temps que met l'argent à entrer ou sortir",
         "synthese": "Un ratio compare le contenu de deux tiroirs ; le classement des comptes décide donc du résultat"},
    {"cartes": [("Ressources longues", "Argent qui reste longtemps à la disposition du FPPN : fonds propres, provisions, emprunts."),
                ("Dettes à court terme", "Argent que le FPPN doit rendre dans l'année : fournisseurs, cliniques, État."),
                ("Emplois longs", "Argent placé pour longtemps : terrains, bâtiments, prêts, avances aux structures liées."),
                ("Trésorerie", "Argent disponible tout de suite ou presque : caisse, comptes en banque, dépôts à terme.")]},
    "Image utilisée pour un lecteur non spécialiste : le bilan est une commode à tiroirs. Le plan comptable (PCMS) donne à chaque compte un numéro qui indique son tiroir : classe 1 pour les ressources longues, classe 2 pour les emplois longs, classe 4 pour les tiers, classe 5 pour la trésorerie. Un ratio de liquidité compare le tiroir « trésorerie » au tiroir « dettes à court terme ». Si un compte est rangé dans le mauvais tiroir, le ratio change sans que la réalité change. C'est pourquoi la suite présente chaque ratio deux fois : avec les comptes rangés comme le FPPN les tient, puis rangés selon leur nature.",
    source="Source : PCMS (structure des classes de comptes) ; Adenka")
contenu({"titre": "Deux lectures : comme le FPPN, puis retraitée", "sous_titre": "Trois groupes de comptes sont rangés dans un tiroir qui ne reflète pas leur nature",
         "synthese": "Adenka calcule d'abord chaque ratio avec les comptes tels qu'ils sont tenus, puis avec trois reclassements motivés"},
        "Le classement « comme le FPPN » suit les numéros de comptes cités par la Décision n° 05/2022 dans ses formules. Le classement retraité ne modifie aucun montant : il déplace seulement des comptes d'un tiroir à un autre, pour une raison écrite. Le détail compte par compte (246 comptes) figure dans l'onglet « 10_classement_comptes » du classeur des calculs.",
        tableau={"titre": tab("Reclassements de la lecture retraitée"), "entetes": ["Comptes", "Tiroir comme le FPPN", "Tiroir retraité", "Raison du reclassement"], "largeurs": [6.6, 5.4, 5.4, 13.8],
                 "lignes": [["41195100, 41195200 : épargne PEP et EFC", "Dette à court terme (classe 4)", "Ressource longue", "L'épargne est un risque long (art. 3 du règlement d'exécution) : elle n'est pas exigible en totalité dans l'année."],
                            ["47111100 à 47111105, 47111201, 485 : avances et créances durables", "Créance à court terme (classe 4)", "Emploi long", "Ces montants financent le CMP, le Grenier et l'Hôpital de la Police depuis plusieurs années, sans remboursement à court terme."],
                            ["10620000 : réserve disponible", "Réserve libre (libellé du FPPN)", "Nature à confirmer", "Le PCMS réserve le compte 106 aux écarts de réévaluation des immeubles, qui ne comptent pas dans la marge de solvabilité."]]},
        disposition="dessous",
        analyse=["**Les deux lectures encadrent la réalité.** La lecture comme le FPPN est prudente pour l'épargne, mais optimiste pour les avances durables ; la lecture retraitée corrige les deux."],
        implication="La réponse définitive dépend de l'échéancier réel de l'épargne : les conditions de retrait et les âges des épargnants trancheront entre les deux lectures.",
        source="Source : PCMS, fiches n° 1 et 25 ; " + REX + ", art. 3 ; balances du FPPN ; classement Adenka")
methode_gab({"titre": "Règle de statut", "sous_titre": "Le statut de chaque ratio suit une règle écrite et vérifiable",
             "synthese": "Le statut compare la valeur calculée à la norme du texte ; il est jugé sur les exercices 2022 à 2025"},
            [("Conforme ou non conforme", "Un ratio est conforme quand sa valeur respecte la norme chiffrée de la Décision n° 05/2022 ; il est non conforme sinon. Pour une fourchette, une valeur au-dessus de la borne haute est signalée à part."),
             ("En attente", "Un ratio est en attente quand une donnée nécessaire n'a pas été transmise : fichiers aller des postes 936 et 8Z6, procès-verbaux, registre des contrôles, cotisation au FNG."),
             ("À apprécier", "Un ratio est à apprécier quand la norme n'est pas chiffrée : délai contractuel, objectif du plan stratégique, solde « avoisinant zéro », taux moyen du marché."),
             ("Période jugée", "La Décision est entrée en vigueur le 1er juin 2022 : 2020 et 2021 sont indicatifs, comme les huit mois provisoires de 2026. Deux exercices non conformes de suite sont signalés, car l'article 84 du règlement n° 07/2009 prévoit alors le retrait d'agrément.")],
            "Le statut ne juge pas la gestion du FPPN : il constate l'écart entre un chiffre et une norme. Le statut global d'un ratio est non conforme s'il est non conforme au moins une année de 2022 à 2025.",
            source="Source : " + DEC + ", article 3 ; " + R07 + ", art. 84 ; Adenka")
gab(29, {"titre": "Quatre familles de ratios", "sous_titre": "Les 36 indicateurs examinés se regroupent en quatre familles, de la plus actuarielle à la gouvernance",
         "synthese": "L'ordre de priorité suit l'importance pour la protection des adhérents : solvabilité d'abord, gouvernance ensuite"},
    {"blocs": [("Sinistralité et solvabilité", "Sinistralité, marge de solvabilité, autonomie financière, excédent net : ces ratios disent si les cotisations suffisent et si le FPPN peut absorber un choc."),
               ("Réserves et fonds propres", "Réserve obligatoire, fonds propres et assimilés, fonds des risques longs : ces ratios disent si le FPPN a mis de côté assez de ressources propres."),
               ("Liquidité et équilibre", "Liquidité, trésorerie à échéance et immédiate, équilibre stable, solvabilité de l'épargne : ces ratios disent si l'argent est disponible à temps."),
               ("Gestion et gouvernance", "Frais, personnel, recouvrement, délais, contrôles, réunions des organes, FNG : ces ratios disent si l'organisation est efficace et contrôlée.")]},
    "Les familles reprennent les catégories de l'annexe de la Décision n° 05/2022 (performance, viabilité, pérennité, risques longs), réordonnées selon la priorité actuarielle demandée pour ce rapport.",
    source="Source : " + DEC + " ; classement Adenka")


def ligne_bord(x, annees=(2022, 2023, 2024, 2025, 2026)):
    vals = [val_ratio(x, a) if a in x["valeurs"] else "n.d." for a in annees]
    nom = x["nom"].replace(" (risques longs)", "")
    return [str(x["priorite"]), nom, x["norme"].replace(" (tableau 1) ; 9 mois au § 1.2.2.5.4", "").replace(" (marge disponible au moins égale à la marge exigée)", ""), *vals, x["statut_global"]]


bord_ent = ["N°", "Ratio", "Norme", "2022", "2023", "2024", "2025", "2026 (8 mois)", "Statut"]
bord_larg = [0.9, 8.4, 4.8, 2.7, 2.7, 2.7, 2.7, 2.9, 3.4]
contenu({"titre": "Tableau de bord (1/4) : solvabilité et réserves", "sous_titre": "La sinistralité globale, les réserves, les fonds propres et la liquidité sont non conformes",
         "synthese": "Ratios de priorité actuarielle 1 à 9, lecture comme le FPPN ; 2026 est indicatif"},
        "Valeurs calculées sur les balances, avec les comptes cités par la Décision n° 05/2022. La sinistralité réglementaire est globale (maladie et décès) : elle est sous la fourchette de 70 % à 80 % parce que la garantie décès consomme peu de cotisations. La marge de solvabilité est conforme dans la lecture comme le FPPN, mais pas une fois les pertes reportées déduites (diapositives suivantes).",
        tableau={"titre": tab("Tableau de bord des ratios prudentiels, priorités 1 à 9"), "entetes": bord_ent, "largeurs": bord_larg, "statut_col": 8, "taille": 9,
                 "lignes": [ligne_bord(RT[c]) for c in ("SIN", "MSE", "ROB", "FEP", "LIQ", "EQS", "AUT", "EXN", "TFP")]},
        disposition="dessous",
        analyse=["**Six des neuf ratios prioritaires sont non conformes plusieurs années de suite.** Seuls l'autonomie financière, le financement propre et la marge de solvabilité calculée comme le FPPN respectent la norme."],
        implication="La non-conformité répétée de la réserve obligatoire, des fonds propres et de la liquidité expose le FPPN aux sanctions prévues pour le non-respect des règles prudentielles sur deux exercices.",
        source=SRC_REG)
contenu({"titre": "Tableau de bord (2/4) : risques longs", "sous_titre": "Aucun ratio propre aux risques longs n'est respecté",
         "synthese": "Ratios de priorité 10 à 16, applicables à l'épargne EFC et PEP qualifiée de risque long"},
        "Section II de l'annexe de la Décision n° 05/2022 et articles 19 et 21 du règlement d'exécution. Le ratio de solvabilité de l'épargne reprend la couverture calculée en partie 6. La réserve de trésorerie vaut zéro car aucun compte de réserve de trésorerie n'existe dans les balances.",
        tableau={"titre": tab("Tableau de bord des ratios des risques longs, priorités 10 à 16"), "entetes": bord_ent, "largeurs": bord_larg, "statut_col": 8, "taille": 9,
                 "lignes": [ligne_bord(RT[c]) for c in ("RL1", "RL2", "RL3", "RL4", "RL5", "RL6", "RL7")]},
        disposition="dessous",
        analyse=["**Les engagements d'épargne ne sont adossés ni à des actifs réservés suffisants, ni à une réserve de trésorerie, ni à des fonds propres positifs.** Les deux derniers indicateurs n'ont pas de seuil chiffré et sont à apprécier."],
        implication="Pour les risques longs, le règlement d'exécution exige aussi une étude actuarielle tous les cinq ans et la certification des comptes : le FPPN doit indiquer s'il les a réalisées.",
        source=SRC_REG)
contenu({"titre": "Tableau de bord (3/4) : gestion", "sous_titre": "Les frais dépassent nettement les normes, alors que le personnel et la fidélisation sont conformes",
         "synthese": "Ratios de gestion de priorité 17 à 25 ; le recouvrement n'est mesurable que pour les retraités"},
        "Les charges de fonctionnement suivent la liste de comptes de l'annexe (612, 613, 618, 62 à 68). La fidélisation compare les matricules précomptés au poste 936 d'une année à l'autre. Le délai de paiement des prestataires est à apprécier, car les délais contractuels des conventions n'ont pas été transmis.",
        tableau={"titre": tab("Tableau de bord des ratios de gestion, priorités 17 à 25"), "entetes": bord_ent, "largeurs": bord_larg, "statut_col": 8, "taille": 9,
                 "lignes": [ligne_bord(RT[c]) for c in ("CHF", "CHB", "CUP", "CHP", "FID", "DPP", "FOR", "CRB")] +
                           [[str(RT["REC"]["priorite"]), RT["REC"]["nom"], RT["REC"]["norme"], "", "", "", pc("REC CGRAE", Q["effectifs"]["recouvrement_cgrae"]), "", RT["REC"]["statut_global"] + " (retraités)"]]},
        disposition="dessous",
        analyse=["**Les trois ratios de frais sont non conformes chaque année.** Les frais de fonctionnement absorbent plus d'un tiers des cotisations, contre 20 % au plus selon la norme."],
        implication="La baisse du coût par unité prestée vient surtout de la hausse des prestations, pas d'une baisse des frais : le FPPN doit présenter un plan de réduction des charges de fonctionnement.",
        source=SRC_REG + " ; fichiers de précomptes")
contenu({"titre": "Tableau de bord (4/4) : données en attente", "sous_titre": "Onze indicateurs attendent des documents de gestion ou de gouvernance",
         "synthese": "Ces indicateurs ne sont pas exclus : ils seront calculés dès réception des pièces demandées"},
        "Les indicateurs de gouvernance sont placés en dernier, conformément à la priorité actuarielle du rapport. Chaque pièce manquante est ajoutée à la liste des documents demandés (PBC) et aux guides d'entretien.",
        tableau={"titre": tab("Indicateurs en attente de documents, priorités 26 à 36"), "entetes": ["N°", "Indicateur", "Norme", "Pièce à obtenir"], "largeurs": [0.9, 11.0, 6.0, 13.3], "taille": 9,
                 "lignes": [[str(RT[c]["priorite"]), RT[c]["nom"], RT[c]["norme"], RT[c]["qualitatif"]] for c in ("RBA", "PEN", "IND", "CAP", "DLQ", "RAT", "DRB", "FNG", "TTR", "OSA", "OSF")]},
        disposition="dessous",
        analyse=["**L'absence de compte du fonds national de garantie dans les balances est le point le plus sensible.** L'adhésion au fonds est obligatoire (règlement n° 07/2009, art. 63), et son défaut est sanctionné (arrêté n° 00115/MEPS/CAB, ligne 8)."],
        source=SRC_REG + " ; liste des documents demandés (PBC)")
contenu({"titre": "Liquidité et équilibre : deux lectures", "sous_titre": "Le classement de l'épargne décide de la conformité de la liquidité et de l'équilibre stable",
         "synthese": f"Liquidité 2025 : {val_ratio(RT['LIQ'], 2025)} comme le FPPN, {val_ratio(RT['LIQ'], 2025, True)} retraitée ; équilibre stable 2025 : {val_ratio(RT['EQS'], 2025)} et {val_ratio(RT['EQS'], 2025, True)}"},
        "Liquidité : norme d'au moins 1. Équilibre stable : norme d'au plus 1. Lecture retraitée : épargne en ressource longue, avances durables en emplois longs (diapositive « Deux lectures »). En 2020, les ressources longues comme le FPPN sont négatives : le ratio d'équilibre stable n'a pas de sens et n'est pas représenté.",
        graphique={"titre": fig("Liquidité et équilibre stable, comme le FPPN et retraités"), "type": "lignes", "categories": [str(a) for a in range(2021, 2027)],
                   "series": [{"nom": "Liquidité comme le FPPN", "valeurs": [round(RT["LIQ"]["valeurs"][a], 2) for a in range(2021, 2027)], "couleur": BLEU},
                              {"nom": "Liquidité retraitée", "valeurs": [round(RT["LIQ"]["valeurs_retraitees"][a], 2) for a in range(2021, 2027)], "couleur": ORANGE},
                              {"nom": "Équilibre stable retraité", "valeurs": [round(RT["EQS"]["valeurs_retraitees"][a], 2) for a in range(2021, 2027)], "couleur": GRIS_F},
                              {"nom": "Norme (1)", "valeurs": [1.0] * 6, "couleur": ROUGE}],
                   "format": "0.00", "unite": "Ratio (sans unité)"},
        analyse=[f"**Comme le FPPN tient ses comptes, la liquidité n'atteint jamais 1.** Elle vaut au mieux {val_ratio(RT['LIQ'], 2021)} en 2021, car l'épargne figure parmi les dettes à court terme.",
                 "**Retraitée, la liquidité devient conforme en 2025 et sur les huit mois de 2026.** La trésorerie a fortement augmenté pendant que les dettes à court terme hors épargne restaient stables.",
                 f"**Retraité, l'équilibre stable devient conforme en 2025** ({val_ratio(RT['EQS'], 2025, True)}), car l'épargne finance alors durablement les emplois longs."],
        implication="La conformité dépend de la part réellement exigible de l'épargne dans l'année : le FPPN doit établir l'échéancier de l'EFC (départs à la retraite) et du PEP (échéances des plans).",
        source=SRC_REG)
a25 = AG[2025]
contenu({"titre": "Calcul pas à pas : sinistralité réglementaire 2025", "sous_titre": "La sinistralité globale est sous la fourchette parce que le décès consomme peu",
         "synthese": f"Charges de prestations {F('SIN num 2025', a25['charges_prestations'])} / cotisations acquises {F('SIN den 2025', a25['cotisations_acquises'])} = {pc('SIN 2025 (pas)', a25['charges_prestations'] / a25['cotisations_acquises'])}"},
        "Formule de la Décision (§ 1.2.1.2.5) : comptes 60 + 6901 - 7901 au numérateur, comptes 70 + 7906 - 6906 au dénominateur. Le FPPN n'a pas de compte 6901, 7901, 6906 ni 7906 : les provisions techniques ne varient pas. Le compte 704 « cartes d'accès aux soins » appartient à la classe 70 et entre donc dans les cotisations acquises.",
        pas_a_pas={"entetes": ["Élément", "Comptes", "Montant 2025"],
                   "lignes": [["Cotisations maladie", "701", F("SIN cot mal", a25["cot_maladie"])],
                              ["Cotisations décès et MEDOC", "702", F("SIN cot dec", a25["cot_deces"])],
                              ["Autres produits techniques", "704", F("SIN cot autres", a25["cotisations_acquises"] - a25["cot_maladie"] - a25["cot_deces"])],
                              ["Cotisations acquises", "70", F("SIN cot acq", a25["cotisations_acquises"])],
                              ["Prestations maladie", "601, 6013, 6018, 60800007", F("SIN prest mal", a25["prest_maladie"])],
                              ["Prestations décès", "6021, 6024", F("SIN prest dec", a25["prest_deces"])],
                              ["Autres prestations", "60800008, 60850000", F("SIN prest autres", a25["charges_prestations"] - a25["prest_maladie"] - a25["prest_deces"])],
                              ["Charges de prestations", "60", F("SIN charges", a25["charges_prestations"])],
                              ["Taux de sinistralité", "60 / 70", pc("SIN taux (t)", a25["charges_prestations"] / a25["cotisations_acquises"])]],
                   "largeurs": [6.2, 6.4, 5.6], "rouge": [(3, 2), (7, 2), (8, 2)],
                   "explications": ["Les cotisations acquises additionnent toute la classe 70, y compris les autres produits techniques.", "Les charges de prestations additionnent toute la classe 60, maladie et décès confondus.", "Le taux global se situe sous la borne basse de 70 % de la norme."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2025\\BALANCE FPPN 2025.xlsx"},
        analyse=["**Une sinistralité trop basse est aussi un écart à la norme.** La fourchette de 70 % à 80 % vise à ce que les cotisations reviennent aux adhérents sous forme de prestations."],
        implication=f"Appliqué à la seule maladie, le taux 2025 ({pc('S/C maladie 2025 (SIN)', s[2025]['sc_maladie'])}) se situe dans la fourchette. L'écart global vient de la garantie décès : son niveau de cotisation mérite d'être revu au vu de la mortalité observée, dès réception des données.",
        source=SRC_COMPTA + " ; " + DEC + ", § 1.2.1.2.5")
m25 = MS[2025]
contenu({"titre": "Calcul pas à pas : marge de solvabilité 2025", "sous_titre": "La marge est conforme si l'on ignore les pertes reportées, non conforme sinon",
         "synthese": f"Marge exigée {F('MSE exigée 2025', m25['marge_exigee'])} ; marge disponible {F('MSE dispo FPPN 2025', m25['marge_disponible_comme_fppn'])} comme le FPPN, {F('MSE dispo nette 2025', m25['marge_disponible_nette'])} nette des pertes"},
        "Formule (Décision, § 1.1.1) : marge exigée = le maximum entre 18 % des cotisations acquises de l'exercice précédent et de 26 % de la moyenne des prestations maladie des trois derniers exercices. Marge disponible = fonds de réserve obligatoire + fonds d'établissement + réserves libres + droits d'adhésion + dons et legs non affectés. Le texte ne mentionne pas les pertes reportées ; la lecture retraitée les déduit, car le PCMS définit les réserves comme des excédents laissés à la disposition de la mutuelle (fiche n° 2).",
        pas_a_pas={"entetes": ["Élément", "Calcul", "Montant"],
                   "lignes": [["18 % des cotisations acquises 2024", "18 % x " + F("MSE cot 2024", AG[2024]["cotisations_acquises"]), F("MSE 18pc", m25["exigee_18pc_cotisations_n_1"])],
                              ["26 % des prestations maladie 2023 à 2025", "26 % x moyenne sur 3 ans", F("MSE 26pc", m25["exigee_26pc_prest_maladie_3ans"])],
                              ["Marge exigée", "le maximum des deux", F("MSE exigée (t)", m25["marge_exigee"])],
                              ["Réserve obligatoire", "compte 111", F("MSE 111", m25["reserve_obligatoire_111"])],
                              ["Réserves libres", "compte 1188", F("MSE 1188", m25["reserves_libres_1188"])],
                              ["Réserve disponible", "compte 106", F("MSE 106", m25["reserve_disponible_106"])],
                              ["Marge disponible comme le FPPN", "111 + 1188 + 106", F("MSE dispo (t)", m25["marge_disponible_comme_fppn"])],
                              ["Pertes reportées et résultat", "129, 121, résultat 2025", F("MSE pertes", m25["marge_disponible_nette"] - m25["marge_disponible_comme_fppn"])],
                              ["Marge disponible nette", "fonds propres", F("MSE nette (t)", m25["marge_disponible_nette"])]],
                   "largeurs": [6.6, 7.0, 5.0], "rouge": [(2, 2), (6, 2), (8, 2)],
                   "explications": ["La marge exigée vient de la règle des 18 % des cotisations, supérieure à celle des prestations.", "Sans déduire les pertes, la marge disponible dépasse largement la marge exigée.", "Une fois les pertes reportées déduites, la marge disponible devient négative."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2024 et 2025 (balances des comptes)"},
        analyse=[f"**La conformité dépend de la lecture des pertes reportées.** Le ratio vaut {val_ratio(RT['MSE'], 2025)} comme le FPPN et {val_ratio(RT['MSE'], 2025, True)} net des pertes."],
        implication=f"Le FPPN n'a pas de fonds d'établissement, alors que la Décision exige un niveau minimal égal au tiers de la marge exigée, soit {F('Fonds établissement min 2025', m25['fonds_etablissement_minimum'])} en 2025.",
        source=SRC_COMPTA + " ; " + DEC + ", § 1.1.1 ; PCMS, fiche n° 2")
besoin6 = a25["charges_prestations"] * 6 / 12
contenu({"titre": "Calcul pas à pas : réserve obligatoire 2025", "sous_titre": "La réserve obligatoire couvre cinq mois de prestations, contre six au minimum",
         "synthese": f"{F('ROB 111 2025', a25['reserves_legales_111'])} x 12 / {F('ROB prest 2025', a25['charges_prestations'])} = {N('ROB 2025 (pas)', RT['ROB']['valeurs'][2025], 2)} mois"},
        "Formule (Décision, tableau 1, ligne 10) : réserves obligatoires x 12 / charges de prestations, norme de 6 à 9 mois ; le § 1.2.2.5.4 de l'annexe indique 9 mois. La réserve obligatoire est le compte 111 « Réserves légales ». Réserve nécessaire pour 6 mois = charges de prestations x 6 / 12.",
        pas_a_pas={"entetes": ["Élément", "Calcul", "Montant"],
                   "lignes": [["Réserve obligatoire", "compte 111", F("ROB 111 (t)", a25["reserves_legales_111"])],
                              ["Charges de prestations", "classe 60", F("ROB prest (t)", a25["charges_prestations"])],
                              ["Ratio de réserve obligatoire", "111 x 12 / 60", N("ROB ratio (t)", RT["ROB"]["valeurs"][2025], 2) + " mois"],
                              ["Réserve nécessaire pour 6 mois", "60 x 6 / 12", F("ROB besoin 6 mois", besoin6)],
                              ["Écart à combler", "besoin - réserve", F("ROB écart", besoin6 - a25["reserves_legales_111"])]],
                   "largeurs": [6.8, 5.2, 6.2], "rouge": [(2, 2), (4, 2)],
                   "explications": ["La réserve obligatoire couvre un peu plus de cinq mois de prestations.", "Il manque ce montant pour atteindre la borne basse de six mois."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2025\\BALANCE FPPN 2025.xlsx"},
        analyse=["**La réserve progresse, mais moins vite que les prestations.** Elle vaut zéro jusqu'en 2022 et couvre " + N("ROB 2023 (a)", RT["ROB"]["valeurs"][2023], 2) + " mois en 2023, puis " + N("ROB 2025 (a)", RT["ROB"]["valeurs"][2025], 2) + " mois en 2025."],
        implication="Comme les prestations augmentent chaque année, le besoin de réserve augmente aussi : la dotation annuelle doit suivre la sinistralité et non seulement le résultat.",
        source=SRC_COMPTA + " ; " + DEC + ", tableau 1 et § 1.2.2.5.4")
contenu({"titre": "Calcul pas à pas : fonds propres et assimilés 2025", "sous_titre": "Les ressources propres et les provisions couvrent un cinquième des dettes",
         "synthese": f"Fonds propres et assimilés {F('FEP num 2025', RT['FEP']['num'][2025])} / dettes {F('FEP den 2025', RT['FEP']['den'][2025])} = {N('FEP 2025 (pas)', RT['FEP']['valeurs'][2025], 2)}"},
        "Formule (Décision, § 1.2.2.5.3) : fonds propres et assimilés (comptes 10 à 14, 190 et 15) / dettes, provisions techniques comprises. Le texte place les provisions techniques à la fois au numérateur et au dénominateur ; Adenka applique la formule à la lettre. Les dettes comprennent toute la classe 4 créditrice, dont l'épargne, les emprunts et les soldes bancaires créditeurs.",
        pas_a_pas={"entetes": ["Élément", "Comptes", "Montant 2025"],
                   "lignes": [["Fonds propres", "10, 11, 12, résultat", F("FEP FP", a25["fonds_propres"])],
                              ["Provisions techniques", "15801000, 15802000", F("FEP prov", a25["provisions_techniques"])],
                              ["Fonds propres et assimilés", "numérateur", F("FEP num (t)", RT["FEP"]["num"][2025])],
                              ["Dettes à court terme", "classe 4 créditrice, 5 créditeur", F("FEP dettes CT", a25["dettes_ct_pcms"])],
                              ["dont épargne PEP et EFC", "41195100, 41195200", F("FEP épargne", a25["epargne"])],
                              ["Provisions et autres dettes longues", "15, 16, 196", F("FEP autres", RT["FEP"]["den"][2025] - a25["dettes_ct_pcms"])],
                              ["Dettes, provisions comprises", "dénominateur", F("FEP den (t)", RT["FEP"]["den"][2025])],
                              ["Ratio des fonds propres et assimilés", "numérateur / dénominateur", N("FEP ratio (t)", RT["FEP"]["valeurs"][2025], 2)]],
                   "largeurs": [6.8, 6.0, 5.4], "rouge": [(0, 2), (4, 2), (7, 2)],
                   "explications": ["Les fonds propres sont négatifs : seules les provisions rendent le numérateur positif.", "L'épargne représente l'essentiel des dettes.", "Le ratio reste très loin de la norme d'au moins 1."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2025\\BALANCE FPPN 2025.xlsx"},
        analyse=["**Même en comptant les provisions comme des fonds propres, le FPPN reste loin de la norme.** Le ratio passe de " + N("FEP 2022 (a)", RT["FEP"]["valeurs"][2022], 2) + " en 2022 à " + N("FEP 2025 (a)", RT["FEP"]["valeurs"][2025], 2) + " en 2025."],
        implication="Ce ratio ne peut s'améliorer que par des excédents mis en réserve pendant plusieurs années : il mesure la profondeur du retard de fonds propres, plus qu'un problème de l'année.",
        source=SRC_COMPTA + " ; " + DEC + ", § 1.2.2.5.3")
contenu({"titre": "Calcul pas à pas : liquidité 2025", "sous_titre": "La trésorerie couvre un huitième des dettes à court terme, épargne comprise",
         "synthese": f"Liquidité {val_ratio(RT['LIQ'], 2025)} comme le FPPN, {val_ratio(RT['LIQ'], 2025, True)} sans l'épargne ; trésorerie à échéance {val_ratio(RT['RL2'], 2025)} et {val_ratio(RT['RL2'], 2025, True)}"},
        "Formules : liquidité = trésorerie actif / (dettes à court terme + variation des provisions techniques) (Décision, § 1.2.2.5.2) ; trésorerie à échéance = (valeurs réalisables + disponibles) / dettes à court terme (section II, § 2.2) ; trésorerie immédiate = disponibles / dettes à court terme (§ 2.3), égale ici à la liquidité car les provisions ne varient pas.",
        pas_a_pas={"entetes": ["Élément", "Comptes", "Montant 2025"],
                   "lignes": [["Trésorerie actif", "classe 5 débitrice", F("LIQ tréso", a25["tresorerie_actif"])],
                              ["Créances à court terme", "classes 3 et 4 débitrices, hors avances durables", F("LIQ créances", a25["creances_ct"])],
                              ["Dettes à court terme", "classe 4 créditrice, 5 créditeur", F("LIQ dettes", a25["dettes_ct_pcms"])],
                              ["dont épargne PEP et EFC", "41195100, 41195200", F("LIQ épargne", a25["epargne"])],
                              ["Dettes à court terme hors épargne", "lecture retraitée", F("LIQ dettes hors ép", a25["dettes_ct_hors_epargne"])],
                              ["Liquidité comme le FPPN", "trésorerie / dettes", val_ratio(RT["LIQ"], 2025)],
                              ["Liquidité retraitée", "trésorerie / dettes hors épargne", val_ratio(RT["LIQ"], 2025, True)]],
                   "largeurs": [6.6, 7.4, 4.8], "rouge": [(3, 2), (5, 2), (6, 2)],
                   "explications": ["L'épargne représente presque toutes les dettes à court terme du plan comptable.", "Dans cette lecture, la trésorerie couvre environ un huitième des dettes.", "Sans l'épargne, la trésorerie couvre toutes les dettes à court terme."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2025\\BALANCE FPPN 2025.xlsx"},
        analyse=["**Toute la question de la liquidité se ramène à l'épargne.** Les autres dettes à court terme sont couvertes par la trésorerie en 2025."],
        implication="Le FPPN doit connaître la part de l'épargne exigible dans l'année : sans cet échéancier, la lecture prudente du plan comptable s'impose et la liquidité est non conforme.",
        source=SRC_COMPTA + " ; " + DEC + ", § 1.2.2.5.2 et section II")
contenu({"titre": "Calcul pas à pas : équilibre stable 2025", "sous_titre": "Les emplois longs dépassent les ressources longues, sauf si l'épargne est une ressource longue",
         "synthese": f"Équilibre stable {val_ratio(RT['EQS'], 2025)} comme le FPPN, {val_ratio(RT['EQS'], 2025, True)} retraité ; norme d'au plus 1"},
        "Formule (Décision, § 1.2.3.1) : emplois moyens et longs / ressources moyennes et longues. Comme le FPPN : classe 2 nette / classes 10 à 19. Retraité : on ajoute aux emplois les avances durables suivies en classe 4 (CMP, Grenier, Hôpital de la Police, projets immobiliers, créance de cession) et aux ressources l'épargne PEP et EFC.",
        pas_a_pas={"entetes": ["Élément", "Comptes", "Montant 2025"],
                   "lignes": [["Actif immobilisé net", "classe 2", F("EQS immo", a25["actif_immobilise_net"])],
                              ["Avances et créances durables", "47111100 à 47111105, 47111201, 485", F("EQS avances", a25["emplois_longs_tiers"])],
                              ["Ressources longues", "10 à 19", F("EQS ress", RT["EQS"]["den"][2025])],
                              ["Épargne PEP et EFC", "41195100, 41195200", F("EQS épargne", a25["epargne"])],
                              ["Équilibre stable comme le FPPN", "classe 2 / classe 1", val_ratio(RT["EQS"], 2025)],
                              ["Équilibre stable retraité", "(2 + avances) / (1 + épargne)", val_ratio(RT["EQS"], 2025, True)]],
                   "largeurs": [6.6, 7.4, 4.8], "rouge": [(1, 2), (4, 2), (5, 2)],
                   "explications": ["Des avances durables supérieures à l'actif immobilisé sont suivies hors de la classe 2.", "Comme le FPPN, les emplois longs valent plus d'une fois et demie les ressources longues.", "Avec l'épargne en ressource longue, l'équilibre est respecté de peu."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2025\\BALANCE FPPN 2025.xlsx"},
        analyse=["**L'épargne des adhérents finance les emplois longs du FPPN.** Sans elle, les ressources longues ne couvrent pas les terrains, les bâtiments et les avances durables."],
        implication="Financer des emplois longs avec de l'épargne n'est soutenable que si cette épargne reste longtemps : c'est précisément ce que l'échéancier de l'EFC et du PEP doit démontrer.",
        source=SRC_COMPTA + " ; " + DEC + ", § 1.2.3.1")
contenu({"titre": "Calcul pas à pas : taux d'excédent net 2025", "sous_titre": "L'excédent dépasse le plafond de 5 % des cotisations, ce qui sert à reconstituer les fonds propres",
         "synthese": f"Excédent net {F('EXN num 2025', a25['resultat_net'])} / cotisations acquises {F('EXN den 2025', a25['cotisations_acquises'])} = {pc('EXN 2025 (pas)', RT['EXN']['valeurs'][2025])}"},
        "Formule (Décision, § 1.2.2.5.5) : excédent net de l'exercice / cotisations acquises, norme d'au plus 5 %. Ce plafond vise à éviter qu'une mutuelle prélève sur ses membres plus que nécessaire. Il est en tension avec les normes de fonds propres quand ceux-ci sont négatifs.",
        pas_a_pas={"entetes": ["Élément", "Comptes", "Montant 2025"],
                   "lignes": [["Produits des activités ordinaires", "classe 7", F("EXN produits", a25["produits_7"])],
                              ["Charges des activités ordinaires", "classe 6", F("EXN charges", a25["charges_6"])],
                              ["Résultat hors activités ordinaires", "classe 8", F("EXN hao", a25["resultat_hao"])],
                              ["Excédent net", "7 - 6 + 8", F("EXN excédent", a25["resultat_net"])],
                              ["Cotisations acquises", "classe 70", F("EXN cot", a25["cotisations_acquises"])],
                              ["Taux d'excédent net", "excédent / cotisations", pc("EXN taux (t)", RT["EXN"]["valeurs"][2025])]],
                   "largeurs": [6.8, 5.6, 5.6], "rouge": [(3, 2), (5, 2)],
                   "explications": ["L'excédent net 2025 est positif, mais il baisse depuis 2021.", "Le taux dépasse le plafond de 5 % fixé par la norme."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2025\\BALANCE FPPN 2025.xlsx"},
        analyse=["**Ce dépassement n'est pas un défaut en soi pour le FPPN.** Ses fonds propres sont négatifs, et les excédents sont le seul moyen de les reconstituer et de doter la réserve obligatoire."],
        implication="Un plan de mise en conformité présenté à l'AIRMS doit justifier ce dépassement temporaire par l'objectif de fonds propres, avec un calendrier de retour sous le plafond.",
        source=SRC_COMPTA + " ; " + DEC + ", § 1.2.2.5.5")
contenu({"titre": "Calcul pas à pas : frais de fonctionnement 2025", "sous_titre": "Les charges de fonctionnement absorbent plus d'un tiers des cotisations",
         "synthese": f"Taux des charges de fonctionnement {pc('CHF 2025 (synth)', RT['CHF']['valeurs'][2025])} (norme 20 %), brut {pc('CHB 2025 (synth)', RT['CHB']['valeurs'][2025])} (norme 15 %), par unité prestée {pc('CUP 2025 (synth)', RT['CUP']['valeurs'][2025])} (norme 15 %)"},
        "Formules (Décision, § 1.2.1.2.1 à 1.2.1.2.4). Charges de fonctionnement : comptes 612, 613, 618, 62 à 68. Charges de fonctionnement brut : comptes 61 à 68, 691 et 697. Le compte 65180000 (prime d'assurance décès de 2020) entre dans la classe 65 et donc dans les charges de fonctionnement de 2020.",
        pas_a_pas={"entetes": ["Élément", "Comptes", "Montant 2025"],
                   "lignes": [["Charges de fonctionnement", "612, 613, 618, 62 à 68", F("CHF num", a25["charges_fonctionnement"])],
                              ["Cotisations acquises", "70", F("CHF den", a25["cotisations_acquises"])],
                              ["Taux des charges de fonctionnement", "norme : au plus 20 %", pc("CHF taux (t)", RT["CHF"]["valeurs"][2025])],
                              ["Charges de fonctionnement brut", "61 à 68, 691, 697", F("CHB num", a25["charges_fonctionnement_brut"])],
                              ["Total des produits", "liste de l'annexe", F("CHB den", a25["total_produits"])],
                              ["Taux brut", "norme : au plus 15 %", pc("CHB taux (t)", RT["CHB"]["valeurs"][2025])],
                              ["Charges de prestations", "60", F("CUP den", a25["charges_prestations"])],
                              ["Coût par unité prestée", "norme : au plus 15 %", pc("CUP taux (t)", RT["CUP"]["valeurs"][2025])],
                              ["Charges de personnel / fonctionnement brut", "66, norme : moins de 85 %", pc("CHP taux (t)", RT["CHP"]["valeurs"][2025])]],
                   "largeurs": [7.2, 6.2, 4.8], "rouge": [(2, 2), (5, 2), (7, 2)],
                   "explications": ["Les charges de fonctionnement représentent plus d'un tiers des cotisations.", "Le taux brut dépasse aussi sa norme de 15 %.", "Le FPPN dépense 63 FCFA de fonctionnement pour 100 FCFA de prestations."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2025\\BALANCE FPPN 2025.xlsx"},
        analyse=["**Le personnel n'explique pas l'écart.** Les charges de personnel ne représentent qu'un quart des charges de fonctionnement brut ; les indemnités des organes, les services extérieurs et les amortissements pèsent ensemble davantage."],
        implication="Pour atteindre 20 % des cotisations de 2025, les charges de fonctionnement doivent baisser de " + F("CHF baisse nécessaire", a25["charges_fonctionnement"] - 0.20 * a25["cotisations_acquises"]) + " : l'effort doit porter sur les postes autres que le personnel.",
        source=SRC_COMPTA + " ; " + DEC + ", § 1.2.1.2.1 à 1.2.1.2.4")
trim = RT["RL4"]["den"][2025]
contenu({"titre": "Calcul pas à pas : réserves des risques longs 2025", "sous_titre": "Aucune réserve de trésorerie n'est constituée pour l'épargne, et les fonds propres sont négatifs",
         "synthese": f"Besoin d'un trimestre d'engagements et de frais : {F('RL4 besoin 2025', trim)} ; réserve de trésorerie constituée : 0 FCFA"},
        "Règlement d'exécution, art. 19 : une réserve de trésorerie ou fonds de roulement au moins égale à un trimestre d'engagements et de charges de fonctionnement, pour chaque risque long. Engagements : débits des comptes 41195100 et 41195200 (sorties d'épargne). Charges de fonctionnement : quote-part au prorata de l'épargne dans le total épargne + cotisations. Art. 21 : fonds propres au moins égaux aux prestations des risques longs des deux derniers exercices.",
        pas_a_pas={"entetes": ["Élément", "Calcul", "Montant 2025"],
                   "lignes": [["Sorties d'épargne", "débits 41195100 et 41195200", F("RL4 sorties", a25["sorties_epargne"])],
                              ["Frais attribués à l'épargne", "prorata épargne / (épargne + cotisations)", F("RL4 frais", 4 * trim - a25["sorties_epargne"])],
                              ["Besoin d'un trimestre", "(sorties + frais) / 4", F("RL4 besoin (t)", trim)],
                              ["Réserve de trésorerie constituée", "aucun compte identifié", "0 FCFA"],
                              ["Actifs réservés à l'épargne (pour comparaison)", "banques, DAT et FCP réservés", F("RL4 actifs réservés", a25["actifs_dedies_epargne"])],
                              ["Fonds propres", "10 à 14 + résultat", F("RL5 FP", a25["fonds_propres"])],
                              ["Sorties d'épargne 2024 et 2025", "art. 21", F("RL5 sorties 2 ans", RT["RL5"]["den"][2025])],
                              ["Fonds propres / sorties de deux ans", "norme : au moins 1", N("RL5 ratio (t)", RT["RL5"]["valeurs"][2025], 2)]],
                   "largeurs": [7.0, 7.0, 4.6], "rouge": [(2, 2), (3, 2), (7, 2)],
                   "explications": ["Le besoin d'un trimestre se calcule à partir des sorties réelles d'épargne.", "Aucun compte de réserve de trésorerie n'existe dans les balances.", "Les fonds propres négatifs rendent le ratio de l'article 21 négatif."],
                   "fichier": "AUDIT ADENKA\\DAFC\\2024 et 2025 (balances des comptes)"},
        analyse=["**Les actifs réservés à l'épargne dépassent le besoin d'un trimestre, mais ils ne forment pas une réserve.** Une réserve est une part des fonds propres affectée, et les fonds propres du FPPN sont négatifs."],
        implication="La constitution de cette réserve passe par l'affectation d'excédents futurs, après reconstitution des fonds propres : elle doit figurer dans le plan de mise en conformité.",
        source=SRC_COMPTA + " ; " + REX + ", art. 19 et 21")
methode_gab({"titre": "Excédent par risque : pourquoi le calculer", "sous_titre": "La norme d'affectation des excédents oblige à connaître le résultat de chaque risque",
             "synthese": "Le FPPN ne présente qu'un résultat global ; Adenka reconstitue le résultat de la maladie et du décès"},
            [("Ce que dit la norme", "La Décision n° 05/2022 impose d'affecter à la réserve obligatoire au moins 30 % de l'excédent net du risque maladie et au moins 50 % de celui des autres risques (§ 1.1.1). Le règlement d'exécution fixe une dotation de 20 à 50 % du résultat excédentaire de l'exercice précédent (art. 18)."),
             ("Ce qui nous y a conduits", "Le résultat global du FPPN est positif, mais la sinistralité maladie monte vite pendant que le décès reste très excédentaire (partie 5). Un excédent global peut donc cacher un déficit de la maladie, ce que la norme veut précisément rendre visible."),
             ("Ce que fait le FPPN", "La comptabilité sépare les cotisations et les prestations de chaque risque, mais pas les frais de fonctionnement. Adenka reprend d'abord la dotation réellement passée en réserve, telle que le FPPN l'a décidée."),
             ("Pourquoi répartir les frais", "L'excédent par risque est la différence entre les produits et toutes les charges affectées au risque (Décision, § 1.1.1). Sans clé fournie par le FPPN, Adenka répartit les frais au prorata des cotisations de chaque risque, la clé la plus neutre.")],
            "La répartition au prorata des cotisations est une hypothèse : une autre clé (nombre de dossiers, temps passé) donnerait d'autres excédents. Le FPPN est invité à fournir sa propre clé, qui remplacera celle-ci.",
            source="Source : " + DEC + ", § 1.1.1 ; " + REX + ", art. 18 ; Adenka")
contenu({"titre": "Excédent par risque", "sous_titre": "Depuis 2025, la maladie est déficitaire et le décès finance l'ensemble",
         "synthese": f"Excédent maladie 2025 : {F('Exc mal 2025', EXR[2025]['excedent_maladie'])} ; excédent décès 2025 : {F('Exc dec 2025', EXR[2025]['excedent_deces'])}"},
        "Excédent d'un risque = cotisations du risque - prestations du risque - frais attribués au prorata des cotisations. La prime d'assurance décès de 2020 (65180000) est rattachée au décès. Les autres éléments (produits financiers, transferts de charges, hors activités ordinaires) ne sont pas répartis et figurent dans la colonne « autres » du classeur des calculs.",
        graphique={"titre": fig("Excédent technique par risque, frais répartis au prorata des cotisations"), "type": "barres", "categories": cats_ex,
                   "series": [{"nom": "Maladie", "valeurs": [round(EXR[a]["excedent_maladie"]) for a in range(2020, 2027)], "couleur": BLEU},
                              {"nom": "Décès", "valeurs": [round(EXR[a]["excedent_deces"]) for a in range(2020, 2027)], "couleur": ORANGE}],
                   "format": FMT_M, "unite": "En millions de FCFA"},
        analyse=[f"**La maladie devient déficitaire en 2025, une fois ses frais pris en compte.** Son excédent passe de {F('Exc mal 2022', EXR[2022]['excedent_maladie'])} en 2022 à {F('Exc mal 2025 (t)', EXR[2025]['excedent_maladie'])} en 2025.",
                 f"**Le décès est excédentaire chaque année depuis 2022.** Il dégage {F('Exc dec 2025 (t)', EXR[2025]['excedent_deces'])} en 2025.",
                 f"**Sur les huit mois de 2026, le déficit maladie atteint {F('Exc mal 2026', EXR[2026]['excedent_maladie'])}.**"],
        implication="Le décès subventionne la maladie : si le barème officiel confirme une clé de 70 % pour la maladie, la cotisation maladie ne couvre plus son coût complet et doit être revue.",
        source=SRC_COMPTA + " ; répartition des frais par Adenka")
contenu({"titre": "Affectation des excédents", "sous_titre": "Depuis 2024, la dotation à la réserve obligatoire atteint le minimum exigé par risque",
         "synthese": "Dotation de la réserve légale (compte 111) rapportée au résultat de l'exercice précédent, puis comparée aux deux normes"},
        "Lecture : la colonne « dotation réelle » est l'augmentation du compte 111 entre deux clôtures ; elle traduit l'affectation décidée par l'assemblée générale. Le minimum du règlement d'exécution est de 20 % du résultat excédentaire de l'exercice précédent (art. 18), exigible dans un délai de trois ans après le 1er juin 2022. Le minimum par risque applique 30 % à l'excédent maladie et 50 % à l'excédent décès (Décision, § 1.1.1).",
        tableau={"titre": tab("Dotation de la réserve obligatoire et normes"), "entetes": ["Affectation en", "Résultat de l'exercice précédent", "Dotation réelle", "Taux", "Minimum de 20 %", "Minimum par risque", "Écart au minimum par risque"],
                 "largeurs": [3.2, 5.0, 4.4, 2.6, 4.4, 4.4, 4.8], "taille": 10,
                 "lignes": [[str(a), F(f"AFF res {a}", AFF[a]["resultat_n_1"]), F(f"AFF dot {a}", AFF[a]["dotation_reserve_legale"]), pc(f"AFF taux {a}", AFF[a]["taux_dotation_legale"]),
                             F(f"AFF min20 {a}", 0.20 * AFF[a]["resultat_n_1"]), F(f"AFF minrisque {a}", AFF[a]["dotation_minimale_par_risque"]), F(f"AFF écart {a}", AFF[a]["ecart_dotation"])] for a in (2023, 2024, 2025)]},
        disposition="dessous",
        analyse=[f"**La dotation passe de {pc('AFF taux 2023 (a)', AFF[2023]['taux_dotation_legale'])} du résultat en 2023 à {pc('AFF taux 2025 (a)', AFF[2025]['taux_dotation_legale'])} en 2025.** En 2023, elle restait sous les deux minimums, pendant le délai de mise en conformité de trois ans.",
                 "**En 2024, une partie des excédents a aussi servi à apurer les pertes reportées et à doter les réserves facultatives.**"],
        implication="Le FPPN respecte désormais la règle d'affectation, mais la réserve reste sous six mois de prestations : la règle d'affectation seule ne suffit pas à rattraper le retard.",
        source=SRC_COMPTA + " ; " + DEC + ", § 1.1.1 ; " + REX + ", art. 18")
contenu({"titre": "Marge de solvabilité", "sous_titre": "La marge disponible ne couvre la marge exigée que si l'on ignore les pertes reportées",
         "synthese": f"2025 : marge exigée {F('MS exigée 2025 (g)', MS[2025]['marge_exigee'])}, disponible {F('MS dispo 2025 (g)', MS[2025]['marge_disponible_comme_fppn'])} comme le FPPN, {F('MS nette 2025 (g)', MS[2025]['marge_disponible_nette'])} nette des pertes"},
        "Marge exigée = le maximum entre 18 % des cotisations acquises de l'exercice précédent et de 26 % de la moyenne des prestations maladie des trois derniers exercices (Décision, § 1.1.1). L'exercice 2020 n'est pas calculé, faute de balances 2018 et 2019.",
        graphique={"titre": fig("Marge de solvabilité exigée et disponible"), "type": "barres", "categories": [str(a) if a < 2026 else "31/08/2026" for a in range(2021, 2027)],
                   "series": [{"nom": "Marge exigée", "valeurs": [round(MS[a]["marge_exigee"]) for a in range(2021, 2027)], "couleur": GRIS_C},
                              {"nom": "Disponible comme le FPPN", "valeurs": [round(MS[a]["marge_disponible_comme_fppn"]) for a in range(2021, 2027)], "couleur": BLEU},
                              {"nom": "Disponible nette des pertes", "valeurs": [round(MS[a]["marge_disponible_nette"]) for a in range(2021, 2027)], "couleur": ORANGE}],
                   "format": FMT_M, "unite": "En millions de FCFA"},
        analyse=[f"**La marge disponible comme le FPPN repose surtout sur la réserve disponible du compte 10620000** ({F('MS 106 (g)', MS[2025]['reserve_disponible_106'])}), dont la nature reste à confirmer.",
                 "**Nette des pertes reportées, la marge disponible est négative chaque année**, même si elle se rapproche de zéro jusqu'en 2025.",
                 f"**La marge exigée augmente avec les cotisations** : elle passe de {F('MS exigée 2021 (g)', MS[2021]['marge_exigee'])} en 2021 à {F('MS exigée 2026 (g)', MS[2026]['marge_exigee'])} en 2026."],
        implication="Si le compte 10620000 est un écart de réévaluation, comme son numéro l'indique dans le PCMS, il ne compte pas dans la marge, et le FPPN est non conforme dans les deux lectures. Le FPPN doit documenter l'origine de ce montant.",
        source=SRC_COMPTA + " ; " + DEC + ", § 1.1.1 ; PCMS, fiche n° 1")
contenu({"titre": "Contrôle du plan comptable", "sous_titre": "Deux écarts au plan comptable touchent directement les calculs actuariels",
         "synthese": "Contrôle limité aux comptes qui alimentent les provisions, les réserves, l'épargne, les cotisations, les prestations et les placements"},
        "Le contrôle ne constitue pas une revue comptable complète : il vérifie seulement que les comptes utilisés dans ce rapport sont rangés comme le prévoit le PCMS, et mesure l'effet d'un écart sur les ratios.",
        tableau={"titre": tab("Contrôle du classement des comptes selon le PCMS"), "entetes": ["Point", "Attendu selon le PCMS", "Constat", "Incidence"], "largeurs": [4.6, 8.2, 9.2, 9.2], "taille": 9,
                 "lignes": [[p["point"], p["attendu_pcms"], p["constat"], p["incidence"]] for p in Q["pcms"]]},
        disposition="dessous",
        analyse=["**Le reclassement des provisions techniques en provisions réglementées depuis 2024 et l'usage du compte 106 sont les deux écarts qui changent la lecture du bilan.** Les autres comptes suivent le plan comptable."],
        source="Source : PCMS (mai 2012), liste des comptes et fiches n° 1, 2, 6 et 25 ; balances du FPPN")
contenu({"titre": "Exigences non chiffrées", "sous_titre": "Plusieurs obligations du règlement d'exécution ne se mesurent pas par un ratio",
         "synthese": "Chaque obligation est comparée aux pièces disponibles ; plusieurs restent à documenter par le FPPN"},
        "Ces exigences complètent les ratios. Elles relèvent du cadre actuariel quand elles portent sur les provisions, les risques longs, les placements et les réserves. Elles sont reprises dans les guides d'entretien.",
        tableau={"titre": tab("Exigences non chiffrées et situation du FPPN"), "entetes": ["Exigence", "Texte", "Situation constatée"], "largeurs": [9.6, 6.2, 15.4], "taille": 9,
                 "lignes": [["Provisions techniques suffisantes à tout moment, calculées dossier par dossier", "Règl. d'exécution, art. 7 à 11", "Montant identique de 2020 à 2026 : la provision n'est pas recalculée (partie 9)."],
                            ["Comptes certifiés par un commissaire aux comptes (risques longs)", "Règl. d'exécution, art. 12", "Aucun rapport de commissaire aux comptes n'a été transmis."],
                            ["États financiers propres à chaque risque long", "Règl. d'exécution, art. 13", "Les balances ne séparent pas l'EFC et le PEP en états distincts."],
                            ["Étude actuarielle tous les cinq ans (risques longs)", "Règl. d'exécution, art. 14", "Aucune étude antérieure n'a été transmise ; la présente mission peut y contribuer."],
                            ["Cantonnement des patrimoines entre risques courts et longs", "Règl. d'exécution, art. 15", "Des comptes bancaires réservés existent, mais ils couvrent seulement 16,28 % de l'épargne."],
                            ["Placements sûrs, rentables, liquides et diversifiés", "Règl. d'exécution, art. 17", "La décision ministérielle qui fixe ces conditions n'a pas été transmise."],
                            ["Pertes des activités sanitaires non couvertes par les cotisations", "Règl. d'exécution, art. 27", "Le FPPN finance des charges de fonctionnement « au CMP » sans comptabilité séparée transmise."],
                            ["Adhésion et cotisation au fonds national de garantie", "Règl. n° 07/2009, art. 63", "Aucun compte du fonds de garantie n'apparaît dans les balances."]]},
        disposition="dessous",
        analyse=["**Les exigences propres aux risques longs ne sont pas documentées.** Elles découlent pourtant directement de la gestion de l'EFC et du PEP."],
        source="Source : " + REX + " ; " + R07 + " ; documents reçus le 15/10/2026")
contenu({"titre": "Exposition aux sanctions", "sous_titre": "Le non-respect des règles prudentielles expose le FPPN à des sanctions graduées",
         "synthese": "Barème de l'arrêté n° 00115/MEPS/CAB du 12/12/2023 pour une mutuelle de plus de 10 000 000 FCFA de cotisations"},
        "Les sanctions sont prononcées par le ministre chargé de la mutualité sociale, sur rapport ou avis motivé de l'AIRMS, après que la mutuelle a été entendue (règlement d'exécution, art. 38 et 39). Ce tableau décrit une exposition théorique : la décision appartient à l'autorité de contrôle.",
        tableau={"titre": tab("Manquements liés aux constats du rapport et sanctions prévues"), "entetes": ["Manquement", "Ligne de l'arrêté", "Sanction", "Peine pécuniaire"], "largeurs": [10.0, 3.4, 10.2, 7.6], "taille": 10,
                 "lignes": [["Non-respect des règles comptables et prudentielles durant un exercice", "6", "Avertissement, injonction de régularisation, suspension des dirigeants ou suspension provisoire de l'agrément", "1 000 000 FCFA"],
                            ["Défaut d'adhésion ou de versement au fonds national de garantie", "8", "Mêmes sanctions que la ligne 6", "1 000 000 FCFA"],
                            ["Non-respect des règles prudentielles sur deux exercices", "13", "Retrait de l'agrément (règlement n° 07/2009, art. 84)", "3 000 000 FCFA"]]},
        disposition="dessous",
        analyse=["**Les peines pécuniaires sont faibles ; la sanction lourde est le retrait d'agrément.** Celui-ci emporte la dissolution de la mutuelle dès sa publication au Journal officiel (règlement n° 07/2009, art. 84)."],
        implication="Un plan de mise en conformité daté et chiffré, présenté à l'AIRMS, est le meilleur moyen de démontrer la volonté de régularisation et d'éviter les sanctions les plus lourdes.",
        source="Source : arrêté n° 00115/MEPS/CAB du 12/12/2023, annexe ; " + R07 + ", art. 83 et 84 ; " + REX + ", art. 38 à 41")
retenir("conformité prudentielle", ("Les normes de solidité financière de l'UEMOA ne sont pas respectées", f"{NB_NC} ratios non conformes sur 36, souvent plusieurs années de suite",
                                    "Le FPPN respecte les ratios d'activité courante, mais pas ceux de réserves, de fonds propres, de liquidité et de frais."),
        f"La réserve obligatoire couvre {N('ROB 2025 (ret)', RT['ROB']['valeurs'][2025], 2)} mois de prestations, les fonds propres et assimilés {N('FEP 2025 (ret)', RT['FEP']['valeurs'][2025], 2)} fois les dettes, et les frais {pc('CHF 2025 (ret)', RT['CHF']['valeurs'][2025])} des cotisations.",
        "Des non-conformités répétées exposent le FPPN au retrait d'agrément prévu par l'article 84 du règlement n° 07/2009.",
        "Le FPPN doit présenter à l'AIRMS un plan de mise en conformité daté, et documenter le compte 106, le fonds de garantie et l'échéancier de l'épargne.",
        "Les ratios sont calculés sur les balances : leur conclusion ne dépend pas des données manquantes, sauf pour la liquidité, qui dépend de l'échéancier de l'épargne.")

# ---------------------------------------------------------------- PARTIE 11 : MORTALITÉ
separateur(11, "Tables de mortalité", "Rôle, choix des tables, indicateurs de référence")
methode_gab({"titre": "Pourquoi des tables de mortalité", "sous_titre": "Les tables servent à mesurer et à évaluer tout engagement lié à la vie humaine",
             "synthese": "Elles préparent deux analyses de la mission : la mortalité des policiers et les engagements viagers"},
            [("qx", "La table donne, pour chaque âge x, la probabilité de décéder avant l'âge x + 1. C'est la brique de base de tous les calculs."),
             ("Lx", "La table donne aussi le nombre de survivants à chaque âge, sur 1 000 000 de naissances. Le rapport de deux Lx donne une probabilité de survie."),
             ("Premier usage", "Comparer la mortalité réelle des policiers à la table de référence dit si la cotisation décès est bien calibrée."),
             ("Second usage", "Les tables permettent de calculer la valeur des engagements qui durent toute la vie, comme une prise en charge santé maintenue aux retraités.")],
            "Ces deux analyses demandent les dates de naissance (PBC, items 9 et 28). Le règlement d'exécution n° 004/2022 définit la table de mortalité comme un tableau qui donne les probabilités de décès ou de survie et l'espérance de vie selon l'âge et le sexe (art. 1).",
            source=SRC_CIMA)
contenu({"titre": "Choix des tables", "sous_titre": "Le principe de prudence fixe la table selon le risque porté",
         "synthese": "CIMA H pour le risque décès, CIMA F pour le risque de survie, TD et TV en comparaison historique"},
        "Principe de prudence : on choisit la table qui défavorise l'organisme. Pour un risque décès (le FPPN paie au décès), une table qui prévoit plus de décès est prudente : CIMA H. Pour un risque de survie (le FPPN paie tant que la personne vit), une table qui prévoit des vies plus longues est prudente : CIMA F. Les fichiers fournis portent eux-mêmes cet usage dans le nom de leur feuille. Les tables TD et TV sont les anciennes tables réglementaires remplacées en 2012 (Code CIMA, article 338 modifié).",
        tableau={"titre": tab("Tables de mortalité CIMA et usage retenu"), "entetes": ["Table", "Feuille du fichier", "Risque", "Usage dans la mission"], "largeurs": [3.8, 5.6, 7.0, 14.8],
                 "lignes": [["CIMA H", MORT["CIMA H"]["feuille"], "Décès (paiement au décès)", "Calibrage de la cotisation décès et comparaison à la mortalité observée"],
                            ["CIMA F", MORT["CIMA F"]["feuille"], "Survie (paiement tant que la personne vit)", "Engagements viagers : prise en charge des retraités, rentes"],
                            ["TD CIMA", MORT["TD CIMA"]["feuille"], "Décès (ancienne table)", "Comparaison historique"],
                            ["TV CIMA", MORT["TV CIMA"]["feuille"], "Survie (ancienne table)", "Comparaison historique"]]},
        disposition="dessous",
        analyse=["**La prudence change de sens selon le risque.** Une même table peut être prudente pour un capital décès et imprudente pour une rente : c'est pourquoi deux tables sont retenues."],
        implication="Le FPPN porte les deux risques (capital décès et, potentiellement, prise en charge viagère des retraités) : chaque engagement sera évalué avec la table qui le défavorise.",
        source=SRC_CIMA)
cs = C["courbes_survie"]
contenu({"titre": "Courbes de survie", "sous_titre": "Les tables de survie prévoient des vies plus longues que les tables de décès",
         "synthese": f"Survivants à 60 ans sur 100 naissances : {N('Survie60 H', 100 * MORT['CIMA H']['survie_60'], 2)} (CIMA H), {N('Survie60 F', 100 * MORT['CIMA F']['survie_60'], 2)} (CIMA F)"},
        "Une courbe de survie donne la part des personnes encore en vie à chaque âge. L'écart entre les courbes mesure l'écart de prudence entre une table de décès et une table de survie.",
        graphique={"titre": fig("Survivants pour 100 naissances, par âge"), "type": "lignes", "categories": [str(a) for a in cs["ages"]],
                   "series": [{"nom": n, "valeurs": [round(v, 2) for v in cs[n]], "couleur": c} for n, c in (("CIMA H", BLEU), ("CIMA F", ORANGE), ("TD CIMA", GRIS_F), ("TV CIMA", GRIS_C))],
                   "format": "0.00", "etiquettes": False, "unite": "En % des naissances"},
        analyse=[f"**L'espérance de vie à la naissance vaut {N('e0 H', MORT['CIMA H']['e0'], 2)} ans avec CIMA H et {N('e0 F', MORT['CIMA F']['e0'], 2)} ans avec CIMA F.**",
                 f"**Les anciennes tables prévoient davantage de décès** : {N('e0 TD', MORT['TD CIMA']['e0'], 2)} ans avec TD et {N('e0 TV', MORT['TV CIMA']['e0'], 2)} ans avec TV.",
                 "**Le choix de la table change donc fortement les montants calculés.**"],
        implication="Un engagement viager évalué avec une ancienne table est sous-estimé : les calculs de la mission utiliseront les tables en vigueur depuis 2012.",
        source=SRC_CIMA)
cq = C["courbes_qx"]
contenu({"titre": "Probabilités de décès", "sous_titre": "La probabilité de décès double environ tous les dix ans aux âges actifs",
         "synthese": f"Probabilité de décès à 40 ans : {N('q40 H', 1000 * MORT['CIMA H']['q40'], 2)} pour mille (CIMA H), {N('q40 F', 1000 * MORT['CIMA F']['q40'], 2)} pour mille (CIMA F)"},
        "qx en pour mille aux âges de 20 à 80 ans. Ce sont les âges de la population du FPPN (actifs puis retraités). La comparaison avec la mortalité observée se fera par le rapport décès observés / décès attendus (voir partie 12).",
        graphique={"titre": fig("Probabilité de décès dans l'année (qx), par âge"), "type": "lignes", "categories": [str(a) for a in cq["ages"]],
                   "series": [{"nom": n, "valeurs": [round(v, 2) for v in cq[n]], "couleur": c} for n, c in (("CIMA H", BLEU), ("CIMA F", ORANGE), ("TD CIMA", GRIS_F), ("TV CIMA", GRIS_C))],
                   "format": "0.00", "etiquettes": False, "unite": "En pour mille"},
        analyse=[f"**À 55 ans, la probabilité de décès vaut {N('q55 H', 1000 * MORT['CIMA H']['q55'], 2)} pour mille avec CIMA H et {N('q55 F', 1000 * MORT['CIMA F']['q55'], 2)} pour mille avec CIMA F.**",
                 f"**Entre 20 et 60 ans, CIMA H prévoit {pc('H vs TD', -C['mortalite_comparaison']['deces_20_60_H_vs_TD'])} de décès de moins que TD**, et CIMA F {pc('F vs TV', -C['mortalite_comparaison']['deces_20_60_F_vs_TV'])} de moins que TV.",
                 "**La table ronde CIMA FANAF de 2011 indique des écarts de 17 % pour CIMA H et d'environ 4 % pour CIMA F, sur l'ensemble des âges.**"],
        implication="Pour une population de policiers, exposée à des risques professionnels, la mortalité réelle peut s'écarter de la table : seul le rapport décès observés / décès attendus le dira.",
        source=SRC_CIMA)
contenu({"titre": "Indicateurs de référence", "sous_titre": "Une rente viagère vaut entre 12 et 18 années de paiement selon l'âge et la table",
         "synthese": f"Valeur d'une rente de 1 FCFA par an à 60 ans, au taux de 3,5 % : {N('a60 F', MORT['CIMA F']['a60'], 2)} (CIMA F), {N('a60 H', MORT['CIMA H']['a60'], 2)} (CIMA H)"},
        "Rente viagère : valeur aujourd'hui de 1 FCFA payé chaque année, au début de l'année, tant que la personne vit, actualisé au taux de 3,5 % (plafond du Code CIMA). Formule : ä_x = Σ (L_(x+k) / L_x) × v^k, avec v = 1 / 1,035. Exemple : si le FPPN prenait en charge 100 000 FCFA par an de soins pour un retraité de 60 ans, l'engagement vaudrait environ 100 000 × ä_60. Ce sont des valeurs de référence, pas des engagements du FPPN.",
        tableau={"titre": tab("Espérance de vie et valeur de la rente viagère (taux de 3,5 %)"), "entetes": ["Table", "e55 (ans)", "e60 (ans)", "ä55", "ä60", "ä65"], "largeurs": [3.8, 3.0, 3.0, 3.0, 3.0, 3.0],
                 "lignes": [[t, N(f"e55 {t}", MORT[t]["e55"], 2), N(f"e60 {t}", MORT[t]["e60"], 2), N(f"a55 {t}", MORT[t]["a55"], 2), N(f"a60 {t}", MORT[t]["a60"], 2), N(f"a65 {t}", MORT[t]["a65"], 2)] for t in ("CIMA H", "CIMA F", "TD CIMA", "TV CIMA")]},
        disposition="droite",
        analyse=[f"**Une rente servie à partir de 60 ans vaut {N('a60 F (t)', MORT['CIMA F']['a60'], 2)} années de paiement avec CIMA F.** Avec CIMA H, elle n'en vaut que {N('a60 H (t)', MORT['CIMA H']['a60'], 2)}.",
                 "**L'écart entre les deux tables dépasse deux années de paiement à 60 ans.** C'est l'ordre de grandeur de l'enjeu du choix de la table."],
        implication="Ces facteurs serviront à évaluer tout engagement viager du FPPN dès que les âges et les coûts annuels seront connus.",
        source=SRC_CIMA)
retenir("tables de mortalité", ("Les tables CIMA sont prêtes ; il manque les dates de naissance pour les appliquer", "Choix prudent : CIMA H pour le décès, CIMA F pour la survie",
                                "Les outils de calcul sont en place, mais ils ne peuvent pas être appliqués aux mutualistes du FPPN sans leurs âges."),
        "Une rente à 60 ans vaut 16,84 années de paiement avec CIMA F, contre 14,26 avec CIMA H. Les anciennes tables sous-estiment la durée de vie.",
        "La mortalité des policiers et la valeur des engagements viagers restent inconnues tant que les dates de naissance manquent.",
        "Le FPPN doit transmettre le fichier des mutualistes avec date de naissance et sexe (PBC, item 9) et la base des décès (item 28).",
        "Les indicateurs présentés sont des références techniques. Ils deviendront des engagements chiffrés du FPPN dès réception des données.")

# ---------------------------------------------------------------- PARTIE 12 : MÉTHODES À VENIR
separateur(12, "Méthodes actuarielles", "À venir : provisions, mortalité, projection")
methode_gab({"titre": "Provisions : méthode PSAP", "sous_titre": "La provision pour sinistres à payer additionne dossiers connus et sinistres non déclarés",
             "synthese": "PSAP = dossiers déclarés non payés + sinistres survenus non encore déclarés (IBNR), par garantie et par clôture"},
            [("Dossiers connus", "Adenka additionne les montants dus des dossiers déclarés et non payés, dossier par dossier (PBC 14 et 29)."),
             ("Sinistres tardifs (IBNR)", "Adenka estime les sinistres survenus mais non encore déclarés à partir des délais observés dans le passé, par la méthode Chain Ladder (PBC 13, 28 et 40)."),
             ("Provision totale", "La provision additionne les dossiers connus et les tardifs, par garantie et par clôture, comme l'exige l'article 10 du règlement d'exécution."),
             ("Comparaison", "Adenka compare la provision calculée aux comptes 15801000 et 15802000, et mesure l'écart (boni ou mali) à chaque clôture.")],
            "Les dossiers connus non payés se comptent dossier par dossier. Les sinistres non déclarés (tardifs) s'estiment à partir des délais observés dans le passé : c'est le rôle de la méthode Chain Ladder.",
            source="Source : Adenka ; " + REX + ", art. 10", attente=True)
gab(39, {"titre": "Méthode Chain Ladder", "sous_titre": "Le triangle des paiements révèle le rythme de règlement des sinistres",
         "synthese": "Quatre étapes, de la base des prestations à la provision pour sinistres tardifs"},
    {"etapes": [("Ranger", "Adenka range les paiements cumulés dans un triangle : une ligne par année de survenance, une colonne par délai de règlement."),
                ("Mesurer", "Adenka calcule pour chaque délai le facteur de développement : la progression moyenne des paiements d'un délai au suivant."),
                ("Compléter", "Adenka complète le triangle avec ces facteurs pour obtenir la charge ultime de chaque année de survenance."),
                ("Provisionner", "La provision est la charge ultime moins les paiements déjà faits ; elle s'ajoute aux dossiers connus.")],
     "attente": True},
    "Hypothèse : le rythme de règlement passé se répète. Aucun chiffre n'est présenté : les données ne sont pas encore reçues.",
    source="Source : Adenka (méthode actuarielle standard)")
contenu({"titre": "Chain Ladder : le triangle", "sous_titre": "Les cases « ? » sont les paiements futurs que la méthode estime",
         "synthese": "Facteur de développement f(j) = Σ C(i, j+1) / Σ C(i, j) ; charge ultime = dernier montant connu x produit des facteurs suivants"},
        "Principe : on range les paiements cumulés dans un triangle (année de survenance × délai). Le facteur de développement mesure la progression moyenne d'un délai au suivant. On complète le triangle avec ces facteurs ; la charge ultime moins les paiements déjà faits donne la provision. Hypothèse : le rythme de règlement passé se répète. Aucun chiffre n'est présenté : les données ne sont pas encore reçues.",
        triangle=True, attente=True, source="Source : Adenka (méthode actuarielle standard)")
methode_gab({"titre": "Mortalité observée", "sous_titre": "Le rapport décès observés / décès attendus mesure l'écart à la table",
             "synthese": "Ratio A/E = décès observés / Σ qx(table) × exposition au risque, par âge"},
            [("Exposition", "Adenka calcule le nombre d'années-personnes passées à chaque âge pendant la période, à partir des dates de naissance, d'entrée et de sortie."),
             ("Décès attendus", "Les décès attendus sont l'exposition de chaque âge multipliée par la probabilité de décès de la table CIMA H, additionnée sur tous les âges."),
             ("Ratio A/E", "Le ratio rapporte les décès observés aux décès attendus. Un ratio supérieur à 100 % signifie une mortalité plus forte que la table."),
             ("Signification", "Un intervalle de confiance fondé sur la loi de Poisson dit si l'écart à la table est significatif ou dû au hasard.")],
            "Exposition : nombre d'années-personnes passées à chaque âge pendant la période. Décès attendus : exposition × qx de la table CIMA H. Données : base des décès (PBC 28), mutualistes avec dates (PBC 9, 10, 11).",
            source="Source : Adenka", attente=True)
gab(41, {"titre": "Projection de l'équilibre technique", "sous_titre": "La projection teste si les cotisations futures couvriront les prestations futures",
         "synthese": "Six étapes, de la population d'aujourd'hui à l'équilibre de chaque année sur 5 à 10 ans"},
    {"etapes": [("Population", "Adenka répartit les mutualistes par âge, sexe et statut (actif, retraité), à partir du fichier des dates de naissance."),
                ("Vieillissement", "Chaque année, la population vieillit d'un an ; les décès suivent la table CIMA H et les départs à la retraite suivent l'âge légal."),
                ("Entrées", "Les promotions de recrues de juillet sont ajoutées selon une hypothèse de recrutement à valider avec le FPPN."),
                ("Cotisations", "Les cotisations futures sont les salaires projetés multipliés par le taux de 5 % du barème officiel."),
                ("Prestations", "Les prestations futures sont les effectifs par âge multipliés par le coût moyen des soins et des décès à chaque âge."),
                ("Équilibre", "L'écart entre cotisations, prestations et frais donne l'équilibre de chaque année, testé par des sensibilités.")],
     "attente": True},
    "Données nécessaires : âges, coût des soins par âge, barèmes, hypothèses de recrutement. Aucune projection chiffrée n'est présentée.",
    source="Source : Adenka")
contenu({"titre": "Travaux en attente des données", "sous_titre": "Ces méthodes seront appliquées dès réception des données listées",
         "synthese": "Voir la liste des documents demandés (PBC pointée) et la partie 14 du rapport"},
        "Aucun chiffre n'est présenté pour ces méthodes : les appliquer sans les données conduirait à des résultats sans valeur. Chaque méthode est reliée aux numéros de la liste des documents demandés.",
        tableau={"titre": tab("Méthodes en attente et données nécessaires"), "entetes": ["Méthode", "Données nécessaires", "PBC"], "largeurs": [7.4, 18.0, 5.8],
                 "lignes": [["PSAP et Chain Ladder", "Bases des prestations avec dates, dossiers non payés à chaque clôture", "13, 14, 28, 29, 40"],
                            ["Mortalité observée (A/E)", "Mutualistes avec date de naissance, sexe, entrée et sortie, base des décès", "9, 10, 11, 28"],
                            ["Ratios de gestion en attente", "Fichiers aller 936 et 8Z6, contrôles a posteriori, procès-verbaux, cotisation au fonds de garantie", "Nouveaux points PBC"],
                            ["Projection de l'équilibre", "Âges, barèmes, coût des soins par âge, hypothèses de recrutement", "1, 2, 3, 9, 40"],
                            ["Engagements d'épargne", "Comptes individuels PEP et EFC, règlement des produits", "16, 17, 18, 19"]]},
        disposition="dessous",
        analyse=["**Toutes les méthodes en attente dépendent de trois fichiers** : les bases de prestations, le fichier des mutualistes avec dates de naissance et les comptes individuels d'épargne."],
        implication="Ces trois fichiers permettront aussi de réaliser l'étude actuarielle quinquennale exigée pour les risques longs (art. 14 du règlement d'exécution).",
        source="Source : liste des documents demandés (PBC) pointée au 18/10/2026")

# ---------------------------------------------------------------- PARTIE 13 : OUTILS
separateur(13, "Outils et traçabilité", "Chaîne de traitement, extraits de code, contrôle des chiffres")
contenu({"titre": "Environnement de travail", "sous_titre": "Les traitements sont écrits en Python et exécutés dans VS Code sous Windows",
         "synthese": "Dix scripts conservés, rejouables à chaque nouvel envoi du FPPN"},
        "Les images des diapositives suivantes reproduisent la fenêtre de VS Code sous Windows et y placent des extraits réels des scripts de la mission. Ce sont des reconstitutions, pas des captures d'écran ; le terminal n'est affiché que lorsque sa sortie a réellement été obtenue.",
        tableau={"titre": tab("Scripts de la mission"), "entetes": ["Script", "Rôle", "Contrôle intégré"], "largeurs": [7.4, 12.4, 11.4],
                 "lignes": [["extraction_fppn.py", "Extraction sans perte, archive découpée pour l'envoi", "Empreintes SHA-256, lignes et cellules"],
                            ["envoi_fichiers_fppn.py", "Envoi de fichiers précis en parties", "Empreinte de chaque partie"],
                            ["base_fppn.py", "Base d'analyse : lecture par en-tête ou par contenu", "Lignes de total, montants doublés"],
                            ["grenier_fppn.py, compta_fppn.py", "Lecture des exports du logiciel de prêts et des balances Sage", "Totaux imprimés retrouvés, débit = crédit"],
                            ["analyses_fppn.py, analyses_complements_fppn.py", "Indicateurs, tests, sensibilités, tables CIMA", "Classeur des calculs"],
                            ["ratios_airms_fppn.py", "Ratios prudentiels, classement des comptes, excédent par risque", "Rapprochement avec les agrégats de la version 2"],
                            ["rapport_v3_fppn.py, classeur_v2_fppn.py", "Rapport, graphiques et tableaux natifs, classeur", "Contrôle automatique des chiffres"]]},
        disposition="dessous",
        analyse=["**Chaque chiffre du rapport peut être recalculé par un tiers.** Les scripts sont conservés dans le dossier de la mission et rejoués à chaque nouvel envoi."],
        source="Source : dépôt de la mission (dossier scripts)")
for etape, (entree, traitement, controle, sortie) in [
        ("Réception", ("4 287 fichiers envoyés par le FPPN, dont des exports de logiciels et des PDF.", "Le script dresse l'inventaire, repère doublons, fichiers vides et formats mal nommés.", "Chaque fichier reçoit une empreinte SHA-256 qui prouve qu'il n'a pas changé.", "Inventaire des fichiers et liste des fichiers à renvoyer.")),
        ("Extraction", ("Tableurs Excel, pages web nommées « .xls » et fichiers texte de la Solde.", "Le script lit les premiers octets pour reconnaître le vrai format, puis convertit chaque feuille en CSV sans perte.", "Les lignes et les cellules sont comptées avant et après conversion : 491 feuilles identiques.", "Fichiers CSV et journal des feuilles converties.")),
        ("Base", ("Fichiers CSV des précomptes, des balances et du Grenier.", "Le script reconnaît les colonnes par leur en-tête ou par leur contenu, puis réunit 4,45 millions de lignes.", "Les lignes de total et les montants doublés sont écartés et enregistrés à part.", "Base d'analyse au format Parquet et journal des rejets.")),
        ("Contrôles", ("Base d'analyse et totaux imprimés par les logiciels.", "Le script recalcule les totaux et rapproche les fichiers aller et retour, poste par poste et mois par mois.", "Les 37 balances âgées retrouvent leurs totaux et les 7 exercices sont équilibrés.", "Tableaux de contrôle et anomalies signalées au FPPN.")),
        ("Analyses", ("Base contrôlée, données externes validées et textes de l'UEMOA.", "Les scripts calculent les indicateurs, les tests statistiques, les tables CIMA et les 36 ratios prudentiels.", "Les agrégats recalculés sont comparés à ceux de la version précédente au franc près.", "Fichiers de résultats et classeur des calculs.")),
        ("Rapport", ("Fichiers de résultats et gabarits Timesaver Adenka.", "Le script construit les graphiques et tableaux natifs, les schémas modifiables et les notes.", "Chaque chiffre affiché est relu dans le PowerPoint produit et comparé à sa valeur exacte.", "Rapport PowerPoint et onglet de traçabilité."))]:
    gab(13, {"titre": f"Chaîne de traitement : {etape.lower()}", "sous_titre": f"L'étape « {etape} » transforme une entrée en sortie, avec un contrôle automatique",
             "synthese": "Entrée, traitement, contrôle et sortie de l'étape ; schéma modifiable, complété au fil des entretiens"},
        {"cartes": [("Entrée", entree), ("Traitement", traitement), ("Contrôle", controle), ("Sortie", sortie)]},
        f"Étape « {etape} » de la chaîne de traitement présentée en partie 2. Chaque étape correspond à un script conservé dans le dossier de la mission.",
        source="Source : scripts Python Adenka (dossier scripts de la mission)")
for nomimg, titre, sous, synth, expl, impl in [
        ("code_extraction", "Code : reconnaître le vrai format", "Le script lit les premiers octets pour reconnaître le vrai format d'un fichier", "C'est ainsi que les 70 fichiers « .xls » du Grenier ont été reconnus comme des pages web lisibles",
         ["**Un fichier Excel moderne commence par les octets « PK », un ancien par d'autres octets.**", "**Une page web contient « <html » ou « <table ».** Le script la lit alors comme un tableau.", "**Le terminal montre la sortie réelle de l'extraction sur le poste de la mission** : 91,5 Mo en 5 parties."],
         "Ce test évite de déclarer illisible un fichier simplement mal nommé : aucun fichier du Grenier n'a été perdu."),
        ("code_total", "Code : écarter les lignes de total", "Une ligne égale à la somme des autres est écartée automatiquement", "Règle appliquée feuille par feuille, avec une tolérance de 1 %",
         ["**Le script calcule la somme de la feuille.**", "**Une ligne dont le montant égale la somme des autres est un total**, pas un mutualiste.", "**Chaque ligne écartée est enregistrée** avec son fichier et son montant."],
         "La règle est automatique et tracée : le FPPN peut vérifier chaque ligne écartée dans le journal des rejets."),
        ("code_double", "Code : corriger les montants doublés", "Les montants écrits deux fois à la suite sont corrigés et signalés", "Correction limitée au poste décès, au-delà d'un million de FCFA",
         ["**Le script découpe le montant en deux moitiés.**", "**Si les deux moitiés sont identiques, il retient la première.**", "**Le montant d'origine est conservé** dans une colonne à part."],
         "La correction est réversible : le montant d'origine reste disponible pour le FPPN."),
        ("code_grenier", "Code : lire les montants du Grenier", "Les montants imprimés avec espaces et « XOF » sont convertis sans perte", "Les totaux recalculés sont égaux aux totaux imprimés dans les 37 balances âgées",
         ["**Le script retire les espaces, y compris insécables.**", "**Il retire aussi la mention « XOF ».**", "**Le nombre obtenu est comparé au total imprimé** par le logiciel."],
         "L'égalité des totaux sur les 37 balances prouve que la lecture n'a rien perdu."),
        ("code_sc", "Code : calculer la sinistralité", "Le ratio S/C est calculé compte par compte, pour chaque exercice", "Chaque compte utilisé est cité dans l'annexe des méthodes",
         ["**Les cotisations viennent des comptes 7010 et 7021.**", "**Les prestations viennent des comptes 601, 6013, 6018, 60800007, 6021 et 6024.**", "**Le résultat est écrit dans le classeur des calculs.**"],
         "Le choix des comptes est explicite dans le code : un lecteur peut le contester et le modifier."),
        ("code_stats", "Code : intervalle de confiance", "L'intervalle de Wilson encadre la proportion de prêts en retard", "Méthode adaptée aux proportions, même faibles",
         ["**La proportion vaut le nombre de prêts en retard divisé par le nombre de prêts.**", "**L'intervalle est centré et élargi** pour rester entre 0 et 1.", "**z vaut 1,96 pour un niveau de confiance de 95 %.**"],
         "L'intervalle dit dans quelle plage se situe la vraie proportion de prêts en retard, compte tenu du hasard."),
        ("code_statut", "Code : statut d'un ratio prudentiel", "Le statut de chaque ratio suit une règle écrite dans le code", "Conforme, non conforme, en attente ou à apprécier, selon la norme du texte",
         ["**Une valeur absente donne « En attente » ; une norme non chiffrée donne « À apprécier ».**", "**Le test compare la valeur à la norme** : fourchette, minimum ou maximum.", "**Les années 2020, 2021 et 2026 portent la mention « indicatif ».** Le terminal montre la sortie réelle du script."],
         "La règle de statut est la même pour les 36 indicateurs : elle ne laisse aucune place à l'appréciation au cas par cas.")]:
    contenu({"titre": titre, "sous_titre": sous, "synthese": synth},
            "Extrait réel du script, placé dans une fenêtre de VS Code sous Windows reconstituée (thème sombre, fil d'Ariane du dossier de la mission, numéros de ligne). Le code complet est dans le dossier scripts de la mission ; il peut être relu et rejoué par un tiers.",
            image={"fichier": os.path.join(IMG, nomimg + ".png"), "legende": fig(f"Extrait du script dans VS Code sous Windows ({titre.split(': ')[1]})")}, analyse=expl, implication=impl,
            source="Source : scripts de la mission (fenêtre VS Code reconstituée, pas une capture d'écran)")
methode_gab({"titre": "Contrôle automatique des chiffres", "sous_titre": "Chaque chiffre du rapport est recalculé et comparé au texte affiché",
             "synthese": "Le contrôle relit le PowerPoint produit et vérifie que chaque chiffre enregistré figure sur sa diapositive"},
            [("Enregistrement", "À la construction, chaque chiffre est produit par une fonction qui garde sa valeur exacte, sa diapositive et le texte affiché."),
             ("Relecture", "Après la construction, un script relit le fichier PowerPoint, textes, tableaux et notes compris."),
             ("Comparaison", "Le script vérifie la présence de chaque texte enregistré sur sa diapositive et compte les chiffres retrouvés."),
             ("Résultat", "Le résultat figure dans l'onglet « Traçabilité » du classeur des calculs, ligne par ligne, avec la valeur exacte de chaque chiffre.")],
            "Principe : aucun chiffre n'est saisi à la main. Le contrôle garantit que le rapport affiche exactement les valeurs calculées, avec deux décimales pour les pourcentages et des montants exacts en FCFA.")

# ---------------------------------------------------------------- PARTIE 14 : ANALYSES EN ATTENTE ET RISQUES
separateur(14, "Analyses en attente", "Documents à obtenir, risques, points à clarifier")
contenu({"titre": "Analyses en attente", "sous_titre": "Les analyses actuarielles centrales dépendent de trois bases non reçues",
         "synthese": "Provisions, mortalité, équilibre technique, contrôles santé et onze indicateurs prudentiels seront réalisés dès réception des documents"},
        "Liste détaillée et à jour dans la PBC pointée.",
        tableau={"titre": tab("Analyses en attente et documents nécessaires"), "entetes": ["Analyse", "Objet", "Documents nécessaires (PBC)"], "largeurs": [7.2, 11.6, 12.4], "taille": 10,
                 "lignes": [["Provisions (PSAP, Chain Ladder)", "Montant restant à payer, y compris non déclaré", "Prestations décès et santé, dossiers non payés (13, 14, 28, 29, 40)"],
                            ["Sinistralité par année de survenance", "Coût réel de chaque année", "Bases des prestations (13, 28, 40)"],
                            ["Mortalité observée (A/E)", "Comparaison à CIMA H", "Mutualistes avec date de naissance (9), décès (28)"],
                            ["Délais et dossiers en attente", "Délais de déclaration et de paiement", "Bases des prestations avec dates (28, 40)"],
                            ["Contrôles santé", "Doublons, bénéficiaires hors fichier, prestations indues", "Factures santé (40), cartes actives (42)"],
                            ["Engagements d'épargne", "Dette envers chaque épargnant et échéancier", "Comptes individuels et règlements (16, 17)"],
                            ["Indicateurs prudentiels en attente", "Recouvrement, contrôles, réunions, fonds de garantie", "Fichiers aller 936 et 8Z6, procès-verbaux, registres"],
                            ["Équilibre technique", "Projection sur 5 à 10 ans", "Âges, barèmes, hypothèses (1, 2, 3, 9)"]]},
        disposition="dessous",
        analyse=["**Les analyses en attente portent sur le cœur de la mission : les engagements et les provisions.**"],
        source="Source : PBC pointée au 18/10/2026")
contenu({"titre": "Points à clarifier", "sous_titre": "Douze points demandent une réponse du FPPN avant la prochaine version",
         "synthese": "Chaque point est relié à un constat chiffré du rapport et repris dans les guides d'entretien"},
        "Ces questions peuvent être posées lors des prochains entretiens ; les réponses seront intégrées à la version suivante du rapport.",
        tableau={"titre": tab("Points à clarifier avec le FPPN"), "entetes": ["N°", "Point", "Question au FPPN"], "largeurs": [1.0, 7.8, 22.4], "taille": 9,
                 "lignes": [["1", "Répartition du taux de 5 %", "Quel document fixe le barème : 3,5 / 1,5 ou 3,77 / 1,23 ? Pourquoi la comptabilité suit-elle une autre clé ?"],
                            ["2", "Provisions figées", "Comment les provisions ont-elles été calculées, et pourquoi n'ont-elles pas changé depuis 2020 ?"],
                            ["3", "Couverture de l'épargne", "Quelles sont les conditions de retrait et les échéances du PEP et de l'EFC ?"],
                            ["4", "Créance sur l'Hôpital de la Police", "Quelle est sa nature, et un remboursement est-il prévu ?"],
                            ["5", "Retards du Grenier début 2026", "Pourquoi la plus grande part de l'encours était-elle en retard fin janvier 2026 ?"],
                            ["6", "Écart échéances / précomptes", "Comment les échéances non précomptées sont-elles réglées ?"],
                            ["7", "Sinistralité 2020 à 2022", "Des factures de ces années ont-elles été comptabilisées plus tard ?"],
                            ["8", "Agrément AIRMS", "Quels sont le numéro, la date et les risques autorisés de l'agrément du FPPN ?"],
                            ["9", "Sigle CMP et réserve disponible", "Que signifie le sigle CMP, et quelle est l'origine du compte 10620000 « Réserve disponible » ?"],
                            ["10", "Fonds national de garantie", "Le FPPN adhère-t-il au FOMUS-CI, et à quel montant de cotisation ?"],
                            ["11", "Risques longs", "Une étude actuarielle et une certification des comptes ont-elles été réalisées ?"],
                            ["12", "Données manquantes", "Quand le FPPN transmettra-t-il les précomptes de 2020 à 2023, les fichiers de décembre 2024 et les grands livres endommagés ?"]]},
        disposition="dessous",
        analyse=["**Les points 8 à 11 sont nouveaux : ils découlent de la lecture des textes de l'AIRMS.**"],
        source="Source : analyses Adenka ; guides d'entretien mis à jour")

# ---------------------------------------------------------------- ANNEXES
SL.append({"type": "separateur", "num": "A", "titre": "Annexes", "sous_titre": "Méthodes, listes, abréviations, références", "notes": ""})
SL.append({"type": "annexe", "titre": "Méthodes et formules", "notes": "Formules exactes utilisées dans le rapport.", "tableau": {"titre": tab("Indicateurs, formules et sources"),
           "entetes": ["Indicateur", "Formule", "Source des données"], "largeurs": [6.0, 14.6, 10.6], "taille": 8,
           "lignes": [["Ratio S/C", "Prestations comptabilisées / cotisations comptabilisées, par garantie et par exercice", "Comptes 601, 6013, 6018, 60800007, 6021, 6024 ; 7010, 7021"],
                      ["Taux de cotisation", "Précompte du mois / salaire de base du mois, par agent", "Fichiers 936 détaillés, 8Z6"],
                      ["Taux de précompte", "Montant retour / montant aller, mois communs", "Fichiers aller et retour"],
                      ["TCAM", "(V_fin / V_début)^(1/n) - 1", "Selon l'indicateur"],
                      ["Tendance", "Régression linéaire par moindres carrés, test de Student sur la pente, IC à 95 %", "Séries mensuelles ou annuelles"],
                      ["Intervalle de Wilson", "(p + z²/2n ± z √(p(1-p)/n + z²/4n²)) / (1 + z²/n), z = 1,96", "Balance âgée du 31/08/2026"],
                      ["Valeur réelle", "Valeur nominale × 100 / indice des prix (base 100 en 2020)", "Banque mondiale, FP.CPI.TOTL.ZG"],
                      ["Couverture de l'épargne", "Actifs réservés / (comptes PEP + EFC)", "Balances des comptes"],
                      ["Sinistralité réglementaire", "(60 + 6901 - 7901) / (70 + 7906 - 6906)", "Décision n° 05/2022, § 1.2.1.2.5"],
                      ["Marge de solvabilité exigée", "max(18 % des cotisations acquises n-1, 26 % des prestations maladie moyennes sur 3 ans)", "Décision n° 05/2022, § 1.1.1"],
                      ["Réserve obligatoire", "Réserves obligatoires x 12 / charges de prestations", "Décision n° 05/2022, tableau 1"],
                      ["Liquidité", "Trésorerie actif / (dettes à court terme + variation des provisions techniques)", "Décision n° 05/2022, § 1.2.2.5.2"],
                      ["Équilibre stable", "Emplois moyens et longs / ressources moyennes et longues", "Décision n° 05/2022, § 1.2.3.1"],
                      ["Rente viagère ä_x", "Σ (L_(x+k) / L_x) × (1,035)^(-k)", "Tables CIMA"],
                      ["Espérance de vie e_x", "Σ L_(x+k) / L_x pour k ≥ 1, + 0,5", "Tables CIMA"]]},
           "source": "Source : Adenka ; " + DEC})
SL.append({"type": "annexe", "titre": "Outils et bibliothèques", "notes": "", "tableau": {"titre": tab("Outils et bibliothèques utilisés"),
           "entetes": ["Outil", "Usage", "Version"], "largeurs": [7.0, 18.2, 6.0],
           "lignes": [["Python", "Langage des traitements, scripts conservés et rejouables", "3.14 (poste de la mission)"],
                      ["Visual Studio Code", "Éditeur et exécution des scripts sous Windows", "dernière"],
                      ["pandas, numpy", "Tables de données, calculs, agrégations", "3.0 et 2.5"],
                      ["statsmodels", "Régressions et tests statistiques", "0.15"],
                      ["python-calamine, openpyxl, lxml", "Lecture sans perte des fichiers Excel et des pages web", "dernière"],
                      ["pyarrow", "Stockage de la base d'analyse (Parquet)", "26.0"],
                      ["pymupdf, pdftotext, tesseract", "Lecture des PDF comptables et des textes scannés (OCR)", "dernière"],
                      ["pygments, Pillow", "Mise en forme des extraits de code", "dernière"],
                      ["python-pptx", "Rapport, graphiques et tableaux natifs", "1.0"]]},
           "source": "Source : Adenka"})
SL.append({"type": "liste_figures", "notes": ""})
SL.append({"type": "liste_tableaux", "notes": ""})
_sig = sorted({(nom, defi) for nom, _, defi in SIGLES}, key=lambda t: t[0].lower())
_moit = (len(_sig) + 1) // 2
SL.append({"type": "annexe", "titre": "Signes et abréviations", "notes": "Les sigles de chaque diapositive sont aussi définis en bas de la diapositive.", "tableau": {"titre": None,
           "entetes": ["Sigle", "Signification", "Sigle", "Signification"], "largeurs": [2.6, 13.0, 2.6, 13.0], "taille": 8, "hauteur_ligne": 0.5,
           "lignes": [[_sig[i][0], _sig[i][1], _sig[i + _moit][0] if i + _moit < len(_sig) else "", _sig[i + _moit][1] if i + _moit < len(_sig) else ""] for i in range(_moit)]},
           "source": None})
SL.append({"type": "annexe", "titre": "Références et sources", "notes": "", "tableau": {"titre": None,
           "entetes": ["Référence", "Contenu utilisé", "Lien ou emplacement"], "largeurs": [10.2, 9.8, 11.2], "taille": 8,
           "lignes": [["Envoi du FPPN du 15/10/2026", "Précomptes, balances, exports du Grenier", "Dossier 04-Documents reçus du client"],
                      ["Grands livres 2021, 2022, 2024, 2025 (PDF)", "Détail des écritures", "Envoi complémentaire du 18/10/2026"],
                      ["Règlement n° 07/2009/CM/UEMOA du 26/06/2009", "Cadre de la mutualité sociale", "airms-ci.com, documents publics"],
                      ["Règlement d'exécution n° 004/2022/COM/UEMOA du 01/06/2022", "Règles prudentielles, garantie, contrôle", "airms-ci.com, documents publics"],
                      ["Décision n° 05/2022/COM/UEMOA du 01/06/2022", "Indicateurs et ratios prudentiels", "airms-ci.com, documents publics"],
                      ["Règlement n° 03/2012/CM/UEMOA et PCMS (mai 2012)", "Plan comptable des mutuelles sociales", "airms-ci.com, documents publics"],
                      ["Arrêté n° 00115/MEPS/CAB du 12/12/2023", "Sanctions administratives", "airms-ci.com, documents publics"],
                      ["Arrêtés n° 092 et 093/MEPS/CAB du 06/09/2023", "Agrément, immatriculation, frais", "airms-ci.com, documents publics"],
                      ["Répertoire AIRMS des mutuelles agréées (15/10/2026)", "Agrément et immatriculation du FPPN", "airms-ci.com, documents publics"],
                      ["Ordonnance n° 2024-1043", "FOMUS-CI", "civlii.laws.africa"],
                      ["Code CIMA, article 338", "Tables et taux technique de 3,5 %", "cima-afrique.org"],
                      ["Tables CIMA H, CIMA F, TD, TV ; table ronde CIMA FANAF (2011)", "Mortalité de référence", "Fichiers fournis le 18/10/2026 ; fanaf.org"],
                      ["BCEAO, communiqués du CPM 2020 à 2025", "Taux directeurs", "bceao.int"],
                      ["Banque mondiale, FP.CPI.TOTL.ZG", "Inflation 2020 à 2025", "donnees.banquemondiale.org"],
                      ["Classeur des calculs v3", "Tableaux, séries, ratios, classement des comptes, traçabilité", "FPPN_Calculs_rapport_20261015_v3.xlsx"]]},
           "source": None})
SL.append({"type": "merci", "notes": ""})

# ===========================================================================
# CONSTRUCTION
# ===========================================================================
GAB = {"couverture": 1, "sommaire": 5, "separateur": 2, "contenu": 5, "annexe": 5, "liste_figures": 5, "liste_tableaux": 5, "merci": 110}
tmp = SORTIE.replace(".pptx", "_tmp.pptx")
selection(TIMESAVER, tmp, [x["gab"] if x["type"] == "gab" else GAB[x["type"]] for x in SL])


def paragraphes(x):
    """Paragraphes d'analyse d'une diapositive : constat et preuves, puis l'implication dans un paragraphe distinct."""
    paras = list(x.get("analyse") or x.get("texte") or [])
    if x.get("implication"):
        paras.append("**Implication pour le FPPN.** " + x["implication"])
    return paras


def longueur(paras):
    return sum(len(p.replace("**", "")) for p in paras)


def disposition_graphique(paras):
    L = longueur(paras)
    if L <= 650:
        return 18.2, 20.2, 12.4
    if L <= 1000:
        return 16.6, 18.6, 14.0
    return 15.0, 17.0, 15.6


CJ = {}
for i, x in enumerate(SL, 1):
    k = str(i)
    if x["type"] == "couverture":
        CJ[k] = {"Google Shape;138;p25": "Rapport actuariel\net statistique\ndu FPPN", "Text Placeholder 4": "Rapport d'étape au 15 octobre 2026, version 3"}
    elif x["type"] == "separateur":
        CJ[k] = {"Rectangle 24": str(x["num"]), "Rectangle 9": x["titre"], "Rectangle 5": x["sous_titre"]}
    elif x["type"] in ("sommaire", "liste_figures", "liste_tableaux", "annexe"):
        CJ[k] = {"@vider": True}
    elif x["type"] == "gab":
        # Les gabarits mis à l'échelle par compresser() ne doivent pas être réduits une première fois par l'en-tête
        CJ[k] = {"@entete": dict(x["entete"], ajuster=x["gab"] not in (9, 13, 29, 31, 39, 41))}
        if x.get("source"):
            CJ[k]["@source"] = x["source"]
    elif x["type"] == "contenu":
        d = {"@vider": True, "@entete": x["entete"]}
        if "graphique" in x:
            lg, _, _ = disposition_graphique(paragraphes(x))
            g = dict(x["graphique"])
            d["@texte1"] = {"texte": g.pop("titre"), "style": "titre_graphique", "x": 1.4, "y": 3.85, "largeur": lg, "hauteur": 0.7}
            if g.get("unite"):
                d["@texte3"] = {"texte": g.pop("unite"), "style": "unite", "x": 1.4, "y": 4.45, "largeur": lg, "hauteur": 0.45}
            g.update({"x": 1.4, "y": 4.95, "largeur": lg, "hauteur": BAS_CONTENU - 4.95})
            d["@graphique"] = g
        if "graphiques" in x:
            hg = (13.4 - 4.95) if x.get("bandeau") else (BAS_CONTENU - 4.95)
            for j, gg in enumerate(x["graphiques"]):
                g = dict(gg)
                xx = 1.4 + j * 16.0
                d[f"@texte1{j}"] = {"texte": g.pop("titre"), "style": "titre_graphique", "x": xx, "y": 3.85, "largeur": 15.2, "hauteur": 0.7}
                if g.get("unite"):
                    d[f"@texte3{j}"] = {"texte": g.pop("unite"), "style": "unite", "x": xx, "y": 4.45, "largeur": 15.2, "hauteur": 0.45}
                g.update({"x": xx, "y": 4.95, "largeur": 15.2, "hauteur": hg})
                d[f"@graphique{j}"] = g
        if x.get("source"):
            d["@source"] = x["source"]
        CJ[k] = d
remplir(tmp, CJ, tmp)

prs = Presentation(tmp)
ts = Presentation(TIMESAVER)
icones = [sh for sh in ts.slides[110].shapes if sh.shape_type == 13 and sh.width < Cm(3)]


def icone_blanche(blob):
    from PIL import Image
    im = Image.open(io.BytesIO(blob)).convert("RGBA")
    blanc = Image.new("RGBA", im.size, (255, 255, 255, 0))
    alpha = im.split()[3]
    if alpha.getextrema() == (255, 255):
        alpha = im.convert("L").point(lambda v: 255 if v < 128 else 0)
    blanc.putalpha(alpha)
    out = io.BytesIO()
    blanc.save(out, "PNG")
    return out.getvalue()


def txt(slide, x, y, l, h, texte, taille=14, gras=False, couleur=NOIR, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(l), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.15)
    for i, ligne in enumerate(str(texte).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ligne
        r.font.name = "Helvetica"
        r.font.size = Pt(taille)
        r.font.bold = gras
        r.font.color.rgb = RGBColor.from_string(couleur)
    return tb


def hauteur_texte(paras, largeur, taille, espace=6):
    cpl = max(1.0, (largeur - 0.3) / (0.0181 * taille))
    lignes = sum(math.ceil(len(p.replace("**", "")) / cpl) for p in paras)
    return lignes * taille * 1.22 * 0.03528 + len(paras) * espace * 0.03528


def corps(slide, x, y, l, h, paras, tmax=14, tmin=9):
    """Texte d'analyse justifié, **gras** accepté ; la taille est la plus grande qui tient dans la zone."""
    if not paras or h < 0.6:
        return None
    taille = tmin
    for t in (14, 13, 12, 11.5, 11, 10.5, 10, 9.5, 9):
        if t <= tmax and hauteur_texte(paras, l, t) <= h:
            taille = t
            break
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(l), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Cm(0.15)
    for i, bloc in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.JUSTIFY
        p.space_after = Pt(6)
        for morceau in re.split(r"(\*\*.+?\*\*)", bloc):
            if not morceau:
                continue
            r = p.add_run()
            gras = morceau.startswith("**") and morceau.endswith("**")
            r.text = morceau[2:-2] if gras else morceau
            r.font.name = "Helvetica"
            r.font.size = Pt(taille)
            r.font.bold = gras
            r.font.color.rgb = RGBColor.from_string(NOIR)
    return tb


def boite(slide, x, y, l, h, texte, fond=BLEU, couleur="FFFFFF", taille=12, gras=True, forme=MSO_SHAPE.ROUNDED_RECTANGLE, contour=None):
    sh = slide.shapes.add_shape(forme, Cm(x), Cm(y), Cm(l), Cm(h))
    if fond:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fond)
    else:
        sh.fill.background()
    if contour:
        sh.line.color.rgb = RGBColor.from_string(contour)
        sh.line.width = Pt(1.5)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Cm(0.15)
    for i, ligne in enumerate(str(texte).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = ligne
        r.font.name = "Helvetica"
        r.font.size = Pt(taille if i == 0 else max(taille - 2, 9))
        r.font.bold = gras and i == 0
        r.font.color.rgb = RGBColor.from_string(couleur)
    return sh


def fleche(slide, x1, y1, x2, y2, couleur=GRIS_F, epaisseur=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = RGBColor.from_string(couleur)
    c.line.width = Pt(epaisseur)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


COUL_STATUT = [("Conforme", BLEU, "FFFFFF"), ("Non conforme", ROUGE, "FFFFFF"), ("Au-dessus", ORANGE, NOIR),
               ("En attente", GRIS_C, NOIR), ("À apprécier", GRIS_C, NOIR)]


def hauteurs_lignes(spec, taille):
    hs = []
    for i, vals in enumerate([spec["entetes"]] + spec["lignes"]):
        h = 0
        for v, w in zip(vals, spec["largeurs"]):
            cpl = max(1.0, (w - 0.3) / (0.0186 * taille))
            n = max(1, math.ceil(len(str(v)) / cpl))
            h = max(h, n * taille * 1.2 * 0.03528 + 0.16)
        hs.append(max(h, spec.get("hauteur_ligne") or 0.5))
    return hs


def tableau(slide, spec, y0=3.85, x0=1.4, taille=None, hauteur_max=None):
    """Tableau natif ; renvoie le bas du tableau (cm). La police diminue si le tableau dépasse la hauteur disponible."""
    y = y0
    if spec.get("titre"):
        txt(slide, x0, y, sum(spec["largeurs"]), 0.6, spec["titre"], 12, False, ORANGE, PP_ALIGN.CENTER)
        y += 0.75
    lignes, larg = spec["lignes"], spec["largeurs"]
    taille = taille or spec.get("taille") or (11 if len(lignes) <= 8 else 10)
    hmax = (hauteur_max or (BAS_CONTENU - y))
    hs = hauteurs_lignes(spec, taille)
    while sum(hs) > hmax and taille > 7.5:
        taille -= 0.5
        hs = hauteurs_lignes(spec, taille)
    n = len(lignes) + 1
    shape = slide.shapes.add_table(n, len(larg), Cm(x0), Cm(y), Cm(sum(larg)), Cm(sum(hs)))
    t = shape.table
    tblPr = t._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblPr.set(attr, "0")
    st = tblPr.find(qn("a:tableStyleId"))
    if st is not None:
        st.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"
    for j, w in enumerate(larg):
        t.columns[j].width = Cm(w)
    for i in range(n):
        t.rows[i].height = Cm(hs[i])
        vals = spec["entetes"] if i == 0 else lignes[i - 1]
        for j, v in enumerate(vals):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Cm(0.12)
            c.margin_top = c.margin_bottom = Cm(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            p.text = ""
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", "0")
            pPr.set("indent", "0")
            pPr.set("lvl", "0")
            for bn in pPr.findall(qn("a:buNone")):
                pPr.remove(bn)
            pPr.insert(0, pPr.makeelement(qn("a:buNone"), {}))
            r = p.add_run()
            r.text = str(v)
            r.font.name = "Helvetica"
            r.font.size = Pt(taille)
            r.font.bold = (i == 0) or (j == 0)
            fond, coul = (NOIR, "FFFFFF") if i == 0 else ((GRIS if i % 2 == 0 else "FFFFFF"), NOIR)
            if i > 0 and spec.get("statut_col") == j:
                for debut, f_, c_ in COUL_STATUT:
                    if str(v).startswith(debut):
                        fond, coul = f_, c_
                        r.font.bold = True
                        break
            r.font.color.rgb = RGBColor.from_string(coul)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string(fond)
    return y + sum(hs)


def tableau_avec_analyse(slide, x):
    spec, paras = dict(x["tableau"]), paragraphes(x)
    if not paras:
        tableau(slide, spec)
        return
    if x.get("disposition") == "droite":
        larg = spec["largeurs"]
        if sum(larg) > 19.6:
            f_ = 19.6 / sum(larg)
            spec["largeurs"] = [round(w * f_, 2) for w in larg]
        bas = tableau(slide, spec)
        xt = 1.4 + sum(spec["largeurs"]) + 0.6
        corps(slide, xt, 3.85, 32.6 - xt, BAS_CONTENU - 3.85, paras)
    else:
        besoin = max(1.6, min(4.2, hauteur_texte(paras, 31.2, 11) + 0.2))
        bas = tableau(slide, spec, hauteur_max=BAS_CONTENU - 3.85 - 0.75 - besoin - 0.3)
        corps(slide, 1.4, bas + 0.3, 31.2, BAS_CONTENU - bas - 0.3, paras, 13)


def pastille(slide, x, y, num):
    return boite(slide, x, y, 0.75, 0.75, str(num), ROUGE, "FFFFFF", 11, True, MSO_SHAPE.OVAL)


def pas_a_pas(slide, spec, paras):
    lignes = spec["lignes"]
    taille = 10 if not spec.get("petit") else 9
    txt(slide, 1.4, 3.75, sum(spec["largeurs"]), 0.5, "Extrait : " + spec["fichier"], 9, False, GRIS_F)
    t_spec = {"titre": None, "entetes": spec["entetes"], "lignes": lignes, "largeurs": spec["largeurs"], "hauteur_ligne": 0.55}
    y_tab = 4.3
    hs = hauteurs_lignes(t_spec, taille)
    bas = tableau(slide, t_spec, y_tab, 1.4, taille)
    larg = spec["largeurs"]
    for k, (ri, cj) in enumerate(spec["rouge"], 1):
        xx = 1.4 + sum(larg[:cj])
        yy = y_tab + sum(hs[:ri + 1])
        cadre = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(xx - 0.05), Cm(yy - 0.03), Cm(larg[cj] + 0.1), Cm(hs[ri + 1] + 0.06))
        cadre.fill.background()
        cadre.line.color.rgb = RGBColor.from_string(ROUGE)
        cadre.line.width = Pt(2.25)
        pastille(slide, xx + larg[cj] - 0.5, yy - 0.45, k)
    xr = 1.4 + sum(larg) + 0.6
    lr = 33.0 - xr - 0.4
    if lr < 7:
        xr, lr = 1.4, sum(larg)
        y = bas + 0.35
    else:
        y = 4.3
    for k, e in enumerate(spec["explications"], 1):
        h = max(0.8, hauteur_texte([e], lr - 1.0, 12) + 0.1)
        pastille(slide, xr, y, k)
        txt(slide, xr + 1.0, y - 0.05, lr - 1.0, h, e, 12)
        y += h + 0.25
    if lr == sum(larg):           # explications sous le tableau : l'analyse suit, pleine largeur
        corps(slide, 1.4, y + 0.1, 31.2, BAS_CONTENU - y - 0.1, paras, 12)
    else:
        y_an = max(y + 0.2, 4.3)
        if BAS_CONTENU - y_an >= 2.0:
            corps(slide, xr, y_an, lr, BAS_CONTENU - y_an, paras, 12)
        else:
            corps(slide, 1.4, bas + 0.3, sum(larg), BAS_CONTENU - bas - 0.3, paras, 12)


def bandeau_attente(slide, y=15.15):
    boite(slide, 1.4, y, 31.2, 0.95, "Travail en attente des données : voir la liste des données manquantes (PBC pointée et partie 14)", ORANGE, NOIR, 13)


def triangle(slide):
    ent = ["Survenance", "Délai 1", "Délai 2", "Délai 3", "Délai 4"]
    lignes = [["Année n-3", "C(1,1)", "C(1,2)", "C(1,3)", "C(1,4)"], ["Année n-2", "C(2,1)", "C(2,2)", "C(2,3)", "?"],
              ["Année n-1", "C(3,1)", "C(3,2)", "?", "?"], ["Année n", "C(4,1)", "?", "?", "?"]]
    tableau(slide, {"titre": None, "entetes": ent, "lignes": lignes, "largeurs": [4.6, 3.6, 3.6, 3.6, 3.6], "hauteur_ligne": 1.1}, 4.3, 1.4, 14)
    txt(slide, 1.4, 10.4, 19.0, 4.4, "C(i, j) : paiements cumulés de l'année de survenance i, au délai j.\n« ? » : montants futurs à estimer.\nFacteur de développement : f(j) = Σ C(i, j+1) / Σ C(i, j).\nCharge ultime = dernier montant connu × produit des facteurs suivants.\nProvision = charge ultime - paiements déjà faits.", 13)
    boite(slide, 21.6, 4.3, 11.0, 4.9, "Hypothèse\nLe rythme de règlement passé se répète dans le futur.", GRIS, NOIR, 15, True, MSO_SHAPE.ROUNDED_RECTANGLE)
    boite(slide, 21.6, 9.6, 11.0, 4.9, "Données\nPaiements par année de survenance et par délai, issus des bases de prestations (PBC 13, 28, 40).", GRIS, NOIR, 15, True, MSO_SHAPE.ROUNDED_RECTANGLE)
    bandeau_attente(slide)


def schema(slide, nom, paras):
    if nom == "fonctionnement":
        boite(slide, 1.4, 6.9, 6.0, 3.2, "Mutualistes\nactifs et retraités", BLEU, "FFFFFF", 16)
        boite(slide, 9.0, 4.4, 6.2, 2.4, "Solde (actifs)", GRIS, NOIR, 14, contour=BLEU)
        boite(slide, 9.0, 10.2, 6.2, 2.4, "CGRAE (retraités)", GRIS, NOIR, 14, contour=BLEU)
        boite(slide, 16.8, 6.5, 6.6, 4.0, "FPPN\nCotisations " + F("Cot (schéma)", s[2025]["cot_maladie"] + s[2025]["cot_deces"]), ORANGE, NOIR, 16)
        sorties = [("Prestations maladie\n" + F("Prest mal (schéma)", s[2025]["prest_maladie"]), 5.0),
                   ("Prestations décès\n" + F("Prest déc (schéma)", s[2025]["prest_deces"]), 7.6),
                   ("Épargne PEP et EFC\n" + F("Épargne (schéma)", b[2026]["epargne_totale"]), 10.2),
                   ("Prêts du Grenier\n" + F("Grenier (schéma)", b[2026]["grenier_financement"]), 12.8),
                   ("CMP et terrains\n" + F("CMP+terrains (schéma)", b[2026]["cmp_investissements"] + b[2026]["terrains"]), 15.1)]
        for t_, yy in sorties:
            boite(slide, 25.2, yy - 1.05, 7.4, 2.1, t_, BLEU, "FFFFFF", 12)
            fleche(slide, 23.4, 8.5, 25.2, yy, GRIS_F)
        fleche(slide, 7.4, 7.9, 9.0, 5.6, GRIS_F)
        fleche(slide, 7.4, 9.1, 9.0, 11.4, GRIS_F)
        fleche(slide, 15.2, 5.6, 16.8, 7.7, GRIS_F)
        fleche(slide, 15.2, 11.4, 16.8, 9.3, GRIS_F)
        corps(slide, 1.4, 13.2, 22.6, BAS_CONTENU - 13.2, paras, 12)
    elif nom == "precomptes":
        etapes_ = [("1. FPPN", "Le FPPN prépare le fichier aller : qui précompter et combien."), ("2. Solde / CGRAE", "La Solde ou la CGRAE retient les montants sur le salaire ou la pension."),
                   ("3. Fichier retour", "Le fichier retour liste les montants réellement retenus."), ("4. FPPN", "Le FPPN rapproche aller et retour, comptabilise et relance les rejets.")]
        for k, (t_, c_) in enumerate(etapes_):
            x = 1.4 + k * 8.0
            boite(slide, x, 4.6, 6.6, 2.2, t_, BLEU if k % 2 == 0 else ORANGE, "FFFFFF" if k % 2 == 0 else NOIR, 16)
            txt(slide, x, 7.1, 6.6, 2.4, c_, 13, False, NOIR, PP_ALIGN.CENTER)
            if k < 3:
                fleche(slide, x + 6.6, 5.7, x + 8.0, 5.7, GRIS_F, 3)
        boite(slide, 1.4, 9.8, 31.2, 1.9, "Contrôle Adenka : taux de précompte = montant retour / montant aller, poste par poste et mois par mois (partie 3)", GRIS, NOIR, 14, contour=BLEU)
        corps(slide, 1.4, 12.2, 31.2, BAS_CONTENU - 12.2, paras, 13)


def compresser(slide, y0, y1, n0, n1, uniforme=False):
    """Ramène dans [n0, n1] les formes du gabarit situées dans [y0, y1] (cm) : place de l'en-tête et des sigles.
    En mode uniforme, chaque forme est réduite dans les deux sens autour du centre de la diapositive."""
    f_ = (n1 - n0) / (y1 - y0)
    for sh in slide.shapes:
        if sh.top is None:
            continue
        top, h = sh.top / 360000, sh.height / 360000
        if h > 18 or top < y0 - 0.05 or top > y1 or sh.name.startswith(("ZoneTexte", "Title", "Adenka_")):
            continue
        centre = top + h / 2
        if uniforme:
            l_, w_ = sh.left / 360000, sh.width / 360000
            cx = 16.93 + (l_ + w_ / 2 - 16.93) * f_
            cy = n0 + (centre - y0) * f_
            sh.width, sh.height = Cm(w_ * f_), Cm(max(h * f_, 0.01))
            sh.left, sh.top = Cm(cx - w_ * f_ / 2), Cm(cy - h * f_ / 2)
            continue
        if sh.shape_type in (1, 9, 17) and not sh.name.startswith("Google Shape"):
            nt, nh = n0 + (top - y0) * f_, h * f_
            sh.top, sh.height = Cm(nt), Cm(max(nh, 0.01))
        else:
            sh.top = Cm(n0 + (centre - y0) * f_ - h / 2)


def ecrire(sh, lignes_, tailles=None, gras=None):
    tf = sh.text_frame
    paras = [p for p in tf.paragraphs]
    for k, ligne in enumerate(lignes_):
        if k < len(paras):
            p = paras[k]
            runs = p.runs
            if runs:
                runs[0].text = ligne
                for r in runs[1:]:
                    r._r.getparent().remove(r._r)
            else:
                p.text = ligne
        else:
            p = tf.add_paragraph()
            p.text = ligne
        for r in p.runs:
            if tailles:
                r.font.size = Pt(tailles[min(k, len(tailles) - 1)])
            if gras is not None:
                r.font.bold = gras[min(k, len(gras) - 1)]
    for p in paras[len(lignes_):]:
        p._p.getparent().remove(p._p)


def remplir_gab(slide, num, d, analyse=None):
    formes = list(slide.shapes)

    def par_pos(pred):
        return sorted([sh for sh in formes if pred(sh)], key=lambda sh: (round(sh.top / 360000), sh.left))

    if num == 49:
        for nm, t_ in zip(("Freeform 7", "Freeform 9", "Freeform 11", "Freeform 13"), d["titres"]):
            for sh in formes:
                if sh.name == nm:
                    ecrire(sh, [t_])
        for sh, v in zip(par_pos(lambda sh: sh.name == "TextBox 31"), d["valeurs"]):
            ecrire(sh, [v])
            sh.text_frame.word_wrap = False
            sh.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
            for r in sh.text_frame.paragraphs[0].runs:
                r.font.size = Pt(21)
        for sh, v in zip(par_pos(lambda sh: sh.name == "TextBox 96"), d["commentaires"]):
            ecrire(sh, [v])
    elif num == 23:
        for sh, (t_, c_) in zip(par_pos(lambda sh: sh.name == "TextBox 49"), d["blocs"]):
            ecrire(sh, [t_, c_])
    elif num == 83:
        noms = {"CAUSE": "CAUSES", "EFFECT": "EFFETS"}
        for sh in formes:
            if sh.has_text_frame:
                t0 = sh.text_frame.text.strip()
                if t0 in noms:
                    ecrire(sh, [noms[t0]])
                elif t0.startswith("CAUSE 0"):
                    ecrire(sh, ["CAUSE " + t0[-1]])
                elif t0.startswith("EFFECT 0"):
                    ecrire(sh, ["EFFET " + t0[-1]])
        dummies = [sh for sh in formes if sh.has_text_frame and sh.text_frame.text.startswith("This is dummy")]
        cols = sorted(dummies, key=lambda sh: sh.left)
        paires = [sorted(cols[i:i + 2], key=lambda sh: sh.top) for i in range(0, len(cols), 2)]
        for (haut, bas), c_, e_ in zip(paires, d["causes"], d["effets"]):
            ecrire(haut, [c_])
            ecrire(bas, [e_])
    elif num in (39, 41):
        compresser(slide, 3.5, 16.2, 4.0, 14.9 if d.get("attente") else 16.1)
        formes = list(slide.shapes)
        nums = {sh.text_frame.text.strip(): sh for sh in formes if sh.has_text_frame and sh.text_frame.text.strip() in ("01", "02", "03", "04", "05", "06")}
        noms_r = ("Rectangle 18", "Rectangle 19", "Rectangle 21", "Rectangle 22", "Rectangle 28", "Rectangle 29")
        rects = [sh for sh in formes if sh.name in noms_r]
        for k, (t_, c_) in enumerate(d["etapes"], 1):
            nsh = nums[f"0{k}"]
            r = min(rects, key=lambda sh: abs(sh.top - nsh.top) + abs(sh.left - nsh.left))
            ecrire(r, [t_, c_], [14, 12] if num == 39 else [13, 11], [True, False])
        if d.get("attente"):
            bandeau_attente(slide)
    elif num == 9:
        for sh in formes:
            if sh.name == "Rectangle 27":
                ecrire(sh, [""])
        compresser(slide, 2.0, 17.2, 3.75, 14.85 if d.get("attente") else 16.25)
        formes = list(slide.shapes)
        blocs = par_pos(lambda sh: sh.name in ("Rectangle 22", "Rectangle 23", "Rectangle 32", "Rectangle 33"))
        for sh, (t_, c_) in zip(blocs, d["blocs"]):
            ecrire(sh, [t_, c_], [13, 11.5], [True, False])
            sh.text_frame.word_wrap = True
        if d.get("attente"):
            bandeau_attente(slide)
    elif num == 13:
        compresser(slide, 2.6, 16.9, 3.9, 16.25, uniforme=True)
        formes = list(slide.shapes)
        titres = par_pos(lambda sh: sh.name in ("Rectangle 3", "Rectangle 23", "Rectangle 39", "Rectangle 44"))
        corps_ = par_pos(lambda sh: sh.name in ("Rectangle 30", "Rectangle 31", "Rectangle 32", "Rectangle 33"))
        for sh, (t_, _) in zip(titres, d["cartes"]):
            ecrire(sh, [t_], [14], [True])
        for sh, (_, c_) in zip(corps_, d["cartes"]):
            ecrire(sh, [c_], [12.5], [False])
            sh.text_frame.word_wrap = True
            for p in sh.text_frame.paragraphs:
                pPr = p._p.get_or_add_pPr()
                for bu in list(pPr):
                    if bu.tag.endswith(("buChar", "buAutoNum")):
                        pPr.remove(bu)
                pPr.insert(0, pPr.makeelement(qn("a:buNone"), {}))
                pPr.set("marL", "0")
                pPr.set("indent", "0")
    elif num == 29:
        compresser(slide, 3.5, 15.6, 4.2, 16.1)
        formes = list(slide.shapes)
        blocs = par_pos(lambda sh: sh.name in ("Rectangle 30", "Rectangle 31", "Rectangle 33", "Rectangle 41"))
        for sh, (t_, c_) in zip(blocs, d["blocs"]):
            ecrire(sh, [t_, c_], [14, 11.5], [True, False])
    elif num == 31:
        for sh in formes:
            if sh.name == "Title 2":
                ecrire(sh, [""])
        compresser(slide, 3.2, 16.1, 4.0, 16.2, uniforme=True)
        formes = list(slide.shapes)
        blocs = par_pos(lambda sh: sh.name in ("Rectangle 5", "Rectangle 7", "Rectangle 9", "Rectangle 12"))
        for sh, (t_, c_) in zip(blocs, d["cartes"]):
            ecrire(sh, [t_, c_], [14, 12], [True, False])
            sh.text_frame.word_wrap = True
    elif num == 107:
        charts = sorted([sh for sh in formes if getattr(sh, "has_chart", False) and sh.has_chart], key=lambda sh: sh.left)
        for sh, (v, lib, cle) in zip(charts, d["anneaux"]):
            cd = CategoryChartData()
            cd.categories = ["Valeur", "Reste"]
            cd.add_series("Série", (round(100 * v, 2), round(100 - 100 * v, 2)))
            sh.chart.replace_data(cd)
            sh.chart.has_title = False
            nh = Cm(7.4)
            delta = sh.top - Cm(4.4) + (sh.height - nh) // 2
            for o in formes:
                if o is not sh and sh.left <= o.left + o.width / 2 <= sh.left + sh.width and sh.top <= o.top + o.height / 2 <= sh.top + sh.height:
                    o.top = o.top - delta
            sh.top, sh.height = Cm(4.4), nh
            pts = sh.chart.plots[0].series[0].points
            for k_, coul in enumerate((ORANGE, BLEU)):
                pts[k_].format.fill.solid()
                pts[k_].format.fill.fore_color.rgb = RGBColor.from_string(coul)
            txt(slide, sh.left / 360000 - 0.3, sh.top / 360000 + sh.height / 360000 + 0.05, sh.width / 360000 + 0.6, 0.9, pc(cle, v), 20, True, BLEU, PP_ALIGN.CENTER)
            txt(slide, sh.left / 360000 - 0.3, sh.top / 360000 + sh.height / 360000 + 0.9, sh.width / 360000 + 0.6, 1.6, lib, 11, False, NOIR, PP_ALIGN.CENTER)
        if analyse:
            corps(slide, 1.4, 14.35, 31.2, BAS_CONTENU - 14.35, [analyse], 13)


def sommaire(slide, entrees):
    txt(slide, 2.7, 0.85, 20, 1.1, "SOMMAIRE", 24, True, BLEU)
    pas = 1.37 if len(entrees) <= 9 else 13.4 / len(entrees)
    taille = 16 if len(entrees) <= 9 else (13 if len(entrees) <= 14 else 12)
    centres = [3.4 + i * pas + pas / 2 for i in range(len(entrees))]
    hb = min(1.01, pas - 0.12)
    ligne = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(2.59), Cm(centres[0]), Cm(2.59), Cm(centres[-1]))
    ligne.line.color.rgb = RGBColor.from_string(BLEU)
    ligne.line.width = Pt(1)
    for i, (intitule, num) in enumerate(entrees):
        cy = centres[i]
        h = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(2.59), Cm(cy), Cm(4.85), Cm(cy))
        h.line.color.rgb = RGBColor.from_string(BLEU)
        h.line.width = Pt(1)
        dd = min(1.2, pas - 0.1)
        dsq = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(4.85 + (1.2 - dd) / 2), Cm(cy - dd / 2), Cm(dd), Cm(dd))
        dsq.fill.solid()
        dsq.fill.fore_color.rgb = RGBColor.from_string(NOIR)
        dsq.line.fill.background()
        if icones:
            ti = dd * 0.6
            slide.shapes.add_picture(io.BytesIO(icone_blanche(icones[i % len(icones)].image.blob)), Cm(4.85 + 0.6 - ti / 2), Cm(cy - ti / 2), Cm(ti), Cm(ti))
        for xx, l, t_, al in ((6.79, 22.98, intitule, PP_ALIGN.LEFT), (29.88, 1.53, str(num), PP_ALIGN.CENTER)):
            cadre = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(xx), Cm(cy - hb / 2), Cm(l), Cm(hb))
            cadre.fill.background()
            cadre.line.color.rgb = RGBColor.from_string(NOIR)
            cadre.line.width = Pt(1.5)
            tf = cadre.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Cm(0.3)
            tf.margin_top = tf.margin_bottom = Cm(0)
            p = tf.paragraphs[0]
            p.alignment = al
            r = p.add_run()
            r.text = t_
            r.font.name = "Helvetica"
            r.font.size = Pt(taille)
            r.font.color.rgb = RGBColor.from_string(NOIR)


entrees = [("Résumé exécutif", 3)]
for i, x in enumerate(SL, 1):
    if x["type"] == "separateur" and x["num"] != "R":
        entrees.append((x["titre"], i))

slides = list(prs.slides)
for i, x in enumerate(SL, 1):
    sl = slides[i - 1]
    COURANTE = i
    if x["type"] == "separateur":
        for sh in sl.shapes:
            if sh.name in ("Rectangle 9", "Rectangle 5"):
                sh.width = Cm(10.6)
            if sh.name == "Rectangle 24" and len(str(x["num"])) > 1:
                sh.text_frame.word_wrap = False
                for p_ in sh.text_frame.paragraphs:
                    for r in p_.runs:
                        r.font.size = Pt(r.font.size.pt * 0.7 if r.font.size else 40)
    if x["type"] == "couverture":
        for sh in sl.shapes:
            if sh.name == "Google Shape;138;p25":
                sh.width = Cm(12.8)
            if sh.name == "Text Placeholder 4":
                sh.left, sh.width = Cm(21.5), Cm(11.0)
    for sh in sl.shapes:
        if getattr(sh, "has_chart", False) and sh.has_chart and x["type"] != "gab":
            try:
                if any(v is not None and v < 0 for pl in sh.chart.plots for se in pl.series for v in se.values):
                    sh.chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
            except Exception:
                pass
    if x["type"] == "sommaire":
        sommaire(sl, entrees)
    elif x["type"] == "gab":
        remplir_gab(sl, x["gab"], x["remplir"], x.get("analyse"))
    elif x["type"] == "contenu":
        paras = paragraphes(x)
        if "tableau" in x:
            tableau_avec_analyse(sl, x)
        if "pas_a_pas" in x:
            pas_a_pas(sl, x["pas_a_pas"], paras)
        if x.get("triangle"):
            triangle(sl)
        if "schema" in x:
            schema(sl, x["schema"], paras)
        if "graphique" in x:
            _, xt, lt = disposition_graphique(paras)
            corps(sl, xt, 4.0, lt, BAS_CONTENU - 4.0, paras)
        if x.get("bandeau"):
            corps(sl, 1.4, 13.75, 31.2, BAS_CONTENU - 13.75, [x["bandeau"]], 13)
        if "image" in x:
            from PIL import Image
            im = Image.open(x["image"]["fichier"])
            larg = 21.8
            haut = larg * im.height / im.width
            if haut > BAS_CONTENU - 4.75:
                haut = BAS_CONTENU - 4.75
                larg = haut * im.width / im.height
            sl.shapes.add_picture(x["image"]["fichier"], Cm(1.4), Cm(4.65), Cm(larg), Cm(haut))
            txt(sl, 1.4, 3.85, larg, 0.7, x["image"]["legende"], 12, False, ORANGE, PP_ALIGN.CENTER)
            corps(sl, 1.4 + larg + 0.6, 4.0, 32.6 - (1.4 + larg + 0.6), BAS_CONTENU - 4.0, paras, 13)
    elif x["type"] == "annexe":
        txt(sl, 1.52, 0.4, 31.24, 1.1, x["titre"], 20, True, NOIR)
        tableau(sl, x["tableau"], 2.0, hauteur_max=15.0)
        if x.get("source"):
            txt(sl, 1.02, 17.15, 31.75, 0.51, x["source"], 7.5, False, "808080")
    elif x["type"] in ("liste_figures", "liste_tableaux"):
        liste = FIG if x["type"] == "liste_figures" else TAB
        nom = "Figure" if x["type"] == "liste_figures" else "Tableau"
        txt(sl, 1.52, 0.4, 31.24, 1.1, "Liste des figures" if nom == "Figure" else "Liste des tableaux", 20, True, NOIR)
        moitie = (len(liste) + 1) // 2
        for col, part in enumerate((liste[:moitie], liste[moitie:])):
            if part:
                tableau(sl, {"titre": None, "entetes": [nom, "Titre", "Diapo."], "largeurs": [1.6, 12.4, 1.4],
                             "lignes": [[str(n_), t_, str(d_)] for n_, t_, d_ in part], "hauteur_ligne": 0.3}, 1.8, 1.4 + col * 16.0, 7, hauteur_max=15.6)
    if x.get("notes"):
        sl.notes_slide.notes_text_frame.text = x["notes"]
COURANTE = None


def textes_visibles(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from textes_visibles(sh.shapes)
            continue
        if sh.has_text_frame:
            yield sh.text_frame.text
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    yield c.text
        if getattr(sh, "has_chart", False) and sh.has_chart:
            for t_ in sh.chart._chartSpace.iter(qn("c:v")):
                yield t_.text or ""


def sans_point_virgule(tf):
    """Charte Adenka : pas de point-virgule dans les phrases ; on coupe en deux phrases."""
    for p_ in tf.paragraphs:
        for r in p_.runs:
            while " ; " in r.text:
                avant, apres = r.text.split(" ; ", 1)
                r.text = avant + ". " + apres[:1].upper() + apres[1:]


for i, (sl, x) in enumerate(zip(prs.slides, SL), 1):
    for sh in sl.shapes:
        if sh.has_text_frame:
            sans_point_virgule(sh.text_frame)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    sans_point_virgule(c.text_frame)
    # Sigles de la diapositive, définis en bas de page (au-dessus de la source)
    if x["type"] in ("contenu", "gab"):
        texte = "\n".join(t_ for t_ in textes_visibles(sl.shapes) if t_)
        sig_ = sigles_du_texte(texte)
        if sig_:
            ligne_s = "Sigles : " + " | ".join(sig_)
            taille_s = 7 if len(ligne_s) < 330 else 6.5
            txt(sl, 1.02, 16.48, 31.75, 0.66, ligne_s, taille_s, False, "595959", PP_ALIGN.LEFT, MSO_ANCHOR.BOTTOM)
prs.save(SORTIE)
finaliser(SORTIE)
os.remove(tmp)

# ---------------------------------------------------------------- contrôle automatique des chiffres
prs2 = Presentation(SORTIE)
textes = []
for sl in prs2.slides:
    morceaux = list(textes_visibles(sl.shapes))
    morceaux.append(sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else "")
    textes.append("\n".join(morceaux))
for t_ in TRACE:
    t_["retrouve"] = t_["texte"] in textes[t_["diapositive"] - 1] if t_["diapositive"] <= len(textes) else False
json.dump({"figures": FIG, "tableaux": TAB, "diapositives": len(SL), "trace": TRACE},
          open(SORTIE.replace(".pptx", "_index.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)
ok = sum(t_["retrouve"] for t_ in TRACE)
print(f"{len(SL)} diapositives ; {len(FIG)} figures ; {len(TAB)} tableaux ; chiffres retrouvés {ok}/{len(TRACE)}")
for t_ in TRACE:
    if not t_["retrouve"]:
        print("  NON RETROUVÉ", t_["diapositive"], t_["indicateur"], t_["texte"])
