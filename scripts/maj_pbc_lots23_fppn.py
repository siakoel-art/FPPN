# -*- coding: utf-8 -*-
"""
FPPN - Pointage de la PBC actuarielle sur le contenu des lots 2 et 3 des nouveaux fichiers (22/10/2026)
=====================================================================================================
Lot 2 (26 fichiers) et lot 3 (29 fichiers) : empreintes SHA-256 conformes aux journaux « journal_lot2.csv » et
« journal_lot3.csv ». Les PDF scannés ont été lus par reconnaissance de caractères (tesseract, français) ; les montants
cités ont été vérifiés sur l'image de la page. Ajoute les items 110 à 114 (demandes issues du rapport v9), des points à
clarifier et l'onglet « Pointage lots 2 et 3 (contenu) ». Un statut n'est jamais abaissé.
Utilisation : python maj_pbc_lots23_fppn.py <pbc.xlsx>   (modifié en place)
"""
import datetime as dt
import sys
from copy import copy

import openpyxl

PBC = sys.argv[1]
R = dt.datetime(2026, 10, 22)
EF = "« ETATS FINANCIERS 2020.pdf » à « ETATS FINANCIERS 2025.pdf » et « ETATS FINANCIERS AU 31-08-2026.pdf »"
RG = ("« RAPPORT DE GESTION 2020 FPPN - BASE - 10 (A).docx », « 1- RAPPORT MORAL 30 juin 2021 définitif (A).docx », « RAPPORT MORAL 2022 (définitif) (A).docx », "
      "« BILAN_DE_L'EXERCICE_CLOS_31.12.2023_(corrigé_et_finalisé)[1].pdf », « RAPPORT DE GESTION DETAILLE EXERCICE 2024.pdf », « RAPPORT GESTION NOUVEAUpdf.pdf », « Rapport de Gestion  2025 Final (2).pdf »")
CGR = "« CONVENTION DE PARTENARIAT ENTRE IPS-CGRAE ET LE FPPN.pdf »"
ORG = "« OrganigrammeFPPN_ORI (A) 2020.docx », « ORGANIGRAMME FPPN 2021(1).pdf », « COURRIER DE NOMINATION DÉFINITIVE.pdf »"
MAJ = {  # item : (statut proposé, constat sur le contenu) ; le statut n'est relevé que s'il est supérieur au statut actuel
    4: ("Partiellement Reçu", "Lu : " + CGR + " (scan) : précompte sur pension des cotisations de prise en charge médicale (art. 3), pas de suspension des soins pour retard de la CGRAE (art. 4.1 e), "
        "précompte dans la limite de la quotité cessible (art. 4.2 a), aucun frais de gestion (art. 5). Annexe 2 (bulletin de souscription) et montant de la cotisation non reçus."),
    79: ("Partiellement Reçu", "La convention IPS-CGRAE confirme la garantie santé des retraités souscripteurs ; durée, cotisation et ayants droit couverts non précisés (" + CGR + ")."),
    108: ("Partiellement Reçu", "La convention IPS-CGRAE organise la couverture santé des policiers retraités ; la décision des organes sur la qualité de membre des retraités reste à recevoir."),
    47: ("Partiellement Reçu", "Lus : " + EF + " (scans). Totaux de l'actif brut, résultat et trésorerie identiques aux balances, sauf 2023 (version antérieure : actif brut inférieur de 10 000 000 FCFA, excédent de 928 699 102 contre 930 629 102 FCFA) "
         "et résultat au 31/08/2026 (-437 148 460 contre -441 256 500 FCFA). Rapports du commissaire aux comptes et états par risque long non reçus."),
    6: ("Partiellement Reçu", "Les états financiers 2020 au 31/08/2026 montrent les provisions au passif sans détail par garantie (" + EF + ")."),
    7: ("Partiellement Reçu", "Lus : rapports de gestion et rapports moraux 2020 à 2025 (" + RG + ") : analyses de l'excédent et de l'épargne, sans calcul actuariel des engagements."),
    112: None,  # item créé plus bas
    12: ("Partiellement Reçu", "Les rapports de gestion donnent les cotisations maladie annuelles (« Cotisation Maladie (3,5 % SB) »), sans détail mensuel ni par catégorie (" + RG + ")."),
    50: ("Partiellement Reçu", "Le rapport de gestion 2023 corrigé propose l'affectation de l'excédent 2023 (36 % aux réserves obligatoires, 20 % aux réserves libres, 44 % au déficit cumulé) ; PV d'AG non reçus."),
    107: ("Partiellement Reçu", "Propositions d'affectation lues dans les rapports de gestion ; PV d'AG d'approbation de 2020 à 2025 non reçus."),
    35: ("Partiellement Reçu", "Lus : « PLAN D'ACTION 2024.pdf » et « Plan d'action 2025.pdf » ; rapport de contrôle de l'AIRMS et plan de redressement non reçus."),
    53: ("Partiellement Reçu", "Lus : listes des administrateurs et des membres de la commission de surveillance et de contrôle d'août 2020 et d'août 2023 (scans) ; calendrier et PV des réunions non reçus."),
    84: ("Partiellement Reçu", "Lus : « Certificat de Propriété Terrain Marcori (Siège).pdf », « Certificat de Propriété Terrain Plateau.pdf » et « Vente ferme par les consorts SIRANSY TOURE SOULEYMANE AU PROFIT DU FPPN.pdf » "
         "(lot n° 40 de 1 469 m²). Scans partiellement lisibles ; situation et valeur de chaque terrain non reçues."),
    23: ("Partiellement Reçu", "Titres de propriété du siège (Marcory) et du terrain du Plateau reçus ; inventaire et évaluation des biens non reçus."),
    85: ("Partiellement Reçu", "Lus : « CONTRAT DE BAIL SONAPIE N° 2124-03002 FPPN.pdf » et « … 2124-03037 FPPN.pdf » (scans peu lisibles, bail de deux ans, loyer trimestriel) ; objet à rapprocher des immeubles d'Aboisso, encaissements non reçus."),
    88: ("Partiellement Reçu", "Lu : « CONVENTION TRANSVIE-CMP.pdf » (soins du CMP aux assurés de Transvie sur lettre de garantie) ; bilan d'ouverture du CMP non reçu."),
    55: ("Partiellement Reçu", "Le rapport de gestion 2023 corrigé nomme la structure : Centre Médical du Policier (CMP). Statut juridique non reçu."),
    26: ("Partiellement Reçu", "Lus : " + ORG + " ; contacts des responsables non reçus."),
    109: ("Partiellement Reçu", "Organigrammes 2020 et 2021 reçus (" + ORG + ") ; délibération approuvant le cadre organique non reçue."),
    103: ("Partiellement Reçu", "Organigrammes reçus ; le statut civil ou policier de chaque agent n'y figure pas. Rattachement aux entités établi par les récapitulatifs de salaires 2024 (CMP) et 2025 (Grenier)."),
    97: ("Partiellement Reçu", "Lot 3 : DISA et DASC 2021 à 2025 scannées et états des travailleurs occasionnels 2024 et 2025, en double des fichiers Excel du lot 1."),
    18: ("Partiellement Reçu", "Le rapport de gestion 2025 ne mentionne aucun taux servi ; aucune écriture d'intérêts sur les comptes d'épargne dans le grand livre 2024."),
}
MAJ = {k: v for k, v in MAJ.items() if v}
NOUVEAUX = [  # (document demandé, priorité, commentaire)
    ("Version des comptes 2023 approuvée par l'Assemblée générale (les états financiers scannés portent un actif brut de 16 157 155 123 FCFA et un excédent de 928 699 102 FCFA, la balance et le rapport de gestion corrigé 16 167 155 123 FCFA et 930 629 102 FCFA) "
     "et explication de l'écart de résultat au 31/08/2026 (états financiers -437 148 460 FCFA, balance -441 256 500 FCFA).", "Urgent", "Demande issue des contrôles croisés du rapport v9 (partie 3)."),
    ("Clé de répartition maladie (936) / décès (8Z6) appliquée en janvier et février 2024 (rapport médian 8Z6 / 936 de 0,3059 et 0,3023, contre 0,3263 ensuite) et liste des dix adhérents d'écart "
     "entre l'étude des ratios et les fichiers 8Z6 de janvier à juin 2026.", "Moyen", "Demande issue des contrôles croisés du rapport v9 (partie 3)."),
    ("Rapport de gestion 2025 rectifié présentant le solde de l'épargne due aux mutualistes au 31/12/2025 (PEP et EFC), distinct des cotisations collectées dans l'année (4 387 823 635 FCFA présentés comme l'épargne au 31/12/2025).",
     "Urgent", "Demande issue du rapport v9 (partie 6)."),
    ("Personnel : âge de départ à la retraite appliqué, nombre d'enfants à charge de chaque salarié, justification des douze dates d'embauche modifiées dans la DISA 2025, rattachement (DG, CMP, Grenier) "
     "des dix-neuf agents recrutés en 2025, note de calcul du compte 196 (complète les items 78 et 98).", "Urgent", "Demande issue du recalcul de la provision du personnel (rapport v9, partie 9)."),
    ("Annexe 2 de la convention IPS-CGRAE (bulletin de souscription), barème de la cotisation des retraités et liste des retraités souscripteurs non précomptés faute de quotité cessible suffisante.",
     "Moyen", "Demande issue de la lecture de la convention IPS-CGRAE (rapport v9, partie 8)."),
]
POINTS = [
    ("Comptes 2023 et 31/08/2026", "Deux versions des comptes 2023 (écart de 10 000 000 FCFA sur l'actif brut et de 1 930 000 FCFA sur l'excédent) ; résultat au 31/08/2026 différent de 4 108 040 FCFA entre états financiers et balance.", "Version approuvée et explication des écarts (item 110)."),
    ("Clé maladie / décès avant mars 2024", "En janvier et février 2024, le précompte décès ne vaut pas 1,23 / 3,77 du précompte maladie ; la clé 3,77 / 1,23 s'applique à plus de 99 % des lignes ensuite.", "Clé appliquée et décision (item 111)."),
    ("Épargne présentée en Assemblée générale", "Le rapport de gestion 2025 présente la collecte de l'année (4 387 823 635 FCFA) comme l'épargne au 31/12/2025, alors que la dette d'épargne de la balance dépasse 13,8 milliards de FCFA.", "Rapport rectifié (item 112)."),
    ("Intérêts servis", "Aucune écriture d'intérêts sur les comptes PEP et EFC dans le grand livre 2024 et aucun compte de charge d'intérêts servis dans les balances 2020 au 31/08/2026.", "Indiquer comment les intérêts promis sont calculés et payés (items 18 et 99)."),
    ("Personnel", "Cinq salariés ont 55 ans ou plus en 2025 et restent en poste ; douze dates d'embauche changent dans la DISA 2025 ; dix-neuf recrues de 2025 sans entité.", "Âge de départ, dates et entités (item 113)."),
    ("Retraités", "La convention IPS-CGRAE ne fixe pas la cotisation des retraités et interdit au FPPN de suspendre les soins en cas de retard de reversement.", "Bulletin, barème et liste des non précomptés (item 114)."),
]


def main():
    wb = openpyxl.load_workbook(PBC)
    ws = wb["PBC Actuariat FPPN"]
    rows, n = {}, 0
    for r in range(11, ws.max_row + 1):
        b = ws.cell(r, 2).value
        if isinstance(b, str) and b.startswith("="):
            n += 1
            rows[n] = r
    # nouveaux items 110 et suivants, sous une section dédiée
    sect = max(r for r in range(11, ws.max_row + 1) if ws.cell(r, 2).value is None and ws.cell(r, 3).value)
    dern, prec = rows[n], rows[n]
    cols = range(1, ws.max_column + 1)
    st_s = {c: copy(ws.cell(sect, c)._style) for c in cols}
    st_i = {c: copy(ws.cell(dern, c)._style) for c in cols}
    r = ws.max_row + 1
    for c in cols:
        ws.cell(r, c)._style = copy(st_s[c])
    ws.cell(r, 3, "Documents reçus le 22/10/2026 (demandes issues du rapport v9)")
    ws.row_dimensions[r].height = ws.row_dimensions[sect].height
    for doc, prio, com in NOUVEAUX:
        r += 1
        for c in cols:
            ws.cell(r, c)._style = copy(st_i[c])
        ws.cell(r, 2, f"=+B{prec}+1")
        ws.cell(r, 3, doc)
        ws.cell(r, 5, prio)
        ws.cell(r, 7, R)
        ws.cell(r, 8, dt.datetime(2026, 10, 29))
        ws.cell(r, 11, com)
        ws.cell(r, 13, "Non Reçu")
        ws.row_dimensions[r].height = 75
        prec = r
        n += 1
        rows[n] = r
    rang = {"Non Reçu": 0, "Partiellement Reçu": 1, "Reçu": 2}
    trace = []
    for item, (statut, texte) in sorted(MAJ.items()):
        r = rows[item]
        avant = ws.cell(r, 13).value
        if rang.get(statut, 0) > rang.get(avant, 0):
            ws.cell(r, 13).value = statut
            ws.cell(r, 9).value = R
        ancien = ws.cell(r, 11).value or ""
        ws.cell(r, 11).value = (ancien + "\n" if ancien else "") + "Lots 2 et 3 lus le 22/10/2026 : " + texte
        trace.append((item, str(ws.cell(r, 3).value)[:160], avant, ws.cell(r, 13).value, texte))
    for k, (doc, _, com) in enumerate(NOUVEAUX):
        item = n - len(NOUVEAUX) + 1 + k
        trace.append((item, doc[:160], "Nouveau", "Non Reçu", com))
    sy = wb["Synthèse du pointage"]
    par = {}
    for item, r in rows.items():
        par.setdefault(ws.cell(r, 13).value, []).append(str(item))
    for k, st in ((6, "Reçu"), (7, "Partiellement Reçu"), (8, "Non Reçu")):
        sy.cell(k, 4).value = ", ".join(par.get(st, []))
    pc = wb["Points à clarifier"]
    d = pc.max_row
    for k, (sujet, constat, demande) in enumerate(POINTS, 1):
        rr = d + k
        for c in range(2, 6):
            pc.cell(rr, c)._style = copy(pc.cell(d, c)._style)
        pc.row_dimensions[rr].height = 60
        pc.cell(rr, 2, int(pc.cell(d, 2).value) + k)
        pc.cell(rr, 3, sujet)
        pc.cell(rr, 4, constat)
        pc.cell(rr, 5, demande)
    nom = "Pointage lots 2 et 3 (contenu)"
    if nom in wb.sheetnames:
        del wb[nom]
    modele = wb["Pointage autres équipes"]
    po = wb.create_sheet(nom)
    po.sheet_view.showGridLines = False
    po["B2"] = "Pointage sur le contenu des lots 2 et 3 des nouveaux fichiers (22/10/2026)"
    po["B2"].font = copy(modele["B2"].font)
    po["B3"] = ("Lot 2 : 26 fichiers, lot 3 : 29 fichiers, empreintes SHA-256 conformes aux journaux « journal_lot2.csv » et « journal_lot3.csv ». Les scans ont été lus par reconnaissance de caractères "
                "et les montants cités vérifiés sur l'image. Reçus sans item de la PBC : budgets de 2020 à 2026, PV d'arrêté de caisse de 2020 au 26/08/2026, convention DEPA. Items 110 à 114 ajoutés.")
    po["B3"].font, po["B3"].alignment = copy(modele["B3"].font), copy(modele["B3"].alignment)
    po.merge_cells("B3:F3")
    po.row_dimensions[3].height = 48
    for j, t in enumerate(["Item", "Document demandé (PBC actuarielle)", "Statut avant", "Statut après", "Constat sur le contenu"], 2):
        po.cell(5, j, t)._style = copy(modele.cell(5, j)._style)
    for i, ligne in enumerate(trace, 6):
        for j, v in enumerate(ligne, 2):
            po.cell(i, j, v)._style = copy(modele.cell(6, j)._style)
        po.row_dimensions[i].height = 90
    for col, w in zip("BCDEF", (7, 48, 16, 16, 100)):
        po.column_dimensions[col].width = w
    po.freeze_panes = "C6"
    po.page_setup.orientation = "landscape"
    po.page_setup.fitToWidth, po.page_setup.fitToHeight = 1, 0
    po.sheet_properties.pageSetUpPr.fitToPage = True
    po.print_title_rows = "5:5"
    wb.save(PBC)
    print(f"{len(MAJ)} items pointés, {sum(1 for t in trace if t[2] != t[3] and t[2] != 'Nouveau')} statuts relevés, {len(NOUVEAUX)} items ajoutés (jusqu'à {n}), {len(POINTS)} points ajoutés")


if __name__ == "__main__":
    main()
