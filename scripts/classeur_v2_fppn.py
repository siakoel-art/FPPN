# -*- coding: utf-8 -*-
"""
FPPN - Classeur des calculs v2 (rapport + compléments + traçabilité des chiffres)
================================================================================

Regroupe les onglets des deux classeurs de calculs et ajoute l'onglet « Traçabilité » :
chaque chiffre écrit dans le PowerPoint, sa diapositive, sa valeur exacte et le résultat du contrôle automatique.
Utilisation : python classeur_v2_fppn.py <calculs_rapport.xlsx> <calculs_complements.xlsx> <rapport_index.json> <sortie.xlsx> [calculs_ratios.xlsx]
Avec le classeur des ratios (version 3), ses onglets « 10_ » sont ajoutés, ainsi qu'un tableau de bord mis en forme.
Version 5 : arguments facultatifs <calculs_provisions.xlsx> <formules.json> : onglets « 12_ » des provisions et onglet
« 11_formules » (registre complet des formules, « Non communiquée » en rouge).
Version 8 : argument facultatif <calculs_sensibilites.xlsx> : onglets « 13 » à formules vivantes (sensibilités et leviers),
copiés avec leurs formules, commentaires, cellules fusionnées et hauteurs de lignes ; recalcul complet à l'ouverture.
Version 9 : argument facultatif <calculs_v9.xlsx> : onglets 14 à 17 à formules (provision du personnel agent par agent,
mortalité brute et sinistralité du FPPN, EFC, contrôles croisés) ; onglet « 00_validite » (périmètre et validité de
chaque analyse) et colonne « Valable pour » dans l'onglet 9_figures_tableaux.
Version 10 : argument facultatif <calculs_v10.xlsx> : onglets 18 à 25 de la revue actuarielle (provisions, décès, épargne,
retraités, Grenier, adossement, projection), chacun ouvert par un bloc d'explication en mots simples ; onglet « 0_guide »
(mode d'emploi pour un lecteur non technique, explication et lien de chaque onglet) ; bloc d'explication en tête des
onglets de valeurs ; texte sans numéro de version.
"""
import json
import sys
from copy import copy

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

RAPPORT, COMPLEMENTS, INDEX, SORTIE = sys.argv[1:5]
RATIOS = sys.argv[5] if len(sys.argv) > 5 else None
PROV = sys.argv[6] if len(sys.argv) > 6 else None
FORMULES = sys.argv[7] if len(sys.argv) > 7 else None
SENSIB = sys.argv[8] if len(sys.argv) > 8 else None
V9X = sys.argv[9] if len(sys.argv) > 9 else None
V10X = sys.argv[10] if len(sys.argv) > 10 else None
BASE = sys.argv[11] if len(sys.argv) > 11 else None
TABLES = sys.argv[12] if len(sys.argv) > 12 else None
EXCLUS = set()
CTRL_LIGNES = None
ENTETE = PatternFill("solid", fgColor="112A40")
_v = __import__("re").search(r"_v(\d+)\.xlsx$", SORTIE)
VER = _v.group(1) if _v else "2"


def copier(src, dst):
    for row in src.iter_rows():
        for c in row:
            d = dst.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                d.font, d.fill, d.border = copy(c.font), copy(c.fill), copy(c.border)
                d.number_format, d.alignment = c.number_format, copy(c.alignment)
            if c.comment:
                d.comment = copy(c.comment)
    for k, dim in src.column_dimensions.items():
        dst.column_dimensions[k].width = dim.width
    for k, dim in src.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[k].height = dim.height
    for m in src.merged_cells.ranges:
        dst.merge_cells(str(m))
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines
    dst.freeze_panes = src.freeze_panes


def entete(ws, titres, largeurs):
    for j, (t, w) in enumerate(zip(titres, largeurs), 1):
        c = ws.cell(row=1, column=j, value=t)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), ENTETE
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


wb = openpyxl.Workbook()
lm = wb.active
lm.title = "0_lisez_moi"
idx = json.load(open(INDEX, encoding="utf-8"))
trace = idx["trace"]
ok = sum(t["retrouve"] for t in trace)
lignes = [
    "Classeur des calculs du rapport actuariel et statistique du FPPN (mission du 01/09/2020 au 31/08/2026, informations exploitées depuis le 01/01/2020).",
    "Onglets 1 à 7 : données, effectifs, sinistralité, bilan, Grenier, produits annexes, placements.",
    "Onglets 8 : inflation et valeurs réelles, taux de croissance, tendances, sensibilités, provisions, tables de mortalité.",
    "Onglet 9_tracabilite : chaque chiffre écrit dans le rapport, avec sa diapositive et sa valeur exacte.",
    f"Contrôle automatique : {ok} chiffres retrouvés sur {len(trace)} enregistrés, {idx['diapositives']} diapositives.",
    "Formats : pourcentages à deux décimales, montants exacts en FCFA sans décimales.",
    "Sources : fichiers reçus du FPPN le 15/10/2026, Banque mondiale (FP.CPI.TOTL.ZG), tables CIMA validées le 18/10/2026.",
]
for i, t in enumerate(lignes, 1):
    lm.cell(row=i, column=1, value=t)
lm["A1"].font = Font(bold=True, size=12)
lm.column_dimensions["A"].width = 120

for fichier in [f for f in (RAPPORT, COMPLEMENTS, RATIOS, PROV) if f]:
    src = openpyxl.load_workbook(fichier)
    for ws in src.worksheets:
        copier(ws, wb.create_sheet(ws.title))

tr = wb.create_sheet("9_tracabilite")
entete(tr, ["Diapositive", "Indicateur", "Valeur exacte", "Texte affiché", "Retrouvé"], [12, 48, 22, 30, 11])
for i, t in enumerate(trace, 2):
    tr.cell(row=i, column=1, value=t["diapositive"])
    tr.cell(row=i, column=2, value=t["indicateur"])
    tr.cell(row=i, column=3, value=t["valeur_exacte"])
    tr.cell(row=i, column=4, value=t["texte"])
    tr.cell(row=i, column=5, value="Oui" if t["retrouve"] else "Non").alignment = Alignment(horizontal="center")

fg = wb.create_sheet("9_figures_tableaux")
entete(fg, ["Type", "Numéro", "Titre", "Diapositive", "Valable pour"], [10, 10, 70, 12, 90])
r = 2
for typ, liste in (("Figure", idx["figures"]), ("Tableau", idx["tableaux"])):
    for l_ in liste:
        n, titre, dia = l_[:3]
        for j, v in enumerate((typ, n, titre, dia, l_[3] if len(l_) > 3 else ""), 1):
            fg.cell(row=r, column=j, value=v).alignment = Alignment(vertical="top", wrap_text=j in (3, 5))
        r += 1
if idx.get("validite"):
    va = wb.create_sheet("00_validite", 1)
    entete(va, ["Partie", "Analyse ou diapositive", "Diapositive", "Valable pour"], [12, 60, 14, 110])
    for i, l_ in enumerate(idx["validite"], 2):
        for j, v in enumerate(l_, 1):
            va.cell(row=i, column=j, value=v).alignment = Alignment(vertical="top", wrap_text=True)
    va.freeze_panes = "A2"

if RATIOS:
    src = openpyxl.load_workbook(RATIOS)["10_ratios"]
    lignes_r = list(src.iter_rows(values_only=True))
    ent = {v: i for i, v in enumerate(lignes_r[0])}
    tb = wb.create_sheet("10_tableau_de_bord", 1)
    tb.sheet_view.showGridLines = False
    tb["B2"] = "Tableau de bord des ratios prudentiels du FPPN (Décision n° 05/2022/COM/UEMOA)"
    tb["B2"].font = Font(bold=True, size=14, color="112A40")
    tb["B3"] = ("Statut : « Conforme » si la valeur respecte la norme du texte ; « Non conforme » sinon ; « En attente » si une donnée manque ; "
                "« À apprécier » si la norme n'est pas chiffrée. Statut global jugé sur 2022 à 2025 ; 2020, 2021 et 2026 indicatifs.")
    tb["B3"].font = Font(italic=True, size=9, color="53565A")
    cols = ["Priorité", "Famille", "Ratio", "Norme", "2022", "2023", "2024", "2025", "2026", "Statut global (2022 à 2025)", "Statut retraité", "Source"]
    larg = [8, 22, 46, 26, 11, 11, 11, 11, 11, 18, 14, 50]
    fin = Side(style="thin", color="BBBCBC")
    for j, (c, w) in enumerate(zip(cols, larg), 2):
        x = tb.cell(row=5, column=j, value=c.replace("Statut global (2022 à 2025)", "Statut 2022-2025"))
        x.font, x.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="060606")
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        tb.column_dimensions[x.column_letter].width = w
    for i, l in enumerate(lignes_r[1:], 6):
        pct = "%" in str(l[ent["Norme"]]) and "mois" not in str(l[ent["Norme"]])
        for j, c in enumerate(cols, 2):
            v = l[ent[c]] if c in ent else None
            x = tb.cell(row=i, column=j, value=v)
            x.border = Border(bottom=fin)
            x.alignment = Alignment(vertical="center", wrap_text=c in ("Ratio", "Norme", "Source"))
            if c in ("2022", "2023", "2024", "2025", "2026") and isinstance(v, (int, float)):
                x.number_format = "0.00%" if pct else "# ##0.00"
    der = 5 + len(lignes_r) - 1
    for colonne in ("K", "L"):
        plage = f"{colonne}6:{colonne}{der}"
        tb.conditional_formatting.add(plage, FormulaRule(formula=[f'LEFT({colonne}6,4)="Conf"'], fill=PatternFill("solid", fgColor="112A40"), font=Font(color="FFFFFF", bold=True)))
        tb.conditional_formatting.add(plage, FormulaRule(formula=[f'LEFT({colonne}6,3)="Non"'], fill=PatternFill("solid", fgColor="C00000"), font=Font(color="FFFFFF", bold=True)))
        tb.conditional_formatting.add(plage, FormulaRule(formula=[f'OR(LEFT({colonne}6,3)="En ",LEFT({colonne}6,1)="À")'], fill=PatternFill("solid", fgColor="F2F2F2"), font=Font(color="53565A")))
    tb.freeze_panes = "E6"
    lm.cell(row=len(lignes) + 1, column=1, value="Onglets 10 : ratios prudentiels, classement des comptes, marge de solvabilité, excédent par risque, contrôle du PCMS, adossement par échéance ; onglet 10_tableau_de_bord mis en forme.")

if FORMULES:
    import sys as _s
    import os as _o
    _s.path.insert(0, _o.path.dirname(_o.path.abspath(__file__)))
    from formules_fppn import COLONNES
    reg = json.load(open(FORMULES, encoding="utf-8"))
    fo = wb.create_sheet("11_formules", 1)
    entete(fo, [t for _, t, _ in COLONNES], [w for _, _, w in COLONNES])
    for i, x in enumerate(reg, 2):
        for j, (k, _, _) in enumerate(COLONNES, 1):
            c = fo.cell(row=i, column=j, value=x.get(k, ""))
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if str(c.value).startswith("Non communiquée"):
                c.font = Font(color="C00000", bold=True)
    fo.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(COLONNES))}{len(reg) + 1}"
    fo.freeze_panes = "D2"
    lm.cell(row=len(lignes) + 2, column=1, value="Onglet 11_formules : registre complet des formules (méthode du FPPN, méthode réglementaire, écart, risques, source). Onglets 12_ : provisions et variantes de sinistralité.")
if V9X:
    src = openpyxl.load_workbook(V9X)
    for ws in src.worksheets:
        copier(ws, wb.create_sheet(ws.title))
    lm.cell(row=len(lignes) + 5, column=1, value="Onglets 14 à 17 : provision du personnel agent par agent à partir des DISA, paramètres modifiables en orange (14a), calcul par agent avec formules (14b), synthèse par entité DG, CMP, Grenier (14c), table CIMA F (14t) ; "
            "mortalité brute et sinistralité du FPPN (15) ; intérêts de l'EFC selon le taux servi (16) ; contrôles croisés (17). Onglet 00_validite : périmètre et validité de chaque analyse.")
if V10X:
    src = openpyxl.load_workbook(V10X)
    for ws in src.worksheets:
        copier(ws, wb.create_sheet(ws.title))
    lm.cell(row=len(lignes) + 6, column=1, value="Onglets 18 à 25 : revue actuarielle complète (provision maladie, allocations décès, risques courts, intérêts d'épargne, santé des retraités, dépréciation du Grenier, adossement actif-passif, projection). Chaque onglet commence par une explication en mots simples ; les cellules orange sont modifiables.")
if SENSIB:
    src = openpyxl.load_workbook(SENSIB)
    for ws in src.worksheets:
        copier(ws, wb.create_sheet(ws.title))
    if BASE:
        # onglets 1 à 12 reconstruits en formules sur des onglets de données brutes (D1 à D8) ; onglet 0_controle
        import os as _os
        sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
        import classeur_formules_fppn as CF
        import charte_excel_fppn as _ch
        from explications_excel_fppn import EXPLIC as _EXPL
        _C = json.load(open(_os.path.join(_os.path.dirname(RAPPORT), "complements.json"), encoding="utf-8"))
        _Q = json.load(open(_os.path.join(_os.path.dirname(RAPPORT), "ratios.json"), encoding="utf-8"))
        _avant = set(wb.sheetnames)
        from exemples_excel_fppn import exemples as _ex
        _R = json.load(open(_os.path.join(_os.path.dirname(RAPPORT), "resultats.json"), encoding="utf-8"))
        for _k, _t in _ex(_R, _C, _Q).items():
            _EXPL.setdefault(_k, {})["Exemple chiffré"] = _t
        _P, _ROW, _AGR, _raw = CF.reconstruire(wb, BASE, _os.path.dirname(RAPPORT), _C, TABLES, _EXPL, _ch.lisible)
        CF.reconstruire_10_12(wb, BASE, _raw, _P, _ROW, _AGR, _Q, _EXPL, _ch.lisible)
        _v10j = _os.path.join(_os.path.dirname(V10X), "v10.json") if V10X else None
        if _v10j and _os.path.exists(_v10j):
            CF.CTRL.extend(tuple(x) for x in json.load(open(_v10j, encoding="utf-8")).get("controles_formules", []))
        CF.relier_existants(wb, _AGR, _P)
        CTRL_LIGNES = CF.onglet_controle(wb)
        EXCLUS = set(CF.CREES)
        json.dump({"controle": CTRL_LIGNES}, open(SORTIE + ".controle.json", "w"))
    wb.calculation.fullCalcOnLoad = True
    sys.path.insert(0, __import__("os").path.dirname(__import__("os").path.abspath(__file__)))
    import charte_excel_fppn
    charte_excel_fppn.appliquer(wb, EXCLUS)
    lm.cell(row=len(lignes) + 4, column=1, value="Charte graphique Adenka : Helvetica, titre au-dessus et source sous chaque tableau, en-tête noir, première colonne grise, lignes alternées, cellules modifiables en orange, statuts conformes en bleu nuit et non conformes en rouge.")
    lm.cell(row=len(lignes) + 3, column=1, value="Onglets 13 : analyses de sensibilité et leviers de conformité, à formules vivantes. Méthode et références (13a), "
            "paramètres modifiables en orange (13b), données (13c), statuts et montants à combler (13d), facteurs récurrents (13e), montants d'équilibre des leviers (13f), "
            "inflation médicale ANStat (13g), choc de 100 000 000 FCFA et faisabilité conjointe (13h), rendement et taux du marché (13i). Un commentaire explique chaque cellule clé.")
# ---- onglet 0_guide : mode d'emploi pour un lecteur non technique, une ligne par onglet avec lien
sys.path.insert(0, __import__("os").path.dirname(__import__("os").path.abspath(__file__)))
from explications_excel_fppn import EXPLIC, GUIDE  # noqa: E402
try:
    import classeur_formules_fppn as _CFm
    _TXT = _CFm.TEXTES
except ImportError:
    _TXT = {}
_RUB = ["À quoi sert cet onglet", "D'où viennent les chiffres", "Comment le calcul est fait", "Exemple chiffré", "Comment lire le résultat", "Valable pour", "Limites", "Calcul Python"]


def _rubriques_onglet(ws):
    """Rubriques écrites en tête d'un onglet (colonne A = nom de la rubrique, colonne B = texte)."""
    t = {}
    for r in range(1, min(ws.max_row, 30) + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.strip() in _RUB and isinstance(ws.cell(r, 2).value, str):
            t[a.strip()] = ws.cell(r, 2).value
    return t


gd = wb.create_sheet("0_guide", 0)
gd.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHI", (26, 34, 40, 46, 40, 40, 34, 30, 40)):
    gd.column_dimensions[col].width = w
gd["A1"] = "Mode d'emploi du classeur des calculs"
gd["A1"].font = Font(name="Helvetica", size=14, bold=True)
aide = ["Ce classeur contient tous les calculs du rapport actuariel et statistique du FPPN, sous forme de formules : en cliquant sur une cellule, on voit le calcul et les cellules qu'il utilise. Chaque onglet commence par une explication en mots simples.",
        "Navigation : la ligne 1 de chaque onglet porte des boutons-liens (⌂ Guide, ◀ Précédent, Suivant ▶, ✓ Contrôle) et, sur fond orange, les onglets d'où viennent les chiffres de l'onglet. Ils marchent sans macro. Les titres des tableaux restent visibles en haut de l'écran quand on descend (tableaux Excel).",
        "Ordre de lecture : onglets de données D1 à D8 (fichiers du FPPN, sans calcul) → 10_agregats (montants comptables, formules sur D1) → onglets 1 à 12 (analyses) → onglets 13 à 25 (sensibilités, provision du personnel, revue actuarielle). L'onglet 0_controle prouve que chaque formule redonne le chiffre du rapport.",
        "Codes couleur : en-tête noir = titres de colonnes ; première colonne grise = libellé de la ligne ; cellule orange = paramètre modifiable (le reste se recalcule) ; fond bleu clair = valeur calculée par Python ; « Non conforme » en rouge, « Conforme » en bleu nuit.",
        "Calculs faits en Python : les lignes individuelles des fichiers de précomptes (4 454 416 lignes, une par agent et par mois) dépassent la limite d'Excel (1 048 576 lignes par feuille). Elles sont lues par le script « base_fppn.py » puis résumées par mois par « analyses_fppn.py » (onglets D3 et D3b). Les grandeurs qui exigent ces lignes (nombre d'agents différents, entrées et sorties, médianes, taux ligne par ligne) sont sur fond bleu clair et renvoient au rapport (méthode et résultats).",
        "Boutons macros (facultatif, 2 minutes) : 1. Alt + F11 ; Fichier > Importer un fichier… ; choisir « FPPN_navigation.bas » livré avec ce classeur ; fermer l'éditeur. 2. Alt + F8 ; choisir « InstallerBoutons » ; Exécuter. 3. Fichier > Enregistrer sous > « Classeur Excel (prenant en charge les macros) (*.xlsm) ». Les boutons Guide, Précédent, Suivant, Tous les onglets et Contrôle apparaissent en haut de chaque onglet.",
        "Montants en FCFA exacts sans décimales ; pourcentages à deux décimales. 2026 = du 01/01 au 31/08/2026 (balance provisoire). Avant de reprendre un chiffre, lire sa ligne « Valable pour » (onglet 00_validite)."]
for i, t in enumerate(aide, 3):
    gd.cell(i, 1, t).alignment = Alignment(wrap_text=True, vertical="top")
    gd.cell(i, 1).font = Font(name="Helvetica", size=10)
    gd.merge_cells(start_row=i, start_column=1, end_row=i, end_column=9)
    gd.row_dimensions[i].height = 44
r0 = 3 + len(aide) + 1
ent = ["Onglet"] + _RUB
for j, t in enumerate(ent, 1):
    c = gd.cell(r0, j, t)
    c.font, c.fill = Font(name="Helvetica", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="000000")
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
rr = r0 + 1
for nom in wb.sheetnames:
    if nom == "0_guide":
        continue
    c = gd.cell(rr, 1, nom)
    c.hyperlink = f"#'{nom}'!A1"
    c.font = Font(name="Helvetica", color="112A40", underline="single", bold=True)
    t = dict(_rubriques_onglet(wb[nom]))
    t.update(_TXT.get(nom, {}))
    t.update(EXPLIC.get(nom, {}))
    if not t and nom in GUIDE:
        t = {"À quoi sert cet onglet": GUIDE[nom][0], "Comment lire le résultat": GUIDE[nom][1]}
    vals = [t.get(k_, "") for k_ in _RUB]
    for j, v in enumerate(vals, 2):
        gd.cell(rr, j, v).alignment = Alignment(wrap_text=True, vertical="top")
        gd.cell(rr, j).font = Font(name="Helvetica", size=9)
        if _RUB[j - 2] == "Calcul Python" and v:
            gd.cell(rr, j).fill = PatternFill("solid", fgColor="DCE6F1")
    gd.row_dimensions[rr].height = min(300, max(30, 12 * (max(len(str(v)) for v in vals) // 42 + 1)))
    rr += 1
gd.freeze_panes = None
manque = [n for n in wb.sheetnames if n not in EXPLIC and n not in GUIDE and n not in _TXT and not _rubriques_onglet(wb[n]) and n not in ("0_guide", "0_lisez_moi")]
if manque:
    print("ONGLETS SANS EXPLICATION :", manque)
# ---- barre de navigation (ligne 1), volets figés réduits, tableaux Excel, module VBA facultatif
import navigation_classeur_fppn as NAV  # noqa: E402
NAV.finaliser(wb)
NAV.ecrire_vba(__import__("os").path.join(__import__("os").path.dirname(SORTIE), "FPPN_navigation.bas"))
if CTRL_LIGNES:
    json.dump({"controle": [CTRL_LIGNES[0] + 1, CTRL_LIGNES[1] + 1]}, open(SORTIE + ".controle.json", "w"))
wb.save(SORTIE)
print(f"{len(wb.sheetnames)} onglets ; {ok}/{len(trace)} chiffres retrouvés")
