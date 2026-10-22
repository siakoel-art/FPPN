# -*- coding: utf-8 -*-
"""
FPPN - Plan de travail du volet actuariel et statistique, pointé au 22/10/2026
==============================================================================

Reprend la structure du plan de travail MAYELIA (description des activités, fait par, interlocuteurs, observations,
livrables, statut, Gantt quotidien) et l'adapte à la mission actuarielle du FPPN :
- les huit phases sont les activités de l'estimation du temps (26 j si les données sont correctes, 39 j si elles
  sont dégradées) ; chaque phase est détaillée en sous-activités dont le temps estimé redonne le total de la phase ;
- chaque sous-activité porte un statut (Effectuée, En cours, Bloquée, Non réalisée) et un taux de réalisation libre ;
- le taux d'une phase et celui de la mission sont des moyennes pondérées par le temps estimé, selon le scénario
  choisi dans la cellule orange (données dégradées par défaut) ;
- le pointage cite, pour chaque sous-activité, la pièce qui le justifie (rapport v8, classeur v8, PBC, comptes
  rendus) et, pour une tâche bloquée, les numéros des items de la PBC attendus.

Charte Adenka : Helvetica, en-têtes noirs à texte blanc, lignes de phase noires, lignes alternées gris clair,
cellule de paramètre orange, statuts en couleur, quadrillage masqué, impression en paysage.

Utilisation : python scripts/plan_travail_fppn.py  (écrit FPPN_Plan_de_travail_actuariel_pointe_20261022.xlsx)
"""
import datetime as dt
import os

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PBC = os.path.join(RACINE, "FPPN_PBC_complete_pointee_20261022.xlsx")
SORTIE = os.path.join(RACINE, "FPPN_Plan_de_travail_actuariel_pointe_20261022.xlsx")

POLICE = "Helvetica"
NOIR, BLEU, ORANGE, GRIS, GRIS_TXT, ROUGE = "060606", "112A40", "EFAC62", "BBBCBC", "53565A", "C00000"
F_ENTETE, F_COL1, F_ALT, BLANC = "000000", "CCCCCC", "E8E8E8", "FFFFFF"
FIN = Side(style="thin", color="000000")
BORD = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)

ARRETE = dt.date(2026, 10, 22)
DEBUT, FIN_MISSION = dt.date(2026, 10, 13), dt.date(2026, 11, 30)
FAIT_PAR = "Actuaire Adenka"
SCENARIOS = ("Données dégradées", "Données correctes")
STATUTS = ("Effectuée", "En cours", "Bloquée", "Non réalisée")
COULEURS_STATUT = {"Effectuée": (BLEU, BLANC), "En cours": (ORANGE, NOIR), "Bloquée": (ROUGE, BLANC), "Non réalisée": (GRIS, NOIR)}
MOIS = ("", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre")
JOURS = ("L", "M", "M", "J", "V", "S", "D")

R8 = "Rapport v8"
C8 = "classeur v8"


def d(j, m):
    return dt.date(2026, m, j)


# Phases : (titre, temps correct, temps dégradé, justification de l'estimation, livrables, sous-activités)
# Sous-activité : (description, interlocuteurs, observations et pièces, j correct, j dégradé, début, fin, statut, taux, items PBC)
PHASES = [
    ("Cadrage : textes, garanties, règles de calcul", 2, 3,
     "Lire les statuts et les règles de chaque prestation, puis fixer les hypothèses",
     "- Comptes rendus de la réunion de lancement\n- PBC actuarielle pointée\n- Guides, comptes rendus et notes d'entretiens\n- Hypothèses retenues (tables, taux, inflation)",
     [
         ("Réunion de lancement : présentation de l'équipe et de la démarche, périmètre (01/01/2020 au 31/08/2026) et quatre questions de la mission",
          "Direction générale, DAF", "Comptes rendus des 13 et 14/10/2026 (FPPN_CR_reunion_de_lancement_Adenka_20261013.pptx, FPPN_Compte_rendu_reunion_de_lancement_20261014.pptx).",
          0.25, 0.25, d(13, 10), d(14, 10), "Effectuée", 1.0, ()),
         ("Élaborer la liste des documents et données (PBC actuarielle), la compléter après chaque entretien et pointer les envois reçus (15/10/2026, lots 1 à 3)",
          "Chef comptable, responsable des cotisations", "PBC pointée au 22/10/2026 (FPPN_PBC_complete_pointee_20261022.xlsx) : 12 items reçus, 39 partiellement reçus, 58 non reçus sur 109. Lot 1 pointé ; lots 2 et 3 préparés. Reste : pointage des lots 2 et 3 et relances.",
          0.5, 0.75, d(14, 10), d(30, 10), "En cours", 0.7, ()),
         ("Lire les statuts, le règlement intérieur, le cahier du régime des prestations sociales et les règles de chaque produit (maladie, décès, EFC, PEP, Grenier, flotte automobile, MEDOC)",
          "Chef de service Assurance, sinistres et prestations", f"{R8}, diapositives 17 à 20 (cadre juridique, statuts et règlement intérieur, garanties et produits). Versions des statuts antérieures à 2023 non reçues.",
          0.5, 0.75, d(15, 10), d(20, 10), "En cours", 0.9, (1,)),
         ("Préparer les guides d'entretien et conduire les entretiens : sinistres et prestations (16/10/2026) ; cotisations, placements et provisions (19/10/2026)",
          "Chef de service Assurance, sinistres et prestations ; responsable des cotisations ; chef comptable", "Guides (FPPN_Guides_entretiens_actuariat_20261016.xlsx et _20261019.xlsx) ; quatre comptes rendus et notes internes (dossier entretiens) ; demandes complémentaires ajoutées à la PBC.",
          0.5, 0.75, d(15, 10), d(20, 10), "Effectuée", 1.0, ()),
         ("Fixer les hypothèses : tables de mortalité CIMA, taux d'actualisation (rendements de l'État de Côte d'Ivoire), inflation médicale (IHPC santé), âge de départ à la retraite, règles de calcul par garantie",
          "Chef comptable", f"Tables retenues ({R8}, diapositives 152 à 158) ; sources externes validées (FPPN_Sources_externes_validees_20261019_v3.xlsx) ; rendements et IHPC santé intégrés au {C8}. Reste : âge de départ par grade et hypothèses de projection (phase 7).",
          0.25, 0.5, d(19, 10), d(23, 10), "En cours", 0.6, (25,)),
     ]),
    ("Contrôle et rapprochement des données", 4, 8,
     "Les fichiers doivent coller avec la comptabilité. Le cas dégradé suppose des relances, des fichiers incomplets et des écarts à expliquer",
     "- Inventaire des fichiers reçus\n- Partie 3 du rapport (données et qualité)\n- Note sur les anomalies de données\n- Onglets 1_ du classeur des calculs",
     [
         ("Extraire et inventorier les fichiers reçus (envoi du 15/10/2026 et nouveaux fichiers du 22/10/2026)",
          "Responsable des cotisations et informatique", "Extraction FPPN_extraction_20261018_1939.zip : 4 287 fichiers, journal journal_fichiers.csv ; inventaire des nouveaux fichiers du 22/10/2026 (scripts extraction_fppn.py, inventaire_nouveaux_fppn.py).",
          0.5, 0.75, d(15, 10), d(22, 10), "Effectuée", 1.0, ()),
         ("Contrôles de qualité : vrai format des fichiers, lignes de total, montants doublés, doublons, fichiers vides ou endommagés",
          "", f"{R8}, diapositives 29 à 36 : 4 454 416 lignes de précomptes contrôlées ; {C8}, onglet 1_rejets.",
          0.75, 1.25, d(16, 10), d(20, 10), "Effectuée", 1.0, ()),
         ("Rapprocher les fichiers aller et retour des précomptes (Solde, CGRAE, GNA) par poste et par mois",
          "Responsable des cotisations", f"{R8}, diapositives 33 et 34 (rapprochement aller et retour, écarts du PEP) ; {C8}, onglet 1_aller_retour.",
          0.5, 1.0, d(16, 10), d(20, 10), "Effectuée", 1.0, ()),
         ("Mesurer la couverture de la période (2020 au 31/08/2026) par poste et par sens, et obtenir les précomptes manquants",
          "Responsable des cotisations", f"Couverture mesurée ({R8}, diapositive 30 ; {C8}, onglet 1_couverture). Précomptes 2020 à 2022 non transmis, sauf PEP. Fichiers aller et retour manquants attendus.",
          0.5, 1.5, d(16, 10), d(6, 11), "Bloquée", 0.5, (51, 61, 62, 63)),
         ("Rapprocher les cotisations, les prestations et l'épargne des fichiers avec la comptabilité (balances Sage 100)",
          "Chef comptable", f"Cotisations comptabilisées et précomptées rapprochées ({R8}, diapositive 46 ; {C8}, onglet 2_cotisations). Reste : prestations (attente des bases) et comptes individuels d'épargne.",
          1.0, 1.5, d(19, 10), d(13, 11), "En cours", 0.5, (17, 19)),
         ("Contrôler les bases de prestations maladie et décès (une ligne par dossier) et les rapprocher des montants comptabilisés",
          "Chef de service Assurance, sinistres et prestations", "Aucune base de prestations extraite du logiciel à ce jour.",
          0.5, 1.5, d(26, 10), d(13, 11), "Bloquée", 0.0, (13, 28, 40, 71)),
         ("Documenter les anomalies de données et les points à clarifier avec le FPPN",
          "Chef comptable, responsable des cotisations", f"Note sur les anomalies de données du 19/10/2026 ; onglet « Points à clarifier » de la PBC ; {R8}, diapositive 179. À tenir à jour à chaque envoi.",
          0.25, 0.5, d(19, 10), d(13, 11), "En cours", 0.6, ()),
     ]),
    ("Évaluation des engagements", 6, 8,
     "Un modèle par garantie, relancé à 8 dates (du 31/12/2019 au 31/08/2026)",
     "- Parties 4 à 8 du rapport (effectifs, cotisations, sinistralité, épargne, Grenier, retraités)\n- Modèle d'évaluation par garantie\n- Engagements aux 8 dates d'arrêté",
     [
         ("Analyser les effectifs (cotisants, entrées et sorties, structure, ancienneté) et les cotisations (taux, valeur réelle, nouvelles recrues)",
          "Responsable des cotisations", f"{R8}, partie 4 (diapositives 37 à 49) ; {C8}, onglets 2_. Limite : fichier des mutualistes partiel, fichier des ayants droit non reçu.",
          1.0, 1.0, d(16, 10), d(21, 10), "Effectuée", 1.0, (9, 10)),
         ("Analyser la sinistralité et le résultat par garantie (prestations sur cotisations, nature des risques, croissance des prestations, frais de gestion)",
          "Chef de service Assurance, sinistres et prestations", f"{R8}, partie 5 (diapositives 50 à 61) ; {C8}, onglets 3_ et 12_sinistralite_variantes.",
          1.0, 1.0, d(16, 10), d(21, 10), "Effectuée", 1.0, ()),
         ("Analyser l'épargne (PEP, EFC), les prêts du Grenier, les retraités CGRAE et les produits annexes",
          "Chef comptable", f"{R8}, parties 6 à 8 (diapositives 62 à 81) ; {C8}, onglets 4_, 5_ et 6_.",
          0.5, 0.5, d(16, 10), d(21, 10), "Effectuée", 1.0, ()),
         ("Construire le modèle d'évaluation par garantie (décès et allocations, maladie, EFC, PEP) : formules, tables, taux, données individuelles",
          "", f"Méthodes posées ({R8}, partie 12, diapositives 159 à 165 ; registre des formules, onglet 11_formules). Modèle non chiffré faute de bases individuelles.",
          1.5, 2.0, d(23, 10), d(5, 11), "En cours", 0.2, ()),
         ("Lancer le modèle aux 8 dates d'arrêté (31/12/2019 à 31/12/2025 et 31/08/2026) et expliquer l'évolution des engagements",
          "", "Attente des fichiers des mutualistes et des ayants droit à chaque date et des bases de prestations.",
          1.5, 2.5, d(3, 11), d(13, 11), "Bloquée", 0.0, (9, 10, 66, 13, 28, 40, 71)),
         ("Conclure sur les engagements envers les adhérents et leurs ayants droit",
          "", f"Conclusion provisoire ({R8}, diapositive 93), à confirmer par le modèle.",
          0.5, 1.0, d(21, 10), d(13, 11), "En cours", 0.3, ()),
     ]),
    ("Revue des provisions", 3, 5,
     "Comparer les provisions au bilan avec mon recalcul, et les provisions passées avec les paiements réels",
     "- Revue des provisions (rapport, partie 9)\n- Recalcul des provisions et tableau de liquidation\n- Conclusion sur la suffisance des provisions",
     [
         ("Inventorier les provisions et dépréciations au bilan par nature et par date, et choisir les formules de contrôle",
          "Chef comptable", f"{R8}, diapositives 84 à 87 ; {C8}, onglets 12_provisions_inventaire, 12_psap_rapprochement et 8_provisions_mois.",
          0.5, 0.75, d(19, 10), d(21, 10), "Effectuée", 1.0, ()),
         ("Recalculer la provision pour prestations à payer maladie (sinistres tardifs par Chain Ladder) et la provision pour allocations décès",
          "Chef comptable ; chef de service Assurance, sinistres et prestations", f"Méthodes présentées ({R8}, diapositives 88, 89 et 162). Triangles non constructibles sans dossiers non payés ni dates de survenance.",
          1.0, 1.75, d(27, 10), d(12, 11), "Bloquée", 0.25, (14, 29, 71, 72, 92)),
         ("Comparer les provisions passées avec les paiements réels (boni et mali de liquidation)",
          "Chef comptable", "Attente des bases de prestations avec dates de survenance et de paiement.",
          0.5, 1.0, d(5, 11), d(13, 11), "Bloquée", 0.0, (13, 28, 71)),
         ("Revoir les autres provisions : dépréciation des prêts du Grenier, compte GNA (flotte automobile), départ à la retraite du personnel",
          "Chef comptable", f"{R8}, diapositives 90 à 92 ; {C8}, onglets 12_grenier_tranches, 12_gna_47111106 et 12_pidr_196. Reste : notes de calcul et fichier du personnel complet.",
          0.5, 0.75, d(20, 10), d(6, 11), "En cours", 0.7, (75, 78, 98, 103)),
         ("Conclure sur la suffisance des provisions",
          "", f"Conclusion provisoire ({R8}, diapositive 94).",
          0.5, 0.75, d(21, 10), d(13, 11), "En cours", 0.3, ()),
     ]),
    ("Adossement et adéquation des placements", 3, 4,
     "Mettre actifs et engagements par échéance, mesurer la couverture et la liquidité",
     "- Adossement actif-passif (rapport, partie 9)\n- Échéancier des actifs et des engagements\n- Conclusion sur l'adéquation des placements",
     [
         ("Décrire la structure de l'actif, les placements, leurs règles de valorisation et le circuit de décision des placements",
          "Chef comptable ; direction générale", f"{R8}, diapositives 95, 98, 104 et 105 ; compte rendu de l'entretien placements du 19/10/2026.",
          0.5, 0.5, d(19, 10), d(21, 10), "Effectuée", 1.0, ()),
         ("Mesurer un adossement de liquidité provisoire (borne favorable) et des scénarios de couverture des engagements",
          "", f"{R8}, diapositives 96, 97, 99 et 100 ; {C8}, onglets 10_actif_mobilisation et 8_sensibilite_retraits.",
          0.75, 0.75, d(20, 10), d(21, 10), "Effectuée", 1.0, ()),
         ("Établir l'échéancier des actifs et des engagements et mesurer la couverture et la liquidité par échéance",
          "Chef comptable", "Attente de l'échéancier et de l'inventaire des placements, et des engagements par échéance (phase 3).",
          1.0, 1.5, d(3, 11), d(13, 11), "Bloquée", 0.1, (20, 80, 81, 82)),
         ("Mesurer le rendement des placements et le comparer aux taux du marché (rendements de l'État de Côte d'Ivoire)",
          "", f"{R8}, diapositives 106 et 150 ; {C8}, onglets 7_rendement_placements et 13i_rendement.",
          0.25, 0.5, d(19, 10), d(21, 10), "Effectuée", 1.0, ()),
         ("Conclure sur l'adossement actif-passif et l'adéquation des placements aux engagements",
          "", f"Conclusions provisoires ({R8}, diapositives 101 et 107).",
          0.5, 0.75, d(21, 10), d(13, 11), "En cours", 0.4, ()),
     ]),
    ("Comparaison avec les indicateurs AIRMS", 1, 2,
     "Calculer les ratios du régulateur à chaque date",
     "- Tableau de bord prudentiel, 2020 au 31/08/2026\n- Analyses de sensibilité et leviers de conformité",
     [
         ("Calculer les ratios prudentiels de la Décision n° 05/2022/COM/UEMOA à chaque date, en deux lectures (comme le FPPN, puis retraitée)",
          "", f"{R8}, diapositives 109 à 129 ; {C8}, onglets 10_ratios et 10_ratios_detail.",
          0.4, 0.75, d(16, 10), d(21, 10), "Effectuée", 1.0, ()),
         ("Contrôler le classement des comptes (PCMS), la marge de solvabilité, l'affectation des excédents et l'excédent par risque",
          "Chef comptable", f"{R8}, diapositives 131 à 134 ; {C8}, onglets 10_controle_pcms, 10_marge_solvabilite, 10_affectation_excedents et 10_excedent_par_risque.",
          0.25, 0.5, d(16, 10), d(21, 10), "Effectuée", 1.0, ()),
         ("Mener les analyses de sensibilité et chiffrer les leviers de conformité (montant à combler, effet de chaque levier)",
          "", f"{R8}, diapositives 138 à 151 ; {C8}, onglets 13a à 13i.",
          0.25, 0.5, d(20, 10), d(21, 10), "Effectuée", 1.0, ()),
         ("Compléter les ratios en attente de données (risques longs, FOMUS-CI, contrôles a posteriori)",
          "Chef comptable", f"{R8}, diapositive 119 (données en attente).",
          0.1, 0.25, d(3, 11), d(13, 11), "Bloquée", 0.0, (47, 48, 49, 52)),
     ]),
    ("Projection de l'équilibre technique", 3, 4,
     "Projeter cotisations, prestations et frais sur 5 à 10 ans, puis tester les hypothèses",
     "- Projection de l'équilibre technique sur 5 à 10 ans\n- Tests de sensibilité des hypothèses",
     [
         ("Fixer les hypothèses de projection : effectifs, cotisations, inflation médicale, mortalité, rendement, frais",
          "Chef comptable", f"Éléments disponibles : tendances et taux de croissance ({C8}, onglets 8_tcam et 8_tendances), inflation médicale ({R8}, diapositive 148), rendements (diapositive 150).",
          0.75, 1.0, d(9, 11), d(12, 11), "En cours", 0.2, ()),
         ("Projeter cotisations, prestations et frais sur 5 à 10 ans par garantie",
          "", "Dépend des phases 3 et 4.",
          1.5, 2.0, d(12, 11), d(19, 11), "Non réalisée", 0.0, ()),
         ("Tester les hypothèses (sensibilités et scénarios) et en tirer les mesures de retour à l'équilibre",
          "", f"Sensibilités ponctuelles déjà menées ({R8}, diapositives 148 à 150), à étendre à la projection.",
          0.75, 1.0, d(18, 11), d(20, 11), "Non réalisée", 0.0, ()),
     ]),
    ("Rapport actuariel et revues", 4, 5,
     "Rédaction, contrôle qualité, prise en compte des remarques",
     "- Note de point d'étape du 21/10/2026\n- Rapport actuariel et statistique provisoire\n- Rapport actuariel et statistique définitif",
     [
         ("Rédiger le rapport actuariel et statistique et le classeur des calculs, avec la traçabilité de chaque chiffre",
          "", f"{R8} (FPPN_Rapport_actuariel_statistique_20261015_v8.pptx, 204 diapositives) et {C8} (onglet 9_tracabilite).",
          1.5, 1.75, d(15, 10), d(20, 11), "En cours", 0.6, ()),
         ("Rédiger la note de point d'étape des travaux (réunion du 21/10/2026)",
          "Direction générale, DAF", "FPPN_Note_point_etape_travaux_20261021.docx.",
          0.25, 0.25, d(20, 10), d(21, 10), "Effectuée", 1.0, ()),
         ("Contrôle qualité : forme, cohérence des chiffres entre sections, niveau de preuve",
          "", "Contrôles menés sur les versions 5 à 8 (mise en page, contrôle automatique des chiffres). À refaire sur le rapport provisoire.",
          0.5, 0.75, d(20, 10), d(23, 11), "En cours", 0.3, ()),
         ("Restitution provisoire au FPPN et recueil des commentaires",
          "Direction générale, DAF, chef comptable", "",
          0.5, 0.5, d(23, 11), d(24, 11), "Non réalisée", 0.0, ()),
         ("Intégrer les commentaires et finaliser le rapport",
          "", "",
          1.0, 1.25, d(25, 11), d(27, 11), "Non réalisée", 0.0, ()),
         ("Transmettre le rapport actuariel et statistique définitif",
          "Direction générale", "",
          0.25, 0.5, d(30, 11), d(30, 11), "Non réalisée", 0.0, ()),
     ]),
]

# Colonnes de l'onglet principal
COLS = ["N°", "Description des activités", "Fait par", "Interlocuteurs FPPN", "Livrables", "Observations et pièces justificatives (pointage au 22/10/2026)",
        "Temps estimé, données correctes (j)", "Temps estimé, données dégradées (j)", "Début prévu", "Fin prévue", "Statut",
        "Taux de réalisation", "Temps réalisé équivalent (j)"]
LARG = [6, 58, 15, 24, 32, 58, 12, 12, 11, 11, 13, 12, 12]
NC = len(COLS)
G0 = NC + 1  # première colonne du Gantt
PARAM = "$F$2"


def police(taille=10, gras=False, italique=False, couleur=NOIR):
    return Font(name=POLICE, size=taille, bold=gras, italic=italique, color=couleur)


def remplir(c, coul):
    c.fill = PatternFill("solid", fgColor=coul)


def cellule(ws, r, col, v=None, fmt=None, gras=False, fond=None, coul=NOIR, h="left", wrap=True, taille=10, bord=True):
    c = ws.cell(r, col, v)
    c.font = police(taille, gras, couleur=coul)
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap)
    if fmt:
        c.number_format = fmt
    if fond:
        remplir(c, fond)
    if bord:
        c.border = BORD
    return c


def lire_pbc():
    """Numéro, intitulé et statut de chaque item de la PBC du 22/10/2026 (numéros recalculés comme les formules)."""
    ws = openpyxl.load_workbook(PBC).worksheets[0]
    items, n = {}, 0
    for r in range(11, 152):
        if ws.cell(r, 2).value is None:
            continue
        n += 1
        items[n] = (str(ws.cell(r, 3).value).strip(), ws.cell(r, 13).value)
    return items


def regles_statut(ws, plage, ancre):
    for s, (fond, txt) in COULEURS_STATUT.items():
        ws.conditional_formatting.add(plage, FormulaRule(formula=[f'{ancre}="{s}"'], fill=PatternFill("solid", fgColor=fond),
                                                         font=Font(name=POLICE, color=txt, bold=True)))


def onglet_plan(wb):
    ws = wb.active
    ws.title = "Plan de travail"
    jours = [DEBUT + dt.timedelta(k) for k in range((FIN_MISSION - DEBUT).days + 1)]
    GN = G0 + len(jours) - 1

    # Titre, sous-titre et paramètre
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NC)
    c = ws.cell(1, 1, "PLAN DE TRAVAIL : AUDIT DU FPPN, 2020 AU 31 AOÛT 2026, VOLET ACTUARIEL ET STATISTIQUE")
    c.font = police(14, True)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26
    ws.merge_cells("A2:D2")
    c = ws.cell(2, 1, f"Pointage au {ARRETE:%d/%m/%Y}. Taux de réalisation pondérés par le temps estimé du scénario choisi ci-contre (cellule orange).")
    c.font = police(9, italique=True, couleur=GRIS_TXT)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    cellule(ws, 2, 5, "Scénario de pondération", gras=True, fond=F_COL1, h="right")
    cellule(ws, 2, 6, SCENARIOS[0], gras=True, fond=ORANGE, h="center")
    dv = DataValidation(type="list", formula1='"' + ",".join(SCENARIOS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("F2")
    cellule(ws, 2, 11, "Taux de la mission", gras=True, fond=F_COL1, h="right")
    ws.row_dimensions[2].height = 30

    # En-têtes (lignes 4 à 6) et Gantt
    for j, (t, w) in enumerate(zip(COLS, LARG), 1):
        ws.merge_cells(start_row=4, start_column=j, end_row=6, end_column=j)
        cellule(ws, 4, j, t, gras=True, fond=F_ENTETE, coul=BLANC, h="center")
        for r in (5, 6):
            ws.cell(r, j).border = BORD
        ws.column_dimensions[L(j)].width = w
    debut_mois = G0
    for k, jr in enumerate(jours):
        col = G0 + k
        ws.column_dimensions[L(col)].width = 3.4
        we = jr.weekday() >= 5
        cellule(ws, 5, col, JOURS[jr.weekday()], gras=True, fond=F_ENTETE if not we else GRIS_TXT, coul=BLANC, h="center", wrap=False, taille=8)
        a = cellule(ws, 6, col, jr, fmt="d", gras=True, fond=F_COL1 if not we else GRIS, h="center", wrap=False, taille=8)
        if jr == ARRETE:
            remplir(a, ORANGE)
            remplir(ws.cell(5, col), ORANGE)
            ws.cell(5, col).font = police(8, True)
        dernier = k == len(jours) - 1 or jours[k + 1].month != jr.month
        if dernier:
            ws.merge_cells(start_row=4, start_column=debut_mois, end_row=4, end_column=col)
            cellule(ws, 4, debut_mois, f"{MOIS[jr.month]} {jr.year}", gras=True, fond=F_ENTETE, coul=BLANC, h="center", wrap=False)
            for cc in range(debut_mois, col + 1):
                ws.cell(4, cc).border = BORD
            debut_mois = col + 1
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 30

    r = 7
    lignes_phase, lignes_sous = [], []
    for p, (titre, jc, jd, justif, livr, sous) in enumerate(PHASES, 1):
        assert abs(sum(s[3] for s in sous) - jc) < 1e-9 and abs(sum(s[4] for s in sous) - jd) < 1e-9, titre
        rp, r1, r2 = r, r + 1, r + len(sous)
        lignes_phase.append(rp)
        # Ligne de phase : sommes et moyennes pondérées des sous-activités
        vals = [p, f"Phase {p} : {titre}", "", "", "", f"Estimation : {justif}.",
                f"=SUM(G{r1}:G{r2})", f"=SUM(H{r1}:H{r2})", f"=MIN(I{r1}:I{r2})", f"=MAX(J{r1}:J{r2})",
                f'=IF(L{rp}>=1,"Effectuée",IF(AND(L{rp}=0,COUNTIF(K{r1}:K{r2},"Bloquée")>0),"Bloquée",IF(L{rp}=0,"Non réalisée","En cours")))',
                f'=IFERROR(M{rp}/IF({PARAM}="{SCENARIOS[1]}",G{rp},H{rp}),0)', f"=SUM(M{r1}:M{r2})"]
        fmts = ["0", None, None, None, None, None, "0.00", "0.00", "dd/mm/yyyy", "dd/mm/yyyy", None, "0%", "0.00"]
        for j, (v, f) in enumerate(zip(vals, fmts), 1):
            cellule(ws, rp, j, v, fmt=f, gras=True, fond=NOIR, coul=BLANC, h="center" if j in (1,) or j >= 7 else "left")
        ws.cell(rp, 6).font = police(9, False, True, BLANC)
        ws.row_dimensions[rp].height = 32
        r += 1
        for k, (desc, interl, obs, sc, sd, deb, fin, statut, taux, items) in enumerate(sous, 1):
            fond = F_ALT if k % 2 == 0 else BLANC
            obs_txt = obs + (f" Items PBC attendus : {', '.join(str(i) for i in items)}." if items else "")
            vals = [f"{p}.{k}", desc, FAIT_PAR, interl, None, obs_txt, sc, sd, deb, fin, statut, taux,
                    f'=IF({PARAM}="{SCENARIOS[1]}",G{r},H{r})*L{r}']
            fmts = ["@", None, None, None, None, None, "0.00", "0.00", "dd/mm/yyyy", "dd/mm/yyyy", None, "0%", "0.00"]
            for j, (v, f) in enumerate(zip(vals, fmts), 1):
                cellule(ws, r, j, v, fmt=f, fond=fond, h="center" if j in (1, 3) or j >= 7 else "left",
                        gras=j in (11, 12))
            ws.cell(r, 6).font = police(9)
            n = max(len(desc) / 55, len(obs_txt) / 55, len(interl) / 22, 1)
            ws.row_dimensions[r].height = max(30, 13 * (int(n) + 1) + 4)
            lignes_sous.append(r)
            r += 1
        # Livrables fusionnés sur les sous-activités, comme dans le modèle
        ws.merge_cells(start_row=r1, start_column=5, end_row=r2, end_column=5)
        c = ws.cell(r1, 5, livr)
        c.font = police(9)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        remplir(c, BLANC)
        for rr in range(r1, r2 + 1):
            ws.cell(rr, 5).border = BORD

    # Total de la mission
    rt = r
    ph = ",".join(str(x) for x in lignes_phase)
    som = lambda col: "=" + "+".join(f"{col}{x}" for x in lignes_phase)
    vals = ["", "Total de la mission actuarielle et statistique", "", "", "", "",
            som("G"), som("H"), f"=MIN(I7:I{rt - 1})", f"=MAX(J7:J{rt - 1})",
            f'=IF(L{rt}>=1,"Effectuée",IF(L{rt}=0,"Non réalisée","En cours"))',
            f'=IFERROR(M{rt}/IF({PARAM}="{SCENARIOS[1]}",G{rt},H{rt}),0)', som("M")]
    fmts = [None, None, None, None, None, None, "0.00", "0.00", "dd/mm/yyyy", "dd/mm/yyyy", None, "0%", "0.00"]
    for j, (v, f) in enumerate(zip(vals, fmts), 1):
        cellule(ws, rt, j, v, fmt=f, gras=True, fond=F_COL1, h="center" if j >= 7 else "left", taille=11)
    ws.row_dimensions[rt].height = 26
    cellule(ws, 2, 12, f"=L{rt}", fmt="0%", gras=True, fond=F_COL1, h="center", taille=12)

    # Gantt : période planifiée (bleu nuit pour les sous-activités, orange pour les phases), week-ends grisés
    for rr in range(7, rt + 1):
        for col in range(G0, GN + 1):
            ws.cell(rr, col).border = Border(left=Side(style="hair", color=GRIS), right=Side(style="hair", color=GRIS),
                                             top=Side(style="hair", color=GRIS), bottom=Side(style="hair", color=GRIS))
    g1, gN = L(G0), L(GN)
    for rr in lignes_phase:
        ws.conditional_formatting.add(f"{g1}{rr}:{gN}{rr}", FormulaRule(formula=[f"AND({g1}$6>=$I{rr},{g1}$6<=$J{rr})"],
                                                                         fill=PatternFill("solid", fgColor=ORANGE), stopIfTrue=True))
    plage_sous = f"{g1}7:{gN}{rt - 1}"
    ws.conditional_formatting.add(plage_sous, FormulaRule(formula=[f'AND($K7="Effectuée",{g1}$6>=$I7,{g1}$6<=$J7)'],
                                                          fill=PatternFill("solid", fgColor=BLEU), stopIfTrue=True))
    ws.conditional_formatting.add(plage_sous, FormulaRule(formula=[f'AND($K7="Bloquée",{g1}$6>=$I7,{g1}$6<=$J7)'],
                                                          fill=PatternFill("solid", fgColor=ROUGE), stopIfTrue=True))
    ws.conditional_formatting.add(plage_sous, FormulaRule(formula=[f"AND({g1}$6>=$I7,{g1}$6<=$J7)"],
                                                          fill=PatternFill("solid", fgColor=GRIS_TXT), stopIfTrue=True))
    ws.conditional_formatting.add(f"{g1}7:{gN}{rt}", FormulaRule(formula=[f"WEEKDAY({g1}$6,2)>5"], fill=PatternFill("solid", fgColor=F_ALT)))
    regles_statut(ws, f"K7:K{rt}", "K7")

    # Saisie guidée : statut en liste, taux entre 0 et 100 %
    dvs = DataValidation(type="list", formula1='"' + ",".join(STATUTS) + '"', allow_blank=True)
    dvt = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                         error="Saisir un taux entre 0 % et 100 %.", errorTitle="Taux de réalisation")
    ws.add_data_validation(dvs)
    ws.add_data_validation(dvt)
    for rr in lignes_sous:
        dvs.add(f"K{rr}")
        dvt.add(f"L{rr}")

    # Légende sous le tableau
    rl = rt + 2
    ws.cell(rl, 2, "Légende du diagramme").font = police(10, True)
    leg = [(BLEU, "Période planifiée, sous-activité effectuée"), (GRIS_TXT, "Période planifiée, sous-activité en cours ou non réalisée"),
           (ROUGE, "Période planifiée, sous-activité bloquée (attente de documents)"), (ORANGE, "Période planifiée de la phase ; colonne orange : date du pointage"),
           (F_ALT, "Week-end")]
    for k, (coul, txt) in enumerate(leg, 1):
        a = ws.cell(rl + k, 1)
        remplir(a, coul)
        a.border = BORD
        ws.cell(rl + k, 2, txt).font = police(9)
    ws.cell(rl + len(leg) + 2, 1, "Source : estimation du temps de la mission actuarielle ; rapport actuariel et statistique v8, classeur des calculs v8, "
                                  "PBC pointée au 22/10/2026, comptes rendus d'entretiens ; pointage Adenka").font = police(8, italique=True, couleur=GRIS_TXT)

    ws.freeze_panes = ws.cell(7, 3)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.tabColor = BLEU
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:6"
    ws.print_area = f"A1:{gN}{rl + len(leg) + 2}"
    ws.oddFooter.left.text = "Adenka : plan de travail actuariel FPPN"
    ws.oddFooter.right.text = "Page &P / &N"
    return ws, lignes_phase, rt


def onglet_synthese(wb, lignes_phase, rt):
    ws = wb.create_sheet("Synthèse")
    P = "'Plan de travail'!"
    ws.cell(1, 1, "Tableau 1 : Avancement par phase au 22/10/2026").font = police(14, True)
    ws.cell(2, 1, "Calculé à partir de l'onglet « Plan de travail » : toute mise à jour des statuts et des taux s'y reporte. "
                  "Le temps restant est le temps estimé du scénario choisi moins le temps réalisé équivalent.").font = police(9, italique=True, couleur=GRIS_TXT)
    ent = ["Phase", "Temps estimé, données correctes (j)", "Temps estimé, données dégradées (j)", "Temps réalisé équivalent (j)",
           "Temps restant (j)", "Taux de réalisation", "Effectuées", "En cours", "Bloquées", "Non réalisées", "Sous-activités"]
    for j, t in enumerate(ent, 1):
        cellule(ws, 4, j, t, gras=True, fond=F_ENTETE, coul=BLANC, h="center")
    ws.row_dimensions[4].height = 42
    bornes = lignes_phase + [rt]
    r = 5
    for k, rp in enumerate(lignes_phase):
        r1, r2 = rp + 1, bornes[k + 1] - 1
        fond = F_ALT if k % 2 else BLANC
        vals = [f"={P}B{rp}", f"={P}G{rp}", f"={P}H{rp}", f"={P}M{rp}",
                f'=IF({P}{PARAM}="{SCENARIOS[1]}",B{r},C{r})-D{r}', f"={P}L{rp}"]
        vals += [f'=COUNTIF({P}K{r1}:K{r2},"{s}")' for s in STATUTS] + [f"=COUNTA({P}K{r1}:K{r2})"]
        fmts = [None, "0.00", "0.00", "0.00", "0.00", "0%"] + ["0"] * 5
        for j, (v, f) in enumerate(zip(vals, fmts), 1):
            cellule(ws, r, j, v, fmt=f, gras=j == 1, fond=F_COL1 if j == 1 else fond, h="left" if j == 1 else "right")
        ws.row_dimensions[r].height = 30
        r += 1
    vals = ["Total de la mission", f"=SUM(B5:B{r - 1})", f"=SUM(C5:C{r - 1})", f"=SUM(D5:D{r - 1})", f"=SUM(E5:E{r - 1})", f"={P}L{rt}"]
    vals += [f"=SUM({L(j)}5:{L(j)}{r - 1})" for j in range(7, 12)]
    for j, (v, f) in enumerate(zip(vals, [None, "0.00", "0.00", "0.00", "0.00", "0%"] + ["0"] * 5), 1):
        cellule(ws, r, j, v, fmt=f, gras=True, fond=F_COL1, h="left" if j == 1 else "right")
    rtot = r
    ws.cell(r + 1, 1, "Source : onglet « Plan de travail » ; pointage Adenka au 22/10/2026").font = police(8, italique=True, couleur=GRIS_TXT)
    for j, w in enumerate([48, 14, 14, 14, 12, 12, 11, 11, 11, 12, 12], 1):
        ws.column_dimensions[L(j)].width = w
    for col, s in zip("GHIJ", STATUTS):
        fond, txt = COULEURS_STATUT[s]
        remplir(ws[f"{col}4"], fond)
        ws[f"{col}4"].font = police(10, True, couleur=txt)

    ch = BarChart()
    ch.type = "bar"
    ch.style = 10
    ch.title = "Taux de réalisation par phase"
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 1
    ch.y_axis.number_format = "0%"
    ch.y_axis.majorGridlines = None
    ch.x_axis.scaling.orientation = "maxMin"
    ch.add_data(Reference(ws, min_col=6, min_row=4, max_row=rtot - 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=rtot - 1))
    ch.series[0].graphicalProperties.solidFill = BLEU
    ch.series[0].graphicalProperties.line.solidFill = BLEU
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.dataLabels.showSerName = ch.dataLabels.showCatName = ch.dataLabels.showLegendKey = False
    ch.legend = None
    ch.height, ch.width = 9, 22
    ws.add_chart(ch, f"A{rtot + 4}")
    ws.cell(rtot + 3, 1, "Figure 1 : Taux de réalisation par phase au 22/10/2026").font = police(10, True)
    ws.cell(rtot + 23, 1, "Source : onglet « Plan de travail » ; pointage Adenka").font = police(8, italique=True, couleur=GRIS_TXT)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = BLEU
    ws.freeze_panes = "B5"


def onglet_bloquants(wb, items):
    ws = wb.create_sheet("Documents attendus")
    ws.cell(1, 1, "Tableau 2 : Documents de la PBC dont dépendent les sous-activités bloquées ou en cours").font = police(14, True)
    ws.cell(2, 1, "Numéros et statuts repris de la PBC pointée au 22/10/2026 (FPPN_PBC_complete_pointee_20261022.xlsx).").font = police(9, italique=True, couleur=GRIS_TXT)
    usage = {}
    for p, ph in enumerate(PHASES, 1):
        for k, s in enumerate(ph[5], 1):
            for i in s[9]:
                usage.setdefault(i, []).append((f"{p}.{k}", s[7]))
    for j, t in enumerate(["Item PBC", "Document attendu", "Statut dans la PBC", "Sous-activités concernées", "Effet sur le plan"], 1):
        cellule(ws, 4, j, t, gras=True, fond=F_ENTETE, coul=BLANC, h="center")
    r = 5
    for k, i in enumerate(sorted(usage)):
        titre, statut = items[i]
        sa = usage[i]
        effet = "Bloque la sous-activité" if any(s == "Bloquée" for _, s in sa) else "Limite la conclusion"
        fond = F_ALT if k % 2 else BLANC
        for j, v in enumerate([i, titre, statut, ", ".join(x for x, _ in sa), effet], 1):
            cellule(ws, r, j, v, fmt="0" if j == 1 else None, gras=j == 1, fond=F_COL1 if j == 1 else fond, h="center" if j in (1, 3, 4) else "left")
        ws.cell(r, 2).font = police(9)
        ws.row_dimensions[r].height = max(28, 13 * (len(titre) // 75 + 1))
        r += 1
    ws.conditional_formatting.add(f"C5:C{r - 1}", FormulaRule(formula=['C5="Non Reçu"'], font=Font(name=POLICE, color=ROUGE, bold=True)))
    ws.cell(r + 1, 1, "Source : PBC pointée au 22/10/2026 ; onglet « Plan de travail »").font = police(8, italique=True, couleur=GRIS_TXT)
    for j, w in enumerate([10, 90, 18, 20, 24], 1):
        ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:E{r - 1}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return len(usage)


def onglet_legende(wb):
    ws = wb.create_sheet("Mode d'emploi")
    ws.cell(1, 1, "Mode d'emploi du plan de travail").font = police(14, True)
    blocs = [
        ("Structure", [
            ("Phases", "Les huit phases sont les activités de l'estimation du temps de la mission (26 j si les données sont correctes, 39 j si elles sont dégradées)."),
            ("Sous-activités", "Chaque phase est détaillée en sous-activités. Leur temps estimé redonne exactement le temps de la phase dans les deux scénarios."),
            ("Pointage", "La colonne « Observations et pièces justificatives » cite la pièce qui justifie le statut (rapport v8, classeur v8, PBC, comptes rendus) et les items de la PBC attendus."),
        ]),
        ("Mise à jour", [
            ("Statut", "Liste déroulante sur chaque sous-activité. Le statut des phases et de la mission est calculé."),
            ("Taux de réalisation", "Saisie libre entre 0 % et 100 % sur chaque sous-activité."),
            ("Scénario", "Cellule orange F2 de l'onglet « Plan de travail » : données dégradées (par défaut) ou données correctes."),
            ("Dates", "Les colonnes « Début prévu » et « Fin prévue » pilotent le diagramme ; le modifier revient à changer ces dates."),
        ]),
        ("Calculs", [
            ("Temps réalisé équivalent", "Temps estimé du scénario choisi multiplié par le taux de réalisation de la sous-activité."),
            ("Taux d'une phase", "Somme des temps réalisés équivalents de la phase divisée par son temps estimé : moyenne des taux pondérée par le temps."),
            ("Taux de la mission", "Somme des temps réalisés équivalents divisée par le temps estimé total du scénario choisi."),
            ("Statut d'une phase", "Effectuée à 100 % ; Bloquée si le taux est nul et qu'une sous-activité est bloquée ; Non réalisée si le taux est nul ; En cours sinon."),
        ]),
    ]
    r = 3
    for titre, lignes in blocs:
        cellule(ws, r, 1, titre, gras=True, fond=F_ENTETE, coul=BLANC)
        cellule(ws, r, 2, "", fond=F_ENTETE)
        r += 1
        for k, (a, b) in enumerate(lignes):
            cellule(ws, r, 1, a, gras=True, fond=F_COL1)
            cellule(ws, r, 2, b, fond=F_ALT if k % 2 else BLANC)
            ws.row_dimensions[r].height = max(18, 14 * (len(b) // 95 + 1))
            r += 1
        r += 1
    cellule(ws, r, 1, "Statuts", gras=True, fond=F_ENTETE, coul=BLANC)
    cellule(ws, r, 2, "", fond=F_ENTETE)
    defs = {"Effectuée": "Travail terminé, pièce justificative disponible.",
            "En cours": "Travail engagé, résultat provisoire ou partiel.",
            "Bloquée": "Travail arrêté ou limité dans l'attente de documents du FPPN (items de la PBC cités).",
            "Non réalisée": "Travail non commencé, selon le calendrier prévu."}
    for s in STATUTS:
        r += 1
        fond, txt = COULEURS_STATUT[s]
        cellule(ws, r, 1, s, gras=True, fond=fond, coul=txt, h="center")
        cellule(ws, r, 2, defs[s])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    items = lire_pbc()
    wb = openpyxl.Workbook()
    _, lignes_phase, rt = onglet_plan(wb)
    onglet_synthese(wb, lignes_phase, rt)
    n = onglet_bloquants(wb, items)
    onglet_legende(wb)
    wb.save(SORTIE)

    # Contrôle : taux attendus calculés en Python (scénario par défaut)
    tot = sum(ph[2] for ph in PHASES)
    real = 0
    for p, ph in enumerate(PHASES, 1):
        rp = sum(s[4] * s[8] for s in ph[5])
        real += rp
        print(f"Phase {p} : {rp / ph[2]:.0%} ({rp:.2f} j sur {ph[2]} j)")
    print(f"Mission : {real / tot:.0%} ({real:.2f} j sur {tot} j) ; {n} items de la PBC cités")
    print("Écrit :", SORTIE)


if __name__ == "__main__":
    main()
