# -*- coding: utf-8 -*-
"""
FPPN - Note interne confidentielle sur l'entretien cotisations du 19/10/2026
Reprend la mise en forme de la note interne de l'entretien sinistres (16/10/2026).
Utilisation : python note_cotisations_fppn.py <note_modele.docx> <sortie.docx>
"""
import sys
from copy import deepcopy

import docx

MODELE, SORTIE = sys.argv[1:3]

ENTETE = [
    "FPPN : note interne sur l'entretien cotisations et précomptes",
    "Document interne Adenka. Confidentiel. Ne pas diffuser au client.",
    "Entretien du 19 octobre 2026 avec M. Kouamé Attuoman, responsable des cotisations, qui est aussi le responsable informatique. "
    "Participants Adenka : Jean-Noël Sialou (conduite de l'entretien), Ruth Miensa et Eliam Affolaby.",
    "Objet : consigner, de façon factuelle et neutre, les éléments sensibles relevés pendant l'entretien et non repris dans le compte rendu destiné au client.",
]

SECTIONS = [
    ("1. Cumul des fonctions et dépendance à une personne", [
        "L'interlocuteur est le responsable informatique. Il gère les cotisations parce qu'il maîtrise le logiciel qui produit les fichiers aller.",
        "Il produit les fichiers aller des produits complémentaires, des retraités (CGRAE) et de la flotte automobile (GNA), et les dépose sur la plateforme de partage.",
        "Il décrit les collaborateurs qui traitent les souscriptions, les rachats et les demandes des mutualistes comme très chargés.",
        "Le suivi des cotisations repose donc sur une seule personne, dont ce n'est pas le métier principal. Aucune procédure écrite n'a été citée pour assurer la continuité en son absence.",
    ]),
    ("2. Absence de contrôle des cotisations reçues", [
        "Aucun état des cotisations n'est tenu. Aucun rapprochement n'est fait entre le fichier aller et le fichier retour, et aucun suivi n'est fait mutualiste par mutualiste.",
        "Le FPPN ne connaissait pas les motifs de rejet. Un état séparé des motifs de rejet n'est reçu que depuis environ deux mois.",
        "Pour la maladie et le décès, le FPPN ne reçoit pas le fichier aller : la direction du personnel de la Police le transmet directement à la Solde, sans plateforme de partage. Le FPPN ne peut donc pas mesurer ce qui aurait dû être précompté.",
        "Des paiements manuels sont possibles et sont enregistrés avec leur date et leur heure. Leur encaissement et leur suivi comptable restent à vérifier avec le chef comptable.",
        "À vérifier : états des motifs de rejet (PBC, item 64), paiements manuels (PBC, item 65), fichiers aller maladie et décès (PBC, item 61).",
    ]),
    ("3. Usage d'un modèle d'intelligence artificielle local sur des données personnelles", [
        "L'interlocuteur indique utiliser un modèle d'intelligence artificielle installé en local pour reporter en masse les mécanos des nouvelles recrues sur les fiches de souscription.",
        "Ce traitement porte sur des données personnelles des policiers (matricule, nom, mécano) et se fait hors de toute procédure documentée.",
        "Risque de conformité : la protection des données à caractère personnel (loi n° 2013-450 du 19 juin 2013) s'applique à ce traitement. Sa conformité est à apprécier par un juriste.",
        "Risque d'erreur : une erreur d'appariement entre matricule et mécano peut faire précompter la cotisation d'un policier sur la paie d'un autre.",
        "Risque de traçabilité : aucun contrôle des reports n'a été cité, et le résultat ne peut pas être revu par une autre personne.",
        "Risque de dépendance : l'outil et la méthode ne sont maîtrisés que par l'interlocuteur.",
        "Recommandation à porter : obtenir de la direction du personnel une table de correspondance matricule et mécano (PBC, item 67), qui supprimerait le report manuel.",
    ]),
    ("4. Règles des produits non écrites", [
        "L'EFC est passée de 6 000 FCFA à 10 000 FCFA pour les nouvelles souscriptions depuis 2021, sans document formel selon l'interlocuteur. Les anciens souscripteurs restent à 6 000 FCFA mais peuvent choisir 10 000 FCFA. Les fichiers confirment la coexistence des deux montants.",
        "L'EFC et le PEP sont présentés aux élèves policiers comme obligatoires dès l'école, sans texte qui le prévoie. Ce point expose le FPPN à des contestations de souscripteurs.",
        "L'interlocuteur cite un délai d'au moins trois mois de cotisation avant de pouvoir bénéficier des prestations. L'entretien du 16/10/2026 indiquait l'absence de délai de carence pour le régime obligatoire : les deux déclarations sont à recouper.",
        "À vérifier : décision de hausse de l'EFC (PBC, item 58), règle des trois mois (PBC, item 59), documentation des produits (PBC, item 57).",
    ]),
    ("5. Flotte automobile, MEDOC et Grenier", [
        "Flotte : les primes vont directement à GNA. Selon l'interlocuteur, GNA met mal à jour les soldes. Le FPPN arrête le précompte à la date de fin, parfois avec un mois de plus, et rembourse les trop-perçus sur demande.",
        "MEDOC : le produit est en pause depuis environ deux mois, en raison de difficultés d'utilisation par les pharmacies. La période de fonctionnement reste à préciser.",
        "Grenier : l'interlocuteur cite trois types de prêts, alors que les fichiers de production en comptent huit. L'achat d'équipements via une plateforme partenaire n'a pas bien fonctionné selon lui. Une nouvelle application est terminée et attend son lancement.",
    ]),
    ("6. Points d'attention pour l'équipe", [
        "Ces éléments émanent d'une seule personne. Ils doivent être recoupés avec les documents demandés et avec les entretiens du chef comptable et de la direction.",
        "Ne pas reprendre ces éléments dans un document destiné au client sans validation de la direction de mission, en particulier l'usage du modèle d'intelligence artificielle.",
        "Points à recouper en priorité : règle des trois mois, circuit du fichier retour maladie et décès, encaissement des paiements manuels, base de calcul du taux maladie et décès.",
    ]),
]

d = docx.Document(MODELE)
corps = d.element.body
paras = d.paragraphs
modele_entete = [deepcopy(p._p) for p in paras[:4]]
modele_titre = deepcopy(paras[4]._p)
modele_puce = deepcopy(paras[5]._p)
for p in paras:
    corps.remove(p._p)
sectpr = corps.find(docx.oxml.ns.qn("w:sectPr"))


def ajouter(modele, texte):
    el = deepcopy(modele)
    sectpr.addprevious(el) if sectpr is not None else corps.append(el)
    p = docx.text.paragraph.Paragraph(el, d._body)
    runs = p.runs
    runs[0].text = texte
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


for m, t in zip(modele_entete, ENTETE):
    ajouter(m, t)
for titre, puces in SECTIONS:
    ajouter(modele_titre, titre)
    for t in puces:
        ajouter(modele_puce, t)
d.save(SORTIE)
print(len(d.paragraphs), "paragraphes")
