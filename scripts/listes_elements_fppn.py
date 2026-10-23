# -*- coding: utf-8 -*-
"""
FPPN - Listes exhaustives des éléments à obtenir pour terminer le rapport
========================================================================
Onglet 1 : éléments à obtenir du FPPN (documents, données, réponses écrites).
Onglet 2 : données externes à rechercher (textes officiels, séries publiques), avec les liens trouvés.
Onglet 3 : points retirés des questions au FPPN, car la réponse figure déjà dans les documents reçus.
Onglet 4 : doublons de la PBC (items qui demandent la même pièce).
Utilisation : python listes_elements_fppn.py <sortie.xlsx>
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SORTIE = sys.argv[1]
NOIR, GRIS = "000000", "D9D9D9"
COUL = {"Urgent": "FF0000", "Moyen": "FFC000", "Faible": "00B050",
        "Non reçu": "FF0000", "Partiellement reçu": "FFC000", "Reçu": "00B050", "Nouvelle demande": "FFFF00",
        "Lien trouvé, non ouvert": "FFC000", "Aucun lien direct trouvé": "FF0000", "Document déjà détenu": "00B050"}
BLANC_SUR = {"FF0000", "00B050"}
fin = Side(style="thin", color="808080")
BORD = Border(left=fin, right=fin, top=fin, bottom=fin)

FPPN = [
    # thème, élément, pourquoi, diapositives, PBC, statut PBC, priorité
    ("Personnel", "Âge de départ à la retraite appliqué au personnel civil du FPPN.", "5 agents ont 55 ans ou plus et sont toujours en poste en 2025, dont un de 60 ans. La provision de 24 446 566 FCFA baisse de 29,71 % si l'âge passe à 60 ans.", "118, 128", "113", "Non reçu", "Urgent"),
    ("Personnel", "Justification des 12 dates d'embauche qui changent entre les DISA 2021 à 2023 et la DISA 2025.", "L'ancienneté de l'agent A25-01 baisse de plus de sept ans. La provision dépend directement de l'ancienneté.", "128", "113", "Non reçu", "Urgent"),
    ("Personnel", "Entité de rattachement (Direction Générale, CMP, Grenier) des 19 agents recrutés en 2025 et nombre d'enfants à charge de chaque agent.", "Les 19 agents ne figurent dans aucun récapitulatif de salaires par entité. L'âge de départ conventionnel augmente d'un an par enfant à charge (au plus 3).", "118, 126, 127", "113", "Non reçu", "Urgent"),
    ("Personnel", "Note de calcul et fichiers de calcul de la provision pour départ à la retraite (compte 196), de 2020 à 2026.", "Le rapport recalcule la provision agent par agent, mais ne peut pas comparer sa méthode à celle du FPPN.", "119", "78, 98", "Non reçu", "Moyen"),
    ("Personnel", "Fichier du personnel de la Direction Générale avec le statut de chaque agent (civil salarié ou policier en service).", "Seuls les civils salariés entrent dans la provision pour départ à la retraite.", "129", "103", "Partiellement reçu", "Moyen"),
    ("Cotisations", "Document qui notifie à la Solde la répartition du taux de 5 % (3,77 % maladie et 1,23 % décès précomptés) et clé appliquée en janvier et février 2024.", "Les statuts (art. 17) fixent 3,5 % et 1,5 %. La Solde précompte 3,77 % et 1,23 % sur 209 670 lignes de huit mois.", "22, 51", "56, 111", "Non reçu", "Urgent"),
    ("Cotisations", "Décision qui a porté l'EFC de 6 000 FCFA à 10 000 FCFA pour les souscriptions depuis 2021.", "Les fichiers montent la part des lignes à 10 000 FCFA de 62,52 % en 2023 à 84,98 % en 2026, sans texte.", "56", "58", "Non reçu", "Moyen"),
    ("Cotisations", "Texte qui fixe le délai d'au moins trois mois de cotisation avant les prestations, avec les produits concernés.", "L'entretien du 16/10/2026 indiquait l'absence de délai de carence, celui du 19/10/2026 un délai de trois mois.", "56, 220", "59", "Non reçu", "Urgent"),
    ("Cotisations", "Nature des références PTI, RPEP, EP, TI, R-PEP et PTQ du poste 998 (remboursements de prêts PEP des comptes 27110000 et 27120000 ou autre).", "Ces références portent 174 461 860 FCFA de précomptes de 2020 à août 2026, à côté de l'épargne PEP.", "80", "Nouvelle", "Nouvelle demande", "Moyen"),
    ("Fichiers de précomptes", "Fichiers retour de décembre 2024 reçus vides (0 octet) : « 122024_effectif_police_936_DECEMBRE 2024.xlsx » (maladie) et « FPPN_ALLOC_DECES_8Z6_DECEMBRE_2024.xlsx » (décès).", "Le mois de décembre 2024 n'a aucune donnée pour la maladie et le décès. Aucun envoi suivant (18/10 et 22/10/2026) ne les contient.", "41, 48", "62", "Non reçu", "Urgent"),
    ("Fichiers de précomptes", "Fichiers retour décès 8Z6 de mai à octobre 2024 et d'avril 2025 (dossier « FPPN-SOLDE\\ALLOCATION DECES »).", "Huit mois de précomptes décès manquent sur la période 2024 à 2026.", "41", "62", "Non reçu", "Urgent"),
    ("Fichiers de précomptes", "Fichier retour Grenier 9X4 de décembre 2024 (série « MTG_LEGRENIER_9X4_… »), fichier retour EFC 870 de mars 2025, fichier retour CGRAE de janvier 2024 et de juin 2024 (le fichier reçu pour juin 2024 est un fichier aller).", "Ces mois manquent dans le rapprochement entre échéances et précomptes.", "41", "63", "Non reçu", "Moyen"),
    ("Fichiers de précomptes", "Fichiers retour maladie 936 et décès 8Z6 de janvier 2020 à décembre 2023.", "Les analyses mensuelles commencent en janvier 2024. Les années 2020 à 2023 ne sont lues qu'en totaux comptables.", "36", "62", "Non reçu", "Urgent"),
    ("Fichiers de précomptes", "Fichiers aller maladie 936 et décès 8Z6 de janvier 2020 à août 2026, détenus par la direction du personnel de la Police.", "Le taux de recouvrement des cotisations (indicateur REC) ne peut pas être calculé sans ces fichiers.", "157, 162, 218", "51, 61", "Non reçu", "Urgent"),
    ("Fichiers de précomptes", "Fichiers aller et retour MEDOC (8Z7) de 2024 à 2026, et fichiers aller et retour de la flotte 9W3 manquants (aller : 03, 05, 06 et 08/2025, 03 et 08/2026 ; retour : 04, 06 et 08/2025, 01/2026).", "Les retours MEDOC de février et avril 2026 n'ont aucune ligne.", "41", "63", "Non reçu", "Faible"),
    ("Fichiers de précomptes", "États des motifs de rejet reçus de la Solde, et liste des paiements manuels de cotisations de 2020 à août 2026.", "Les écarts entre fichiers aller et retour traduisent des rejets dont le motif est inconnu.", "39", "64, 65", "Non reçu", "Moyen"),
    ("Fichiers de précomptes", "Procédure écrite du circuit des fichiers aller et retour par produit.", "Le circuit présenté repose sur la seule déclaration du responsable des cotisations (entretien du 19/10/2026).", "25, 26", "Nouvelle", "Nouvelle demande", "Faible"),
    ("Prestations", "Base des prestations maladie, une ligne par facture : date des soins, date de dépôt, date de paiement, montant.", "Base de la provision pour sinistres à payer maladie, des délais de paiement et des contrôles santé. Seule l'année de la facture est lisible dans son numéro.", "108, 218, 221", "40, 71", "Non reçu", "Urgent"),
    ("Prestations", "Base des sinistres décès, une ligne par dossier, et liste des dossiers composant le solde du compte 41591000 au 31/08/2026 (avec la date de dépôt de chaque demande).", "La dette décès de 489 419 669 FCFA ne peut pas être triée entre dossiers payables et dossiers forclos.", "112, 129", "28, 115", "Partiellement reçu", "Urgent"),
    ("Prestations", "Liste des factures et dossiers non payés au 31/12 de chaque année de 2019 à 2025 et au 31/08/2026.", "Elle permet de vérifier la dette envers les cliniques (compte 40110000) dossier par dossier.", "109, 218", "14, 29", "Non reçu", "Urgent"),
    ("Prestations", "Explication de l'écart décès 2021 : 156 020 000 FCFA de sinistres selon l'étude des ratios du FPPN, contre 401 856 666 FCFA dans les comptes.", "Les deux sources concordent au franc près de 2022 à 2025, pas en 2021.", "66", "Nouvelle", "Nouvelle demande", "Moyen"),
    ("Prestations", "Validation, ligne par ligne, des montants décès supérieurs à 1 000 000 FCFA corrigés par Adenka dans les fichiers de précomptes.", "La correction est limitée au poste décès et à ce seuil pour ne modifier aucun montant réel.", "38", "Nouvelle", "Nouvelle demande", "Faible"),
    ("Prestations", "Texte qui fonde l'exclusion du suicide (quatre décès valorisés à 0 FCFA dans l'état 2025).", "Le règlement intérieur reçu ne mentionne pas cette exclusion.", "111", "116", "Non reçu", "Moyen"),
    ("Mutualistes", "Fichier des mutualistes avec la date de naissance, et fichier des ayants droit.", "Sans âge, ni la mortalité observée (décès observés / décès attendus), ni les engagements viagers, ni la projection technique ne peuvent être calculés.", "197, 201, 218", "9, 10, 66", "Partiellement reçu", "Urgent"),
    ("Mutualistes", "Âge de départ à la retraite des policiers par grade, avec le texte de référence.", "Il fixe l'échéance de versement de l'EFC de chaque titulaire.", "129", "25", "Non reçu", "Moyen"),
    ("Retraités", "Fichiers CGRAE avec la date de naissance de tous les retraités, prestations payées aux retraités, annexe 2 de la convention IPS-CGRAE (bulletin de souscription) et décision fixant la garantie santé des retraités.", "Seuls les fichiers retour de février à avril 2024 portent la date de naissance, pour 105 retraités. La durée de la garantie (à vie ou limitée) fixe la provision mathématique.", "98, 99, 105", "79, 108, 114, 120", "Partiellement reçu", "Urgent"),
    ("Épargne", "Décision signée fixant le taux servi sur le PEP et l'EFC (les documents reçus citent 3,5 %, 5,5 % et 14 %).", "L'engagement d'épargne avec intérêts (19 022 429 629 FCFA au taux du manuel) change avec le taux retenu.", "82, 129, 134", "18, 99", "Partiellement reçu", "Urgent"),
    ("Épargne", "Comptes individuels d'épargne (une ligne par mutualiste et par année) et rapprochement avec les comptes 41195100 et 41195200.", "Aucun échéancier individuel : l'épargne de 15 743 516 391 FCFA n'est pas ventilée par échéance.", "129, 134, 218", "17, 19", "Partiellement reçu", "Urgent"),
    ("Épargne", "Rapport de gestion 2025 rectifié présentant le solde de l'épargne due au 31/12/2025.", "Le rapport de gestion 2025 présente 4 387 823 635 FCFA de cotisations de l'année comme l'épargne.", "86", "112", "Non reçu", "Moyen"),
    ("Provisions", "Toute pièce sur l'origine des provisions héritées (2 867 602 878 FCFA maladie et 1 961 695 412 FCFA obsèques) et procès-verbal de l'Assemblée générale approuvant leur reclassement en comptes 15801000 et 15802000 en 2024.", "Ces montants sont identiques de 2020 au 31/08/2026. Le chef comptable les date d'environ 2004 (entretien du 19/10/2026).", "106, 109, 222", "72, 95, 106", "Partiellement reçu", "Urgent"),
    ("Placements", "Nature du DAT BOA (compte 52210200) : placement du FPPN ou placement affecté à l'EFC, avec conventions et relevés.", "L'intitulé du compte le rattache à l'EFC. Le classement change la couverture de l'épargne.", "222", "80", "Partiellement reçu", "Urgent"),
    ("Placements", "Bulletin de souscription et relevés du FCP géré par United Capital, inventaire de tous les placements de 2016 au 31/08/2026, délibérations du Conseil d'administration autorisant les placements.", "Les statuts (art. 85) exigent une autorisation du Conseil pour chaque placement.", "146, 147", "81, 82, 105", "Non reçu", "Urgent"),
    ("Placements", "Rapport de réévaluation du siège et décision de comptabilisation au compte 10620000, et origine du compte 106 « Réserve disponible ».", "Le grand livre 2022 intitule le compte « Réévaluation matériels ». Un écart de réévaluation ne compte pas dans la marge de solvabilité.", "176, 177", "50, 86", "Non reçu", "Urgent"),
    ("Placements", "Dossier de la créance sur l'hôpital de la Police (3 407 519 850 FCFA, compte 47111100) et échéancier de remboursement.", "Montant inchangé depuis 2020. Sa valeur conditionne l'actif retenu dans l'adossement.", "134", "87", "Non reçu", "Moyen"),
    ("Placements", "Projet de Cocody : budget, plan de financement, calendrier, décision des organes, part financée par l'épargne. Décision de désignation ou de prolongation du directeur général.", "Le projet engagerait l'épargne exigible dans un actif long, alors que la liquidité ne couvre pas l'année qui vient.", "148", "117, 118", "Non reçu", "Urgent"),
    ("Placements", "Comptabilité séparée de la Clinique médicale du policier (CMP) et résultats annuels de 2020 à 2025.", "5 803 025 455 FCFA d'investissements et 1 280 300 860 FCFA de charges en comptes 47, sans valeur de marché.", "134", "88, 119", "Partiellement reçu", "Moyen"),
    ("Flotte automobile", "Convention d'apporteur d'affaires avec GNA, décomptes annuels des commissions, détail des mouvements du compte 47111106.", "La commission comptabilisée en 2024 ne vaut pas 20 % des précomptes 9W3 de l'année (7 913 002 FCFA).", "116", "68, 75, 89", "Non reçu", "Moyen"),
    ("Grenier", "État des impayés par tranche de retard au 31/12/2025 et au 31/08/2026, politique de dépréciation, et mode de règlement des échéances non précomptées.", "Aucune dépréciation passée depuis 2022 (compte 49 non mouvementé).", "94, 221", "73", "Partiellement reçu", "Moyen"),
    ("MEDOC", "Documentation de MEDOC : conventions avec les pharmacies, date de démarrage, décision de suspension, état des prises en charge.", "Le produit porte des cotisations (compte 7022) sans règle écrite transmise.", "100, 220", "60, 69", "Non reçu", "Faible"),
    ("Conformité", "Décision d'agrément AIRMS (numéro, date, risques autorisés), attestation d'adhésion au FOMUS-CI et preuves de paiement, étude actuarielle des risques longs, rapports du commissaire aux comptes 2020 à 2025.", "Aucun compte FNG dans les balances. Les risques longs exigent une étude actuarielle (règlement d'exécution n° 004/2022, art. 14).", "162, 219", "44, 47, 48, 49", "Partiellement reçu", "Urgent"),
    ("Conformité", "Version des comptes 2023 approuvée par l'Assemblée générale et explication de l'écart de résultat au 31/08/2026.", "Les états financiers scannés et la balance diffèrent de 10 000 000 FCFA à l'actif brut 2023.", "43", "110", "Non reçu", "Urgent"),
    ("Conformité", "Procès-verbaux des organes 2022 à 2025, registre des contrôles a posteriori, liste des prestataires conventionnés, fichier des cartes actives, liste des fraudes, effectif total de la Police nationale.", "Onze indicateurs de la Décision n° 05/2022 restent non calculables.", "162", "39, 41, 42, 52, 53", "Partiellement reçu", "Moyen"),
    ("Conformité", "État des cotisations retenues par l'État et non reversées, et état des arriérés apurés en 2026.", "Ces montants entrent dans le recouvrement des cotisations et les dettes envers les prestataires.", "162", "15, 22", "Non reçu", "Faible"),
    ("Historique", "Contrats et décomptes de gestion de la garantie décès avec les assureurs (LMAI, Sanlam), de novembre 2013 à 2020.", "En 2020, l'assureur portait le risque : les cotisations et sinistres du FPPN diffèrent des comptes.", "66", "32, 94", "Non reçu", "Faible"),
]

EXT = [
    # donnée, utilisation, diapositives, document exact, liens, ce qu'il faut transmettre, statut lien
    ("Montant du SMIG mensuel en vigueur au 31/12/2025", "Plafond de l'indemnité de départ à la retraite : 25 fois le SMIG annuel (convention collective interprofessionnelle du 19/07/1977, art. 40). L'indemnité projetée la plus élevée atteint 104 542 202 FCFA.", "128",
     "Décret n° 2022-986 du 21/12/2022 portant revalorisation du SMIG (75 000 FCFA par mois à compter du 01/01/2023, selon le texte publié). Vérifier qu'aucun décret postérieur ne l'a modifié avant le 31/12/2025.",
     "https://civlii.laws.africa/en/akn/ci/act/decree/2022/986/fra@2023-07-18/source\nhttps://www.financialafrik.com/2022/12/23/cote-divoire-le-smig-passe-de-60-000-a-75-000-fcfa-a-partir-de-janvier-2023/",
     "Le PDF du décret, ou du Journal officiel qui le publie.", "Lien trouvé, non ouvert"),
    ("Barème de dépréciation des prêts en retard des SFD", "Dépréciation des prêts du Grenier par tranche de retard (formule F-07-03).", "134, 221, 240",
     "Référentiel comptable spécifique des SFD de l'UMOA (comptes 2991 à 2993, provisions par ancienneté : moins de 6 mois, 6 à 12 mois, plus de 12 mois) et texte BCEAO qui fixe les taux. À titre de comparaison bancaire : instruction BCEAO n° 026-11-2016.",
     "https://www.cb-umoa.org/sites/default/files/2022-05/R%C3%A9f%C3%A9rentiel%20comptable%20sp%C3%A9cifique%20des%20SFD_1.pdf\nhttps://www.bceao.int/sites/default/files/2017-11/instruction_no_026-11-2016_-_relative_a_la_comptabilisation_et_a_l_evaluation_des_engagements_en_souffrance.pdf",
     "Les deux PDF.", "Lien trouvé, non ouvert"),
    ("Conditions de placement des mutuelles sociales en Côte d'Ivoire", "Règles de valorisation des actifs et conformité des placements (règlement d'exécution n° 004/2022, art. 17 : conditions fixées par décision du ministre en charge de la mutualité sociale).", "135, 178",
     "Décision (ou arrêté) du ministre de l'Emploi et de la Protection sociale fixant les conditions de sécurité, de rendement, de liquidité et de diversification des placements des mutuelles sociales.",
     "Aucun lien direct trouvé. Pistes : rubrique des textes du site de l'AIRMS (https://www.airms-ci.com) et Journal officiel de Côte d'Ivoire. Le texte est aussi demandé au FPPN (PBC, item 54).",
     "Le texte signé, avec sa date.", "Aucun lien direct trouvé"),
    ("Série ANStat des prix de la santé après octobre 2024 et pondérations", "Inflation médicale pondérée par les prestations du FPPN (onglet 13g, formule F-10-97). L'extraction du 19/10/2026 s'arrête en octobre 2024.", "191",
     "Bulletins mensuels de l'IHPC de novembre 2024 à août 2026 (tableau par fonction, ligne « 06 Santé » et postes de santé), et note méthodologique de l'IHPC base 2023 (pondérations issues de l'EHCVM 2021/2022).",
     "https://www.anstat.ci/publication-details/55c4c7ae14c38d1ecaed614da678aa712852639695b30b24f462d9124111c0e5deebf9d2855e2ae9b53b189b7f5c55fe647dd307ecde6d4b59add75cb26918deXB_UDE6G92M5goSnbIldpUK9_LOvn_voIwRYAhW0l1g (janvier 2026)\nhttps://www.anstat.ci/publication-details/53725cd7d740aa1b3ef826f9baba865482fa3ff96dc9c1e1c20b7497982478cdcb916fa86853ce67fb26ebad0f5b23affe7e9b98f0595c91947932101ec40dacFtIzOtL47TwzzUZYLFGSqmQ8MQfB5aYRgXxDIlv_lK4urck8HZ8jalYqVqQQcOes (juillet 2025)\nhttps://www.anstat.ci/assets/publications/files/File_val_indicateur1744026915.pdf (mars 2025)",
     "Les bulletins mensuels, ou une nouvelle extraction des mêmes postes que le 19/10/2026 prolongée jusqu'en août 2026, et la note sur les pondérations.", "Lien trouvé, non ouvert"),
    ("Texte du Code CIMA (provision mathématique et sinistres tardifs)", "Formules F-09-03 (sinistres tardifs, méthode CIMA) et F-09-04 (provision mathématique, art. 338 : tables et taux de 3,5 %). L'article 334-8 et la circulaire n° 00230/CIMA/CRCA/PDT/2005 sont déjà détenus.", "236 à 241",
     "Code des assurances CIMA, édition 2019, livre III, titre III (articles 334 à 338).",
     "https://cima-afrique.org/wp-content/uploads/2023/06/CODE-CIMA-2019.pdf",
     "Le PDF du code.", "Lien trouvé, non ouvert"),
    ("Cotisation des mutuelles au FOMUS-CI", "Indicateur n° 33 de la Décision n° 05/2022 (taux de règlement de la cotisation au fonds national de garantie). Le règlement n° 07/2009, art. 63, déjà détenu, renvoie à un texte d'exécution.", "162",
     "Ordonnance n° 2024-1043 du 05/12/2024 portant attributions, organisation et fonctionnement du FOMUS-CI, et texte d'application qui fixe le montant ou le taux de la contribution des mutuelles.",
     "https://agp.africanlii.org/en/akn/ci/act/ord/2024/1043/fra@2025-03-21\nhttps://civlii.laws.africa/en/akn/ci/act/ord/2024/1043/fra@2025-03-21/source",
     "L'ordonnance et le texte qui fixe la contribution.", "Lien trouvé, non ouvert"),
]

RESOLUS = [
    ("Signification du sigle CMP", "Clinique médicale du policier (indication de l'équipe Adenka du 23/10/2026). Item 55 de la PBC pointé « Reçu ».", "Question 9 retirée. Sigle corrigé dans tout le rapport."),
    ("Barème décès des officiers", "« CARTOGRAPHIE DES RISQUES COUVERTS ET DES PRESTATIONS OFFERTES.docx » (reçu le 22/10/2026) : officier de police 1 600 000 FCFA, sous-officier 1 400 000 FCFA.", "Question 19 retirée."),
    ("Compte de la provision de fin d'année des allocations décès", "En 2024 et 2025, le crédit de l'année au compte 41591000 est exactement le total de l'état de fin d'année. Le grand livre 2024 porte l'écriture du 31/12/2024 « PROVISION ALLO DECES 2024 » de 140 800 000 FCFA.", "Question 18 retirée."),
    ("Types de prêts du Grenier", "Les fichiers de production (« Production_de_la_periode 2021.xls » à « production_de_la_periode AOUT 2026 CRS.xls ») portent 8 types : ACTIVITES GENERATRICE DE REVENUS (470 prêts), CREDIT DE SANTE (993), CREDIT GRENIER SHOP (41), CREDIT HABITAT (2 023), CREDIT MARIAGE (121), CRÉDIT FLASH (8 497), CRÉDIT SCOLAIRE (3 216), EQUIPEMENT DOMESTIQUE (210). Item 70 de la PBC pointé « Reçu ».", "Question 17 retirée. Les 8 types sont cités dans le rapport."),
    ("Grands livres endommagés", "Tous ont été renvoyés en version lisible le 22/10/2026 (dossier « 02-Nouveaux_GL et Balances ») : « GRAND LIVRE DES COMPTES 2022.xlsx » (01/01 au 31/12/2022, 192 comptes), « GRAND LIVRE 2021.xlsx », « GRAND LIVRE DES TIERS 2021.xlsx », « GRAND LIVRE DES TIERS FPPN 2023.xlsx », « GRAND LIVRE FPPN 2024.xlsx », « GRAND LIVRE DES TIERS 2024.xlsx », « GRAND LIVRES DES TIERS 2025.xlsx ».", "Retirés de la question 12. Diapositive « Fichiers vides ou endommagés » corrigée."),
    ("Canal de réception des fichiers retour maladie et décès", "Les fichiers retour 936 (« …_effectif_police_936_….xlsx », dossier « FPPN-SOLDE\\EFFECTIF POLICE (MALADIE) ») et 8Z6 (« FPPN_ALLOC_DECES_8Z6_….xlsx », dossier « FPPN-SOLDE\\ALLOCATION DECES ») ont été transmis à Adenka le 15/10/2026. Le canal ne change aucun résultat.", "Question 14 retirée. Seuls les mois manquants sont demandés (onglet 1)."),
    ("Provisions figées", "Montants identiques depuis le report à nouveau du 01/09/2020 (« BALANCE FORMAT EXCEL DE 09-2020 AU 31-12-2020.xlsx ») jusqu'au 31/08/2026 : 2 867 602 878 FCFA (maladie) et 1 961 695 412 FCFA (obsèques). Seul mouvement : reclassement de 19011000 et 19012000 vers 15801000 et 15802000 en 2024, au même montant. Le chef comptable les date d'environ 2004 (entretien du 19/10/2026, archives perdues lors d'un sinistre en 2010).", "Questions 2 et 23 fusionnées : il s'agit des provisions héritées."),
    ("Factures 2020 à 2022 comptabilisées plus tard", "Mesuré par le journal SIN des grands livres des tiers 2023 à 2025 (factures d'une année enregistrées l'année suivante).", "Question 7 retirée. La base des prestations reste demandée (onglet 1) pour les délais."),
]

DOUBLONS = [
    ("51 et 61", "Fichiers aller maladie 936 et décès 8Z6 de la direction du personnel.", "Fusionner dans l'item 61 (période 2020 à 2026)."),
    ("10 et 66", "Fichier des ayants droit.", "Fusionner dans l'item 66 (format du modèle)."),
    ("13, 28, 40 et 71", "Bases des prestations (13 : toutes prestations ; 28 : décès ; 40 et 71 : santé).", "Garder 28 pour le décès et 71 pour la santé. Clore 13 et 40 en renvoyant à ces items."),
    ("14 et 29", "Dossiers de prestations non payés.", "Garder 14 (toutes garanties, toutes dates) avec le détail décès de 29."),
    ("18 et 99", "Décision fixant le taux servi sur l'épargne.", "Fusionner dans l'item 99."),
    ("16, 46 et 102", "Règlement et manuels des produits PEP et EFC.", "Items reçus : garder 102, clore 16 et 46."),
    ("3, 45 et 56", "Barème des cotisations et répartition du taux de 5 %.", "45 reçu (statuts, art. 17). Garder 56 pour la notification à la Solde. Clore 3."),
    ("2, 33, 100 et 101", "Règles et barèmes des prestations.", "Tous reçus : regrouper sous l'item 100."),
    ("30, 93 et 101", "Barème des allocations décès.", "93 et 101 reçus. Garder 30 pour les barèmes antérieurs à 2020."),
    ("50 et 86", "Origine du compte 10620000 et rapport de réévaluation.", "Fusionner dans l'item 86."),
    ("76 et 103", "Fichier du personnel.", "Fusionner dans l'item 103 (avec le statut civil ou policier)."),
    ("77 et 104", "Convention collective et contrat type du personnel.", "Fusionner dans l'item 104."),
    ("78 et 98", "Note et fichiers de calcul de la provision retraite du personnel.", "Fusionner dans l'item 78."),
    ("79, 108 et 114", "Garantie santé des retraités.", "Garder 79 (règles) et 114 (bulletin et barème). Clore 108."),
    ("37, 68 et 89", "Convention et états de GNA (flotte automobile).", "Fusionner dans l'item 89."),
    ("32 et 94", "Contrats avec les assureurs de la garantie décès.", "Fusionner dans l'item 94."),
    ("57, 60 et 69", "Documentation des produits, dont MEDOC.", "Garder 57 pour les règlements et 69 pour MEDOC. Clore 60."),
    ("88 et 119", "Clinique médicale du policier (CMP).", "Fusionner dans l'item 119."),
    ("6 et 72", "Note de calcul des provisions techniques.", "Fusionner dans l'item 72."),
]


def remplir(c, txt):
    c.value = txt
    if txt in COUL:
        c.fill = PatternFill("solid", fgColor=COUL[txt])
        c.font = Font(name="Helvetica", size=9, bold=True, color="FFFFFF" if COUL[txt] in BLANC_SUR else "000000")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def feuille(ws, titre, sous_titre, entetes, largeurs, lignes, colorees, legendes):
    ws.sheet_view.showGridLines = False
    ws["A1"] = titre
    ws["A1"].font = Font(name="Helvetica", size=14, bold=True)
    ws["A2"] = sous_titre
    ws["A2"].font = Font(name="Helvetica", size=9, italic=True)
    for j, (e, w) in enumerate(zip(entetes, largeurs), 1):
        c = ws.cell(4, j, e)
        c.font = Font(name="Helvetica", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NOIR)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORD
        ws.column_dimensions[c.column_letter].width = w
    for i, l in enumerate(lignes, 5):
        for j, v in enumerate(l, 1):
            c = ws.cell(i, j)
            c.font = Font(name="Helvetica", size=9)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BORD
            if j in colorees:
                remplir(c, v)
            else:
                c.value = v
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    ws.auto_filter.ref = f"A4:{ws.cell(4, len(entetes)).column_letter}{4 + len(lignes)}"
    r = 6 + len(lignes)
    for titre_leg, items in legendes:
        ws.cell(r, 2, titre_leg).font = Font(name="Helvetica", size=10, bold=True)
        r += 1
        for code, sens in items:
            c = ws.cell(r, 2)
            remplir(c, code)
            c.border = BORD
            d = ws.cell(r, 3, sens)
            d.font = Font(name="Helvetica", size=9)
            d.alignment = Alignment(wrap_text=True, vertical="center")
            d.border = BORD
            r += 1
        r += 1


LEG_PRIO = ("Code couleur : priorité", [("Urgent", "Bloque une conclusion centrale du rapport (provisions, engagements, conformité)."),
                                        ("Moyen", "Précise un résultat déjà établi ou un indicateur."),
                                        ("Faible", "Complète la documentation sans changer les conclusions.")])
LEG_STAT = ("Code couleur : statut dans la PBC du 22/10/2026", [("Non reçu", "Aucun document reçu pour cet item."),
                                                               ("Partiellement reçu", "Une partie seulement a été reçue."),
                                                               ("Nouvelle demande", "Élément absent de la PBC, à y ajouter.")])
wb = Workbook()
ws = wb.active
ws.title = "1_A obtenir du FPPN"
feuille(ws, "Éléments à obtenir du FPPN pour terminer le rapport",
        "Liste exhaustive tirée du rapport et de la PBC pointée au 22/10/2026. Numéros de diapositive du rapport du 22/10/2026. Une ligne regroupe les items de PBC qui demandent la même pièce.",
        ["N°", "Thème", "Élément à obtenir", "Pourquoi : constat ou analyse concerné", "Diapositives", "Items PBC", "Statut PBC", "Priorité"],
        [5, 16, 60, 60, 12, 12, 16, 11],
        [(i, *x) for i, x in enumerate(FPPN, 1)], {7, 8}, [LEG_PRIO, LEG_STAT])
ws = wb.create_sheet("2_Données externes")
feuille(ws, "Données externes à rechercher et à transmettre à Adenka",
        "Textes officiels et séries publiques absents des documents du client et des documents déjà collectés. Les liens viennent d'une recherche web du 23/10/2026 : l'environnement de travail d'Adenka ne peut pas les ouvrir, il faut donc les vérifier à l'ouverture.",
        ["N°", "Donnée", "Utilisation dans le rapport", "Diapositives", "Document exact à récupérer", "Liens", "À transmettre", "Statut du lien"],
        [5, 26, 48, 11, 52, 70, 30, 16],
        [(i, *x) for i, x in enumerate(EXT, 1)], {8},
        [("Code couleur : statut du lien", [("Lien trouvé, non ouvert", "Adresse trouvée par recherche web, à ouvrir et vérifier."),
                                            ("Aucun lien direct trouvé", "Texte à obtenir auprès de l'AIRMS, du Journal officiel ou du FPPN.")])])
ws = wb.create_sheet("3_Points déjà résolus")
feuille(ws, "Questions retirées : la réponse figure dans les documents reçus",
        "Ces points ne seront plus posés au FPPN. La preuve est citée avec le nom exact du fichier.",
        ["N°", "Point", "Preuve dans les documents", "Effet sur le rapport"], [5, 30, 90, 45],
        [(i, *x) for i, x in enumerate(RESOLUS, 1)], set(), [])
ws = wb.create_sheet("4_Doublons PBC")
feuille(ws, "Doublons de la PBC du 22/10/2026",
        "Items qui demandent la même pièce. La fusion proposée sera appliquée dans la prochaine version de la PBC.",
        ["N°", "Items concernés", "Pièce demandée", "Fusion proposée"], [5, 18, 60, 60],
        [(i, *x) for i, x in enumerate(DOUBLONS, 1)], set(), [])
for w in wb.worksheets:
    for row in w.iter_rows(min_row=5):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value.split("\n")[0].split(" ")[0]
                c.font = Font(name="Helvetica", size=9, color="0563C1", underline="single")
wb.save(SORTIE)
print(SORTIE, len(FPPN), len(EXT), len(RESOLUS), len(DOUBLONS))
