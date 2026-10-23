# -*- coding: utf-8 -*-
"""
FPPN - Logigrammes modifiables pour PowerPoint (formes natives et connecteurs attachés)
======================================================================================

Formes standard des logigrammes : début / fin (terminateur), étape (processus), décision (losange),
document. Les connecteurs sont attachés aux formes : déplacer une forme dans PowerPoint entraîne la flèche.
Des couloirs (« swimlanes ») matérialisent les acteurs. Une forme « alerte » (contour rouge) signale
une étape où un contrôle manque.

Utilisation :
    from logigramme_fppn import Logigramme
    lg = Logigramme(slide)
    lg.couloirs(["Acteur 1", "Acteur 2"], y0=3.9, hauteur=2.4)
    lg.noeud("a", "debut", "Texte", x=6, y=4.3)
    lg.lien("a", "b", "oui")
"""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

NOIR, GRIS, BLEU, ORANGE, ROUGE, GRIS_F, GRIS_C = "060606", "F2F2F2", "112A40", "EFAC62", "C00000", "53565A", "BBBCBC"
FORMES = {"debut": MSO_SHAPE.FLOWCHART_TERMINATOR, "fin": MSO_SHAPE.FLOWCHART_TERMINATOR,
          "etape": MSO_SHAPE.FLOWCHART_PROCESS, "decision": MSO_SHAPE.FLOWCHART_DECISION,
          "document": MSO_SHAPE.FLOWCHART_DOCUMENT, "alerte": MSO_SHAPE.FLOWCHART_PROCESS,
          "donnee": MSO_SHAPE.FLOWCHART_DATA}
STYLE = {  # fond, texte, contour
    "debut": (BLEU, "FFFFFF", None), "fin": (BLEU, "FFFFFF", None), "etape": ("FFFFFF", NOIR, BLEU),
    "decision": (ORANGE, NOIR, None), "document": (GRIS, NOIR, GRIS_F), "alerte": ("FFFFFF", ROUGE, ROUGE),
    "donnee": (GRIS, NOIR, GRIS_F)}
# Points de connexion des formes de logigramme : 0 haut, 1 gauche, 2 bas, 3 droite
COTE = {"haut": 0, "gauche": 1, "bas": 2, "droite": 3}


def _texte(tf, texte, taille, couleur, gras=False, align=PP_ALIGN.CENTER):
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Cm(0.08)
    tf.margin_top = tf.margin_bottom = Cm(0.03)
    for i, ligne in enumerate(str(texte).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ligne
        r.font.name = "Helvetica"
        r.font.size = Pt(taille)
        r.font.bold = gras or (i == 0 and "\n" in str(texte))
        r.font.color.rgb = RGBColor.from_string(couleur)


class Logigramme:
    def __init__(self, slide, taille=9):
        self.s = slide
        self.taille = taille
        self.n = {}

    def couloirs(self, noms, y0=3.9, hauteur=2.4, x0=1.4, largeur=31.2, entete=3.9):
        """Couloirs horizontaux : un par acteur, nom à gauche sur fond bleu nuit."""
        self.lanes = {}
        for k, nom in enumerate(noms):
            y = y0 + k * hauteur
            fond = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x0), Cm(y), Cm(largeur), Cm(hauteur))
            fond.fill.solid()
            fond.fill.fore_color.rgb = RGBColor.from_string("FFFFFF" if k % 2 else "F7F7F7")
            fond.line.color.rgb = RGBColor.from_string(GRIS_C)
            fond.line.width = Pt(0.75)
            fond.shadow.inherit = False
            tete = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x0), Cm(y), Cm(entete), Cm(hauteur))
            tete.fill.solid()
            tete.fill.fore_color.rgb = RGBColor.from_string(BLEU)
            tete.line.fill.background()
            tete.shadow.inherit = False
            _texte(tete.text_frame, nom, 10, "FFFFFF", True)
            self.lanes[k] = (y, hauteur)
        return self

    def centre_couloir(self, k, h):
        y, H = self.lanes[k]
        return y + (H - h) / 2

    def noeud(self, cle, type_, texte, x, y, l=3.3, h=1.55, taille=None):
        if type_ == "decision":
            h = max(h, 1.9)
        sh = self.s.shapes.add_shape(FORMES[type_], Cm(x), Cm(y), Cm(l), Cm(h))
        fond, coul, contour = STYLE[type_]
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fond)
        if contour:
            sh.line.color.rgb = RGBColor.from_string(contour)
            sh.line.width = Pt(2.25 if type_ == "alerte" else 1.25)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        _texte(sh.text_frame, texte, taille or (self.taille - 0.5 if type_ == "decision" else self.taille), coul,
               type_ in ("debut", "fin", "alerte"))
        sh.name = f"Logigramme {cle}"
        self.n[cle] = sh
        return sh

    def lien(self, a, b, etiquette=None, de=None, vers=None, couleur=GRIS_F, pointille=False):
        """Flèche attachée de la forme a à la forme b. Côtés choisis automatiquement si non précisés."""
        A, B = self.n[a], self.n[b]
        ax, ay = A.left + A.width / 2, A.top + A.height / 2
        bx, by = B.left + B.width / 2, B.top + B.height / 2
        if de is None or vers is None:
            if abs(ay - by) < Cm(0.3):
                de, vers = ("droite", "gauche") if bx > ax else ("gauche", "droite")
            elif abs(ax - bx) < Cm(0.3):
                de, vers = ("bas", "haut") if by > ay else ("haut", "bas")
            else:
                de = "droite" if bx > ax else "gauche"
                vers = "haut" if by > ay else "bas"
        droit = (de, vers) in (("droite", "gauche"), ("gauche", "droite"), ("bas", "haut"), ("haut", "bas")) and \
            (abs(ay - by) < Cm(0.3) or abs(ax - bx) < Cm(0.3))
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT if droit else MSO_CONNECTOR.ELBOW, 0, 0, Cm(1), Cm(1))
        c.begin_connect(A, COTE[de])
        c.end_connect(B, COTE[vers])
        c.line.color.rgb = RGBColor.from_string(couleur)
        c.line.width = Pt(1.75)
        ln = c.line._get_or_add_ln()
        if pointille:
            ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        c.name = f"Flèche {a} vers {b}"
        if etiquette:
            bx_, by_ = c.begin_x / 360000, c.begin_y / 360000
            dx, dy, al = {"droite": (0.12, -0.5, PP_ALIGN.LEFT), "gauche": (-2.52, -0.5, PP_ALIGN.RIGHT),
                          "bas": (0.15, 0.05, PP_ALIGN.LEFT), "haut": (0.15, -0.5, PP_ALIGN.LEFT)}[de]
            tb = self.s.shapes.add_textbox(Cm(bx_ + dx), Cm(by_ + dy), Cm(2.4), Cm(0.45))
            _texte(tb.text_frame, etiquette, self.taille - 1.5, GRIS_F, True, al)
            tb.text_frame.vertical_anchor = MSO_ANCHOR.TOP
        return c

    def note(self, texte, x, y, l, h, couleur=GRIS_F, taille=None):
        tb = self.s.shapes.add_textbox(Cm(x), Cm(y), Cm(l), Cm(h))
        _texte(tb.text_frame, texte, taille or self.taille - 1, couleur, False, PP_ALIGN.LEFT)
        return tb

    def legende(self, x, y, types=("debut", "etape", "decision", "document", "alerte")):
        noms = {"debut": "Début ou fin", "etape": "Étape", "decision": "Décision", "document": "Document",
                "alerte": "Contrôle absent", "donnee": "Fichier"}
        for k, t in enumerate(types):
            xx = x + k * 4.2
            sh = self.s.shapes.add_shape(FORMES[t], Cm(xx), Cm(y), Cm(0.9), Cm(0.5))
            fond, _, contour = STYLE[t]
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor.from_string(fond)
            if contour:
                sh.line.color.rgb = RGBColor.from_string(contour)
                sh.line.width = Pt(1.5 if t == "alerte" else 1)
            else:
                sh.line.fill.background()
            sh.shadow.inherit = False
            self.note(noms[t], xx + 1.0, y - 0.05, 3.1, 0.6, NOIR, 8)


# ---------------------------------------------------------------------------
# Logigrammes de la mission
# ---------------------------------------------------------------------------
def circuit_precomptes(slide, y0=3.9, h_couloir=2.5):
    """Circuit des précomptes déclaré lors de l'entretien du 19/10/2026 (responsable des cotisations)."""
    lg = Logigramme(slide, 8.5)
    lg.couloirs(["Direction du personnel de la Police", "Solde, CGRAE, GNA", "Plateforme de partage",
                 "FPPN : responsable des cotisations"], y0, h_couloir)
    L, H = 3.4, 1.7
    X = [5.6 + k * 3.85 for k in range(7)]
    c = lg.centre_couloir
    # Produits obligatoires : maladie (936) et décès (8Z6), transmission directe sans partage
    lg.noeud("a1", "document", "Fichier aller\n936 et 8Z6", X[0], c(0, H), L, H)
    lg.noeud("a2", "etape", "Précompte\nsur salaire (Solde)", X[0], c(1, H), L, H)
    lg.noeud("a3", "document", "Fichier retour reçu\nen direct, canal à préciser", X[1], c(3, H), L, H)
    lg.noeud("a4", "decision", "Fichier aller\nau FPPN ?", X[2], c(3, 1.9), L, 1.9)
    lg.noeud("a5", "alerte", "Rapprochement\nimpossible", X[3], c(3, H) - 0.02, L, H)
    lg.lien("a1", "a2", "transmission directe")
    lg.lien("a2", "a3", "sans partage", de="droite", vers="haut")
    lg.lien("a3", "a4")
    lg.lien("a4", "a5", "non")
    # Produits complémentaires (Solde), CGRAE et GNA : via la plateforme de partage
    lg.noeud("b1", "debut", "Souscriptions\nEFC, PEP, prêts, GNA, MEDOC", X[4], c(3, H), L, H)
    lg.noeud("b2", "document", "Fichier aller mensuel\nmécano, code, montant", X[5], c(3, H), L, H)
    lg.noeud("b3", "etape", "Dépôt du\nfichier aller", X[5], c(2, H), L, H)
    lg.noeud("b4", "etape", "Précompte\nsur salaire ou pension", X[5], c(1, H), L, H)
    lg.noeud("b5", "document", "Fichier retour\net état des rejets récent", X[6], c(2, H), L, H)
    lg.noeud("b6", "alerte", "Retour conservé,\nsans rapprochement", X[6], c(3, H), L, H)
    lg.lien("b1", "b2")
    lg.lien("b2", "b3")
    lg.lien("b3", "b4")
    lg.lien("b4", "b5", de="droite", vers="haut")
    lg.noeud("b7", "document", "Convention IPS-CGRAE :\nquotité cessible ; soins\nmaintenus si retard", X[3], c(1, H), L, H)
    lg.lien("b7", "b4", "retraités", pointille=True)
    lg.lien("b5", "b6")
    lg.note("Produits obligatoires", X[0], y0 - 0.02, 7.6, 0.45, ROUGE, 8)
    lg.note("Produits complémentaires, retraités et flotte automobile", X[4], y0 - 0.02, 11.4, 0.45, ROUGE, 8)
    return lg


def nouvelles_recrues(slide, y0=4.0):
    """Parcours d'une nouvelle recrue, de l'école de police au premier précompte (entretien du 19/10/2026)."""
    lg = Logigramme(slide, 9)
    L, H = 3.6, 1.6
    X = [1.6 + k * 4.45 for k in range(7)]
    y1, y2 = y0 + 0.6, y0 + 4.4
    lg.noeud("r1", "debut", "Élève policier\nmatricule sans mécano", X[0], y1, L, H)
    lg.noeud("r2", "document", "Fiches de souscription\nEFC et PEP", X[1], y1, L, H)
    lg.noeud("r3", "etape", "Attente\ndu mécano", X[2], y1, L, H)
    lg.noeud("r4", "decision", "Mécano\nattribué ?", X[3], y1 - 0.08, L, 1.75)
    lg.noeud("r5", "alerte", "Report du mécano\nsur chaque fiche, en masse", X[4], y1, L, H)
    lg.noeud("r6", "etape", "Ajout au fichier\naller du mois", X[5], y1, L, H)
    lg.noeud("r7", "decision", "Précompte\nau 1er mois ?", X[6], y1 - 0.08, L, 1.75)
    lg.noeud("r10", "fin", "Précompte effectif\népargne en cours", X[6], y2, L, H)
    lg.noeud("r8", "alerte", "Précompte décalé\nau mois suivant", X[5], y2, L, H)
    lg.noeud("r9", "etape", "Paiement manuel\npossible, avec date", X[4], y2, L, H)
    lg.lien("r1", "r2")
    lg.lien("r2", "r3")
    lg.lien("r3", "r4")
    lg.lien("r4", "r5", "oui")
    lg.lien("r4", "r3", "non", de="haut", vers="haut", pointille=True)
    lg.lien("r5", "r6")
    lg.lien("r6", "r7")
    lg.lien("r7", "r8", "non", de="bas", vers="haut")
    lg.lien("r8", "r9", "ou", pointille=True)
    lg.lien("r7", "r10", "oui", de="droite", vers="droite")
    return lg


def chaine_traitement(slide, y0=3.95):
    """Chaîne de traitement Adenka, du fichier reçu au chiffre du rapport, avec ses contrôles."""
    lg = Logigramme(slide, 8.5)
    L, H = 3.5, 1.5
    X = [1.5 + k * 4.45 for k in range(7)]
    y1, y2, y3 = y0 + 0.25, y0 + 4.1, y0 + 7.9
    lg.noeud("c1", "debut", "Envois du FPPN : 4 287\nfichiers (15/10), 124\nfichiers (22/10)", X[0], y1, L, H)
    lg.noeud("c2", "etape", "Inventaire\nformat réel, doublons", X[1], y1, L, H)
    lg.noeud("c3", "decision", "Fichier\nlisible ?", X[2], y1 - 0.12, L, 1.75)
    lg.noeud("c4", "document", "Liste des fichiers\nà renvoyer (PBC)", X[2], y2, L, H)
    lg.noeud("c5", "etape", "Extraction sans perte,\nempreinte numérique ;\nscans lus par OCR", X[3], y1, L, H)
    lg.noeud("c6", "decision", "Lignes\nidentiques ?", X[4], y1 - 0.12, L, 1.75)
    lg.noeud("c7", "etape", "Base d'analyse\n4,45 millions de lignes", X[5], y1, L, H)
    lg.noeud("c8", "decision", "Total ou\nmontant doublé ?", X[6], y1 - 0.12, L, 1.75)
    lg.noeud("c9", "document", "Journal des rejets\net corrections", X[6], y2, L, H)
    lg.noeud("c10", "etape", "Contrôles : totaux,\naller et retour,\ncontrôles croisés", X[5], y2, L, H)
    lg.noeud("c11", "etape", "Analyses\nindicateurs, tests, ratios", X[4], y2, L, H)
    lg.noeud("c12", "document", "Classeur des calculs :\nformules et contrôle\nau franc près", X[3], y2, L, H)
    lg.noeud("c13", "etape", "Rapport\ngraphiques natifs", X[3], y3, L, H)
    lg.noeud("c14", "decision", "Chiffre\nexact ?", X[4], y3 - 0.12, L, 1.75)
    lg.noeud("c15", "fin", "Livraison\net traçabilité", X[5], y3, L, H)
    lg.lien("c1", "c2")
    lg.lien("c2", "c3")
    lg.lien("c3", "c4", "non")
    lg.lien("c3", "c5", "oui")
    lg.lien("c5", "c6")
    lg.lien("c6", "c5", "non : refaire", de="haut", vers="haut", pointille=True)
    lg.lien("c6", "c7", "oui")
    lg.lien("c7", "c8")
    lg.lien("c8", "c9", "oui", de="droite", vers="droite")
    lg.lien("c8", "c10", "non", de="bas", vers="haut")
    lg.lien("c9", "c10")
    lg.lien("c10", "c11")
    lg.lien("c11", "c12")
    lg.lien("c12", "c13")
    lg.lien("c13", "c14")
    lg.lien("c14", "c15", "oui")
    lg.lien("c14", "c13", "non : corriger", de="haut", vers="haut", pointille=True)
    return lg


def decision_placement(slide, y0=3.9, h_couloir=2.35):
    """Circuit de décision d'un placement décrit par le chef comptable (entretien du 19/10/2026)."""
    lg = Logigramme(slide, 8.5)
    lg.couloirs(["Banque", "Direction générale", "DAF", "Comptabilité (chef comptable)"], y0, h_couloir)
    L, H = 3.4, 1.6
    X = [5.6 + k * 3.85 for k in range(7)]
    c = lg.centre_couloir
    lg.noeud("p1", "debut", "Proposition\nde placement", X[0], c(0, H), L, H)
    lg.noeud("p2", "etape", "Analyse de\nla proposition", X[1], c(1, H), L, H)
    lg.noeud("p3", "etape", "Demande du point\nde trésorerie", X[2], c(2, H), L, H)
    lg.noeud("p4", "etape", "Point de\ntrésorerie", X[2], c(3, H), L, H)
    lg.noeud("p5", "decision", "Investir ?", X[3], c(1, 1.9), L, 1.9)
    lg.noeud("p6", "fin", "Pas de\nplacement", X[3], c(0, H), L, H)
    lg.noeud("p7", "document", "Lettre d'ouverture\nà la banque", X[4], c(2, H), L, H)
    lg.noeud("p8", "etape", "Prélèvements\nmensuels", X[5], c(0, H), L, H)
    lg.noeud("p9", "document", "Relevés\nbancaires", X[6], c(0, H), L, H)
    lg.noeud("p10", "etape", "Comptabilisation\nsur relevés", X[6], c(3, H), L, H)
    lg.noeud("p11", "fin", "Débouclage, dont\npaiement de l'EFC", X[5], c(3, H), L, H)
    lg.noeud("p12", "alerte", "Ni comité, ni\nprocédure écrite", X[4], c(3, H), L, H)
    lg.noeud("p13", "alerte", "Mandat du directeur\ngénéral documenté\npour 2020-2023 seulement", X[0], c(1, H), L, H)
    lg.noeud("p14", "alerte", "Projet immobilier de\nCocody : décision des\norganes non reçue", X[0], c(2, H), L, H)
    lg.lien("p1", "p2", de="droite", vers="haut")
    lg.lien("p2", "p3", de="droite", vers="haut")
    lg.lien("p3", "p4")
    lg.lien("p4", "p5", de="droite", vers="bas")
    lg.lien("p5", "p6", "non", de="haut", vers="bas")
    lg.lien("p5", "p7", "oui", de="droite", vers="haut")
    lg.lien("p7", "p8", de="droite", vers="bas")
    lg.lien("p8", "p9")
    lg.lien("p9", "p10")
    lg.lien("p10", "p11")
    return lg


def provision_deces(slide, y0=4.0):
    """Calcul de fin d'année des allocations décès restant à payer, décrit par le chef comptable (19/10/2026)."""
    lg = Logigramme(slide, 9)
    L, H = 3.6, 1.6
    X = [1.6 + k * 4.45 for k in range(7)]
    y1, y2 = y0 + 0.6, y0 + 4.2
    lg.noeud("d1", "debut", "Liste des décès\n(direction du personnel)", X[0], y1, L, H)
    lg.noeud("d2", "etape", "Informaticien :\narrêt du précompte", X[1], y1, L, H)
    lg.noeud("d3", "document", "Liste annuelle\nau chef comptable", X[2], y1, L, H)
    lg.noeud("d4", "etape", "Pointage avec les\nallocations payées", X[3], y1, L, H)
    lg.noeud("d5", "decision", "Allocation\npayée ?", X[4], y1 - 0.15, L, 1.9)
    lg.noeud("d6", "fin", "Charge de\nl'exercice", X[5], y1, L, H)
    lg.noeud("d7", "etape", "Barème : 1,4 / 1,6 / 2,2 M\nFCFA ; conjoint, enfant ;\nsuicide valorisé à 0", X[4], y2, L, H)
    lg.noeud("d8", "fin", "Dette envers les\nbénéficiaires (41591000)", X[5], y2, L, H)
    lg.noeud("d9", "alerte", "Décès non encore\ndéclarés au 31/12\n(tardifs)", X[0], y2, L, H)
    lg.noeud("d10", "alerte", "Forclusion (RI, art. 58)\nnon vérifiable : date\nde dépôt non saisie", X[1], y2, L, H)
    lg.noeud("d11", "alerte", "2024 : 102 allocations\npayées par les comptes\nbancaires de l'épargne", X[2], y2, L, H)
    lg.noeud("d12", "fin", "Frais funéraires :\ndette au 41592000", X[3], y2, L, H)
    lg.noeud("d13", "alerte", "Provision obsèques\n15802000 inchangée :\ndouble emploi", X[6], y2, L, H)
    lg.lien("d8", "d13", "en plus", pointille=True)
    lg.lien("d1", "d2")
    lg.lien("d2", "d3")
    lg.lien("d3", "d4")
    lg.lien("d4", "d5")
    lg.lien("d5", "d6", "oui")
    lg.lien("d5", "d7", "non", de="bas", vers="haut")
    lg.lien("d7", "d8")
    lg.lien("d1", "d9", "hors liste", de="bas", vers="haut", pointille=True)
    return lg


def circuit_facture(slide, y0=4.0):
    """Circuit d'une prestation maladie, de la facture au paiement (règlement intérieur, art. 52, 57 et 60)."""
    lg = Logigramme(slide, 9)
    L, H = 3.6, 1.6
    X = [1.6 + k * 4.45 for k in range(7)]
    y1, y2 = y0 + 0.6, y0 + 4.2
    lg.noeud("f1", "debut", "Soins dans une\nclinique conventionnée", X[0], y1, L, H)
    lg.noeud("f2", "document", "Facture de la clinique\n(numéro, date)", X[1], y1, L, H)
    lg.noeud("f3", "etape", "Dépôt au FPPN\n(date de dépôt)", X[2], y1, L, H)
    lg.noeud("f4", "etape", "Enregistrement au\n40110000 (journal SIN)", X[3], y1, L, H)
    lg.noeud("f5", "decision", "Payée dans\nles 3 mois ?", X[4], y1 - 0.15, L, 1.9)
    lg.noeud("f6", "fin", "Paiement de la clinique\n(RI, art. 52)", X[5], y1, L, H)
    lg.noeud("f7", "alerte", "Arriérés : délai moyen\n53 mois en 2020,\n2,62 mois en 2025", X[5], y2, L, H)
    lg.noeud("f8", "debut", "Soins payés par\nle mutualiste", X[0], y2, L, H)
    lg.noeud("f9", "etape", "Demande de rembourse-\nment au plus 3 mois\naprès la facture (art. 60)", X[1], y2, L, H)
    lg.noeud("f10", "fin", "Remboursement au plus\n6 mois après la\ndemande (art. 57)", X[2], y2, L, H)
    lg.noeud("f11", "alerte", "Date des soins\nnon saisie en\ncomptabilité", X[3], y2, L, H)
    lg.noeud("f12", "document", "Provision maladie : soins\nde l'année facturés\naprès le 31/12", X[6], y2, L, H)
    lg.lien("f1", "f2")
    lg.lien("f2", "f3")
    lg.lien("f3", "f4")
    lg.lien("f4", "f5")
    lg.lien("f5", "f6", "oui")
    lg.lien("f5", "f7", "non", de="bas", vers="haut")
    lg.lien("f8", "f9")
    lg.lien("f9", "f10")
    lg.lien("f4", "f11", de="bas", vers="haut", pointille=True)
    lg.lien("f6", "f12", "facture déposée après le 31/12", de="droite", vers="haut", pointille=True)
    return lg
