# -*- coding: utf-8 -*-
"""
FPPN - Compte rendu de l'entretien du 19/10/2026 avec le responsable des cotisations (PowerPoint, charte Adenka)
==============================================================================================================

Les déclarations de l'interlocuteur sont présentées comme telles (« à confirmer par document ») et confrontées
aux fichiers de précomptes reçus le 15/10/2026. Les propos sensibles figurent dans la note interne confidentielle.
Utilisation : python cr_cotisations_fppn.py <dossier_base> <skill_adenka_powerpoint> <sortie.pptx>
"""
import os
import sys

import pandas as pd

BASE, SKILL, SORTIE = sys.argv[1:4]
sys.path.insert(0, os.path.join(SKILL, "scripts"))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx_tools import selection, remplir, finaliser  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.util import Cm, Pt  # noqa: E402
import logigramme_fppn as LG  # noqa: E402
from outils_slides_fppn import (BAS_CONTENU, BLEU, GRIS, GRIS_F, NOIR, ORANGE, corps, ligne_sigles, tableau,  # noqa: E402
                                txt)

TIMESAVER = os.path.join(SKILL, "assets", "Timesaver_Adenka.pptx")
SRC = "Source : entretien du 19/10/2026 avec le responsable des cotisations ; déclarations à confirmer par document"
SRC_D = SRC + " ; fichiers de précomptes reçus le 15/10/2026, traitements Adenka"
M = ["janv.", "févr.", "mars", "avr.", "mai", "juin", "juil.", "août", "sept.", "oct.", "nov.", "déc."]


def nb(x, dec=0):
    return f"{x:,.{dec}f}".replace(",", " ").replace(".", ",")


def pc(x):
    return nb(100 * x, 2) + " %"


# ---------------------------------------------------------------------------
# Confrontation aux données
# ---------------------------------------------------------------------------
pre = pd.read_parquet(os.path.join(BASE, "precomptes.parquet"))
pre = pre[pre.retenu & (pre.sens == "retour") & (pre.annee >= 2020) & ((pre.annee < 2026) | (pre.mois <= 8))]
efc = pre[pre.poste == "870"]
part10 = {a: (g.montant == 10000).mean() for a, g in efc.groupby("annee")}
part6 = {a: (g.montant == 6000).mean() for a, g in efc.groupby("annee")}
pep = pre[pre.poste == "998"]
pep_m5 = (pep.montant % 5000 == 0).mean()
pep_lt10 = (pep.montant < 10000).mean()
pep_max = pep.montant.max()
pep_med = {a: g.montant.median() for a, g in pep.groupby("annee")}
gna = pre[pre.poste == "9W3"].montant.value_counts()
medoc = pre[pre.poste == "8Z7"].groupby(["annee", "mois"]).size()

inv = pd.read_parquet(os.path.join(BASE, "sources_mois.parquet"))
inv = inv[inv.retenu]
MOIS = [(a, m) for a in range(2020, 2027) for m in range(1, 13) if not (a == 2026 and m > 8)]


def plages(liste):
    out, s, p = [], None, None
    for x in liste:
        i = MOIS.index(x)
        if s is None:
            s = p = i
        elif i == p + 1:
            p = i
        else:
            out.append((s, p))
            s = p = i
    if s is not None:
        out.append((s, p))
    f = lambda i: f"{M[MOIS[i][1] - 1]} {MOIS[i][0]}"  # noqa: E731
    return ", ".join(f(a) if a == b else f"{f(a)} à {f(b)}" for a, b in out) or "aucun"


manq = {}
for po in ("936", "8Z6", "870", "998", "9X4", "9W3", "8Z7", "CGRAE"):
    for se in ("aller", "retour"):
        x = inv[(inv.poste == po) & (inv.sens == se)]
        have = set(zip(x.annee, x.mois))
        manq[(po, se)] = (len(have & set(MOIS)), plages([m for m in MOIS if m not in have]))
MANQ = manq

# ---------------------------------------------------------------------------
# Contenu
# ---------------------------------------------------------------------------
SL = []


def diapo(titre, sous_titre, synthese, rendu, notes="", source=SRC):
    SL.append({"entete": {"titre": titre, "sous_titre": sous_titre, "synthese": synthese}, "rendu": rendu, "notes": notes, "source": source})


def r_infos(s):
    tiles = [("19 octobre 2026", "Date de l'entretien"), ("Cotisations", "Circuit des précomptes et produits"),
             ("Kouamé Attuoman", "responsable des cotisations et informatique"), ("Documentation", "retour annoncé le jour même")]
    tw = (31.2 - 3 * 0.4) / 4
    for i, (a, b) in enumerate(tiles):
        x = 1.4 + i * (tw + 0.4)
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(3.85), Cm(tw), Cm(1.8))
        r.fill.solid()
        r.fill.fore_color.rgb = RGBColor.from_string(GRIS)
        r.line.fill.background()
        r.shadow.inherit = False
        txt(s, x + 0.3, 3.95, tw - 0.6, 0.8, a, 15, True, ORANGE)
        txt(s, x + 0.3, 4.85, tw - 0.6, 0.6, b, 9, False, GRIS_F)
    bas = tableau(s, {"titre": None, "entetes": ["Structure", "Nom", "Fonction", "Rôle dans l'entretien"], "largeurs": [4.6, 7.0, 10.6, 9.0],
                      "lignes": [["FPPN", "Kouamé Attuoman", "Responsable des cotisations et responsable informatique", "Interlocuteur"],
                                 ["Adenka", "Jean-Noël Sialou", "Consultant actuaire", "Conduite de l'entretien"],
                                 ["Adenka", "Ruth Miensa", "Équipe Adenka", "Participante"],
                                 ["Adenka", "Eliam Affolaby", "Équipe Adenka", "Participante, échanges avec l'interlocuteur"]]}, 6.2, taille=11)
    corps(s, 1.4, bas + 0.5, 31.2, BAS_CONTENU - bas - 0.5,
          ["**Objet de l'entretien.** Comprendre comment les cotisations sont précomptées et suivies, comment le barème est fixé et quelles règles s'appliquent à chaque produit.",
           "**Méthode.** Chaque déclaration est présentée comme telle et confrontée aux fichiers de précomptes reçus le 15/10/2026. Les points non documentés sont repris dans la demande complémentaire de documents (PBC, items 56 à 70)."], 12)


def r_organisation(s):
    blocs = [("Fonction", "Le responsable des cotisations est aussi le responsable informatique. Il gère les cotisations parce qu'il maîtrise le logiciel qui produit les fichiers."),
             ("Équipe", "Des collaborateurs traitent chaque jour les souscriptions, les rachats et les demandes des mutualistes. L'interlocuteur les décrit comme très chargés."),
             ("Barème", "L'interlocuteur renvoie à la comptabilité pour la répartition du taux entre maladie et décès. Le barème doit être déposé sur la plateforme de partage."),
             ("Suivi", "Il n'existe ni état des cotisations, ni rapprochement des fichiers aller et retour, ni suivi des précomptes mutualiste par mutualiste.")]
    for k, (t, c) in enumerate(blocs):
        x, y = 1.4 + (k % 2) * 15.9, 3.9 + (k // 2) * 4.3
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(15.3), Cm(3.9))
        b.fill.background()
        b.line.color.rgb = RGBColor.from_string(BLEU)
        b.line.width = Pt(1.5)
        b.shadow.inherit = False
        h = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(15.3), Cm(0.9))
        h.fill.solid()
        h.fill.fore_color.rgb = RGBColor.from_string(BLEU)
        h.line.fill.background()
        txt(s, x + 0.3, y + 0.1, 14.7, 0.7, t, 13, True, "FFFFFF")
        txt(s, x + 0.4, y + 1.1, 14.5, 2.7, c, 12)
    corps(s, 1.4, 12.6, 31.2, BAS_CONTENU - 12.6,
          ["**Le suivi des cotisations repose sur une seule personne, dont ce n'est pas le métier principal.** Les fichiers sont produits et déposés chaque mois, mais leur retour n'est ni contrôlé ni rapproché.",
           "**Implication pour la mission.** Les écarts entre ce qui est demandé et ce qui est précompté ne sont pas expliqués par le FPPN : Adenka les mesure à partir des fichiers (rapport, partie 3), mais leur justification doit venir du FPPN."], 12)


def r_circuit(s):
    LG.circuit_precomptes(s, 3.9, 2.4)
    corps(s, 1.4, 13.75, 31.2, BAS_CONTENU - 13.75,
          ["**Deux circuits coexistent.** Pour la maladie et le décès, la direction du personnel de la Police transmet directement le fichier aller à la Solde, sans plateforme de partage, et le FPPN ne reçoit que le retour. Pour les produits complémentaires, les retraités et la flotte automobile, le responsable des cotisations dépose le fichier aller sur la plateforme et l'entité concernée y dépose le retour.",
           "**Implication pour le FPPN.** Pour la maladie et le décès, le FPPN ne dispose pas de ce qui aurait dû être précompté : il ne peut pas mesurer ses pertes de cotisations. Pour les autres produits, il dispose des deux fichiers mais ne les rapproche pas."], 11)


def r_circuits_produits(s):
    bas = tableau(s, {"titre": "Circuit des fichiers par produit, selon l'interlocuteur", "entetes": ["Produits", "Fichier aller", "Destinataire", "Fichier retour", "Plateforme de partage"],
                      "largeurs": [6.8, 8.4, 3.6, 8.0, 4.4],
                      "lignes": [["Maladie (936) et décès (8Z6)", "Transmis directement par la direction du personnel de la Police", "Solde", "Transmis directement au FPPN, canal à préciser", "Non"],
                                 ["EFC (870), PEP (998), Grenier (9X4), MEDOC (8Z7)", "Produit et déposé par le responsable des cotisations", "Solde", "Déposé par la Solde", "Oui"],
                                 ["Retraités", "Produit et déposé par le responsable des cotisations", "CGRAE", "Déposé par la CGRAE", "Oui"],
                                 ["Flotte automobile (9W3)", "Produit et déposé par le responsable des cotisations", "GNA", "Déposé par GNA", "Oui"]]}, 3.85, taille=11)
    corps(s, 1.4, bas + 0.4, 31.2, BAS_CONTENU - bas - 0.4,
          [f"**Ce circuit explique l'absence totale des fichiers aller maladie et décès dans l'envoi du 15/10/2026** : {MANQ[('936', 'aller')][0]} mois reçus sur 80 pour la maladie et {MANQ[('8Z6', 'aller')][0]} pour le décès. Le FPPN ne les produit pas.",
           "**Le fichier aller contient le mécano, le nom, le code du poste, une référence et le montant.** La Solde retrouve le mutualiste par son mécano et précompte le montant indiqué.",
           "**Implication pour la mission.** Les fichiers aller maladie et décès doivent être demandés à la direction du personnel de la Police ou à la Solde (PBC, item 61). Sans eux, le taux de recouvrement des cotisations exigé par la Décision n° 05/2022 (au moins 90 %) ne peut pas être calculé pour ces postes."], 12)


def r_regles(s):
    bas = tableau(s, {"titre": "Règles des produits déclarées et confrontées aux fichiers de précomptes", "entetes": ["Produit", "Règle déclarée", "Ce que montrent les fichiers (2020 à août 2026)", "À confirmer"],
                      "largeurs": [3.6, 10.0, 10.6, 7.0], "taille": 10,
                      "lignes": [["Maladie (936)", "Pourcentage du salaire de base. La répartition avec le décès relève de la comptabilité.", "La Solde retient 3,77 % du salaire de base, pour tous les agents.", "Barème officiel et document qui le notifie."],
                                 ["Décès (8Z6)", "Pourcentage du salaire de base.", "La Solde retient 1,23 % du salaire de base. Le total avec la maladie fait 5,00 %.", "Barème officiel (3,5 % et 1,5 % indiqués le 16/10/2026)."],
                                 ["EFC (870)", "6 000 FCFA par mois avant 2021, 10 000 FCFA pour les souscriptions depuis 2021. Les anciens restent à 6 000 FCFA et peuvent passer à 10 000 FCFA.",
                                  f"Les deux montants coexistent : 10 000 FCFA dans {pc(part10[2023])} des lignes en 2023 et {pc(part10[2026])} en 2026, 6 000 FCFA dans {pc(part6[2023])} puis {pc(part6[2026])}.", "Aucun document formel selon l'interlocuteur."],
                                 ["PEP (998)", "À partir de 10 000 FCFA, par paliers de 5 000 FCFA. Durée à vérifier.",
                                  f"{pc(pep_m5)} des montants sont des multiples de 5 000 FCFA et {pc(pep_lt10)} sont inférieurs à 10 000 FCFA. Montant maximal : {nb(pep_max)} FCFA.", "Durée des plans et conditions de retrait (règlement du PEP)."],
                                 ["EFC et PEP", "Présentés aux élèves policiers comme obligatoires dès l'école, sans texte.", "Non mesurable dans les fichiers.", "Aucun texte selon l'interlocuteur."],
                                 ["Prestations", "Au moins 3 mois de cotisation avant de pouvoir bénéficier des prestations.", "Non mesurable sans base de prestations.", "Produits concernés et texte. L'entretien du 16/10/2026 indiquait l'absence de délai de carence pour le régime obligatoire."]]}, 3.85)
    corps(s, 1.4, bas + 0.35, 31.2, BAS_CONTENU - bas - 0.35,
          ["**Les fichiers confirment les déclarations sur l'EFC et le PEP, mais aucune règle n'est appuyée par un texte transmis.** La part de l'EFC à 10 000 FCFA augmente chaque année, ce qui correspond à l'arrivée de nouveaux souscripteurs depuis 2021.",
           "**Implication pour la mission.** Les engagements d'épargne se calculeront sur les montants réellement précomptés. Le règlement de chaque produit reste indispensable pour connaître les échéances et les conditions de retrait."], 11)


def r_rejets(s):
    bas = tableau(s, {"titre": "Rejets, contrôles et paiements hors précompte", "entetes": ["Sujet", "Ce qu'indique l'interlocuteur", "Conséquence pour la mission", "Suite"],
                      "largeurs": [4.2, 11.8, 10.2, 5.0], "taille": 10,
                      "lignes": [["Rejets", "Le fichier retour compte moins de personnes que le fichier aller (radiations, suspensions, autres motifs). Les motifs n'étaient pas connus jusqu'à récemment.", "Les écarts mesurés par Adenka ne peuvent pas être justifiés pour les mois antérieurs.", "PBC, item 64"],
                                 ["États des motifs", "La Solde transmet depuis environ deux mois un état séparé des motifs de rejet.", "Une explication des rejets devient possible pour les mois récents.", "PBC, item 64"],
                                 ["Rapprochement", "Aucun rapprochement des fichiers aller et retour n'est fait par le service, faute de temps.", "Le taux de précompte n'est suivi par personne au FPPN.", "Recommandation"],
                                 ["Premier précompte", "Le premier précompte d'un nouveau souscripteur intervient souvent le mois suivant la souscription.", "Les versements d'une année peuvent porter sur le mois précédent.", "Analyse des fichiers"],
                                 ["Paiements manuels", "Un mutualiste non précompté peut payer manuellement. Le paiement est enregistré avec sa date et son heure.", "Une partie des cotisations échappe aux fichiers de précomptes.", "Chef comptable, PBC, item 65"],
                                 ["Flotte GNA", "Les primes vont directement à GNA. GNA met mal à jour les soldes. Le FPPN arrête le précompte à la date de fin, parfois avec un mois de plus, et rembourse les trop-perçus sur demande.", "Risque de trop-perçus et de réclamations des adhérents.", "PBC, item 68"]]}, 3.85)
    corps(s, 1.4, bas + 0.35, 31.2, BAS_CONTENU - bas - 0.35,
          ["**Le FPPN ne contrôle pas ce qu'il reçoit.** Les fichiers retour sont conservés sans vérification, alors que les écarts mensuels mesurés par Adenka atteignent parfois plus de 20 % du montant demandé (rapport, partie 3).",
           "**Implication pour le FPPN.** Un rapprochement mensuel simple (aller, retour, motifs) permettrait de relancer les précomptes manqués et de justifier chaque écart."], 11)


def r_recrues(s):
    LG.nouvelles_recrues(s, 3.9)
    corps(s, 1.4, 11.5, 31.2, BAS_CONTENU - 11.5,
          ["**Les nouvelles recrues souscrivent avant d'avoir un mécano.** Les élèves policiers n'ont qu'un matricule. Le FPPN attend l'attribution des mécanos, puis reporte chaque mécano sur les fiches de souscription avant de les inscrire au fichier aller. Le premier précompte peut être décalé d'un mois.",
           "**Risques relevés.** Le report en masse des mécanos peut produire des erreurs d'appariement, et donc des précomptes sur la mauvaise personne. Il crée un retard entre la souscription et le premier versement, et il manipule des données personnelles en grand volume.",
           "**Implication pour le FPPN.** Une table de correspondance matricule et mécano, fournie par la direction du personnel, éviterait le report manuel et sécuriserait le premier précompte."], 11)


def r_autres(s):
    gtyp = "activités génératrices de revenus, crédit de santé, Grenier shop, habitat, mariage, flash, scolaire, équipement domestique"
    bas = tableau(s, {"titre": "Autres produits évoqués", "entetes": ["Produit", "Ce qu'indique l'interlocuteur", "Ce que montrent les fichiers", "Point à clarifier"],
                      "largeurs": [4.0, 11.0, 9.4, 6.8], "taille": 10,
                      "lignes": [["Flotte automobile GNA (9W3)", "La prime annuelle est précomptée par mois et versée directement à GNA.", f"Les montants mensuels les plus fréquents sont {nb(gna.index[0])} FCFA et {nb(gna.index[1])} FCFA.", "Convention, primes reversées, soldes par adhérent."],
                                 ["MEDOC (8Z7)", "Complémentaire pour les médicaments : le policier paie 20 % en pharmacie. Produit en pause depuis environ deux mois, à cause de difficultés d'utilisation par les pharmacies.",
                                  f"Précomptes de 5 000 FCFA retournés de décembre 2025 à juillet 2026, sauf février et avril 2026 (retours vides).", "Période de fonctionnement, conventions avec les pharmacies, décision de suspension."],
                                 ["Prêts du Grenier (9X4)", "Trois types : prêt ordinaire à partir de 100 000 FCFA selon le salaire, prêt flash de 10 000 FCFA par téléphone, achat d'équipements via un partenaire. Une nouvelle application est terminée et attend son lancement.",
                                  f"Les fichiers de production comptent huit types de crédit : {gtyp}.", "Liste et conditions de chaque type de prêt."]]}, 3.85)
    corps(s, 1.4, bas + 0.35, 31.2, BAS_CONTENU - bas - 0.35,
          ["**La description orale des prêts est plus simple que ce que montrent les fichiers.** Huit types de crédit apparaissent dans la production du Grenier, contre trois cités.",
           "**Implication pour la mission.** La politique de crédit lisible et la documentation de MEDOC et de la flotte sont nécessaires pour apprécier les engagements que ces produits font porter au FPPN."], 11)


def r_manquants(s):
    noms = {"936": "Maladie (936)", "8Z6": "Décès (8Z6)", "870": "EFC (870)", "998": "PEP (998)", "9X4": "Grenier (9X4)", "9W3": "Flotte GNA (9W3)", "8Z7": "MEDOC (8Z7)", "CGRAE": "Retraités CGRAE"}
    lignes = [[noms[po], str(MANQ[(po, "aller")][0]), MANQ[(po, "aller")][1], str(MANQ[(po, "retour")][0]), MANQ[(po, "retour")][1]] for po in noms]
    bas = tableau(s, {"titre": "Fichiers de précomptes reçus et mois manquants, janvier 2020 à août 2026 (80 mois)", "entetes": ["Poste", "Aller reçus", "Mois sans fichier aller", "Retour reçus", "Mois sans fichier retour"],
                      "largeurs": [4.0, 2.2, 11.4, 2.2, 11.4], "taille": 9, "lignes": lignes}, 3.85)
    corps(s, 1.4, bas + 0.3, 31.2, BAS_CONTENU - bas - 0.3,
          ["**Une partie des mois manquants peut correspondre à des périodes où le produit n'existait pas encore**, notamment MEDOC avant novembre 2025, la flotte avant 2023 et les retraités avant 2024. La date de démarrage de chaque produit est demandée (PBC, item 60).",
           "**Implication pour la mission.** Les fichiers aller maladie et décès sont à demander à la direction du personnel de la Police ou à la Solde. Les autres fichiers manquants sont à demander au responsable des cotisations, qui les dépose sur la plateforme."], 11)


def r_documents(s):
    bas = tableau(s, {"titre": "Documents et points d'attention", "entetes": ["Sujet", "Constat ou engagement", "Conséquence pour la mission", "Suite"],
                      "largeurs": [5.0, 10.4, 10.0, 5.8], "taille": 10,
                      "lignes": [["Barème maladie et décès", "À déposer sur la plateforme de partage. Non reçu au 19/10/2026.", "La répartition 3,77 % et 1,23 % reste à valider.", "PBC, item 56"],
                                 ["Documentation des produits", "Retour annoncé le jour même de l'entretien.", "Règles de l'EFC, du PEP, de MEDOC, de GNA et du Grenier.", "PBC, items 57 à 60"],
                                 ["Fichiers aller maladie et décès", "Transmis directement par la direction du personnel à la Solde.", "Recouvrement des cotisations non mesurable.", "PBC, item 61"],
                                 ["Fichiers manquants", "Mois sans fichier aller ou retour, par poste.", "Couverture incomplète de la période.", "PBC, items 62 et 63"],
                                 ["Motifs de rejet", "États reçus depuis environ deux mois.", "Justification des écarts récents.", "PBC, item 64"],
                                 ["Ayants droit", "Fichier attendu, sans âges.", "Évaluation des engagements santé.", "PBC, item 66"],
                                 ["Paiements manuels", "Encaissement à préciser avec le chef comptable.", "Cotisations hors fichiers.", "Guide, entretien placements"]]}, 3.85)
    corps(s, 1.4, bas + 0.35, 31.2, BAS_CONTENU - bas - 0.35,
          ["**L'entretien éclaire le circuit des cotisations, mais ne remplace pas les documents.** Les déclarations seront intégrées au rapport avec la mention « à confirmer par document » jusqu'à réception des pièces."], 12)


diapo("Informations de l'entretien", "Entretien avec le responsable des cotisations", "Un entretien consacré au circuit des précomptes, aux règles des produits et aux données des mutualistes", r_infos,
      "Interlocuteur : Kouamé Attuoman, responsable des cotisations et responsable informatique du FPPN.", "Source : entretien du 19/10/2026")
diapo("Organisation du suivi des cotisations", "Les cotisations sont suivies par le responsable informatique, sans contrôle des retours",
      "Le responsable des cotisations produit les fichiers aller, mais aucun état des cotisations n'est tenu", r_organisation)
diapo("Circuit des précomptes", "La maladie et le décès suivent un circuit direct que le FPPN ne maîtrise pas",
      "Les produits complémentaires passent par la plateforme de partage, la maladie et le décès sont transmis en direct", r_circuit,
      "Logigramme modifiable : chaque forme et chaque flèche peuvent être déplacées dans PowerPoint. Les formes à contour rouge signalent une étape où un contrôle manque.")
diapo("Circuit des fichiers par produit", "Le FPPN ne produit pas les fichiers aller de la maladie et du décès",
      "Quatre circuits, dont trois passent par la plateforme de partage", r_circuits_produits, source=SRC_D)
diapo("Règles des produits", "Les fichiers confirment les montants de l'EFC et du PEP, sans texte qui les fonde",
      "Maladie et décès en pourcentage du salaire de base, EFC et PEP en montants fixes, sans règle écrite", r_regles, source=SRC_D)
diapo("Rejets et contrôles", "Les rejets n'étaient pas expliqués, et les retours ne sont pas rapprochés",
      "Les motifs de rejet ne sont transmis que depuis environ deux mois", r_rejets, source=SRC_D)
diapo("Nouvelles recrues", "Les recrues souscrivent avant d'avoir un mécano, ce qui retarde le premier précompte",
      "Le report des mécanos sur les fiches de souscription se fait en masse, après leur attribution", r_recrues,
      "Logigramme modifiable. Les volumes cités en exemple pendant l'entretien ne sont pas repris.")
diapo("Autres produits", "La flotte, MEDOC et le Grenier fonctionnent avec des règles peu documentées",
      "MEDOC est en pause, et le Grenier compte plus de types de prêts que ceux cités", r_autres, source=SRC_D)
diapo("Fichiers manquants", "Les fichiers aller maladie et décès manquent sur toute la période",
      "Mois sans fichier aller ou retour, par poste, sur les 80 mois de la période d'analyse", r_manquants,
      "Calcul sur l'inventaire des fichiers lisibles reçus le 15/10/2026. Les fichiers vides sont comptés comme manquants.",
      "Source : fichiers reçus le 15/10/2026 ; inventaire et traitements Adenka")
diapo("Documents et points d'attention", "La documentation des produits conditionne la suite des travaux",
      "Les documents promis et les points relevés sont repris dans la demande complémentaire", r_documents)

# ---------------------------------------------------------------------------
# Construction
# ---------------------------------------------------------------------------
tmp = SORTIE.replace(".pptx", "_tmp.pptx")
selection(TIMESAVER, tmp, [1] + [5] * len(SL) + [110])
CJ = {"1": {"Google Shape;138;p25": "Entretien\nCotisations et\nprécomptes", "Text Placeholder 4": "19 octobre 2026"}}
for i, x in enumerate(SL, 2):
    CJ[str(i)] = {"@vider": True, "@entete": x["entete"], "@source": x["source"]}
CJ[str(len(SL) + 2)] = {}
remplir(tmp, CJ, tmp)
prs = Presentation(tmp)
slides = list(prs.slides)
for sh in slides[0].shapes:
    if sh.name == "Google Shape;138;p25":
        sh.width = Cm(12.8)
for i, x in enumerate(SL, 1):
    s = slides[i]
    x["rendu"](s)
    if x["notes"]:
        s.notes_slide.notes_text_frame.text = x["notes"]
    textes = []
    for sh in s.shapes:
        if sh.has_text_frame:
            textes.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            textes += [c.text for r in sh.table.rows for c in r.cells]
    ligne_sigles(s, textes)
prs.save(SORTIE)
finaliser(SORTIE)
os.remove(tmp)
print(len(SL) + 2, "diapositives")
