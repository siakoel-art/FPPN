"""Plan de travail du volet actuariel : pointage au 23/10/2026.

Met à jour les statuts, les taux de réalisation et les pièces justificatives sur le rapport du 23/10/2026,
renumérote les documents attendus selon la PBC du 23/10/2026 (doublons fusionnés), remplace les liens
externes « [1]PT_Actuariat! » par des liens internes et recrée la figure du taux de réalisation par phase.
Les temps estimés (26 jours dans le scénario « données dégradées ») ne changent pas.

Usage : python3 maj_plan_fppn.py plan.xlsx FPPN_PBC_complete_pointee_20261023.xlsx pbc_map.json sortie.xlsx
"""
import json
import sys
from copy import copy

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

PLAN, PBC, MAP, SORTIE = sys.argv[1:5]
M = {int(k): v for k, v in json.load(open(MAP, encoding="utf-8"))["mapping"].items()}
RAP = "rapport du 23/10/2026"
XL = "classeur FPPN_Calculs_rapport_20261023.xlsx"


def n(*anciens):
    """Numéros de la PBC du 23/10/2026 pour des numéros de la PBC du 22/10/2026."""
    return ", ".join(str(x) for x in sorted({M[a] for a in anciens}))


# PBC du 23/10/2026 : numéro, intitulé, statut
wp = openpyxl.load_workbook(PBC)["PBC Actuariat FPPN"]
ITEMS = {}
for r in range(11, wp.max_row + 1):
    if wp.cell(r, 13).value and wp.cell(r, 3).value:
        ITEMS[len(ITEMS) + 1] = (wp.cell(r, 3).value.split("\n")[0].strip(), wp.cell(r, 13).value)
assert len(ITEMS) == 90, len(ITEMS)
compte = {s: sum(1 for _, x in ITEMS.values() if x == s) for s in ("Reçu", "Partiellement Reçu", "Non Reçu")}

# sous-activité : (statut, taux, pièces justificatives et pointage)
POINT = {
    "1.1": ("Effectuée", 1, None),
    "1.2": ("En cours", 0.8, f"PBC pointée au 23/10/2026 (FPPN_PBC_complete_pointee_20261023.xlsx) : 90 items après fusion des doublons, dont {compte['Reçu']} reçus, {compte['Partiellement Reçu']} partiellement reçus et {compte['Non Reçu']} non reçus. Questions et documents à transmettre (FPPN_Demandes_au_FPPN_20261023.xlsx : 26 questions, 84 documents). Reste : relances."),
    "1.3": ("En cours", 0.9, f"{RAP.capitalize()}, diapositives 29 à 32 (cadre juridique, statuts et règlement intérieur, garanties et produits). Versions des statuts antérieures à 2023 non reçues. Item PBC attendu : {n(1)}."),
    "1.4": ("Effectuée", 1, None),
    "1.5": ("En cours", 0.7, f"Tables retenues ({RAP}, diapositives 218 à 224). SMIG de 75 000 FCFA (décret n° 2022-986 du 21/12/2022) appliqué au plafond de l'indemnité de départ (diapositive 142). Rendements lus dans les grands livres (diapositive 162). Reste : âge de départ du personnel et hypothèses de projection (phase 7). Items PBC attendus : {n(25, 78)}."),
    "2.1": ("Effectuée", 1, "Extraction FPPN_extraction_20261018_1939.zip : 4 287 fichiers, journal journal_fichiers.csv. Envois du 22/10/2026 inventoriés. Fichiers vides ou endommagés recensés (" + RAP + ", diapositive 53)."),
    "2.2": ("Effectuée", 1, f"{RAP.capitalize()}, diapositives 46 à 54 (contrôles, qualité des bases variable par variable). Vérification indépendante des chiffres (FPPN_Verification_coherence_20261023.xlsx) : 506 feuilles de précomptes, 43 102 485 135 FCFA, 7 balances identiques."),
    "2.3": ("Effectuée", 1, f"{RAP.capitalize()}, diapositives 51 et 52 (rapprochement aller et retour, écarts du PEP) ; {XL}, onglet 1_aller_retour. Circuits par produit d'après les logigrammes du FPPN, contrôles absents signalés (diapositives 35 à 38 ; FPPN_Logigrammes_cotisations_mis_a_jour_20261023.pptx)."),
    "2.4": ("En cours", 0.8, f"Couverture mesurée de janvier 2020 à août 2026 ({RAP}, diapositive 48). Restent vides : fichiers maladie 936 et décès 8Z6 de décembre 2024. Fichier retour 9X4 de décembre 2024 manquant (diapositive 53). Items PBC attendus : {n(51, 61, 62, 63)}."),
    "2.5": ("En cours", 0.6, f"Cotisations comptabilisées et précomptées rapprochées ({RAP}, diapositive 68). Contrôles croisés états financiers, effectifs et rapprochements bancaires (diapositives 56 à 58). Reste : comptes individuels d'épargne. Items PBC attendus : {n(17, 19)}."),
    "2.6": ("Bloquée", 0.25, f"Aucune base de prestations extraite du logiciel. Le journal SIN et les états décès de fin d'année servent de base provisoire ({RAP}, diapositives 119 à 127). Items PBC attendus : {n(13, 28, 40, 71)}."),
    "2.7": ("En cours", 0.8, f"Note sur les anomalies de données du 19/10/2026. Vérification de cohérence du 23/10/2026 (FPPN_Verification_coherence_20261023.xlsx, onglet 1_Anomalies). Questions au FPPN ({RAP}, diapositives 243 et 244). À tenir à jour à chaque envoi."),
    "3.1": ("Effectuée", 1, f"{RAP.capitalize()}, partie 4 (diapositives 59 à 71). Limite : dates de naissance absentes des fichiers de précomptes (diapositive 54), fichier des ayants droit non reçu. Items PBC attendus : {n(9, 10)}."),
    "3.2": ("Effectuée", 1, f"{RAP.capitalize()}, partie 5 (diapositives 72 à 86) et ratios par garantie (diapositive 175) ; {XL}, onglets 3_, 12_sinistralite_variantes et 28_ratios_garantie."),
    "3.3": ("Effectuée", 1, f"{RAP.capitalize()}, parties 6 à 8 (diapositives 87 à 113) ; {XL}, onglets 4_, 5_ et 6_."),
    "3.4": ("En cours", 0.5, f"Méthodes retenues chiffrées sur les documents reçus : provision maladie, allocations décès, intérêts dus sur l'épargne, provision du personnel agent par agent ({RAP}, diapositives 98, 123, 127 et 131 à 142). Méthodes écartées en annexe (diapositives 246 à 253)."),
    "3.5": ("Bloquée", 0.1, f"Calculs par âge impossibles : dates de naissance absentes (" + RAP + ", diapositive 54). Analyses en attente et méthode retenue pour chacune (diapositive 242). Items PBC attendus : " + n(9, 10, 66, 13, 28, 40, 71) + "."),
    "3.6": ("En cours", 0.6, f"Conclusion établie sur les documents reçus ({RAP}, diapositive 143). Les analyses en attente la préciseront sans en changer le sens (diapositive 242)."),
    "4.1": ("Effectuée", 1, f"{RAP.capitalize()}, diapositives 114 à 118 ; {XL}, onglets 12_provisions_inventaire, 12_psap_rapprochement et 8_provisions_mois."),
    "4.2": ("En cours", 0.7, f"Provision maladie mesurée sur les factures arrivées après la clôture et provision décès sur les dossiers connus ({RAP}, diapositives 119 à 128). Reste : base des prestations datée et liste des dossiers non payés. Items PBC attendus : {n(14, 29, 71, 72, 92)}."),
    "4.3": ("En cours", 0.4, f"Provisions maladie et obsèques identiques depuis le 01/09/2020 : comparaison avec le recalcul ({RAP}, diapositive 145). Boni et mali par exercice en attente des bases avec dates de survenance et de paiement. Items PBC attendus : {n(13, 28, 71)}."),
    "4.4": ("En cours", 0.9, f"{RAP.capitalize()}, diapositives 129 à 142 (Grenier, compte GNA, provision du personnel avec et sans plafond conventionnel) ; {XL}, onglets 12_grenier_tranches, 12_gna_47111106, 14b_pidr_individuel et 27_pidr_plafond. Reste : note de calcul de la provision du personnel. Items PBC attendus : {n(75, 78, 98, 103)}."),
    "4.5": ("En cours", 0.7, f"Conclusion établie sur les documents reçus ({RAP}, diapositives 144 à 146)."),
    "5.1": ("Effectuée", 1, f"{RAP.capitalize()}, diapositives 148, 151 et 162 à 164 ; compte rendu de l'entretien placements du 19/10/2026."),
    "5.2": ("Effectuée", 1, f"{RAP.capitalize()}, diapositives 149 à 157 ; {XL}, onglets 10_actif_mobilisation et 8_sensibilite_retraits."),
    "5.3": ("En cours", 0.5, f"Écoulement de l'épargne et écarts de liquidité année par année ({RAP}, diapositives 154 et 155). Encours des placements lus dans les grands livres (diapositive 162). Reste : inventaire et échéancier signés des placements. Items PBC attendus : {n(20, 80, 81, 82)}."),
    "5.4": ("Effectuée", 1, f"Encours moyen et rendement lus dans les grands livres ({RAP}, diapositives 162, 165 et 216) ; {XL}, onglets 7_rendement_placements, 26_placements_gl et D9_placements_gl."),
    "5.5": ("En cours", 0.7, f"Conclusions établies sur les documents reçus ({RAP}, diapositives 158 et 166)."),
    "6.1": ("Effectuée", 1, f"{RAP.capitalize()}, diapositives 168 à 193, avec la définition et la lecture de chaque ratio ; {XL}, onglets 10_ratios et 10_ratios_detail."),
    "6.2": ("Effectuée", 1, f"{RAP.capitalize()}, diapositives 175, 176 et 194 à 201 (ratios par garantie, excédent par risque, affectation des excédents, marge de solvabilité, dispositif actuariel, plan comptable) ; {XL}, onglets 10_controle_pcms, 10_marge_solvabilite, 10_affectation_excedents et 10_excedent_par_risque."),
    "6.3": ("Effectuée", 1, f"{RAP.capitalize()}, diapositives 204 à 217 ; {XL}, onglets 13a à 13i."),
    "6.4": ("Bloquée", 0, f"{RAP.capitalize()}, diapositive 183 (données en attente). Items PBC attendus : {n(47, 48, 49, 52)}."),
    "7.1": ("En cours", 0.5, f"Hypothèses de la projection agrégée ({RAP}, diapositives 225 et 226) : tendances, inflation médicale (diapositive 214), rendements (diapositive 216)."),
    "7.2": ("En cours", 0.4, f"Projection agrégée en deux scénarios ({RAP}, diapositives 227 et 228). Reste : projection par garantie, qui dépend des phases 3 et 4."),
    "7.3": ("En cours", 0.3, f"Tests de résistance et sensibilités ({RAP}, diapositives 156 et 204 à 217), à étendre à la projection par garantie."),
    "8.1": ("En cours", 0.8, f"Rapport du 23/10/2026 (FPPN_Rapport_actuariel_statistique_20261023.pptx, 291 diapositives) et {XL} (onglet 9_tracabilite)."),
    "8.2": ("Effectuée", 1, None),
    "8.3": ("En cours", 0.6, "Contrôle automatique : 2 885 chiffres du rapport retrouvés dans les calculs. Vérification indépendante sur les fichiers du client (FPPN_Verification_coherence_20261023.xlsx, 18 anomalies corrigées). À refaire sur le rapport provisoire."),
}

wb = openpyxl.load_workbook(PLAN)
ws = wb["PT_Actuariat"]
ws["A2"].value = ws["A2"].value.replace("22/10/2026", "23/10/2026")
for c in list(ws[4]) + list(ws[6]):
    if isinstance(c.value, str):
        c.value = c.value.replace("22/10/2026", "23/10/2026")
vus = set()
for r in range(7, 56):
    k = str(ws.cell(r, 1).value)
    if k in POINT:
        st, tx, obs = POINT[k]
        ws.cell(r, 11).value = st
        ws.cell(r, 12).value = tx
        if obs:
            ws.cell(r, 6).value = obs
        vus.add(k)
assert vus == set(POINT), set(POINT) - vus

# synthèse : liens internes à la place des liens vers un classeur externe
s2 = wb["Synthèse-PT_Actuariat"]
for ligne in s2.iter_rows():
    for c in ligne:
        if isinstance(c.value, str):
            c.value = c.value.replace("[1]PT_Actuariat!", "PT_Actuariat!").replace("22/10/2026", "23/10/2026")
wb._external_links = []
ch = BarChart()
ch.type = "bar"
ch.style = 10
ch.title = None
ch.y_axis.scaling.min = 0
ch.y_axis.scaling.max = 1
ch.y_axis.number_format = "0%"
ch.y_axis.majorGridlines = None
ch.x_axis.scaling.orientation = "maxMin"
ch.legend = None
ch.add_data(Reference(s2, min_col=6, min_row=4, max_row=12), titles_from_data=True)
ch.set_categories(Reference(s2, min_col=1, min_row=5, max_row=12))
ch.series[0].graphicalProperties.solidFill = "EFAC62"
ch.series[0].graphicalProperties.line.solidFill = "EFAC62"
ch.dataLabels = DataLabelList()
ch.dataLabels.showVal = True
ch.dataLabels.showCatName = False
ch.dataLabels.showSerName = False
ch.dataLabels.showLegendKey = False
ch.dataLabels.showPercent = False
ch.dataLabels.numFmt = "0%"
ch.height, ch.width = 9.5, 18
s2._charts = []  # le graphique d'origine perd sa mise en forme à la relecture : il est remplacé
s2.add_chart(ch, "B17")

# documents attendus : numéros et statuts de la PBC du 23/10/2026, doublons fusionnés
s3 = wb["Documents attendus"]
s3["A2"].value = "Numéros et statuts repris de la PBC pointée au 23/10/2026 (FPPN_PBC_complete_pointee_20261023.xlsx). Doublons fusionnés : un numéro par document."
anciens = []
for r in range(4, s3.max_row + 1):
    v = s3.cell(r, 1).value
    if isinstance(v, (int, float)) or (isinstance(v, str) and v.isdigit()):
        anciens.append((int(v), s3.cell(r, 4).value, s3.cell(r, 5).value))
styles = [copy(s3.cell(4, j)._style) for j in range(1, 6)]
source_style = copy(s3.cell(s3.max_row, 1)._style)
fus = {}
for a, sous, effet in anciens:
    nv = M[a]
    d = fus.setdefault(nv, {"sous": set(), "effet": set()})
    d["sous"].update(x.strip() for x in str(sous).split(","))
    d["effet"].add(effet)
for r in range(4, s3.max_row + 1):
    for j in range(1, 6):
        s3.cell(r, j).value = None
r = 4
for nv in sorted(fus):
    lib, st = ITEMS[nv]
    if st == "Reçu":
        continue
    eff = "Bloque la sous-activité" if "Bloque la sous-activité" in fus[nv]["effet"] else "Limite la conclusion"
    sous = ", ".join(sorted(fus[nv]["sous"], key=lambda x: [int(p) for p in x.split(".")]))
    for j, v in enumerate((nv, lib, st, sous, eff), 1):
        c = s3.cell(r, j, v)
        c._style = copy(styles[j - 1])
    r += 1
c = s3.cell(r, 1, "Source : PBC pointée au 23/10/2026 ; onglet « Plan de travail »")
c._style = source_style
me = wb["Mode d'emploi"]
me["B6"].value = me["B6"].value.replace("(rapport v8, classeur v8, PBC, comptes rendus)", "(rapport et classeur des calculs du 23/10/2026, PBC, comptes rendus)")
wb.save(SORTIE)
print(SORTIE, "| documents attendus :", r - 4, "| PBC :", compte)
