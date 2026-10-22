# -*- coding: utf-8 -*-
"""
FPPN - Données de la version 9 (documents reçus le 22/10/2026, lots 1 à 3)
=========================================================================
Calcule, à partir des fichiers reçus et des bases déjà constituées :
  1. la provision pour indemnités de départ à la retraite (PIDR) agent par agent, à partir des DISA (CNPS),
     séparée entre Direction Générale, Centre Médical du Policier (CMP) et Grenier ;
  2. la mortalité brute des membres participants (étude des ratios du FPPN) et ses contrôles ;
  3. la comparaison de la méthode de sinistralité du FPPN et de la méthode réglementaire ;
  4. la sensibilité de la charge d'intérêts de l'EFC au taux servi (3,5 %, 5,5 %, 14 %) ;
  5. des contrôles croisés : états financiers contre balances, continuité des à-nouveaux, effectifs du FPPN contre
     fichiers 8Z6, rapprochements bancaires contre balances, rapport 8Z6 / 936 par année.
Sorties : un JSON pour le rapport et un classeur à formules (onglets 14 à 18).
Utilisation : python donnees_v9_fppn.py <dossier_lot1> <dossier_base> <resultats.json> <provisions.json>
              <dossier_tables_cima> <dossier_investing> <sortie.json> <sortie.xlsx>
"""
import csv
import datetime as dt
import glob
import json
import os
import re
import sys
import unicodedata

import openpyxl
import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from python_calamine import CalamineWorkbook

LOT1, BASE, RES, PROV, TABLES, INVEST, SORTIE_J, SORTIE_X = sys.argv[1:9]
PERS = os.path.join(LOT1, "04-Nouveaux Personnel et paie", "ETAT 301 ET DISA PERSONNEL CIVIL")
R = json.load(open(RES, encoding="utf-8"))
PV = json.load(open(PROV, encoding="utf-8"))
OUT = {}


def feuille(f, nom=None):
    w = CalamineWorkbook.from_path(f)
    return w.get_sheet_by_name(nom or w.sheet_names[0]).to_python()


def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z ]", " ", s).split()


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def date_fr(v):
    if isinstance(v, (dt.date, dt.datetime)):
        return dt.date(v.year, v.month, v.day)
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(v).strip())
    return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None


# =====================================================================================================================
# 1. PIDR à partir des DISA
# =====================================================================================================================
DISA = {2021: "2021/DISA_256927_2021.xls", 2022: "2022/DISA_256927_2022.xls", 2023: "2023/DISA 2023/DISA_256927_2023.xls",
        2025: "2025/DISA 2025/DISA_256927_2025.xls"}
agents = {}
for an, rel in DISA.items():
    t = feuille(os.path.join(PERS, rel))
    for r in t[3:]:
        if str(r[0]).strip().upper().startswith("TOTAL") or not str(r[1]).strip():
            continue
        if "TRAVAILLEUR" in str(r[1]).upper():      # ligne globale des travailleurs occasionnels : exclue
            continue
        brut, mois = float(r[8] or 0), float(r[9] or 0)
        agents.setdefault(an, []).append({
            "annee": an, "ordre": int(float(r[0])), "nom": str(r[1]).strip(), "prenom": str(r[2]).strip(),
            "cnps": str(r[3]).strip(), "naissance": int(float(r[4])), "embauche": date_fr(r[5]),
            "depart": date_fr(r[6]), "brut": brut, "mois": mois, "fichier": os.path.basename(rel)})
OUT["disa_effectifs"] = {a: len(v) for a, v in agents.items()}

# Rattachement aux entités : listes nominatives des récapitulatifs de salaires
f24 = os.path.join(PERS, "2024", "ETAT 301 & DISA 2024 FPPN.xls")
f25 = os.path.join(PERS, "2025", "ETAT 301 & DISA 2025 FPPN.xls")
cmp_noms = []
for r in feuille(f24, "RECAPSALAIRES 2024 CMP")[2:]:
    nom = r[1] if isinstance(r[1], str) and r[1].strip() else r[0]
    if str(r[0]).upper().startswith("TOTAL"):
        break
    if str(nom).strip():
        cmp_noms.append(str(nom).strip())
gre_noms = [str(r[1]).strip() for r in feuille(f25, "RECAPSALAIRES 2025 GRENIER")[2:] if str(r[0]).replace(".0", "").isdigit()]
tous24 = [str(r[1]).strip() for r in feuille(f24, "RECAPSALAIRES 2024")[2:] if str(r[0]).replace(".0", "").isdigit() and str(r[1]).strip()]


def correspond(nom_recap, a):
    """Le nom du récapitulatif (tronqué à 18 caractères) correspond-il à l'agent de la DISA ?"""
    jt = norm(a["nom"] + " " + a["prenom"])
    rt = [x for x in norm(nom_recap) if len(x) >= 2]
    if not rt:
        return 0
    ok = sum(1 for x in rt if any(y.startswith(x) or x.startswith(y) for y in jt))
    # le nom de famille (premier mot du récapitulatif) doit figurer dans la DISA ; variantes d'orthographe tolérées sur un mot sur trois
    if not any(y.startswith(rt[0]) or rt[0].startswith(y) for y in jt):
        return 0
    return ok / len(rt)


def entite(a):
    for liste, nom_e, src in ((gre_noms, "Grenier", "RECAPSALAIRES 2025 GRENIER"), (cmp_noms, "CMP", "RECAPSALAIRES 2024 CMP")):
        if any(correspond(n, a) >= 0.66 for n in liste):
            return nom_e, src
    if any(correspond(n, a) >= 0.66 for n in tous24):
        return "Direction Générale", "RECAPSALAIRES 2024, hors listes CMP et Grenier"
    return "Non documentée", "absent des récapitulatifs de salaires 2024 et 2025"


# Paramètres observés : rotation et croissance des salaires
par_cnps = {a: {x["cnps"]: x for x in v if x["cnps"]} for a, v in agents.items()}
rot = []
for a0, a1 in ((2021, 2022), (2022, 2023), (2023, 2025)):
    present0, present1 = set(par_cnps[a0]), set(par_cnps[a1])
    sortis = present0 - present1
    taux = len(sortis) / len(present0)
    annuel = 1 - (1 - taux) ** (1 / (a1 - a0))
    rot.append({"de": a0, "a": a1, "presents": len(present0), "sortis": len(sortis), "taux_periode": taux, "taux_annuel": annuel})
ROT = sum(x["taux_annuel"] for x in rot) / len(rot)
crois = []
for c, x25 in par_cnps[2025].items():
    x21 = par_cnps[2021].get(c)
    if x21 and x21["mois"] == 12 and x25["mois"] == 12 and x21["brut"] > 0:
        crois.append((x25["brut"] / x21["brut"]) ** (1 / 4) - 1)
crois.sort()
CROIS = crois[len(crois) // 2] if len(crois) % 2 else (crois[len(crois) // 2 - 1] + crois[len(crois) // 2]) / 2
# Dates d'embauche déclarées différemment d'une DISA à l'autre
chg = []
for c in par_cnps[2025]:
    dates = {a: par_cnps[a][c]["embauche"] for a in DISA if c in par_cnps[a]}
    if len(set(dates.values())) > 1:
        chg.append({"cnps": c, "nom": par_cnps[2025][c]["nom"] + " " + par_cnps[2025][c]["prenom"],
                    "dates": {str(a): d.strftime("%d/%m/%Y") for a, d in dates.items()}})
OUT["embauche_modifiee"] = chg

# Taux d'actualisation : rendement de l'obligation de l'État de Côte d'Ivoire à 10 ans (Investing.com), moyenne 2025
cours = []
f10 = [f for f in glob.glob(os.path.join(INVEST, "*.csv")) if "10-Year" in f][0]
with open(f10, encoding="utf-8-sig") as h:
    for l in csv.DictReader(h):
        d = date_fr(l["Date"])
        if d and d.year == 2025:
            cours.append(float(l["Dernier"].replace(",", ".")) / 100)
TAUX10 = sum(cours) / len(cours)

# Table de survie CIMA F (prudente pour un engagement payé si l'agent est en vie au départ)
lx = {int(r[0]): float(r[1]) for r in feuille(os.path.join(TABLES, "CIMA_F.xlsx"))[1:] if r[0] != ""}
AGE_DEP = 55          # convention collective interprofessionnelle du 19/07/1977, art. 40


def bareme(anc):
    return 0.30 * min(anc, 5) + 0.35 * min(max(anc - 5, 0), 5) + 0.40 * max(anc - 10, 0)


def evaluer(a, an, age_dep=AGE_DEP, i=None, g=None, t=None):
    i = TAUX10 if i is None else i
    g = CROIS if g is None else g
    t = ROT if t is None else t
    d_ev = dt.date(an, 12, 31)
    sm = a["brut"] / a["mois"] if a["mois"] else 0
    age = an - a["naissance"]
    mois_anc = (d_ev.year - a["embauche"].year) * 12 + d_ev.month - a["embauche"].month - (1 if d_ev.day < a["embauche"].day else 0)
    anc = max(mois_anc, 0) / 12
    n = max(age_dep - age, 0)
    acquis = sm * bareme(anc)
    proj = sm * (1 + g) ** n * bareme(anc + n)
    prorata = anc / (anc + n) if anc + n > 0 else 0
    survie = lx.get(age + n, 0) / lx[age] if lx.get(age) else 0
    dbo = proj * prorata * survie * (1 - t) ** n / (1 + i) ** n
    return {"salaire_mensuel": sm, "age": age, "anciennete": anc, "annees_restantes": n, "indemnite_acquise": acquis,
            "indemnite_projetee": proj, "prorata": prorata, "survie": survie, "dbo": dbo}


indiv = []
for an in DISA:
    for k, a in enumerate(sorted(agents[an], key=lambda x: x["ordre"]), 1):
        e, src = entite(a)
        v = evaluer(a, an)
        indiv.append({"annee": an, "id": f"A{an % 100:02d}-{k:02d}", **{x: a[x] for x in ("nom", "prenom", "cnps", "naissance", "brut", "mois")},
                      "embauche": a["embauche"].strftime("%d/%m/%Y"), "entite": e, "source_entite": src, **v})
OUT["pidr_individuel"] = indiv
ENT = ("Direction Générale", "CMP", "Grenier", "Non documentée")
synth = []
for an in DISA:
    lig = [x for x in indiv if x["annee"] == an]
    ligne = {"annee": an, "agents": len(lig), "dbo": sum(x["dbo"] for x in lig), "acquis": sum(x["indemnite_acquise"] for x in lig),
             "age_55_atteint": sum(1 for x in lig if x["age"] >= AGE_DEP),
             "provision_196": next(p["provision_196"] for p in PV["pidr"] if p["exercice"] == an)}
    for e in ENT:
        le = [x for x in lig if x["entite"] == e]
        ligne[e] = {"agents": len(le), "dbo": sum(x["dbo"] for x in le), "acquis": sum(x["indemnite_acquise"] for x in le),
                    "masse": sum(x["brut"] for x in le)}
    synth.append(ligne)
OUT["pidr_synthese"] = synth
lig25 = [x for x in indiv if x["annee"] == 2025]
a25 = sorted(agents[2025], key=lambda x: x["ordre"])
sens = []
for lib, kw in (("Hypothèses centrales", {}), ("Départ à 58 ans (trois enfants à charge)", {"age_dep": 58}), ("Départ à 60 ans", {"age_dep": 60}),
                ("Taux d'actualisation de 3,5 %", {"i": 0.035}), ("Sans rotation du personnel", {"t": 0.0}), ("Salaires + 1 point par an", {"g": CROIS + 0.01})):
    sens.append({"hypothese": lib, "dbo": sum(evaluer(a, 2025, **kw)["dbo"] for a in a25)})
OUT["pidr_sensibilites"] = sens
OUT["pidr_parametres"] = {"age_depart": AGE_DEP, "taux_actualisation": TAUX10, "cotations_2025": len(cours), "croissance_salaires": CROIS,
                          "agents_croissance": len(crois), "rotation": ROT, "rotation_detail": rot, "cmp_liste": len(cmp_noms), "grenier_liste": len(gre_noms)}
OUT["pidr_max_individuel_2025"] = max(x["indemnite_projetee"] for x in lig25)

# =====================================================================================================================
# 2. Mortalité brute et sinistralité selon le FPPN (ETUDE DES RATIO MALADIE DECES 2.xlsx)
# =====================================================================================================================
fr = os.path.join(LOT1, "05-Nouveaux Prestations", "ETUDE DES RATIO MALADIE DECES 2.xlsx")
tb, hi = feuille(fr, "Tableau"), feuille(fr, "Historique 2015-2021")


def lig_(t, lib):
    return next(r for r in t if str(r[1]).strip().upper().startswith(lib))


mort = []
for t, cols in ((hi, {2015: 3, 2016: 4, 2017: 5, 2018: 6}), (tb, {2019: 3, 2020: 4, 2021: 5, 2022: 6, 2023: 7, 2024: 8, 2025: 9})):
    for an, c in cols.items():
        adh, dad = float(lig_(t, "NB D'ADHERENTS")[c]), float(lig_(t, "NB DECES ADHERENTS")[c])
        cot = float(lig_(t, "COTISATIONS PAYEES" if t is tb else "PRIMES PAYEES")[c])
        sin = float(lig_(t, "TOTAL SINISTRES")[c])
        mort.append({"annee": an, "adherents": adh, "deces_total": float(lig_(t, "NB TOTAL DECES")[c]), "deces_adherents": dad,
                     "deces_conjoints": float(lig_(t, "NB DECES CONJOINTS")[c]), "deces_enfants": float(lig_(t, "NB DECES ENFANTS")[c]),
                     "taux_brut": dad / adh, "cotisations": cot, "sinistres": sin, "sp_fppn": sin / cot,
                     "capital_deces": float(lig_(t, "CAPITAL DECES")[c]), "feuille": "Tableau" if t is tb else "Historique 2015-2021"})
OUT["mortalite_brute"] = mort
cum26 = {"adherents": float(lig_(tb, "NB D'ADHERENTS")[16]), "deces_total": float(lig_(tb, "NB TOTAL DECES")[16]),
         "cotisations": float(lig_(tb, "COTISATIONS PAYEES")[16]), "sinistres": float(lig_(tb, "TOTAL SINISTRES")[16])}
OUT["mortalite_2026_s1"] = cum26
mens26 = [{"mois": m, "adherents": float(lig_(tb, "NB D'ADHERENTS")[9 + m]), "cotisations": float(lig_(tb, "COTISATIONS PAYEES")[9 + m])} for m in range(1, 7)]

# Effectifs et cotisations du FPPN contre les fichiers 8Z6 de la Solde
pre = pd.read_parquet(os.path.join(BASE, "precomptes.parquet"), columns=["poste", "sens", "annee", "mois", "mat6", "montant", "retenu"])
p8 = pre[(pre.poste == "8Z6") & (pre.sens == "retour") & pre.retenu]
g8 = p8.groupby(["annee", "mois"]).agg(lignes=("mat6", "size"), montant=("montant", "sum")).reset_index()
ctl = []
for an, m, adh, src in [(2024, 11, mort[-2]["adherents"], "Tableau, colonne 2024"), (2025, 12, mort[-1]["adherents"], "Tableau, colonne 2025")] + \
        [(2026, x["mois"], x["adherents"], f"Tableau, colonne {x['mois']:02d}/2026") for x in mens26]:
    r8 = g8[(g8.annee == an) & (g8.mois == m)]
    ctl.append({"annee": an, "mois": m, "adherents_fppn": adh, "lignes_8z6": int(r8.lignes.iloc[0]) if len(r8) else None,
                "montant_8z6": float(r8.montant.iloc[0]) if len(r8) else None, "source": src,
                "cotisations_fppn": next((x["cotisations"] for x in mens26 if an == 2026 and x["mois"] == m), None)})
OUT["controle_adherents_8z6"] = ctl

# Rapport 8Z6 / 936 par année (contrôle de la répartition 3,77 / 1,23)
p9 = pre[(pre.poste == "936") & (pre.sens == "retour") & pre.retenu]
j = p8.groupby(["annee", "mois", "mat6"]).montant.sum().reset_index().merge(
    p9.groupby(["annee", "mois", "mat6"]).montant.sum().reset_index(), on=["annee", "mois", "mat6"], suffixes=("_8z6", "_936"))
j = j[j.montant_936 > 0]
j["r"] = j.montant_8z6 / j.montant_936
cible = 1.23 / 3.77
rap = []
for an, x in j.groupby("annee"):
    rap.append({"annee": int(an), "mois": int(x.mois.nunique()), "lignes": int(len(x)), "mediane": float(x.r.median()),
                "part_cle": float(((x.r - cible).abs() <= 0.005).mean())})
OUT["rapport_8z6_936_annee"] = rap
OUT["rapport_8z6_936_mois"] = [{"annee": int(a_), "mois": int(m_), "lignes": int(len(x)), "mediane": float(x.r.median()),
                                "part_cle": float(((x.r - cible).abs() <= 0.005).mean())} for (a_, m_), x in j.groupby(["annee", "mois"])]
OUT["rapport_cible"] = cible

# =====================================================================================================================
# 3. Balances : à-nouveaux, classe 5, épargne
# =====================================================================================================================
bal = pd.read_parquet(os.path.join(BASE, "compta_balances.parquet"))
bal = bal[bal.nature == "comptes"].copy()
bal["ouv"] = bal.ouv_d.fillna(0) - bal.ouv_c.fillna(0)
bal["sol"] = bal.sol_d.fillna(0) - bal.sol_c.fillna(0)
an_ = []
for an in range(2021, 2027):
    bil_ = bal[bal.compte.str.match(r"^[1-5]") & ~bal.compte.str.match(r"^1[23]")]   # comptes de bilan hors résultat et report
    p, c = bil_[bil_.exercice == an - 1].set_index("compte").sol, bil_[bil_.exercice == an].set_index("compte").ouv
    comptes = sorted(set(p.index) | set(c.index))
    ec = [(k, float(c.get(k, 0) - p.get(k, 0))) for k in comptes if abs(c.get(k, 0) - p.get(k, 0)) > 0.5]
    an_.append({"exercice": an, "comptes_compares": len(comptes), "comptes_en_ecart": len(ec),
                "ecart_absolu": sum(abs(e) for _, e in ec),
                "principaux": [{"compte": k, "intitule": str(bal[bal.compte == k].intitule.iloc[-1]), "ecart": e} for k, e in sorted(ec, key=lambda z: -abs(z[1]))[:6]]})
OUT["continuite_a_nouveaux"] = an_
ep = []
for an in range(2020, 2027):
    b = bal[bal.exercice == an].set_index("compte")
    ep.append({"exercice": an, **{f"{k}_{c}": float(b.loc[c, col]) if c in b.index else None
                                  for c in ("41195100", "41195200") for k, col in (("ouv_c", "ouv_c"), ("mvt_d", "mvt_d"), ("mvt_c", "mvt_c"), ("sol_c", "sol_c"))}})
OUT["epargne_mouvements"] = ep
cl5 = {int(an): float(x[x.compte.str.match(r"^5") & (x.sol > 0)].sol.sum()) for an, x in bal.groupby("exercice")}   # trésorerie-actif : soldes débiteurs
OUT["classe5_debit_net"] = cl5

# États financiers (relevés sur les pages « Bilan Actif » et « Bilan Passif » des PDF reçus le 22/10/2026)
EF = {2020: {"brut": 14446097044, "tresorerie": 881310155, "resultat": 1018579699},
      2021: {"brut": 16195461159, "tresorerie": 2809574473, "resultat": 2083908202},
      2022: {"brut": 15510713126, "tresorerie": 1233051977, "resultat": 2060332677},
      2023: {"brut": 16157155123, "tresorerie": 1242945485, "resultat": 928699102},
      2024: {"brut": 19396236478, "tresorerie": 979838445, "resultat": 792069641},
      2025: {"brut": 21722068914, "tresorerie": 1964171543, "resultat": 437899911},
      2026: {"brut": 23692620745, "tresorerie": 3359100618, "resultat": -437148460}}
bil = {b["exercice"]: b for b in R["bilan"]}
OUT["ef_contre_balances"] = [{"exercice": a, "ef_brut": e["brut"], "balance_brut": bil[a]["total_actif"], "ef_resultat": e["resultat"],
                              "balance_resultat": bil[a]["resultat_exercice"], "ef_tresorerie": e["tresorerie"], "balance_classe5": cl5.get(a)}
                             for a, e in EF.items()]

# =====================================================================================================================
# 4. Rapprochements bancaires
# =====================================================================================================================
erb = []
for f in sorted(glob.glob(os.path.join(LOT1, "07-Nouveaux Trésorerie et placements", "*", "*.xlsx"))):
    w = CalamineWorkbook.from_path(f)
    for sh in w.sheet_names:
        t = w.get_sheet_by_name(sh).to_python()
        bloc = None
        for r in t:
            r = list(r) + [""] * (8 - len(r))
            txt = " ".join(str(x) for x in r[:8])
            m = re.search(r"solde au (\d{1,2})/(\d{1,2})/(\d{2,4})", txt, re.I)
            if m:
                y = int(m.group(3))
                y = y + 2000 if y < 100 else y
                bloc = {"fichier": os.path.basename(f), "compte": sh.strip(), "date": dt.date(y, int(m.group(2)), int(m.group(1))),
                        "compta": num(r[2]) - num(r[3]), "banque": num(r[7]) - num(r[6]), "suspens": []}
                erb.append(bloc)
                continue
            if bloc and str(r[1]).strip().upper().startswith("TOTAL"):
                bloc = None
                continue
            if bloc and isinstance(r[4], (dt.date, dt.datetime)) and num(r[6]):
                bloc["suspens"].append({"date": date_fr(r[4]), "libelle": str(r[5]), "montant": num(r[6])})
fin_annee = [b for b in erb if b["date"].month == 12 and b["date"].day == 31]
res_erb = []
for an in range(2020, 2026):
    bl = {b["compte"]: b for b in fin_annee if b["date"].year == an}
    vieux = [(b["compte"], s) for b in bl.values() for s in b["suspens"] if (b["date"] - s["date"]).days > 365]
    b52 = bal[(bal.exercice == an) & bal.compte.str.startswith("52")].sol.round().tolist()
    res_erb.append({"exercice": an, "comptes": len(bl), "concordants": sum(1 for b in bl.values() if round(b["compta"]) in b52), "solde_compta": sum(b["compta"] for b in bl.values()),
                    "solde_banque": sum(b["banque"] for b in bl.values()), "suspens_plus_un_an": len(vieux),
                    "montant_suspens_plus_un_an": sum(s["montant"] for _, s in vieux),
                    "plus_ancien": min((s["date"] for _, s in vieux), default=None)})
OUT["rapprochements"] = res_erb

# =====================================================================================================================
# 5. Sinistralité : méthode du FPPN contre méthode réglementaire ; EFC : sensibilité au taux servi
# =====================================================================================================================
si = {s["exercice"]: s for s in R["sinistralite"]}
OUT["sinistralite_comparee"] = [{"annee": m["annee"], "fppn_cotisations": m["cotisations"], "fppn_sinistres": m["sinistres"], "fppn_sp": m["sp_fppn"],
                                 "compta_cot_deces": si[m["annee"]]["cot_deces"], "compta_prest_deces": si[m["annee"]]["prest_deces"],
                                 "compta_prime_deces": si[m["annee"]]["prime_assurance_deces"], "compta_sc_deces": si[m["annee"]]["sc_deces"],
                                 "compta_sc_deces_avec_prime": si[m["annee"]]["sc_deces_avec_prime"]}
                                for m in mort if m["annee"] in si]
efc = []
for e in ep:
    if e["sol_c_41195200"] is None:
        continue
    pf = next((x["produits_financiers"] for x in R["rendement_placements"] if x["exercice"] == e["exercice"]), None)
    efc.append({"exercice": e["exercice"], "encours_efc": e["sol_c_41195200"], "produits_financiers": pf,
                **{f"interets_{k}": e["sol_c_41195200"] * t for k, t in (("3_5", 0.035), ("5_5", 0.055), ("14", 0.14))}})
OUT["efc_taux"] = efc

# Grand livre 2024 : nature des écritures des comptes d'épargne (recherche d'intérêts servis)
gl = feuille(os.path.join(LOT1, "02-Nouveaux_GL et Balances", "Nouveaux fichiers", "GRAND LIVRE FPPN 2024.xlsx"), "Grand-livre des comptes")
cur, mv = None, {}
for r in gl:
    a = str(r[0]).strip()
    if re.fullmatch(r"\d{8}", a):
        cur = a
        continue
    if cur in ("41195100", "41195200") and r[1] != "RAN" and str(r[5]).strip():
        lib = str(r[5]).strip().upper()
        nat = ("Intérêts" if re.search(r"INT[EÉ]R", lib) else "Précomptes et cotisations" if re.match(r"(PRECOMPTE|EPARGNE|EFC|PEP)", lib)
               else "Règlements EFC" if lib.startswith("REGLT") else "Remboursements PEP" if lib.startswith("REMB") else "Rachats" if lib.startswith("RACHAT") else "Autres")
        k = (cur, nat)
        x = mv.setdefault(k, {"compte": cur, "nature": nat, "ecritures_debit": 0, "debit": 0.0, "ecritures_credit": 0, "credit": 0.0})
        d_, c_ = num(r[19]), num(r[21])
        x["ecritures_debit"] += d_ != 0
        x["debit"] += d_
        x["ecritures_credit"] += c_ != 0
        x["credit"] += c_
OUT["gl2024_epargne"] = sorted(mv.values(), key=lambda z: (z["compte"], -(z["debit"] + z["credit"])))
OUT["comptes_interets_servis"] = sorted({str(c) for c in bal[bal.compte.str.match(r"^67") & bal.intitule.str.contains("pargn|EFC|PEP", case=False, regex=True)].compte})

json.dump(OUT, open(SORTIE_J, "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)

# =====================================================================================================================
# Classeur à formules
# =====================================================================================================================
NOIR, GRIS, ORANGE = "060606", "E8E8E8", "EFAC62"
wb = openpyxl.Workbook()


def titre(ws, t, src, larg):
    ws.sheet_view.showGridLines = False
    ws["A1"], ws["A2"] = t, src
    ws["A1"].font, ws["A2"].font = Font(name="Helvetica", bold=True, size=13, color="112A40"), Font(name="Helvetica", italic=True, size=9, color="53565A")
    for j, w in enumerate(larg, 1):
        ws.column_dimensions[L(j)].width = w


def entete(ws, r, cols):
    for j, c in enumerate(cols, 1):
        x = ws.cell(r, j, c)
        x.font, x.fill = Font(name="Helvetica", bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NOIR)
        x.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")


# 14a : paramètres modifiables
pa = wb.active
pa.title = "14a_pidr_parametres"
titre(pa, "PIDR : paramètres (cellules orange modifiables)", "Sources : décret n° 2017-210 du 30/03/2017, art. 4 et 5 ; convention collective interprofessionnelle du 19/07/1977, art. 40 ; DISA 2021 à 2025 ; Investing.com ; table CIMA F", [46, 18, 90])
entete(pa, 4, ["Paramètre", "Valeur", "Justification"])
PAR = [("Date d'évaluation", dt.date(2025, 12, 31), "Clôture de l'exercice 2025, dernière DISA reçue."),
       ("Âge de départ à la retraite", AGE_DEP, "Convention collective du 19/07/1977, art. 40 : 55 ans, reculé d'un an par enfant à charge (maximum 3) ; enfants non déclarés dans les fichiers reçus."),
       ("Taux d'actualisation", TAUX10, f"Moyenne 2025 du rendement de l'obligation de l'État de Côte d'Ivoire à 10 ans ({len(cours)} cotations, Investing.com) ; IAS 19, § 83 (taux des obligations d'État à défaut de marché profond)."),
       ("Croissance annuelle des salaires", CROIS, f"Médiane observée sur {len(crois)} agents présents douze mois en 2021 et en 2025 (DISA 2021 et 2025)."),
       ("Rotation annuelle du personnel", ROT, "Moyenne des taux annuels de sortie observés entre les DISA 2021, 2022, 2023 et 2025."),
       ("Barème jusqu'à 5 ans", 0.30, "Décret n° 2017-210, art. 4"), ("Barème de 6 à 10 ans", 0.35, "Décret n° 2017-210, art. 4"), ("Barème au-delà de 10 ans", 0.40, "Décret n° 2017-210, art. 4")]
for k, (lib, v, j_) in enumerate(PAR, 5):
    pa.cell(k, 1, lib)
    c = pa.cell(k, 2, v)
    c.fill = PatternFill("solid", fgColor=ORANGE)
    c.number_format = "dd/mm/yyyy" if isinstance(v, dt.date) else ("0" if isinstance(v, int) else "0.00%")
    pa.cell(k, 3, j_).alignment = Alignment(wrap_text=True)
P = {n: f"'14a_pidr_parametres'!$B${k}" for k, n in enumerate(("date", "age", "i", "g", "t", "b1", "b2", "b3"), 5)}

# 14t : table CIMA F
tt = wb.create_sheet("14t_table_cima_f")
titre(tt, "Table de mortalité CIMA F (nombre de survivants lx)", "Source : fichier CIMA_F.xlsx fourni le 18/10/2026 (Code CIMA, art. 338)", [10, 16])
entete(tt, 4, ["Âge", "lx"])
for k, a in enumerate(sorted(lx), 5):
    tt.cell(k, 1, a)
    tt.cell(k, 2, lx[a]).number_format = "#,##0"
TL = f"'14t_table_cima_f'!$A$5:$B${4 + len(lx)}"

# 14b : calcul agent par agent (2025), formules vivantes
ib = wb.create_sheet("14b_pidr_individuel")
cols = ["Identifiant", "Nom", "Prénom", "Entité", "Rattachement (source)", "Année de naissance", "Date d'embauche", "Salaire brut annuel", "Mois travaillés",
        "Salaire mensuel moyen", "Âge au 31/12/2025", "Ancienneté (années)", "Années avant le départ", "Indemnité acquise", "Indemnité projetée au départ",
        "Prorata des droits", "Probabilité de survie", "Facteur de rotation", "Facteur d'actualisation", "PIDR (unités de crédit projetées)"]
titre(ib, "PIDR agent par agent au 31/12/2025 (méthode des unités de crédit projetées)",
      "Source : DISA_256927_2025.xls (CNPS) ; rattachement : ETAT 301 & DISA 2024 FPPN.xls (RECAPSALAIRES 2024 CMP) et ETAT 301 & DISA 2025 FPPN.xls (RECAPSALAIRES 2025 GRENIER) ; paramètres : onglet 14a",
      [11, 18, 22, 18, 30, 10, 12, 15, 9, 14, 9, 11, 10, 15, 15, 10, 10, 10, 11, 16])
entete(ib, 4, cols)
B = lambda c: f"(0.3*MIN({c},5)+0.35*MIN(MAX({c}-5,0),5)+0.4*MAX({c}-10,0))"  # noqa: E731
for k, x in enumerate(lig25, 5):
    vals = [x["id"], x["nom"], x["prenom"], x["entite"], x["source_entite"], x["naissance"],
            dt.datetime.strptime(x["embauche"], "%d/%m/%Y").date(), x["brut"], x["mois"]]
    for j, v in enumerate(vals, 1):
        ib.cell(k, j, v)
    ib.cell(k, 7).number_format = "dd/mm/yyyy"
    f = {10: f"=IF(I{k}>0,H{k}/I{k},0)", 11: f"=YEAR({P['date']})-F{k}",
         12: f"=MAX(DATEDIF(G{k},{P['date']},\"m\"),0)/12", 13: f"=MAX({P['age']}-K{k},0)",
         14: f"=J{k}*{B(f'L{k}')}", 15: f"=J{k}*(1+{P['g']})^M{k}*{B(f'(L{k}+M{k})')}",
         16: f"=IF(L{k}+M{k}>0,L{k}/(L{k}+M{k}),0)", 17: f"=VLOOKUP(K{k}+M{k},{TL},2,0)/VLOOKUP(K{k},{TL},2,0)",
         18: f"=(1-{P['t']})^M{k}", 19: f"=(1+{P['i']})^-M{k}", 20: f"=O{k}*P{k}*Q{k}*R{k}*S{k}"}
    for j, v in f.items():
        c = ib.cell(k, j, v)
        c.number_format = "#,##0" if j in (10, 14, 15, 20) else ("0.00" if j in (12,) else ("0" if j in (11, 13) else "0.0000"))
    for j in (8,):
        ib.cell(k, j).number_format = "#,##0"
der = 4 + len(lig25)
ib.cell(der + 1, 1, "Total")
for j in (8, 14, 15, 20):
    ib.cell(der + 1, j, f"=SUM({L(j)}5:{L(j)}{der})").number_format = "#,##0"
ib.cell(der + 1, 1).font = Font(bold=True)
ib.freeze_panes = "D5"
ib.cell(4, 20).comment = Comment("PIDR = indemnité projetée × prorata ancienneté actuelle / ancienneté au départ × survie (CIMA F) × (1 - rotation)^n × (1 + taux)^-n. "
                                 "Indemnité = salaire mensuel moyen × barème du décret n° 2017-210 (art. 4) appliqué à l'ancienneté.", "Adenka")

# 14c : synthèse par entité et par année
sy = wb.create_sheet("14c_pidr_synthese")
titre(sy, "PIDR par entité et par année, comparée au compte 196", "Sources : DISA 2021, 2022, 2023, 2025 ; balances du FPPN (compte 196) ; calculs Adenka (onglet 14b pour 2025)", [26, 12, 18, 18, 18, 18, 18, 18])
entete(sy, 4, ["Exercice", "Agents", "PIDR Direction Générale", "PIDR CMP", "PIDR Grenier", "PIDR rattachement non documenté", "PIDR totale", "Compte 196 (FPPN)"])
for k, s in enumerate(synth, 5):
    sy.cell(k, 1, s["annee"])
    sy.cell(k, 2, s["agents"])
    for j, e in enumerate(ENT, 3):
        if s["annee"] == 2025:
            sy.cell(k, j, f"=SUMIF('14b_pidr_individuel'!$D$5:$D${der},\"{e}\",'14b_pidr_individuel'!$T$5:$T${der})")
        else:
            sy.cell(k, j, s[e]["dbo"])
    sy.cell(k, 7, f"=SUM(C{k}:F{k})")
    sy.cell(k, 8, s["provision_196"])
    for j in range(3, 9):
        sy.cell(k, j).number_format = "#,##0"

# 15 : mortalité brute et sinistralité du FPPN
mo = wb.create_sheet("15_mortalite_sinistralite")
titre(mo, "Mortalité brute des membres participants et sinistralité décès selon le FPPN",
      "Source : ETUDE DES RATIO MALADIE DECES 2.xlsx (feuilles Tableau et Historique 2015-2021) ; balances du FPPN ; calculs Adenka", [10, 14, 12, 12, 16, 18, 18, 14, 18, 18, 16])
entete(mo, 4, ["Année", "Adhérents", "Décès adhérents", "Décès total", "Taux brut (pour mille)", "Cotisations payées (FPPN)", "Sinistres déclarés (FPPN)", "S/P du FPPN",
               "Cotisations décès (comptes)", "Prestations décès (comptes)", "Prime d'assurance décès (comptes)"])
for k, m in enumerate(mort, 5):
    s_ = si.get(m["annee"])
    vals = [m["annee"], m["adherents"], m["deces_adherents"], m["deces_total"], f"=C{k}/B{k}*1000", m["cotisations"], m["sinistres"], f"=G{k}/F{k}",
            s_["cot_deces"] if s_ else None, s_["prest_deces"] if s_ else None, s_["prime_assurance_deces"] if s_ else None]
    for j, v in enumerate(vals, 1):
        c = mo.cell(k, j, v)
        c.number_format = "0.00" if j == 5 else ("0.00%" if j == 8 else "#,##0")

# 16 : EFC
ef_ = wb.create_sheet("16_efc_taux_servi")
titre(ef_, "EFC : intérêts annuels dus sur l'encours selon le taux servi",
      "Sources : balances du FPPN (compte 41195200) ; MANUEL DE PRODUIT D'EPARGNE FPPN VF VF.docx (5,5 %) ; CAHIER DE CHARGEEFC & PEP 2021.pdf (14 %) ; entretien du chef comptable (3,5 %)", [10, 18, 14, 14, 14, 18, 18, 18, 18])
entete(ef_, 4, ["Exercice", "Encours EFC (41195200)", "Taux 1", "Taux 2", "Taux 3", "Intérêts au taux 1", "Intérêts au taux 2", "Intérêts au taux 3", "Produits financiers"])
for k, e in enumerate(efc, 5):
    vals = [e["exercice"], e["encours_efc"], 0.035, 0.055, 0.14, f"=B{k}*C{k}", f"=B{k}*D{k}", f"=B{k}*E{k}", e["produits_financiers"]]
    for j, v in enumerate(vals, 1):
        c = ef_.cell(k, j, v)
        c.number_format = "0.00%" if j in (3, 4, 5) else "#,##0"
        if j in (3, 4, 5):
            c.fill = PatternFill("solid", fgColor=ORANGE)

# 17 : contrôles croisés
cc = wb.create_sheet("17_controles_croises")
titre(cc, "Contrôles croisés entre documents", "Sources : ETATS FINANCIERS 2020 à 2025 et AU 31-08-2026.pdf ; balances du FPPN ; ETUDE DES RATIO MALADIE DECES 2.xlsx ; fichiers 8Z6 de la Solde ; ETAT DE RAPPROCHEMENT BANCAIRE FPPN 2020 à 2026.xlsx", [10, 20, 20, 16, 20, 20, 16])
r = 4
entete(cc, r, ["Exercice", "Total actif brut (états financiers)", "Total actif brut (balances)", "Écart", "Résultat (états financiers)", "Résultat (balances)", "Écart", "Trésorerie actif (états financiers)", "Soldes débiteurs classe 5 (balances)", "Écart"])
for x in OUT["ef_contre_balances"]:
    r += 1
    for j, v in enumerate([x["exercice"], x["ef_brut"], x["balance_brut"], f"=B{r}-C{r}", x["ef_resultat"], x["balance_resultat"], f"=E{r}-F{r}", x["ef_tresorerie"], x["balance_classe5"], f"=H{r}-I{r}"], 1):
        cc.cell(r, j, v).number_format = "#,##0"
r += 2
entete(cc, r, ["Année", "Mois", "Adhérents (FPPN)", "Lignes 8Z6 (Solde)", "Écart"])
for x in ctl:
    r += 1
    for j, v in enumerate([x["annee"], x["mois"], x["adherents_fppn"], x["lignes_8z6"], f"=C{r}-D{r}"], 1):
        cc.cell(r, j, v).number_format = "#,##0"
r += 2
entete(cc, r, ["Année", "Mois couverts", "Lignes", "Rapport médian 8Z6 / 936", "Part des lignes à 1,23 / 3,77"])
for x in rap:
    r += 1
    for j, v in enumerate([x["annee"], x["mois"], x["lignes"], x["mediane"], x["part_cle"]], 1):
        cc.cell(r, j, v).number_format = "0.0000" if j == 4 else ("0.00%" if j == 5 else "#,##0")
r += 2
entete(cc, r, ["Exercice", "Comptes comparés", "Comptes dont l'à-nouveau diffère du solde N-1", "Écart absolu cumulé"])
for x in an_:
    r += 1
    for j, v in enumerate([x["exercice"], x["comptes_compares"], x["comptes_en_ecart"], x["ecart_absolu"]], 1):
        cc.cell(r, j, v).number_format = "#,##0"
r += 2
entete(cc, r, ["Exercice", "Comptes bancaires rapprochés", "Soldes égaux à la balance", "Solde comptable (rapprochements)", "Solde bancaire (rapprochements)", "Chèques en suspens > 1 an", "Montant", "Plus ancien"])
for x in res_erb:
    r += 1
    for j, v in enumerate([x["exercice"], x["comptes"], x["concordants"], x["solde_compta"], x["solde_banque"], x["suspens_plus_un_an"], x["montant_suspens_plus_un_an"], x["plus_ancien"]], 1):
        cc.cell(r, j, v).number_format = "dd/mm/yyyy" if j == 8 else "#,##0"
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font and not c.font.bold:
                c.font = Font(name="Helvetica", size=c.font.size or 10, color=c.font.color, italic=c.font.italic)
wb.calculation.fullCalcOnLoad = True
wb.save(SORTIE_X)
print(f"PIDR 2025 : {sum(x['dbo'] for x in lig25):,.0f} FCFA ({len(lig25)} agents) ; compte 196 : {synth[-1]['provision_196']:,.0f}")
print(f"taux {TAUX10:.4%} ; croissance {CROIS:.4%} ; rotation {ROT:.4%}")
