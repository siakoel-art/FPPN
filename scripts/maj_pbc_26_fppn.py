"""PBC du 26/10/2026 : derniers doublons de sens fusionnés, numérotation continue, synthèse et correspondance à jour.

Part de la PBC du 23/10/2026 (90 items) :
- ancien 49 (impayés du Grenier par tranche et dépréciation) fusionné dans l'ancien 10 (prêts du Grenier) ;
- ancien 79 (dossiers du solde du compte 41591000) fusionné dans l'ancien 64 (allocations décès restant à payer) ;
- ancien 7 limité aux prestations santé (les allocations décès non payées sont à l'ancien 64) ;
- ancien 33 limité à la dénomination et au statut du CMP (la comptabilité séparée est à l'ancien 83) ;
- ancien 35 limité aux dates de démarrage et de suspension des produits (la documentation des produits est aux anciens 46, 50 et 71).
Les onglets de pointage historiques gardent la numérotation de leur date.

Usage : python3 maj_pbc_26_fppn.py <PBC du 23/10/2026.xlsx> <sortie.xlsx>
"""
import copy
import sys

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

SRC, SORTIE = sys.argv[1:3]
wb = openpyxl.load_workbook(SRC)
ws = wb["PBC Actuariat FPPN"]
PREM, DERN = 11, 119
DATE = "26/10/2026"
lignes = [r for r in range(PREM, DERN + 1) if ws.cell(r, 2).value is not None]
assert len(lignes) == 90
L = {n: lignes[n - 1] for n in range(1, 91)}  # ancien n° -> ligne


def ajout(n, texte):
    c = ws.cell(L[n], 11)
    c.value = (c.value or "").rstrip() + "\n" + texte


def texte_court(n, k=110):
    t = str(ws.cell(L[n], 3).value).strip()
    return t if len(t) <= k else t[:k].rstrip() + "…"


# ---- fusions
ws.cell(L[10], 3).value = ("Prêts aux mutualistes (« Le Grenier du Policier ») au 31/12/2025 et au 31/08/2026 : encours par emprunteur, échéancier de remboursement, "
                           "impayés par tranche de retard, et décision fixant la dépréciation des prêts en retard (taux appliqués).")
ajout(10, f"Fusion du {DATE} : reprend l'ancien item 49 (« {texte_court(49)} »), statut « {ws.cell(L[49], 13).value} ». "
          "La politique de crédit (« POLITIK DE CREDIT DU GRENIER VFVF.doc », chapitre XI.3) fixe 40 %, 80 % et 100 % selon l'ancienneté du retard. "
          "Aucune dépréciation comptabilisée depuis 2022, et le FPPN cite « garanties et 1 % » : la décision appliquée reste à transmettre.")
ws.cell(L[10], 5).value = "Moyen"

ws.cell(L[64], 3).value = ("Allocations décès restant à payer : fichiers des décès transmis par la direction du personnel, calcul de fin d'année de 2020 à 2025, "
                           "et liste des dossiers composant le solde des comptes 41591000 et 41592000 au 31/08/2026 (489 419 669 FCFA) : nom, grade, "
                           "date du décès, date de dépôt de la demande, montant, motif du non-paiement.")
ajout(64, f"Fusion du {DATE} : reprend l'ancien item 79 (« {texte_court(79)} »), statut « {ws.cell(L[79], 13).value} ». "
          "États de fin d'année 2021 à 2025 reçus. Manquent : fichiers de la direction du personnel et liste détaillée au 31/08/2026 avec la date de dépôt.")

# ---- items recentrés pour supprimer les recoupements
ws.cell(L[7], 3).value = ("Liste des prestations santé non payées à chacune des dates suivantes : 31/12/2019, 31/12/2020, 31/12/2021, 31/12/2022, 31/12/2023, "
                          "31/12/2024, 31/12/2025 et 31/08/2026, factures des prestataires de santé et remboursements aux mutualistes. Une ligne par dossier : "
                          "matricule, type de prestation, date des soins, date de dépôt, montant dû. Les allocations décès non payées sont demandées à l'item « Allocations décès restant à payer ».")
ajout(7, f"Recentrage du {DATE} : la partie décès (ancien item 29) est reprise par l'item « Allocations décès restant à payer ».")

ws.cell(L[33], 3).value = "Dénomination complète et statut juridique de la structure désignée « CMP » dans les comptes 47111104 et 47111105."
ajout(33, f"Recentrage du {DATE} : la comptabilité séparée du CMP est demandée à l'item « Comptabilité séparée du Centre médical du FPPN ».")

ws.cell(L[35], 3).value = ("Date de démarrage et, le cas échéant, de suspension de chaque produit : EFC, PEP, prêts du Grenier, flotte automobile, MEDOC, "
                           "cotisations des retraités (CGRAE).")
ws.cell(L[35], 13).value = "Non Reçu"
ajout(35, f"Recentrage du {DATE} : la documentation des produits est reçue et suivie aux items des manuels d'épargne, de la politique de crédit et des types de prêts du Grenier. "
          "MEDOC est suivi à son item. Reste la date de démarrage et de suspension de chaque produit (ancien item 60), non reçue.")

# ---- suppression des deux lignes fusionnées (de bas en haut), hauteurs de ligne conservées
hauteurs = {r: ws.row_dimensions[r].height for r in range(PREM - 1, DERN + 1)}
supprimees = sorted([L[79], L[49]], reverse=True)
for r in supprimees:
    ws.delete_rows(r)
garde = [r for r in range(PREM - 1, DERN + 1) if r not in supprimees]
for i, r in enumerate(garde):
    ws.row_dimensions[PREM - 1 + i].height = hauteurs[r]
for r in range(DERN - 1, DERN + 1):
    ws.row_dimensions[r].height = None
DERN2 = DERN - len(supprimees)

# ---- numérotation continue (formules) et correspondance
nouv = [r for r in range(PREM, DERN2 + 1) if ws.cell(r, 2).value is not None]
assert len(nouv) == 88
for i, r in enumerate(nouv):
    ws.cell(r, 2).value = "=1" if i == 0 else f"=+B{nouv[i - 1]}+1"
anciens = [n for n in range(1, 91) if n not in (49, 79)]
corr = {a: k + 1 for k, a in enumerate(anciens)}
corr[49], corr[79] = corr[10], corr[64]

# ---- plages des listes déroulantes, des mises en forme conditionnelles et des formules
ws.data_validations.dataValidation = []
for col, liste in (("E", '"Urgent,Moyen,Faible"'), ("J", '"Physique,Electronique,Physique et Electronique"'), ("M", '"Reçu,Non Reçu,Non Applicable,Partiellement Reçu"')):
    dv = DataValidation(type="list", formula1=liste, allow_blank=True)
    dv.add(f"{col}{PREM}:{col}{DERN2}")
    ws.add_data_validation(dv)
cf = ws.conditional_formatting
regles = {str(k.sqref): list(v) for k, v in cf._cf_rules.items()}
cf._cf_rules.clear()
for plage, rs in regles.items():
    col = plage[0]
    for regle in rs:
        cf.add(f"{col}{PREM}:{col}{DERN2}", regle)
ws["M7"].value = f'=IFERROR(COUNTIF($M${PREM}:$M${DERN2},"Reçu")/(COUNTA($M${PREM}:$M${DERN2})-COUNTIF($M${PREM}:$M${DERN2},"Non Applicable")),0)'
ws["C5"].value = f"DEMANDE DES DOCUMENTS (LISTE COMPLÈTE) ET POINTAGE AU {DATE}"

# ---- synthèse du pointage
statut = {k + 1: ws.cell(r, 13).value for k, r in enumerate(nouv)}
s = wb["Synthèse du pointage"]
s["B3"].value = (f"Pointage au {DATE} : envois des 15/10, 18/10, 19/10, 20/10, 21/10, 22/10 et 23/10/2026, vérification de cohérence du 23/10/2026. "
                 "Doublons fusionnés le 23/10 et le 26/10/2026 (onglets « Correspondance des numéros » et « Correspondance 26-10 »).")
for ligne_, st in ((6, "Reçu"), (7, "Partiellement Reçu"), (8, "Non Reçu")):
    s.cell(ligne_, 3).value = f"=COUNTIF('PBC Actuariat FPPN'!$M${PREM}:$M${DERN2},\"{st}\")"
    s.cell(ligne_, 4).value = ", ".join(str(n) for n, v in statut.items() if v == st)
s["D9"].value = "1 à 88"

# ---- correspondance des numéros du 23/10 au 26/10/2026
src = wb["Correspondance des numéros"]
c = wb.copy_worksheet(src)
c.title = "Correspondance 26-10"
for row in c.iter_rows(min_row=6, max_row=c.max_row):
    for cell in row:
        cell.value = None
c["A2"].value = "Correspondance des numéros d'items (26/10/2026)"
c["A3"].value = "Numérotation de la PBC du 23/10/2026 et numérotation du 26/10/2026. Un item fusionné renvoie à l'item qui reprend sa demande."
c["A5"].value, c["B5"].value, c["C5"].value = "N° du 23/10", "N° du 26/10", "Traitement"
modele = [copy.copy(src.cell(6, k)._style) for k in (1, 2, 3)]
trait = {49: "Fusionné avec l'ancien item 10 (doublon : impayés et dépréciation des prêts du Grenier)",
         79: "Fusionné avec l'ancien item 64 (doublon : allocations décès restant à payer)",
         10: "Conservé, reprend l'ancien item 49", 64: "Conservé, reprend l'ancien item 79",
         7: "Conservé, limité aux prestations santé", 33: "Conservé, limité à la dénomination et au statut du CMP",
         35: "Conservé, limité aux dates de démarrage et de suspension des produits"}
for i, a in enumerate(range(1, 91)):
    r = 6 + i
    for k, v in enumerate((str(a), str(corr[a]), trait.get(a, "Conservé"))):
        cell = c.cell(r, k + 1, v)
        cell._style = copy.copy(modele[k])
wb.move_sheet(c, offset=wb.sheetnames.index("Correspondance des numéros") + 1 - (len(wb.sheetnames) - 1))

wb.save(SORTIE)
compte = {st: sum(1 for v in statut.values() if v == st) for st in ("Reçu", "Partiellement Reçu", "Non Reçu")}
print(SORTIE, "items :", len(statut), compte)
