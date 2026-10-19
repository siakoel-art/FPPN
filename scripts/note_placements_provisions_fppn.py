# -*- coding: utf-8 -*-
"""
FPPN - Notes internes confidentielles sur les entretiens du 19/10/2026 avec M. Koffi, chef comptable
(placements à 14h, provisions à 16h). Mise en forme de la note interne de l'entretien sinistres (16/10/2026).
Utilisation : python note_placements_provisions_fppn.py <note_modele.docx> <dossier_sortie>
"""
import os
import sys
from copy import deepcopy

import docx

MODELE, SORTIE = sys.argv[1:3]

NOTES = {
    "FPPN_Note_interne_entretien_placements_20261019_CONFIDENTIEL.docx": (
        ["FPPN : note interne sur l'entretien placements",
         "Document interne Adenka. Confidentiel. Ne pas diffuser au client.",
         "Entretien du 19 octobre 2026, à 14h, avec M. Koffi, chef comptable. Participants Adenka : Emmanuel Dacoury, Marlène Mody, Cédric Kouamé et Jean-Noël Sialou.",
         "Objet : consigner, de façon factuelle et neutre, les éléments sensibles relevés pendant l'entretien et non repris tels quels dans le compte rendu destiné au client."],
        [("1. Concentration des décisions", [
            "Selon l'interlocuteur, toutes les décisions de placement viennent de la direction générale. La comptabilité ne fournit qu'un point de trésorerie.",
            "Il n'existe, à sa connaissance, ni comité de placement, ni avis du conseil d'administration, ni procédure écrite.",
            "Les placements ne sont donc pas choisis selon une politique validée par les organes du FPPN. Ce point touche à la gouvernance et doit être recoupé avec la direction générale et le DAF.",
        ]),
         ("2. Hôpital de la Police", [
            "L'interlocuteur indique que le FPPN a équipé l'hôpital de la Police, payé des médecins et fourni des médicaments, puis que la direction des services de santé de la Police s'en est approprié la gestion.",
            "La créance correspondante (compte 47111100, 3 407 519 850 FCFA) est inchangée depuis 2020. Des démarches auprès de l'État seraient en cours.",
            "Le recouvrement de cette créance est incertain : elle gonfle l'actif et les fonds propres. À vérifier : dossier, correspondances, décision d'autonomie de l'hôpital.",
        ]),
         ("3. Clinique du FPPN (CMP) et Grenier", [
            "Selon l'interlocuteur, l'ancien directeur général voulait donner à la clinique une identité distincte du FPPN. En attendant, toutes ses dépenses sont suivies en comptes 47 ; un cabinet prépare un bilan d'ouverture.",
            "Au 31/08/2026, les investissements au CMP atteignent 5 803 025 455 FCFA et les charges de fonctionnement 1 280 300 860 FCFA (comptes 47111104 et 47111105).",
            "Le Grenier, idée de l'ancien directeur général, n'a pas de personnalité morale. Le FPPN a emprunté à la banque pour le doter et a remboursé l'emprunt lui-même ; le Grenier n'a encore rien remboursé. Selon l'interlocuteur, il a aussi servi à attirer des souscriptions au PEP et à l'EFC.",
        ]),
         ("4. Contrôle comptable et bancaire", [
            "Une écriture validée peut être modifiée ou supprimée, et tous les utilisateurs ont les mêmes droits dans le logiciel comptable. Aucune saisie n'est validée par une seconde personne.",
            "Les rapprochements bancaires ne sont plus supervisés. Les chèques en suspens depuis plus de trois ans sont annulés au changement de signataires, sans analyse. Des comptes bancaires inactifs ne sont pas fermés.",
            "Ces points relèvent des collègues chargés du contrôle interne ; ils fragilisent la fiabilité des balances utilisées par l'actuariat.",
        ]),
         ("5. Incident Orabank", [
            "L'interlocuteur cite un DAT chez Orabank dont la banque a arrêté les prélèvements mensuels après deux versements ; le capital et les intérêts ont été reversés au terme et le DAT n'a pas été reconduit.",
            "Il situe l'incident en 2021 ; les balances montrent un DAT Orabank ouvert en 2022 (58 400 000 FCFA) et soldé en 2023 (60 084 840 FCFA).",
        ]),
         ("6. Points d'attention pour l'équipe", [
            "Ces propos émanent d'une seule personne. Ils doivent être recoupés avec les documents demandés et avec la direction générale.",
            "Ne pas reprendre ces éléments dans un document destiné au client sans validation de la direction de mission, en particulier sur l'hôpital et la gouvernance des placements.",
        ])]),
    "FPPN_Note_interne_entretien_provisions_20261019_CONFIDENTIEL.docx": (
        ["FPPN : note interne sur l'entretien provisions",
         "Document interne Adenka. Confidentiel. Ne pas diffuser au client.",
         "Entretien du 19 octobre 2026, à 16h, avec M. Koffi, chef comptable. Participants Adenka : Issa Bamba et Jean-Noël Sialou.",
         "Objet : consigner, de façon factuelle et neutre, les éléments sensibles relevés pendant l'entretien et non repris tels quels dans le compte rendu destiné au client."],
        [("1. Provisions héritées", [
            "Les provisions maladie (2 867 602 878 FCFA) et obsèques (1 961 695 412 FCFA) auraient été passées vers 2004, avant le PCMS, selon une méthode de l'époque, en fin d'année après les travaux du commissaire aux comptes.",
            "Plusieurs tentatives d'annulation n'ont pas abouti. Les archives auraient été perdues lors d'un sinistre en 2010.",
            "Personne ne peut aujourd'hui justifier ces montants. Leur reprise ou leur maintien a un effet direct sur le résultat et les fonds propres.",
        ]),
         ("2. Absence de provision maladie", [
            "Le FPPN ne calcule aucune provision maladie. Les factures reçues jusqu'en mars de l'année suivante sont rattachées à l'exercice ; les soins facturés plus tard ne sont couverts par rien.",
        ]),
         ("3. Allocations décès", [
            "Le calcul de fin d'année dépend d'une liste transmise par la direction du personnel à l'informaticien, puis au chef comptable. Le compte où il est passé n'a pas été précisé.",
            "Le pointage était fait auparavant par M. Soumahoro ; selon l'interlocuteur, le dossier lui a été retiré (transcription peu lisible, à confirmer).",
            "Le département qui devrait traiter les dossiers d'allocation ne fonctionne pas actuellement ; la comptabilité supplée.",
            "Les décès de conjoints et d'enfants ne sont connus que s'ils sont déclarés ; les délégués présents dans les préfectures doivent relayer l'information.",
        ]),
         ("4. Gestion déléguée et reclassement", [
            "En 2013, faute de pouvoir gérer les allocations sans accumuler de dette, le FPPN aurait confié la garantie à LMAI, puis à Sanlam, avant de la reprendre en interne au même barème.",
            "Le reclassement des provisions du compte 19 vers le compte 158 aurait été fait à la demande de l'AIRMS ; la date et le courrier sont à obtenir.",
        ]),
         ("5. Points d'attention pour l'équipe", [
            "Ces propos émanent d'une seule personne. Ils doivent être recoupés avec les documents demandés, la direction générale et le commissaire aux comptes.",
            "Ne pas reprendre ces éléments dans un document destiné au client sans validation de la direction de mission, en particulier sur l'origine des provisions et le retrait du dossier décès.",
        ])]),
}

for nom, (ENTETE, SECTIONS) in NOTES.items():
    d = docx.Document(MODELE)
    corps = d.element.body
    paras = d.paragraphs
    modele_entete = [deepcopy(p._p) for p in paras[:4]]
    modele_titre = deepcopy(paras[4]._p)
    modele_puce = deepcopy(paras[5]._p)
    for p in paras:
        corps.remove(p._p)
    sectpr = corps.find(docx.oxml.ns.qn("w:sectPr"))

    def ajouter(modele, texte):
        el = deepcopy(modele)
        sectpr.addprevious(el) if sectpr is not None else corps.append(el)
        p = docx.text.paragraph.Paragraph(el, d._body)
        runs = p.runs
        runs[0].text = texte
        for r in runs[1:]:
            r._r.getparent().remove(r._r)

    for m, t in zip(modele_entete, ENTETE):
        ajouter(m, t)
    for titre, puces in SECTIONS:
        ajouter(modele_titre, titre)
        for t in puces:
            ajouter(modele_puce, t)
    d.save(os.path.join(SORTIE, nom))
    print(nom, len(d.paragraphs), "paragraphes")
