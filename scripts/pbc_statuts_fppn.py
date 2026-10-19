# -*- coding: utf-8 -*-
"""
FPPN - Demande complémentaire après la lecture des statuts et du règlement intérieur du 15/06/2023 (items 100 à 109),
reportée dans la PBC générale (section et points à clarifier).
Utilisation : python pbc_statuts_fppn.py <pbc_modele.xlsx> <pbc_generale.xlsx> <sortie_complementaire.xlsx>
(la PBC générale est modifiée en place)
"""
import datetime as dt
import sys
from copy import copy

import openpyxl
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation

MODELE, GENERALE, SORTIE = sys.argv[1:4]
D_DEM, D_LIM = dt.datetime(2026, 10, 19), dt.datetime(2026, 10, 26)
SECTIONS = [
    ("Textes du FPPN cités par les statuts et le règlement intérieur", [
        ("Cahier des charges du régime des prestations sociales du FPPN en vigueur, et versions applicables depuis 2020 (cité aux articles 21 et 40 à 42 du règlement intérieur).", "Urgent", "Physique et Electronique"),
        ("Tableau des allocations en vigueur de 2020 au 31/08/2026 : allocation décès du membre, allocation décès des ayants droit, allocation d'invalidité par corps (articles 43, 44, 58 et 59 du règlement intérieur ; complète l'item 93).", "Urgent", "Physique et Electronique"),
        ("Manuels de procédure du Plan Epargne Plus et de l'Entraide de Fin de Carrière : prestations, rémunération servie, conditions de sortie et de rachat (articles 69 et 71 du règlement intérieur).", "Urgent", "Electronique"),
    ]),
    ("Personnel de la Direction Générale (périmètre de la provision pour départ à la retraite)", [
        ("Fichier du personnel de la Direction Générale avec une colonne de statut : personnel civil salarié ou fonctionnaire de police en service (article 31 du règlement intérieur ; complète l'item 76).", "Urgent", "Electronique"),
        ("Contrat de travail type du personnel civil, convention collective appliquée et tout avenant fixant l'âge de départ à la retraite (complète l'item 77).", "Moyen", None),
    ]),
    ("Décisions des organes", [
        ("Délibérations du Conseil d'Administration autorisant les placements et les ouvertures de comptes (DAT BOA, DAT Orabank, FCP géré par United Capital), comme l'exige l'article 85 des statuts.", "Urgent", None),
        ("Procès-verbal de l'Assemblée Générale approuvant le reclassement des provisions techniques en comptes 15801000 et 15802000 en 2024 (article 75 des statuts).", "Moyen", None),
        ("Procès-verbaux des Assemblées Générales d'affectation des résultats de 2020 à 2025 (article 90 des statuts : réserve légale de 20 % à 50 % de l'excédent net).", "Moyen", None),
        ("Décision fixant la situation des policiers retraités : maintien de la qualité de membre (article 7 des statuts), restitution de la carte (article 38 du règlement intérieur), cotisation et garantie santé (complète l'item 79).", "Urgent", None),
        ("Délibération du Conseil d'Administration approuvant le cadre organique des emplois de la Direction Générale et la grille des salaires et indemnités (articles 61 et 66 des statuts).", "Faible", None),
    ]),
]
POINTS = [
    ("Policiers retraités", "Les statuts (art. 7) font du policier retraité un membre participant, mais le règlement intérieur (art. 38) lui fait restituer sa carte au départ à la retraite, et seule la cotisation du policier en activité est prévue (statuts, art. 80).", "Préciser la situation des retraités (item 108)."),
    ("Réserve légale 2023", "La dotation à la réserve légale de 2023 vaut 8,94 % de l'excédent de 2022, sous le minimum de 20 % de l'article 90 des statuts.", "Transmettre le procès-verbal d'affectation (item 107)."),
    ("Personnel de la Direction Générale", "Deux catégories : personnel civil salarié et fonctionnaires de police percevant des indemnités (règlement intérieur, art. 31). Seul le personnel civil relève a priori de l'indemnité de départ du décret n° 2017-210.", "Transmettre le fichier du personnel avec le statut (item 103)."),
    ("Cahier des charges", "Les montants des prestations, de l'EFC et du PEP sont renvoyés au cahier des charges et au tableau des allocations, non transmis.", "Transmettre les deux documents (items 100 et 101)."),
]
COLS = range(2, 14)


# ---------------------------------------------------------------- PBC complémentaire
wb = openpyxl.load_workbook(MODELE)
ws = wb.worksheets[0]
ws.title = "PBC Textes et personnel"
ws["C4"] = "AUDIT DU FPPN 2020 – 31 AOÛT 2026 : DEMANDE COMPLÉMENTAIRE TEXTES ET PERSONNEL"
ws["C5"] = "SUITE À LA LECTURE DES STATUTS ET DU RÈGLEMENT INTÉRIEUR ÉDITÉS LE 15/06/2023"
st_vide = {c: copy(ws.cell(9, c)._style) for c in COLS}
st_sect = {c: copy(ws.cell(10, c)._style) for c in COLS}
st_item = {c: copy(ws.cell(11, c)._style) for c in COLS}
for r in range(10, ws.max_row + 1):
    for c in COLS:
        ws.cell(r, c).value = None
        ws.cell(r, c)._style = copy(st_vide[c])
    ws.row_dimensions[r].height = 13.5
ws.data_validations.dataValidation = []
r, prec, premier = 10, None, None
for titre, items in SECTIONS:
    for c in COLS:
        ws.cell(r, c)._style = copy(st_sect[c])
    ws.cell(r, 3, titre)
    ws.row_dimensions[r].height = 18
    r += 1
    for texte, prio, support in items:
        for c in COLS:
            ws.cell(r, c)._style = copy(st_item[c])
        ws.cell(r, 2, "=100" if prec is None else f"=+B{prec}+1")
        ws.cell(r, 3, texte)
        ws.cell(r, 5, prio)
        ws.cell(r, 7, D_DEM)
        ws.cell(r, 8, D_LIM)
        ws.cell(r, 10, support)
        ws.cell(r, 13, "Non Reçu")
        ws.row_dimensions[r].height = 21 if len(texte) < 110 else 36
        premier = premier or r
        prec = r
        r += 1
    r += 1
ws["M7"] = f'=IFERROR(COUNTIF($M${premier}:$M${prec},"Reçu")/(COUNTA($M${premier}:$M${prec})-COUNTIF($M${premier}:$M${prec},"Non Applicable")),0)'
for col, liste in (("E", '"Urgent,Moyen,Faible"'), ("J", '"Physique,Electronique,Physique et Electronique"'),
                   ("M", '"Reçu,Non Reçu,Non Applicable,Partiellement Reçu"')):
    dv = DataValidation(type="list", formula1=liste, allow_blank=True)
    dv.add(f"{col}{premier}:{col}{prec}")
    ws.add_data_validation(dv)
for nom in [n for n in wb.sheetnames if n != ws.title]:
    del wb[nom]
wb.save(SORTIE)

# ---------------------------------------------------------------- PBC générale
wb = openpyxl.load_workbook(GENERALE)
ws = wb["PBC Actuariat FPPN"]
num = [x for x in range(11, ws.max_row + 1) if str(ws.cell(x, 2).value or "").startswith("=")]
assert len(num) == 99, len(num)
dern = num[-1]
sect = next(x for x in range(dern, 9, -1) if ws.cell(x, 3).value and not ws.cell(x, 2).value)
st_s = {c: copy(ws.cell(sect, c)._style) for c in COLS}
st_i = {c: copy(ws.cell(dern, c)._style) for c in COLS}
r, prec = dern + 1, dern
for titre, items in SECTIONS:
    r += 1
    for c in COLS:
        ws.cell(r, c)._style = copy(st_s[c])
    ws.cell(r, 3, titre + " (lecture des textes du 15/06/2023)")
    ws.row_dimensions[r].height = ws.row_dimensions[sect].height
    for texte, prio, support in items:
        r += 1
        for c in COLS:
            ws.cell(r, c)._style = copy(st_i[c])
        ws.cell(r, 2, f"=+B{prec}+1")
        ws.cell(r, 3, texte)
        ws.cell(r, 5, prio)
        ws.cell(r, 7, D_DEM)
        ws.cell(r, 8, D_LIM)
        ws.cell(r, 10, support)
        ws.cell(r, 11, "Demande issue de la lecture des statuts et du règlement intérieur du 15/06/2023 (PBC complémentaire textes et personnel).")
        ws.cell(r, 13, "Non Reçu")
        ws.row_dimensions[r].height = 21 if len(texte) < 110 else 36
        prec = r
    r += 1
plage = f"$M$11:$M${prec}"
ws["M7"] = f'=IFERROR(COUNTIF({plage},"Reçu")/(COUNTA({plage})-COUNTIF({plage},"Non Applicable")),0)'
for dv in ws.data_validations.dataValidation:
    col = str(dv.sqref).split(":")[0][0]
    dv.sqref = MultiCellRange(f"{col}11:{col}{prec}")
sy = wb["Synthèse du pointage"]
for k, st in ((6, "Reçu"), (7, "Partiellement Reçu"), (8, "Non Reçu")):
    sy.cell(k, 3, f"=COUNTIF('PBC Actuariat FPPN'!{plage},\"{st}\")")
sy["D8"] = sy["D8"].value + ", " + ", ".join(str(n) for n in range(100, 110))
sy["D9"] = "1 à 109"
pc = wb["Points à clarifier"]
d = pc.max_row
for k, (sujet, constat, demande) in enumerate(POINTS, 1):
    rr = d + k
    for c in range(2, 6):
        pc.cell(rr, c)._style = copy(pc.cell(d, c)._style)
    pc.row_dimensions[rr].height = pc.row_dimensions[d].height
    pc.cell(rr, 2, int(pc.cell(d, 2).value) + k)
    pc.cell(rr, 3, sujet)
    pc.cell(rr, 4, constat)
    pc.cell(rr, 5, demande)
wb.save(GENERALE)
print(f"complémentaire : items 100 à {100 + sum(len(i) for _, i in SECTIONS) - 1} ; générale : jusqu'à la ligne {prec}, points jusqu'à {pc.max_row}")
