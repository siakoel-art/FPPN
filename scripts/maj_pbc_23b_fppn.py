"""PBC du 23/10/2026 : pointage des documents reçus le 23/10/2026 (barème tarifaire du FPPN, barème de la CGRAE, logigrammes des cotisations).

Le fichier est mis à jour sur place : statut (colonne M), date de réception (colonne I), commentaire complété (colonne K),
puis liste des numéros par statut de l'onglet « Synthèse du pointage ».

Usage : python3 maj_pbc_23b_fppn.py FPPN_PBC_complete_pointee_20261023.xlsx
"""
import datetime as dt
import sys

import openpyxl

PBC = sys.argv[1]
RECU = dt.datetime(2026, 10, 23)
POINT = {  # n° PBC du 23/10/2026 : (nouveau statut, complément de commentaire)
    2: ("Partiellement Reçu", "Reçu le 23/10/2026 : barème de la cotisation des retraités (« BAREME_CGRAE.xlsx ») : 2,5 % de (pension de base - MUGEFCI) + 5 000 FCFA (sous-officier), 8 000 FCFA (officier), 10 000 FCFA (commissaire de 2e classe à principal) ou 15 000 FCFA (commissaire divisionnaire à général). Manquent : avenants, annexe 2 et liste des souscripteurs non précomptés."),
    24: ("Partiellement Reçu", "Reçu le 23/10/2026 : barème tarifaire (« Bar_me_tarif_du_FPPN.pdf », 4 pages scannées) : tarifs du CMP pour le FPPN (part FPPN 80 %, part mutualiste 20 %), tarifs proposés aux structures privées, plafond annuel de 1 500 000 FCFA par bénéficiaire et 3 000 000 FCFA par famille. Document non daté, mention « tarifications proposées ». Manquent : conventions signées avec chaque prestataire et date d'application du barème."),
    52: ("Partiellement Reçu", "Reçu le 23/10/2026 : barème de la cotisation des retraités (« BAREME_CGRAE.xlsx »). Manquent : durée de la garantie (à vie ou limitée), ayants droit couverts et décision qui la fonde."),
    88: ("Partiellement Reçu", "Reçu le 23/10/2026 : logigrammes « FLOW_CHART_COTISATION_ET_PRODUITS_DEPARGNE_1.pptx » (EFC et PEP, flotte automobile GNA, pension CGRAE, maladie et décès). Manquent : Grenier (9X4), MEDOC (8Z7) et procédure de rapprochement des fichiers aller et retour."),
}

wb = openpyxl.load_workbook(PBC)
ws = wb["PBC Actuariat FPPN"]
lignes = [r for r in range(11, ws.max_row + 1) if ws.cell(r, 13).value and ws.cell(r, 3).value]
assert len(lignes) == 90, len(lignes)
for n, (st, com) in POINT.items():
    r = lignes[n - 1]
    ws.cell(r, 13).value = st
    if ws.cell(r, 9).value is None:
        ws.cell(r, 9).value = RECU
    ancien = ws.cell(r, 11).value or ""
    if com not in ancien:
        ws.cell(r, 11).value = (ancien + "\n" if ancien else "") + com
syn = wb["Synthèse du pointage"]
par = {"Reçu": [], "Partiellement Reçu": [], "Non Reçu": []}
for n, r in enumerate(lignes, 1):
    par[ws.cell(r, 13).value].append(str(n))
for row in range(6, 9):
    syn.cell(row, 4).value = ", ".join(par[syn.cell(row, 2).value])
syn["B3"].value = syn["B3"].value.replace(" et 22/10/2026,", ", 22/10 et 23/10/2026,")
wb.save(PBC)
print(PBC, {k: len(v) for k, v in par.items()})
