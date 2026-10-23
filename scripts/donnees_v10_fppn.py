# -*- coding: utf-8 -*-
"""
FPPN - Revue actuarielle complète : provisions, engagements, adossement actif-passif, projection
===============================================================================================
Calcule et documente, pour un lecteur non technique, chaque calcul de la revue :
  18 : provision pour prestations à payer (PSAP) maladie, méthode courte (délais de règlement, factures tardives) ;
  19 : allocations décès : états de fin d'année du FPPN, dette au compte 41591000, ancienneté, décès déclarés tardivement,
       comparaison avec la provision obsèques ;
  20 : provision pour risques courts (statuts, art. 88) ;
  21 : engagement d'épargne avec les intérêts dus non comptabilisés ;
  22 : santé des policiers retraités ;
  23 : dépréciation des prêts du Grenier selon sa politique de crédit ;
  24 : adossement actif-passif : écoulement de l'épargne, écarts de liquidité par échéance, tests de résistance ;
  25 : projection agrégée de l'équilibre technique (avec l'avertissement sur ses limites).
Sorties : un JSON pour le rapport et un classeur à formules vivantes, commentées.
Utilisation : python donnees_v10_fppn.py <dossier_lot1> <dossier_lot4> <dossier_base> <resultats.json> <provisions.json>
              <ratios.json> <donnees_v9.json> <sortie.json> <sortie.xlsx>
"""
import collections
import datetime as dt
import glob
import json
import os
import re
import sys

import openpyxl
import pandas as pd
from python_calamine import CalamineWorkbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from excel_pedago_fppn import ENT, FCFA, NB1, NB2, PCT, note, onglet, sous_titre, tableau  # noqa: E402

LOT1, LOT4, BASE, RES, PROV, RAT, V9J, SORTIE_J, SORTIE_X = sys.argv[1:10]
R = json.load(open(RES, encoding="utf-8"))
PV = json.load(open(PROV, encoding="utf-8"))
Q = json.load(open(RAT, encoding="utf-8"))
V9 = json.load(open(V9J, encoding="utf-8"))
OUT = {}
GLDIR = os.path.join(LOT1, "02-Nouveaux_GL et Balances", "Nouveaux fichiers")
PDEC = os.path.join(LOT4, "08-Nouveaux Dossiers Compta", "PROVISION ALLOCATION DECES ET FRAIS FUNERAIRES")
ANS = list(range(2020, 2027))
MOIS = {a: (8 if a == 2026 else 12) for a in ANS}
LIB = {a: (str(a) if a < 2026 else "31/08/2026") for a in ANS}
SRC_BAL = "balances des comptes du FPPN (Sage 100) : « BALANCE FORMAT EXCEL DE 09-2020 AU 31-12-2020.xlsx », « BALANCE FPPN FORMAT EXCEL DU 01-01-2021 AU 31-12-2026.xlsx » (exercice 2021), « BALANCE FPPN 2022.xlsx » à « BALANCE FPPN 2025.xlsx », « BALANCE FPPN DU 01-01-2026 AU 31-08-2026.xlsx » (provisoire)"


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def feuille(f, nom=None):
    w = CalamineWorkbook.from_path(f)
    return w.get_sheet_by_name(nom or w.sheet_names[0]).to_python()


def jour(v):
    if isinstance(v, (dt.date, dt.datetime)):
        return dt.date(v.year, v.month, v.day)
    m = re.search(r"(\d{1,2})[-/](\d{1,2})[-/](\d{4})", str(v))
    if m:
        try:
            return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:          # date mal saisie dans le fichier : ignorée
            return None
    return None


bal = pd.read_parquet(os.path.join(BASE, "compta_balances.parquet"))
bal = bal[bal.nature == "comptes"].copy()
for c in ("ouv_d", "ouv_c", "mvt_d", "mvt_c", "sol_d", "sol_c"):
    bal[c] = bal[c].fillna(0.0)
bal["sol"] = bal.sol_d - bal.sol_c


def cpt(an, compte, col):
    x = bal[(bal.exercice == an) & (bal.compte == compte)]
    return float(x[col].sum()) if len(x) else 0.0


SIN = {s["exercice"]: s for s in R["sinistralite"]}
BIL = {b["exercice"]: b for b in R["bilan"]}

# Formules vers les onglets de données du classeur final (D1_balances, D3, D4, 10_agregats) : aucun montant comptable
# n'est saisi en dur ; chaque cellule ainsi calculée est contrôlée dans l'onglet 0_controle (valeur publiée).
D1Q, D3Q, D4Q = "'D1_balances'", "'D3_precomptes_mensuels'", "'D4_grenier_prets'"
CTRLV = []


def FB(compte, an, col):
    """Montant d'un compte exact de la balance : net_c (crédit - débit), mvt_c, mvt_d, ouv_c (à-nouveaux crédit - débit)."""
    b = f'{D1Q}!$A:$A,"{compte}",{D1Q}!$C:$C,{an}'
    return {"net_c": f"-SUMIFS({D1Q}!$L:$L,{b})", "mvt_c": f"SUMIFS({D1Q}!$I:$I,{b})", "mvt_d": f"SUMIFS({D1Q}!$H:$H,{b})",
            "ouv_c": f"(SUMIFS({D1Q}!$G:$G,{b})-SUMIFS({D1Q}!$F:$F,{b}))"}[col]


def AG(cle, an):
    return f"INDEX('10_agregats'!$D:$J,MATCH(\"{cle}\",'10_agregats'!$A:$A,0),{an - 2019})"


def MOB(cloture, col):
    return f"INDEX('10_actif_mobilisation'!${col}:${col},MATCH(\"{cloture}\",'10_actif_mobilisation'!$A:$A,0))"


def ctl(ws, d, col, valeurs):
    for i, v in enumerate(valeurs):
        if v is not None:
            CTRLV.append((ws.title, f"{col}{d + i}", float(v)))

# =====================================================================================================================
# 18. PSAP maladie : délais de règlement et factures enregistrées tardivement
# =====================================================================================================================
psap_m = []
for a in ANS:
    psap_m.append({"exercice": a, "mois": MOIS[a], "dettes_cliniques": cpt(a, "40110000", "sol_c") - cpt(a, "40110000", "sol_d"),
                   "factures_enregistrees": cpt(a, "40110000", "mvt_c"), "paiements": cpt(a, "40110000", "mvt_d"),
                   "prestations_maladie": SIN[a]["prest_maladie"], "provision_maladie": PV["psap_maladie"][str(a)]})
OUT["psap_maladie_delais"] = psap_m


def annee_ref(s):
    s = str(s).upper()
    for p in (r"\bE(\d{2})-", r"FA ?(\d{2})\d{3,}", r"/(\d{4})\b", r"/(\d{2})\b", r"-(20\d{2})\b", r"\b(20\d{2})[-/]"):
        m = re.search(p, s)
        if m:
            y = int(m.group(1))
            y = y + 2000 if y < 100 else y
            if 2015 <= y <= 2026:
                return y
    return None


GLT = {2023: "GRAND LIVRE DES TIERS FPPN 2023.xlsx", 2024: "GRAND LIVRE DES TIERS 2024.xlsx", 2025: "GRAND LIVRES DES TIERS 2025.xlsx"}
tardives = []
for an, f in GLT.items():
    t = feuille(os.path.join(GLDIR, f))
    cur, tot, ident, prec, n_prec = None, 0.0, 0.0, 0.0, 0
    for r in t:
        a = str(r[0]).strip()
        if re.fullmatch(r"\d{8}", a) and r[3]:
            cur = a
            continue
        if cur and cur.startswith("401") and isinstance(r[0], (dt.date, dt.datetime)) and str(r[1]) == "SIN" and num(r[22]) > 0:
            m = num(r[22])
            tot += m
            y = annee_ref(str(r[10]) + " " + str(r[5]))
            if y:
                ident += m
            if y == an - 1:
                prec += m
                n_prec += 1
    tardives.append({"cloture": an - 1, "grand_livre": f, "annee_enregistrement": an, "factures_sinistres": tot, "references_identifiees": ident,
                     "references_annee_precedente": prec, "nombre_annee_precedente": n_prec})
OUT["psap_maladie_tardives"] = tardives

# =====================================================================================================================
# 19. Allocations décès : états de fin d'année (lot du 22/10/2026), compte 41591000, tardifs
# =====================================================================================================================
ETATS = [  # (exercice, fichier, feuille, colonne date de décès, colonne montant, colonne date d'enregistrement, total imprimé)
    (2021, "PROV ALLOCATION DECES 2021.xlsx", "IMPAYEES GLOBAL 2021", 6, 12, 1, 112820000),
    (2022, "PROV ALLOCATION DECES 2022.xlsx", "ALLOCATION DECES ", 4, 5, None, 88000000),
    (2023, "ALLOC DECES 2023 PAYER ET A PROVISIONNER .xlsx", "ALL DECES  A PROVISIONNER 2023", 6, 11, 1, 124690000),
    (2024, "PROV ALLOCATION DECES 2024.xlsx", "PROVISIONS DECES 2024 (2)", 5, 8, None, 140800000),
    (2025, "PROV ALLOCATIONS DECES ET FRAIS FUNERAIRES 2025.xlsx", "PROV ALLO DECES 2025", 5, 7, None, 130600000),
]
etats = []
for an, f, sh, cd, cm, ce, total in ETATS:
    t = feuille(os.path.join(PDEC, f), sh)
    lig = [r for r in t if str(r[0]).replace(".0", "").strip().isdigit() and len(r) > cm]
    pos = [r for r in lig if num(r[cm]) > 0]
    etats.append({"exercice": an, "fichier": f, "feuille": sh, "total_imprime": total, "lignes": len(lig), "dossiers_valorises": len(pos),
                  "somme_lignes": sum(num(r[cm]) for r in pos), "credit_41591000": cpt(an, "41591000", "mvt_c"),
                  "montants": dict(collections.Counter(int(num(r[cm])) for r in pos))})
ff25 = feuille(os.path.join(PDEC, "PROV ALLOCATIONS DECES ET FRAIS FUNERAIRES 2025.xlsx"), "PROV FRAIS FUNERAIRES 2025")
OUT["deces_frais_funeraires_2025"] = {"total_imprime": 9300000, "credit_41592000": cpt(2025, "41592000", "mvt_c")}
OUT["deces_etats"] = etats
g23 = feuille(os.path.join(PDEC, "ALLOC DECES 2023 PAYER ET A PROVISIONNER .xlsx"), "ALL DECES FPPN GLOBAL")
OUT["deces_2023_global"] = {"total_imprime": 141490000, "prestations_6021_2023": cpt(2023, "60210000", "sol_d") - cpt(2023, "60210000", "sol_c")}
# décès d'une année enregistrés l'année suivante (fichiers 2021 et 2023, qui portent les deux dates)
tard_d = []
for an, f, sh, cd, cm, ce, total in ETATS:
    if ce is None:
        continue
    t = feuille(os.path.join(PDEC, f), sh)
    for r in t:
        if not str(r[0]).replace(".0", "").strip().isdigit() or num(r[cm]) <= 0:
            continue
        dd, de = jour(r[cd]), jour(r[ce])
        if dd and de and de.year == dd.year + 1:
            tard_d.append({"etat": an, "deces": dd.isoformat(), "enregistrement": de.isoformat(), "montant": num(r[cm])})
OUT["deces_tardifs"] = tard_d
stock = []
for a in ANS:
    stock.append({"exercice": a, **{f"{k}_{c}": (cpt(a, c, "ouv_c") - cpt(a, c, "ouv_d") if k == "ouv" else cpt(a, c, k))
                                    for c in ("41591000", "41592000") for k in ("ouv", "mvt_c", "mvt_d")},
                  "provision_obseques": PV["psap_deces"][str(a)]})
OUT["deces_stock"] = stock
# ancienneté du stock au 31/08/2026 : premier entré, premier payé (hypothèse)
vint = [(2021, cpt(2021, "41591000", "mvt_c"))] + [(a, cpt(a, "41591000", "mvt_c")) for a in range(2022, 2027)]
paye = sum(cpt(a, "41591000", "mvt_d") for a in range(2021, 2027))
rest = []
x = paye
for a, m in vint:
    p = min(x, m)
    x -= p
    rest.append({"millesime": a, "dotation": m, "paye_fifo": p, "reste": m - p})
OUT["deces_anciennete"] = rest

# =====================================================================================================================
# 21. Épargne : intérêts dus non comptabilisés
# =====================================================================================================================
ep = []
for a in ANS:
    lig = {"exercice": a, "mois": MOIS[a]}
    for c, k in (("41195100", "pep"), ("41195200", "efc")):
        fin = cpt(a, c, "sol_c") - cpt(a, c, "sol_d")
        deb = fin - cpt(a, c, "mvt_c") + cpt(a, c, "mvt_d")
        lig.update({f"{k}_debut": deb, f"{k}_fin": fin, f"{k}_versements": cpt(a, c, "mvt_c"), f"{k}_sorties": cpt(a, c, "mvt_d")})
    ep.append(lig)
OUT["epargne"] = ep

# =====================================================================================================================
# 23. Grenier : dépréciation selon la politique de crédit (XI.3)
# =====================================================================================================================
BORNES = {"À jour": (-1, 0), "1 à 30 jours": (0, 30), "31 à 90 jours": (30, 90), "91 à 180 jours": (90, 180), "181 à 365 jours": (180, 365), "Plus de 365 jours": (365, 100000)}
TAUX_GR = {"À jour": 0.0, "1 à 30 jours": 0.0, "31 à 90 jours": 0.0, "91 à 180 jours": 0.40, "181 à 365 jours": 0.80, "Plus de 365 jours": 1.0}
gr = [dict(t, taux=TAUX_GR.get(t["tranche"])) for t in PV["grenier_tranches"]]
OUT["grenier_depreciation"] = gr
OUT["grenier_depreciation_comptable"] = {a: sum(cpt(a, c, "sol_c") - cpt(a, c, "sol_d") for c in bal[(bal.exercice == a) & bal.compte.str.match(r"^29")].compte.unique()) for a in (2025, 2026)}

# =====================================================================================================================
# 24. Adossement : écoulement de l'épargne, écarts de liquidité, tests de résistance
# =====================================================================================================================
sorties = []
for e in ep:
    if e["exercice"] in (2021, 2022, 2023, 2024, 2025):
        sorties.append({"exercice": e["exercice"], "pep_taux": e["pep_sorties"] / e["pep_debut"], "efc_taux": e["efc_sorties"] / e["efc_debut"],
                        "total_taux": (e["pep_sorties"] + e["efc_sorties"]) / (e["pep_debut"] + e["efc_debut"])})
OUT["epargne_sorties"] = sorties
OUT["mobilisation"] = Q["mobilisation"]
croiss = []
for a in range(2021, 2026):
    croiss.append({"exercice": a, "prest_maladie": SIN[a]["prest_maladie"] / SIN[a - 1]["prest_maladie"] - 1})
OUT["croissance_prestations"] = croiss


# =====================================================================================================================
# CLASSEUR PÉDAGOGIQUE
# =====================================================================================================================
wb = openpyxl.Workbook()
wb.remove(wb.active)

# ---------------------------------------------------------------- 18a
ws, r = onglet(wb, "18a_psap_maladie_delais", "Provision maladie : en combien de temps le FPPN paie-t-il les cliniques ?", {
    "À quoi sert cet onglet": "Vérifier si les factures des cliniques sont payées dans les 3 mois (règlement intérieur, art. 52). Si c'est le cas, la provision pour prestations à payer en maladie peut être calculée sur quelques mois, et non sur plusieurs années.",
    "D'où viennent les chiffres": "Compte 40110000 « Fournisseurs cliniques privées » et provision maladie (comptes 19011000 puis 15801000) des " + SRC_BAL + " ; prestations maladie : comptes 601, 6013, 6018 et 60800007 des mêmes balances.",
    "Comment le calcul est fait": "1) Factures enregistrées dans l'année = crédits du compte 40110000. 2) Factures moyennes d'un mois = factures de l'année / nombre de mois. 3) Délai moyen de paiement (en mois) = factures non payées en fin d'année / factures moyennes d'un mois. Exemple : 300 de factures non payées pour 100 de factures par mois donnent 3 mois d'attente en moyenne.",
    "Comment lire le résultat": "Un délai de 3 mois ou moins respecte l'article 52. Au-delà, des factures attendent plus longtemps que la règle. La dernière colonne compare la provision maladie inscrite au bilan aux prestations d'un mois.",
    "Valable pour": "Factures des cliniques enregistrées au compte 40110000, exercices 2020 à 2025 et du 01/01 au 31/08/2026. Ne couvre pas les remboursements aux mutualistes ni les bons de prise en charge payés hors de ce compte.",
    "Limites": "Mesure moyenne : elle ne dit pas facture par facture qui a attendu combien de temps. La date des soins n'est pas dans les comptes : on mesure le délai entre l'enregistrement de la facture et son paiement."},
    [16, 18, 18, 18, 14, 18, 18, 16, 18])
d, f = tableau(ws, r, ["Exercice", "Factures enregistrées (crédits 40110000)", "Paiements (débits 40110000)", "Factures non payées en fin de période", "Mois couverts",
                       "Factures moyennes d'un mois", "Délai moyen de paiement (mois)", "Prestations maladie de l'année", "Provision maladie au bilan (en mois de prestations)"],
               [[LIB[x["exercice"]], "=" + FB("40110000", x["exercice"], "mvt_c"), "=" + FB("40110000", x["exercice"], "mvt_d"), "=" + FB("40110000", x["exercice"], "net_c"),
                 x["mois"], "=B{r}/E{r}", "=D{r}/F{r}", "=" + AG("prest_maladie", x["exercice"]), f"={AG('prov_maladie', x['exercice'])}/(H{{r}}/E{{r}})"] for x in psap_m],
               [None, FCFA, FCFA, FCFA, ENT, FCFA, NB1, FCFA, NB1],
               commentaires={(0, 6): "Exemple de lecture : si la valeur vaut 3, les factures non payées représentent en moyenne trois mois de factures."})
for col, k in (("B", "factures_enregistrees"), ("C", "paiements"), ("D", "dettes_cliniques"), ("H", "prestations_maladie")):
    ctl(ws, d, col, [x[k] for x in psap_m])
note(ws, f + 2, "Lecture : de 2020 à 2022, les factures non payées représentent plusieurs années de factures (arriérés). De 2023 au 31/08/2026, le délai moyen s'approche de la règle des 3 mois sans la respecter chaque année.", 9)
# ---------------------------------------------------------------- 18b
ws, r = onglet(wb, "18b_psap_maladie_tardives", "Provision maladie : les factures d'une année arrivent-elles l'année suivante ?", {
    "À quoi sert cet onglet": "Mesurer les soins d'une année facturés et enregistrés seulement l'année suivante. Ce sont les « sinistres survenus mais non déclarés » que la provision doit couvrir (statuts, art. 89 ; règlement d'exécution n° 004/2022, art. 10).",
    "D'où viennent les chiffres": "Grands livres des tiers « GRAND LIVRE DES TIERS FPPN 2023.xlsx », « GRAND LIVRE DES TIERS 2024.xlsx », « GRAND LIVRES DES TIERS 2025.xlsx » : écritures du journal SIN (sinistres) au crédit des comptes de cliniques (401…), colonne « N° facture ».",
    "Comment le calcul est fait": "Pour chaque facture enregistrée, l'année est lue dans son numéro (par exemple « E24-000042 » ou « B.F.A N°23/24 » pour 2024). Une facture enregistrée en 2024 avec un numéro de 2023 est une facture de 2023 arrivée en retard. Son montant est ce que la provision au 31/12/2023 aurait dû couvrir.",
    "Comment lire le résultat": "Plus la part des factures de l'année précédente est faible, plus les factures arrivent vite et plus la provision nécessaire est petite.",
    "Valable pour": "Factures de cliniques enregistrées au journal SIN en 2023, 2024 et 2025, dont l'année est lisible dans le numéro. Ne couvre pas les factures sans année lisible (colonne « Part non identifiée »).",
    "Limites": "L'année du numéro de facture est celle de la facture, pas toujours celle des soins. Les factures dont le numéro ne contient pas d'année ne sont pas classées."},
    [14, 34, 18, 18, 14, 18, 14, 18])
d, f = tableau(ws, r, ["Clôture concernée", "Grand livre lu", "Factures du journal SIN dans l'année suivante", "Dont année lisible", "Part non identifiée",
                       "Dont factures de l'année de clôture (arrivées en retard)", "Nombre de ces factures", "Part des factures en retard"],
               [[f"31/12/{x['cloture']}", x["grand_livre"], x["factures_sinistres"], x["references_identifiees"], "=1-D{r}/C{r}", x["references_annee_precedente"],
                 x["nombre_annee_precedente"], "=F{r}/C{r}"] for x in tardives],
               [None, None, FCFA, FCFA, PCT, FCFA, ENT, PCT])
# ---------------------------------------------------------------- 18c
ws, r = onglet(wb, "18c_psap_maladie_estimation", "Provision maladie : montant nécessaire et comparaison avec le bilan", {
    "À quoi sert cet onglet": "Estimer la provision maladie dont le FPPN a besoin, et la comparer à celle de son bilan (2 867 602 878 FCFA, inchangée depuis 2020).",
    "D'où viennent les chiffres": "Onglets 18a et 18b ; provision maladie des balances (comptes 19011000 puis 15801000).",
    "Comment le calcul est fait": "Provision nécessaire = factures déjà reçues et non payées + soins de l'année facturés l'année suivante. Les factures déjà reçues sont déjà inscrites au passif dans le compte 40110000 : elles ne doivent pas être provisionnées une deuxième fois. La provision au compte 158 ne doit donc couvrir que les soins facturés plus tard. Estimation élargie = factures en retard mesurées / part des factures dont l'année est lisible × (prestations maladie / factures de cliniques), pour couvrir aussi les factures sans année lisible et les prestations payées hors du compte 40110000.",
    "Comment lire le résultat": "Si la provision du bilan dépasse largement les soins facturés plus tard, elle est excédentaire : une partie est une réserve qui ne dit pas son nom.",
    "Valable pour": "Clôtures 2022, 2023 et 2024 (années mesurées dans l'onglet 18b), factures de cliniques au journal SIN.",
    "Limites": "Les remboursements aux mutualistes (jusqu'à 9 mois de délai, art. 57) et les prestations payées hors du compte 40110000 ne sont pas mesurés : la base des prestations avec les dates des soins, du dépôt et du paiement reste nécessaire (PBC, item 71)."},
    [16, 18, 18, 14, 16, 18, 18, 18, 12])
lig = []
for x in tardives:
    a = x["cloture"]
    pa = next(p for p in psap_m if p["exercice"] == a + 1)
    lig.append([f"31/12/{a}", "=" + FB("40110000", a, "net_c"), x["references_annee_precedente"],
                x["references_identifiees"] / x["factures_sinistres"], f"={AG('prest_maladie', a + 1)}/{FB('40110000', a + 1, 'mvt_c')}", "=C{r}/D{r}*E{r}",
                "=" + AG("prov_maladie", a), "=G{r}-F{r}", "=G{r}/F{r}"])
d18c, _ = tableau(ws, r, ["Clôture", "Factures reçues non payées (déjà au passif, 40110000)", "Soins de l'année facturés l'année suivante (18b, années lisibles)", "Part des factures dont l'année est lisible",
                "Prestations maladie / factures de cliniques de l'année suivante", "Provision nécessaire estimée (élargie)", "Provision maladie au bilan", "Écart (bilan - nécessaire)", "Bilan / nécessaire (fois)"],
        lig, [None, FCFA, FCFA, PCT, NB2, FCFA, FCFA, FCFA, NB1],
        commentaires={(0, 5): "Estimation élargie : les factures en retard mesurées sont rapportées à la part des factures dont l'année est lisible, puis à l'ensemble des prestations maladie (y compris celles payées hors du compte 40110000). C'est une estimation haute."})
ctl(ws, d18c, "B", [next(p["dettes_cliniques"] for p in psap_m if p["exercice"] == x["cloture"]) for x in tardives])
ctl(ws, d18c, "G", [PV["psap_maladie"][str(x["cloture"])] for x in tardives])

# ---------------------------------------------------------------- 19a
ws, r = onglet(wb, "19a_deces_etats_fin_annee", "Allocations décès : le calcul de fin d'année du FPPN", {
    "À quoi sert cet onglet": "Montrer comment le FPPN calcule chaque 31 décembre les allocations décès restant à payer, et où il les enregistre.",
    "D'où viennent les chiffres": "Fichiers « PROV ALLOCATION DECES 2021.xlsx », « PROV ALLOCATION DECES 2022.xlsx », « ALLOC DECES 2023 PAYER ET A PROVISIONNER .xlsx », « PROV ALLOCATION DECES 2024.xlsx », « PROV ALLOCATIONS DECES ET FRAIS FUNERAIRES 2025.xlsx » (reçus le 22/10/2026) ; crédits du compte 41591000 « Clients créditeurs / ALLOCATION DECES » dans les balances ; « GRAND LIVRE FPPN 2024.xlsx » (écriture du 31/12/2024 « PROVISION ALLO DECES 2024 »).",
    "Comment le calcul est fait": "Le FPPN liste chaque décès non encore payé et lui applique le barème du grade (sous-officier 1 400 000 FCFA, officier 1 600 000, commissaire 2 200 000 ; conjoint 600 000 à 700 000 ; enfant 120 000 à 150 000). Le total est passé au crédit du compte 41591000. On compare le total du fichier au crédit de l'année.",
    "Comment lire le résultat": "Un écart nul prouve que le fichier est bien celui comptabilisé. Un écart signale une différence à expliquer.",
    "Valable pour": "États de fin d'année 2021 à 2025 tels que reçus ; le fichier 2024 porte la mention « BILAN NON EXHAUSTIF ».",
    "Limites": "Le compte 41591000 a été ouvert en 2021 avec un premier crédit qui reprend aussi des dossiers antérieurs : l'écart de 2021 n'est pas une erreur du fichier."},
    [12, 50, 30, 16, 14, 16, 18, 18, 16])
d19a, _ = tableau(ws, r, ["Exercice", "Fichier", "Feuille", "Total imprimé dans le fichier", "Dossiers valorisés", "Somme des lignes", "Crédit de l'année au 41591000", "Écart (crédit - total imprimé)", "Écart lignes - total"],
        [[x["exercice"], x["fichier"], x["feuille"], x["total_imprime"], x["dossiers_valorises"], x["somme_lignes"], "=" + FB("41591000", x["exercice"], "mvt_c"), "=G{r}-D{r}", "=F{r}-D{r}"] for x in etats],
        [ENT, None, None, FCFA, ENT, FCFA, FCFA, FCFA, FCFA])
ctl(ws, d19a, "G", [x["credit_41591000"] for x in etats])
# ---------------------------------------------------------------- 19b
ws, r = onglet(wb, "19b_deces_dette_et_provision", "Allocations décès : dette réelle et provision obsèques", {
    "À quoi sert cet onglet": "Suivre la dette des allocations décès et frais funéraires restant à payer, année par année, et la comparer à la provision obsèques du bilan.",
    "D'où viennent les chiffres": "Comptes 41591000 (allocations décès) et 41592000 (frais funéraires), provision obsèques (19012000 puis 15802000) : " + SRC_BAL + ".",
    "Comment le calcul est fait": "Dette en fin d'année = dette du début + allocations ajoutées (crédits) - allocations payées (débits). Dette décès totale = allocations + frais funéraires. Écart = provision obsèques - dette décès totale.",
    "Comment lire le résultat": "La dette calculée par le FPPN lui-même est très inférieure à la provision obsèques, qui ne bouge pas depuis 2020. La provision ne correspond à aucun calcul transmis : elle double, au moins en partie, une dette déjà inscrite.",
    "Valable pour": "Exercices 2020 à 2025 et situation au 31/08/2026 (balance provisoire).",
    "Limites": "La dette du 41591000 repose sur la liste du FPPN : un décès que le FPPN ignore n'y figure pas (voir 19c)."},
    [12, 16, 16, 16, 16, 16, 16, 16, 18, 18])
d, f = tableau(ws, r, ["Exercice", "Allocations : début", "Allocations ajoutées", "Allocations payées", "Allocations : fin", "Frais funéraires : fin", "Dette décès totale", "Provision obsèques au bilan", "Écart provision - dette", "Provision / dette"],
               [[LIB[x["exercice"]], "=" + FB("41591000", x["exercice"], "ouv_c"), "=" + FB("41591000", x["exercice"], "mvt_c"), "=" + FB("41591000", x["exercice"], "mvt_d"), "=B{r}+C{r}-D{r}",
                 f"={FB('41592000', x['exercice'], 'ouv_c')}+{FB('41592000', x['exercice'], 'mvt_c')}-{FB('41592000', x['exercice'], 'mvt_d')}", "=E{r}+F{r}",
                 "=" + AG("prov_obseques", x["exercice"]), "=H{r}-G{r}", "=IF(G{r}>0,H{r}/G{r},\"-\")"] for x in stock],
               [None, FCFA, FCFA, FCFA, FCFA, FCFA, FCFA, FCFA, FCFA, NB1])
for col, k in (("B", "ouv_41591000"), ("C", "mvt_c_41591000"), ("D", "mvt_d_41591000"), ("H", "provision_obseques")):
    ctl(ws, d, col, [x[k] for x in stock])
ctl(ws, d, "F", [x["ouv_41592000"] + x["mvt_c_41592000"] - x["mvt_d_41592000"] for x in stock])
r = sous_titre(ws, f + 2, "Ancienneté de la dette au 31/08/2026 (règle « premier entré, premier payé »)", 9)
ws.cell(r, 1, "Allocations payées du 01/01/2021 au 31/08/2026 (débits 41591000)")
ws.cell(r, 2, "=" + "+".join(FB("41591000", a, "mvt_d") for a in range(2021, 2027))).number_format = FCFA
CTRLV.append((ws.title, f"B{r}", float(paye)))
PAYE = f"$B${r}"
r += 2
d0 = r + 1
dfi, _ = tableau(ws, r, ["Année d'inscription", "Allocations inscrites", "Payées (premier entré, premier payé)", "Restant dû", "Plus de 3 ans au 31/08/2026 ?"],
        [[x["millesime"], "=" + FB("41591000", x["millesime"], "mvt_c"), (f"=MAX(0,MIN(B{{r}},{PAYE}))" if k == 0 else f"=MAX(0,MIN(B{{r}},{PAYE}-SUM(B${d0}:B{d0 + k - 1})))"),
          "=B{r}-C{r}", "Oui" if x["millesime"] <= 2022 else "Non"] for k, x in enumerate(rest)],
        [ENT, FCFA, FCFA, FCFA, None])
ctl(ws, dfi, "B", [x["dotation"] for x in rest])
ctl(ws, dfi, "C", [x["paye_fifo"] for x in rest])
# ---------------------------------------------------------------- 19c
ws, r = onglet(wb, "19c_deces_declares_tard", "Allocations décès : décès connus après la clôture", {
    "À quoi sert cet onglet": "Mesurer les décès d'une année que le FPPN n'enregistre qu'après le 31 décembre. Ils manquent dans la liste de fin d'année et doivent être ajoutés à la provision (statuts, art. 89).",
    "D'où viennent les chiffres": "Colonnes « DATE D'ENREGISTREMENT » et « DATE DECES » de « PROV ALLOCATION DECES 2021.xlsx » (feuille IMPAYEES GLOBAL 2021) et de « ALLOC DECES 2023 PAYER ET A PROVISIONNER .xlsx » (feuille ALL DECES  A PROVISIONNER 2023).",
    "Comment le calcul est fait": "On retient les dossiers dont le décès date d'une année et l'enregistrement de l'année suivante, puis on additionne leurs montants.",
    "Comment lire le résultat": "Ces montants donnent l'ordre de grandeur de ce qui manque chaque 31 décembre dans la liste du FPPN.",
    "Valable pour": "Fichiers 2021 et 2023, les seuls qui portent les deux dates.",
    "Limites": "Deux années seulement : l'ordre de grandeur est indicatif. Les dossiers sans date d'enregistrement ne sont pas comptés."},
    [12, 16, 18, 18])
d, f = tableau(ws, r, ["État", "Date du décès", "Date d'enregistrement", "Montant"], [[x["etat"], x["deces"], x["enregistrement"], x["montant"]] for x in tard_d],
               [ENT, None, None, FCFA], total=["Total", "", "", "=SUM(D{d}:D{f})"])

# ---------------------------------------------------------------- 20
ws, r = onglet(wb, "20_risques_courts", "Provision pour risques courts (statuts, article 88)", {
    "À quoi sert cet onglet": "Calculer la provision qui couvre les prestations et les frais entre la date de clôture et la prochaine échéance de cotisation.",
    "D'où viennent les chiffres": "Statuts du FPPN édités le 15/06/2023, art. 88 ; prestations et frais de gestion de l'année : " + SRC_BAL + ".",
    "Comment le calcul est fait": "Provision = (prestations + frais de l'année) / 365 × nombre de jours non couverts par des cotisations. Les cotisations sont précomptées chaque mois sur la solde pour le mois en cours : la cotisation de décembre couvre décembre, la prochaine échéance est le 1er janvier. Le nombre de jours non couverts vaut donc 0.",
    "Comment lire le résultat": "Une provision nulle est cohérente avec des cotisations mensuelles qui couvrent leur mois. La cellule orange permet de tester un décalage (par exemple 30 jours si une cotisation arrivait avec un mois de retard).",
    "Valable pour": "Cotisations maladie et décès précomptées chaque mois par la Solde (postes 936 et 8Z6) et par la CGRAE.",
    "Limites": "Si des cotisations étaient encaissées d'avance pour plusieurs mois, la provision ne serait plus nulle : aucun cas n'a été relevé dans les fichiers de précomptes."},
    [12, 20, 20, 20, 16, 20])
tableau(ws, r, ["Exercice", "Prestations maladie et décès", "Frais de gestion", "Coût d'un jour", "Jours non couverts", "Provision pour risques courts"],
        [[LIB[a], f"={AG('prest_maladie', a)}+{AG('prest_deces', a)}", "=" + AG("frais_gestion", a), f"=(B{{r}}+C{{r}})/{MOIS[a] * 365 // 12}", 0, "=D{r}*E{r}"] for a in ANS],
        [None, FCFA, FCFA, FCFA, ENT, FCFA], orange=(4,))
ctl(ws, r + 1, "B", [SIN[a]["prest_maladie"] + SIN[a]["prest_deces"] for a in ANS])
ctl(ws, r + 1, "C", [SIN[a]["frais_gestion"] for a in ANS])

# ---------------------------------------------------------------- 21
ws, r = onglet(wb, "21_epargne_interets", "Épargne : ce que le FPPN doit vraiment, intérêts compris", {
    "À quoi sert cet onglet": "Le solde comptable de l'épargne ne contient aucun intérêt. Cet onglet calcule les intérêts dus depuis le 01/01/2020 selon chaque taux écrit, puis l'engagement réel.",
    "D'où viennent les chiffres": "Comptes 41195100 (PEP) et 41195200 (EFC) : " + SRC_BAL + ". Taux : « MANUEL DE PRODUIT D'EPARGNE FPPN VF VF.docx » (PEP 3,5 % par an, EFC 5,5 % par an, intérêts non capitalisés), « CARTOGRAPHIE DES RISQUES COUVERTS ET DES PRESTATIONS OFFERTES.docx » (EFC 5,5 % à partir de 2024), entretien du chef comptable du 19/10/2026 (3,5 %).",
    "Comment le calcul est fait": "Encours moyen de l'année = (encours du début + encours de fin) / 2. Intérêts de l'année = encours moyen × taux × mois / 12. Les intérêts ne sont pas capitalisés (manuel) : ils s'additionnent d'année en année sans produire eux-mêmes d'intérêts.",
    "Comment lire le résultat": "L'engagement réel = solde comptable + intérêts dus cumulés. Les cellules orange permettent de changer les taux.",
    "Valable pour": "Intérêts dus du 01/01/2020 au 31/08/2026 : c'est une borne basse, car les intérêts dus avant 2020 ne sont pas comptés.",
    "Limites": "Le taux réellement promis n'est pas arrêté par une décision (3,5 %, 5,5 % ou 14 %). On ne sait pas si des intérêts ont été payés hors comptabilité. Le taux de 14 % (intérêt simple, période non précisée) n'est pas repris ici."},
    [12, 16, 16, 16, 16, 16, 16, 16, 16, 16])
r = sous_titre(ws, r, "Taux retenus (modifiables)", 9)
ws.cell(r, 1, "Taux PEP")
ws.cell(r, 2, 0.035).number_format = PCT
ws.cell(r + 1, 1, "Taux EFC")
ws.cell(r + 1, 2, 0.055).number_format = PCT
for rr in (r, r + 1):
    ws.cell(rr, 2).fill = openpyxl.styles.PatternFill("solid", fgColor="EFAC62")
TP, TE = f"$B${r}", f"$B${r + 1}"
r += 3
d, f = tableau(ws, r, ["Exercice", "PEP : début", "PEP : fin", "Intérêts PEP de l'année", "EFC : début", "EFC : fin", "Intérêts EFC de l'année", "Intérêts de l'année", "Intérêts cumulés", "Engagement réel (solde + intérêts cumulés)"],
               [[LIB[e["exercice"]], f"=C{{r}}-{FB('41195100', e['exercice'], 'mvt_c')}+{FB('41195100', e['exercice'], 'mvt_d')}", "=" + FB("41195100", e["exercice"], "net_c"),
                 f"=(B{{r}}+C{{r}})/2*{TP}*{e['mois']}/12", f"=F{{r}}-{FB('41195200', e['exercice'], 'mvt_c')}+{FB('41195200', e['exercice'], 'mvt_d')}", "=" + FB("41195200", e["exercice"], "net_c"),
                 f"=(E{{r}}+F{{r}})/2*{TE}*{e['mois']}/12",
                 "=D{r}+G{r}", "=SUM(H$" + str(r + 1) + ":H{r})", "=C{r}+F{r}+I{r}"] for e in ep],
               [None, FCFA, FCFA, FCFA, FCFA, FCFA, FCFA, FCFA, FCFA, FCFA])
for col, k in (("B", "pep_debut"), ("C", "pep_fin"), ("E", "efc_debut"), ("F", "efc_fin")):
    ctl(ws, d, col, [e[k] for e in ep])

# ---------------------------------------------------------------- 22
ws, r = onglet(wb, "22_sante_retraites", "Santé des policiers retraités : ce que l'on peut mesurer", {
    "À quoi sert cet onglet": "Décrire la garantie santé des retraités : nombre de retraités qui cotisent, cotisation moyenne, évolution.",
    "D'où viennent les chiffres": "Fichiers retour de la CGRAE reçus le 15/10/2026 (février 2024 à septembre 2026) ; « CONVENTION DE PARTENARIAT ENTRE IPS-CGRAE ET LE FPPN.pdf » ; « ETUDE DES RATIO MALADIE DECES 2.xlsx », feuille Notes (cotisations des retraités au 30/06/2026 : 41 520 789 FCFA).",
    "Comment le calcul est fait": "Cotisation moyenne du mois = montant précompté par la CGRAE / nombre de retraités précomptés.",
    "Comment lire le résultat": "La population de retraités cotisants augmente vite. Leurs prestations ne sont pas isolées dans les comptes : on ne peut pas dire si la garantie est équilibrée.",
    "Valable pour": "Retraités précomptés par la CGRAE, février 2024 à septembre 2026.",
    "Limites": "Sans les âges des retraités ni leurs prestations, aucune provision pour cette garantie ne peut être calculée (voir l'annexe des méthodes non retenues)."},
    [14, 16, 18, 18])
d, f = tableau(ws, r, ["Mois", "Retraités précomptés", "Montant précompté", "Cotisation moyenne"],
        [[c["libelle"], f'=SUMIFS({D3Q}!$G:$G,{D3Q}!$A:$A,"CGRAE",{D3Q}!$B:$B,"retour",{D3Q}!$C:$C,{c["annee"]},{D3Q}!$D:$D,{c["mois"]},{D3Q}!$I:$I,"Oui")',
          f'=SUMIFS({D3Q}!$F:$F,{D3Q}!$A:$A,"CGRAE",{D3Q}!$B:$B,"retour",{D3Q}!$C:$C,{c["annee"]},{D3Q}!$D:$D,{c["mois"]},{D3Q}!$I:$I,"Oui")', "=C{r}/B{r}"] for c in R["cgrae"]], [None, ENT, FCFA, FCFA])
ctl(ws, d, "B", [c["retraites"] for c in R["cgrae"]])
ctl(ws, d, "C", [c["montant"] for c in R["cgrae"]])
# dates de naissance des retraités : présentes seulement dans les fichiers retour de la CGRAE de février à avril 2024
_pr = pd.read_parquet(os.path.join(BASE, "precomptes.parquet"), columns=["poste", "sens", "annee", "mois", "matricule_cgrae", "date_naissance"])
_dn0 = _pr[(_pr.poste == "CGRAE") & (_pr.sens == "retour") & _pr.date_naissance.notna()]
_dn = _dn0.drop_duplicates("matricule_cgrae").copy()
_dn["naissance"] = pd.to_datetime(_dn.date_naissance).map(lambda x: x.replace(year=x.year - 100) if x.year > 2026 else x)   # année lue sur deux chiffres
_dn["age"] = (pd.Timestamp("2026-09-30") - _dn.naissance).dt.days / 365.25
_mois_dn = sorted({(int(a), int(m)) for a, m in zip(_dn0.annee, _dn0.mois)})
OUT["retraites_naissance"] = {"nb": int(len(_dn)), "mois": [f"{m:02d}/{a}" for a, m in _mois_dn], "age_moyen": float(_dn.age.mean()), "age_min": float(_dn.age.min()), "age_max": float(_dn.age.max()),
                              "retraites_fin": int(R["cgrae"][-1]["retraites"])}
r = f + 3
r = sous_titre(ws, r, "Retraités dont la date de naissance figure dans les fichiers de la CGRAE", 4)
r = note(ws, r, "Seuls les fichiers retour de la CGRAE de février à avril 2024 portent la date de naissance. L'année y est écrite sur deux chiffres : une année lue après 2026 est ramenée au siècle précédent (par exemple 2062 devient 1962). Âge calculé au 30/09/2026.", 4)
d2, f2 = tableau(ws, r, ["Retraité (matricule CGRAE)", "Date de naissance", "Âge au 30/09/2026"], [[str(m), n.to_pydatetime(), '=(DATE(2026,9,30)-B{r})/365.25'] for m, n in zip(_dn.matricule_cgrae, _dn.naissance)], [None, "DD/MM/YYYY", NB1])
ws.cell(f2 + 1, 1, "Nombre et âge moyen").font = openpyxl.styles.Font(bold=True)
ws.cell(f2 + 1, 2, f"=COUNTA(A{d2}:A{f2})")
ws.cell(f2 + 1, 3, f"=AVERAGE(C{d2}:C{f2})").number_format = NB1

# ---------------------------------------------------------------- 23
ws, r = onglet(wb, "23_grenier_depreciation", "Prêts du Grenier : dépréciation selon sa propre politique de crédit", {
    "À quoi sert cet onglet": "Calculer la perte probable sur les prêts en retard, avec les taux fixés par la politique de crédit du Grenier.",
    "D'où viennent les chiffres": "Balances âgées du logiciel de prêts du Grenier reçues le 15/10/2026 (arrêtés du 31/12/2025 et du 31/08/2026) ; « POLITIK DE CREDIT DU GRENIER VFVF.doc », chapitre XI.3 : 40 % de 3 à 6 mois de retard, 80 % de 6 à 12 mois, 100 % de 12 à 24 mois.",
    "Comment le calcul est fait": "Dépréciation = encours de la tranche de retard × taux de la politique. Les tranches du logiciel (91 à 180 jours, 181 à 365 jours, plus de 365 jours) correspondent aux trois paliers.",
    "Comment lire le résultat": "Le total est le montant que le FPPN aurait dû constater en perte probable. Les comptes ne portent aucune dépréciation des prêts.",
    "Valable pour": "Prêts du Grenier aux arrêtés du 31/12/2025 et du 31/08/2026.",
    "Limites": "Les tranches du logiciel sont en jours de retard sur l'échéance la plus ancienne ; la politique parle de créance impayée « de plus de 3 mois », mesure voisine."},
    [14, 22, 12, 18, 12, 18])
tableau(ws, r, ["Arrêté", "Tranche de retard", "Prêts", "Encours", "Taux de la politique", "Dépréciation"],
        [[x["arrete"], x["tranche"], f'=COUNTIFS({D4Q}!$A:$A,DATE({x["arrete"][:4]},{int(x["arrete"][5:7])},{int(x["arrete"][8:10])}),{D4Q}!$E:$E,">{BORNES[x["tranche"]][0]}",{D4Q}!$E:$E,"<={BORNES[x["tranche"]][1]}")',
          f'=SUMIFS({D4Q}!$D:$D,{D4Q}!$A:$A,DATE({x["arrete"][:4]},{int(x["arrete"][5:7])},{int(x["arrete"][8:10])}),{D4Q}!$E:$E,">{BORNES[x["tranche"]][0]}",{D4Q}!$E:$E,"<={BORNES[x["tranche"]][1]}")',
          x["taux"], "=D{r}*E{r}"] for x in gr], [None, None, ENT, FCFA, PCT, FCFA], orange=(4,))
ctl(ws, r + 1, "C", [x["prets"] for x in gr])
ctl(ws, r + 1, "D", [x["encours"] for x in gr])

# ---------------------------------------------------------------- 24a
ws, r = onglet(wb, "24a_ecoulement_epargne", "Adossement : à quel rythme l'épargne sort-elle ?", {
    "À quoi sert cet onglet": "L'épargne n'a pas d'échéancier. On mesure la part de l'épargne qui sort chaque année (remboursements, règlements EFC, rachats) pour estimer ce qui sortira dans l'année qui vient.",
    "D'où viennent les chiffres": "Débits et soldes des comptes 41195100 (PEP) et 41195200 (EFC) : " + SRC_BAL + ".",
    "Comment le calcul est fait": "Taux de sortie de l'année = sorties de l'année / épargne du début d'année. C'est la méthode de l'« écoulement conventionnel » utilisée pour les dépôts sans échéance (Vintzel, Gestion des risques bancaires, diapositives 33 à 36).",
    "Comment lire le résultat": "Un taux de 20 % veut dire qu'un cinquième de l'épargne du début d'année est sorti dans l'année. La moyenne sert au scénario central, le maximum au test de résistance.",
    "Valable pour": "Exercices 2021 à 2025 (années complètes).",
    "Limites": "Le passé ne garantit pas l'avenir : une perte de confiance peut accélérer les retraits (Vintzel, diapositives 37 et 38)."},
    [12, 18, 18, 18, 14, 14, 14])
d, f = tableau(ws, r, ["Exercice", "Épargne du début", "Sorties de l'année", "Épargne de fin", "Taux de sortie PEP", "Taux de sortie EFC", "Taux de sortie total"],
               [[e["exercice"], f"=D{{r}}-{FB('41195100', e['exercice'], 'mvt_c')}+{FB('41195100', e['exercice'], 'mvt_d')}-{FB('41195200', e['exercice'], 'mvt_c')}+{FB('41195200', e['exercice'], 'mvt_d')}",
                 f"={FB('41195100', e['exercice'], 'mvt_d')}+{FB('41195200', e['exercice'], 'mvt_d')}", f"={FB('41195100', e['exercice'], 'net_c')}+{FB('41195200', e['exercice'], 'net_c')}",
                 f"={FB('41195100', e['exercice'], 'mvt_d')}/({FB('41195100', e['exercice'], 'net_c')}-{FB('41195100', e['exercice'], 'mvt_c')}+{FB('41195100', e['exercice'], 'mvt_d')})",
                 f"={FB('41195200', e['exercice'], 'mvt_d')}/({FB('41195200', e['exercice'], 'net_c')}-{FB('41195200', e['exercice'], 'mvt_c')}+{FB('41195200', e['exercice'], 'mvt_d')})",
                 "=C{r}/B{r}"] for e in ep if 2021 <= e["exercice"] <= 2025],
               [ENT, FCFA, FCFA, FCFA, PCT, PCT, PCT],
               total=["Moyenne", "", "", "", "=AVERAGE(E{d}:E{f})", "=AVERAGE(F{d}:F{f})", "=AVERAGE(G{d}:G{f})"])
MOYS, MAXS = f"'24a_ecoulement_epargne'!$G${f + 1}", f"MAX('24a_ecoulement_epargne'!$G${d}:$G${f})"
_e5 = [e for e in ep if 2021 <= e["exercice"] <= 2025]
ctl(ws, d, "B", [e["pep_debut"] + e["efc_debut"] for e in _e5])
ctl(ws, d, "C", [e["pep_sorties"] + e["efc_sorties"] for e in _e5])
ctl(ws, d, "D", [e["pep_fin"] + e["efc_fin"] for e in _e5])

# ---------------------------------------------------------------- 24b
ws, r = onglet(wb, "24b_ecarts_liquidite", "Adossement : l'argent disponible couvre-t-il ce qu'il faudra payer dans l'année ?", {
    "À quoi sert cet onglet": "Comparer l'actif disponible à moins d'un an et ce que le FPPN devra payer dans l'année : dettes à court terme et sorties d'épargne prévues.",
    "D'où viennent les chiffres": "Classement des comptes du bilan par délai de mobilisation (onglet 10_actif_mobilisation) ; taux de sortie de l'onglet 24a.",
    "Comment le calcul est fait": "Sorties d'épargne prévues = épargne × taux de sortie. Besoin d'un an = dettes à moins d'un an + sorties d'épargne prévues. Écart = actif à moins d'un an - besoin. Couverture = actif à moins d'un an / besoin. Trois scénarios : central (taux moyen), résistance (taux maximal observé), résistance sévère (taux maximal + pic de prestations maladie : hausse annuelle la plus forte observée en 2024 et 2025). Les hausses de 2021 à 2023 ne sont pas retenues : les prestations maladie comptabilisées en 2021 et 2022 (101 et 115 millions de FCFA pour 1,6 à 1,7 milliard de cotisations maladie) sont anormalement basses, alors que les factures de cliniques non payées dépassaient 500 millions ; la hausse de 2023 (+352 %) reflète ce rattrapage, pas un risque.",
    "Comment lire le résultat": "Une couverture inférieure à 100 % signifie qu'il faudra vendre des actifs longs (terrains, clinique) ou utiliser les cotisations futures pour payer.",
    "Valable pour": "Bilans au 31/12/2025 et au 31/08/2026 (balance provisoire).",
    "Limites": "Les dettes à moins d'un an et les actifs sont classés d'après leur nature comptable, sans échéancier contractuel."},
    [24, 18, 18, 18])
mob = Q["mobilisation"]
pic = max(c["prest_maladie"] for c in croiss if c["exercice"] >= 2024)   # années complètes et comparables (voir l'onglet)
OUT["pic_prestations"] = pic
lignes = []
for m in mob:
    lignes.append([m["cloture"], f"={MOB(m['cloture'], 'B')}+{MOB(m['cloture'], 'C')}", "=" + MOB(m["cloture"], "F"), "=" + MOB(m["cloture"], "G")])
d, f = tableau(ws, r, ["Clôture", "Actif à moins d'un an", "Dettes à moins d'un an", "Épargne due"], lignes, [None, FCFA, FCFA, FCFA])
ctl(ws, d, "B", [m["immediat"] + m["moins_1_an"] for m in mob])
ctl(ws, d, "C", [m["passif_moins_1_an"] for m in mob])
ctl(ws, d, "D", [m["epargne_echeance_a_etablir"] for m in mob])
r = f + 2
ws.cell(r, 1, "Hausse maximale observée des prestations maladie (2024 et 2025)")
ws.cell(r, 2, f"=MAX({AG('prest_maladie', 2024)}/{AG('prest_maladie', 2023)}-1,{AG('prest_maladie', 2025)}/{AG('prest_maladie', 2024)}-1)").number_format = PCT
CTRLV.append((ws.title, f"B{r}", float(pic)))
PIC = f"$B${r}"
ws.cell(r + 1, 1, "Prestations maladie 2025")
ws.cell(r + 1, 2, "=" + AG("prest_maladie", 2025)).number_format = FCFA
CTRLV.append((ws.title, f"B{r + 1}", float(SIN[2025]["prest_maladie"])))
PM25 = f"$B${r + 1}"
r += 3
lig2 = []
for i, m in enumerate(mob):
    rr = d + i
    for nom, taux, choc in (("Central (taux moyen)", MOYS, "0"), ("Résistance (taux maximal)", MAXS, "0"), ("Résistance sévère (taux maximal et pic maladie)", MAXS, f"{PM25}*{PIC}")):
        lig2.append([m["cloture"] + " : " + nom, f"=B{rr}", f"=C{rr}+D{rr}*{taux}+{choc}", "=B{r}-C{r}", "=B{r}/C{r}"])
tableau(ws, r, ["Scénario", "Actif à moins d'un an", "Besoin de l'année", "Écart", "Couverture"], lig2, [None, FCFA, FCFA, FCFA, PCT])

# ---------------------------------------------------------------- 24c
ws, r = onglet(wb, "24c_ecarts_par_annee", "Adossement : écarts de liquidité année par année (31/08/2026)", {
    "À quoi sert cet onglet": "Voir, année après année pendant cinq ans, si l'argent qui devient disponible suffit à payer les dettes et les sorties d'épargne prévues.",
    "D'où viennent les chiffres": "Bilan au 31/08/2026 classé par délai de mobilisation (onglet 10_actif_mobilisation) ; taux de sortie moyen de l'onglet 24a ; " + SRC_BAL + ".",
    "Comment le calcul est fait": "Sorties d'épargne de l'année k = épargne × (1 - taux moyen) puissance (k - 1) × taux moyen : chaque année sort la même part de ce qui reste. Actifs disponibles : actifs à moins d'un an en année 1 ; actifs de 1 à 5 ans répartis à parts égales sur les années 2 à 5. Écart de l'année = actifs disponibles - dettes - sorties d'épargne. Écart cumulé = somme des écarts depuis l'année 1. Méthode de l'écart de liquidité statique (Gestion des risques bancaires, diapositives 86 et 87) : aucun nouveau versement n'est compté.",
    "Comment lire le résultat": "Un écart cumulé négatif est l'argent qui manquerait si le FPPN ne recevait plus de nouveaux versements d'épargne ni d'excédents.",
    "Valable pour": "Bilan provisoire au 31/08/2026 ; taux de sortie moyen observé de 2021 à 2025.",
    "Limites": "Les actifs sont classés d'après leur nature comptable, sans échéancier contractuel ; la répartition des actifs de 1 à 5 ans à parts égales est une hypothèse (modifiable en orange)."},
    [18, 20, 18, 18, 18, 18])
m26 = next(m for m in mob if m["cloture"] == "31/08/2026")
r = sous_titre(ws, r, "Paramètres", 6)
dp, fp = tableau(ws, r, ["Paramètre", "Valeur"], [["Actifs à moins d'un an", f"={MOB('31/08/2026', 'B')}+{MOB('31/08/2026', 'C')}"], ["Actifs de 1 à 5 ans", "=" + MOB("31/08/2026", "D")],
                                               ["Actifs à plus de 5 ans ou délai incertain", "=" + MOB("31/08/2026", "E")],
                                               ["Dettes à moins d'un an", "=" + MOB("31/08/2026", "F")], ["Épargne due", "=" + MOB("31/08/2026", "G")], ["Taux de sortie moyen", "=" + MOYS], ["Part des actifs de 1 à 5 ans disponible chaque année (années 2 à 5)", 0.25]],
                 [None, FCFA], orange=())
for k in (dp + 5, dp + 6):
    ws.cell(k, 2).number_format = PCT
ws.cell(dp + 6, 2).fill = openpyxl.styles.PatternFill("solid", fgColor="EFAC62")
A1, A15, A5, DET, EPG, TX, PART = (f"$B${dp + i}" for i in range(7))
ctl(ws, dp, "B", [m26["immediat"] + m26["moins_1_an"], m26["1_a_5_ans"], m26["plus_5_ans_ou_incertain"], m26["passif_moins_1_an"], m26["epargne_echeance_a_etablir"]])
r = fp + 2
lg = []
for k in range(1, 6):
    act = f"={A1}" if k == 1 else f"={A15}*{PART}"
    lg.append([f"Année {k}", act, f"={DET}" if k == 1 else 0, f"={EPG}*(1-{TX})^({k}-1)*{TX}", "=B{r}-C{r}-D{r}", "=E{r}" if k == 1 else "=F{p}+E{r}".replace("{p}", str(r + k - 1))])
lg.append(["Au-delà de 5 ans", f"={A5}", 0, f"={EPG}*(1-{TX})^5", "=B{r}-C{r}-D{r}", "=F{p}+E{r}".replace("{p}", str(r + 5))])
tableau(ws, r, ["Échéance", "Actifs disponibles", "Dettes à payer", "Sorties d'épargne", "Écart de l'année", "Écart cumulé"], lg, [None, FCFA, FCFA, FCFA, FCFA, FCFA])

# ---------------------------------------------------------------- 25
ws, r = onglet(wb, "25_projection_equilibre", "Projection agrégée de l'équilibre technique (2026 à 2030)", {
    "À quoi sert cet onglet": "Projeter, année par année, les cotisations, les prestations, les frais et le résultat technique, pour voir si le FPPN reste équilibré.",
    "D'où viennent les chiffres": "Cotisations, prestations et frais de 2021 à 2025 : " + SRC_BAL + " (onglet 3_sinistralite_resultat).",
    "Comment le calcul est fait": "Chaque poste de 2025 est multiplié chaque année par (1 + son taux de croissance annuel moyen). Scénario A, « tendance » : taux observés de 2023 à 2025, seule période où les prestations sont comparables (les prestations maladie de 2021 et 2022 sont anormalement basses, voir l'onglet 24b). Scénario B, « stabilisation » : les prestations croissent au même rythme que les cotisations. Résultat technique = cotisations - prestations - frais.",
    "Comment lire le résultat": "Un résultat technique négatif signifie que les cotisations ne couvrent plus les prestations et les frais : il faut puiser dans les réserves.",
    "Valable pour": "Garanties maladie et décès prises ensemble, tendances 2023 à 2025 prolongées (scénario A) ou prestations alignées sur les cotisations (scénario B).",
    "Limites": "AVERTISSEMENT : la projection est seulement agrégée. Une projection actuarielle complète part de la population par âge, la fait vieillir d'un an chaque année, applique la mortalité (table CIMA) et les départs en retraite, puis calcule les cotisations et les prestations de chaque âge. Elle bloque dès la première étape : aucun fichier reçu ne donne la date de naissance des mutualistes (les précomptes et le fichier des effectifs ne portent que la date de recrutement). Sans âge, on ne peut ni faire vieillir la population, ni appliquer la mortalité, ni dater les départs en retraite et donc les versements de l'EFC. La projection prolonge donc les tendances passées, sans les expliquer."},
    [14, 18, 18, 18, 18])


def tcam(k0, k1):
    return (k1 / k0) ** (1 / 4) - 1


cot = {a: SIN[a]["cot_maladie"] + SIN[a]["cot_deces"] for a in range(2021, 2026)}
pre = {a: SIN[a]["prest_maladie"] + SIN[a]["prest_deces"] for a in range(2021, 2026)}
fra = {a: SIN[a]["frais_gestion"] for a in range(2021, 2026)}
OUT["tcam"] = {"cotisations": (cot[2025] / cot[2023]) ** 0.5 - 1, "prestations": (pre[2025] / pre[2023]) ** 0.5 - 1, "frais": (fra[2025] / fra[2023]) ** 0.5 - 1}
r = sous_titre(ws, r, "Historique 2021 à 2025 et taux de croissance annuels moyens de 2023 à 2025", 5)
d, f = tableau(ws, r, ["Exercice", "Cotisations", "Prestations", "Frais", "Résultat technique"],
               [[a, f"={AG('cot_maladie_7010', a)}+{AG('cot_deces_7021', a)}", f"={AG('prest_maladie', a)}+{AG('prest_deces', a)}", "=" + AG("frais_gestion", a), "=B{r}-C{r}-D{r}"] for a in range(2021, 2026)],
               [ENT, FCFA, FCFA, FCFA, FCFA],
               total=["Croissance annuelle moyenne 2023 à 2025", "=(B{f}/B" + str(r + 3) + ")^(1/2)-1", "=(C{f}/C" + str(r + 3) + ")^(1/2)-1", "=(D{f}/D" + str(r + 3) + ")^(1/2)-1", ""])
ctl(ws, d, "B", [cot[a] for a in range(2021, 2026)])
ctl(ws, d, "C", [pre[a] for a in range(2021, 2026)])
ctl(ws, d, "D", [fra[a] for a in range(2021, 2026)])
g = f + 1
for c in "BCD":
    ws[f"{c}{g}"].number_format = PCT
proj = {}
for nom, taux_p in (("A", f"C${g}"), ("B", f"B${g}")):
    r = sous_titre(ws, r if nom == "B" else g + 2, "Scénario " + nom + (" : tendance 2023 à 2025" if nom == "A" else " : prestations alignées sur les cotisations"), 5)
    d2 = r + 1
    lig = []
    for k, a in enumerate(range(2026, 2031)):
        p_ = f if k == 0 else d2 + k - 1
        lig.append([a, f"=B{p_}*(1+B${g})", f"=C{p_}*(1+{taux_p})", f"=D{p_}*(1+D${g})", "=B{r}-C{r}-D{r}"])
    _, f2 = tableau(ws, r, ["Exercice", "Cotisations", "Prestations", "Frais", "Résultat technique"], lig, [ENT, FCFA, FCFA, FCFA, FCFA])
    r = f2 + 2
# valeurs pour le rapport (même calcul qu'Excel)
tc, tp, tf = OUT["tcam"]["cotisations"], OUT["tcam"]["prestations"], OUT["tcam"]["frais"]
for nom, tpp in (("A", tp), ("B", tc)):
    c_, p_, f_ = cot[2025], pre[2025], fra[2025]
    proj[nom] = []
    for a in range(2026, 2031):
        c_, p_, f_ = c_ * (1 + tc), p_ * (1 + tpp), f_ * (1 + tf)
        proj[nom].append({"exercice": a, "cotisations": c_, "prestations": p_, "frais": f_, "resultat": c_ - p_ - f_})
OUT["projection"] = proj

wb.calculation.fullCalcOnLoad = True
wb.save(SORTIE_X)
OUT["controles_formules"] = CTRLV
json.dump(OUT, open(SORTIE_J, "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)

print("PSAP maladie :", [(x["cloture"], round(x["references_annee_precedente"])) for x in tardives])
print("décès :", [(x["exercice"], x["total_imprime"], round(x["credit_41591000"])) for x in etats], "tardifs", len(tard_d), sum(x["montant"] for x in tard_d))
print("sorties épargne :", [(s["exercice"], round(s["total_taux"], 4)) for s in sorties], "pic maladie", round(pic, 4), "tcam", {k: round(v, 4) for k, v in OUT["tcam"].items()})
