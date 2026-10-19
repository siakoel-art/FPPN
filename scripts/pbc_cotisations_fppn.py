# -*- coding: utf-8 -*-
"""
FPPN - Demande complémentaire de documents après l'entretien cotisations du 19/10/2026 (items 56 à 70)
Reprend la mise en forme de la demande complémentaire sinistres et prestations (16/10/2026).
Utilisation : python pbc_cotisations_fppn.py <pbc_modele.xlsx> <dossier_base> <sortie.xlsx>
"""
import datetime as dt
import os
import sys
from copy import copy

import openpyxl
import pandas as pd
from openpyxl.worksheet.datavalidation import DataValidation

MODELE, BASE, SORTIE = sys.argv[1:4]
D_DEM, D_LIM = dt.datetime(2026, 10, 19), dt.datetime(2026, 10, 22)
M = ["janv.", "févr.", "mars", "avr.", "mai", "juin", "juil.", "août", "sept.", "oct.", "nov.", "déc."]
MOIS = [(a, m) for a in range(2020, 2027) for m in range(1, 13) if not (a == 2026 and m > 8)]

# Mois sans fichier, par poste et par sens (inventaire des fichiers reçus le 15/10/2026)
inv = pd.read_parquet(os.path.join(BASE, "sources_mois.parquet"))
inv = inv[inv.retenu]


def plages(liste):
    out, s, p = [], None, None
    for x in liste:
        i = MOIS.index(x)
        if s is None:
            s = p = i
        elif i == p + 1:
            p = i
        else:
            out.append((s, p))
            s = p = i
    if s is not None:
        out.append((s, p))
    f = lambda i: f"{M[MOIS[i][1] - 1]} {MOIS[i][0]}"  # noqa: E731
    return ", ".join(f(a) if a == b else f"{f(a)} à {f(b)}" for a, b in out) or "aucun"


POSTES = {"936": ("Maladie", "Direction du personnel de la Police ou Solde"), "8Z6": ("Décès", "Direction du personnel de la Police ou Solde"),
          "870": ("EFC", "Responsable des cotisations"), "998": ("PEP", "Responsable des cotisations"), "9X4": ("Grenier", "Responsable des cotisations"),
          "9W3": ("Flotte automobile GNA", "Responsable des cotisations"), "8Z7": ("MEDOC", "Responsable des cotisations"),
          "CGRAE": ("Retraités CGRAE", "Responsable des cotisations")}
MANQ = {}
for po in POSTES:
    for se in ("aller", "retour"):
        x = inv[(inv.poste == po) & (inv.sens == se)]
        have = set(zip(x.annee, x.mois)) & set(MOIS)
        MANQ[(po, se)] = (len(have), len(MOIS) - len(have), plages([m for m in MOIS if m not in have]))

SECTIONS = [
    ("Barème et règles des produits", [
        ("Barème des cotisations maladie (poste 936) et décès (poste 8Z6) appliqué de 2020 à 2026, avec la base de calcul (salaire de base), la répartition entre maladie et décès, et le document qui notifie ces taux à la Solde.", "Urgent", "Electronique"),
        ("Documentation de chaque produit : EFC, PEP, MEDOC, prêts du Grenier, flotte automobile (règlement, conditions de souscription, montants, conditions de sortie et de remboursement).", "Urgent", "Electronique"),
        ("Décision ou note ayant porté l'EFC de 6 000 FCFA à 10 000 FCFA pour les nouvelles souscriptions depuis 2021, et règle applicable aux anciens souscripteurs qui choisissent 10 000 FCFA.", "Moyen", None),
        ("Texte fixant le délai d'au moins trois mois de cotisation avant le bénéfice des prestations, avec les produits concernés.", "Urgent", None),
        ("Date de démarrage (et, le cas échéant, de suspension) de chaque produit : EFC, PEP, MEDOC, prêts du Grenier, flotte automobile, cotisations des retraités (CGRAE).", "Moyen", None),
    ]),
    ("Fichiers de précomptes", [
        ("Fichiers aller maladie (poste 936) et décès (poste 8Z6), de janvier 2020 à août 2026, transmis par la direction du personnel de la Police à la Solde (aucun fichier aller reçu sur 80 mois, voir l'onglet « Fichiers manquants »).", "Urgent", "Electronique"),
        ("Fichiers retour maladie et décès manquants (voir l'onglet « Fichiers manquants »), et description du canal par lequel le FPPN les reçoit.", "Urgent", "Electronique"),
        ("Fichiers aller et retour manquants de l'EFC, du PEP, du Grenier, de MEDOC, de la flotte automobile et des retraités (CGRAE), mois par mois (voir l'onglet « Fichiers manquants »).", "Urgent", "Electronique"),
        ("États des motifs de rejet reçus de la Solde (environ deux derniers mois), et tout état antérieur disponible.", "Moyen", "Electronique"),
        ("Liste des paiements manuels de cotisations de 2020 à août 2026 : matricule, produit, mois couvert, montant, date et heure d'enregistrement.", "Moyen", "Electronique"),
    ]),
    ("Mutualistes et ayants droit", [
        ("Fichier des ayants droit au 31/08/2026, au format de l'onglet « Modèle ayants droit ».", "Urgent", "Electronique"),
        ("Table de correspondance entre matricule et mécano des policiers, au format de l'onglet « Modèle matricule mécano », fournie par la direction du personnel de la Police.", "Moyen", "Electronique"),
    ]),
    ("Autres produits", [
        ("États de GNA sur la flotte automobile : primes reçues, soldes par adhérent, trop-perçus remboursés, de 2020 à août 2026, et convention avec GNA.", "Moyen", "Electronique"),
        ("Documentation de MEDOC : conventions avec les pharmacies, période de fonctionnement, décision de suspension et état des prises en charge.", "Moyen", None),
        ("Liste des types de prêts du Grenier et conditions de chaque type (montant, durée, taux, remboursement), en rapport avec les huit types de crédit présents dans les fichiers de production.", "Moyen", None),
    ]),
]

wb = openpyxl.load_workbook(MODELE)
ws = wb.worksheets[0]
ws.title = "PBC Cotisations"
ws["C4"] = "AUDIT DU FPPN 2020 – 31 AOÛT 2026 : DEMANDE COMPLÉMENTAIRE COTISATIONS ET PRÉCOMPTES"
ws["C5"] = "SUITE À L'ENTRETIEN DU 19/10/2026 AVEC LE RESPONSABLE DES COTISATIONS"
COLS = range(2, 14)
st_vide = {c: copy(ws.cell(9, c)._style) for c in COLS}
st_sect = {c: copy(ws.cell(10, c)._style) for c in COLS}
st_item = {c: copy(ws.cell(11, c)._style) for c in COLS}
for r in range(10, ws.max_row + 1):
    for c in COLS:
        ws.cell(r, c).value = None
        ws.cell(r, c)._style = copy(st_vide[c])
    ws.row_dimensions[r].height = 13.5
ws.data_validations.dataValidation = []

r, prec, premier = 10, None, None
for titre, items in SECTIONS:
    for c in COLS:
        ws.cell(r, c)._style = copy(st_sect[c])
    ws.cell(r, 3, titre)
    ws.row_dimensions[r].height = 18
    r += 1
    for texte, prio, support in items:
        for c in COLS:
            ws.cell(r, c)._style = copy(st_item[c])
        ws.cell(r, 2, "=56" if prec is None else f"=+B{prec}+1")
        ws.cell(r, 3, texte)
        ws.cell(r, 5, prio)
        ws.cell(r, 7, D_DEM)
        ws.cell(r, 8, D_LIM)
        ws.cell(r, 10, support)
        ws.cell(r, 13, "Non Reçu")
        ws.row_dimensions[r].height = 21 if len(texte) < 110 else 36 if len(texte) < 230 else 51
        premier = premier or r
        prec = r
        r += 1
    r += 1
dern = prec
ws["M7"] = f'=IFERROR(COUNTIF($M${premier}:$M${dern},"Reçu")/(COUNTA($M${premier}:$M${dern})-COUNTIF($M${premier}:$M${dern},"Non Applicable")),0)'
for col, liste in (("E", '"Urgent,Moyen,Faible"'), ("J", '"Physique,Electronique,Physique et Electronique"'),
                   ("M", '"Reçu,Non Reçu,Non Applicable,Partiellement Reçu"')):
    dv = DataValidation(type="list", formula1=liste, allow_blank=True)
    dv.add(f"{col}{premier}:{col}{dern}")
    ws.add_data_validation(dv)

# Onglets annexes : on garde la mise en forme d'un onglet modèle
mod = wb["Modèle données décès"]
st = {k: copy(mod[k]._style) for k in ("B2", "B3", "B5", "B6", "C6", "D6")}
larg = {k: mod.column_dimensions[k].width for k in "BCD"}


def annexe(nom, titre, note, entetes, lignes, largeurs=None):
    a = wb.create_sheet(nom)
    a.sheet_view.showGridLines = mod.sheet_view.showGridLines
    a["B2"], a["B3"] = titre, note
    a["B2"]._style, a["B3"]._style = copy(st["B2"]), copy(st["B3"])
    for j, e in enumerate(entetes):
        x = a.cell(5, 2 + j, e)
        x._style = copy(st["B5"])
    for i, l in enumerate(lignes, 6):
        for j, v in enumerate(l):
            x = a.cell(i, 2 + j, v)
            x._style = copy(st["B6"] if j == 0 else st["C6"])
    for j, w in enumerate(largeurs or [larg["B"], larg["C"], larg["D"]]):
        a.column_dimensions[openpyxl.utils.get_column_letter(2 + j)].width = w
    return a


lignes = [[f"{POSTES[po][0]} ({po})" if po != "CGRAE" else POSTES[po][0], MANQ[(po, "aller")][0], MANQ[(po, "aller")][2],
           MANQ[(po, "retour")][0], MANQ[(po, "retour")][2], POSTES[po][1]] for po in POSTES]
annexe("Fichiers manquants", "Fichiers de précomptes reçus et mois manquants, janvier 2020 à août 2026 (80 mois)",
       "Source : inventaire Adenka des fichiers reçus le 15/10/2026. Une partie des mois manquants peut correspondre à des périodes où le produit n'existait pas encore (PBC, item 60).",
       ["Poste", "Mois aller reçus", "Mois sans fichier aller", "Mois retour reçus", "Mois sans fichier retour", "Fichier aller à demander à"],
       lignes, [24, 14, 60, 14, 60, 40])
annexe("Modèle ayants droit", "Modèle ayants droit",
       "Une ligne par ayant droit. Une colonne par information. Dates au format JJ/MM/AAAA.",
       ["Colonne", "Contenu attendu", "Exemple"],
       [["Matricule", "Matricule du policier mutualiste", "123456A"],
        ["Mecano", "Mécano du policier mutualiste", "Code mécano"],
        ["Lien", "Conjoint, enfant, autre", "Enfant"],
        ["Rang", "Rang de l'ayant droit (1er enfant, 2e enfant, etc.)", "2"],
        ["Sexe", "Sexe de l'ayant droit", "F"],
        ["Date_naissance", "Date de naissance de l'ayant droit, si elle est connue", "20/09/2012"],
        ["Date_inscription", "Date d'inscription comme ayant droit", "01/03/2021"],
        ["Date_sortie", "Date de sortie (décès, majorité, divorce), vide si toujours couvert", ""],
        ["Statut", "Couvert ou sorti au 31/08/2026", "Couvert"]])
annexe("Modèle matricule mécano", "Modèle matricule mécano",
       "Une ligne par policier. Une colonne par information. Dates au format JJ/MM/AAAA.",
       ["Colonne", "Contenu attendu", "Exemple"],
       [["Matricule", "Matricule du policier", "123456A"],
        ["Mecano", "Mécano attribué par la Solde", "Code mécano"],
        ["Date_attribution", "Date d'attribution du mécano", "15/02/2025"],
        ["Date_entree_ecole", "Date d'entrée à l'école de police", "01/09/2024"],
        ["Statut", "Élève, actif ou retraité au 31/08/2026", "Actif"]])
for nom in ("Modèle données décès", "Modèle données santé"):
    del wb[nom]
wb.active = 0
wb.save(SORTIE)
print(f"items 56 à {56 + sum(len(i) for _, i in SECTIONS) - 1}, lignes {premier}-{dern}")
